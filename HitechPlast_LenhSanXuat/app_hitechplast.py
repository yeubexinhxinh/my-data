# -*- coding: utf-8 -*-
"""
PHAN MEM TAO LENH SAN XUAT - HitechPlast (ban Python, giao dien nhieu tab)
Viet lai tu file Excel VBA "HitechPlast order creation software in 2026".

Chay:  python app_hitechplast.py
Yeu cau: Python 3.8+ (Tkinter co san). Xuat Excel can: pip install openpyxl
Du lieu: data.json (don hang) + sheets.json (cac sheet) cung thu muc.
"""
import json
import os
import sys
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import time
import threading

# MA HOA FILE DU LIEU (json / qc_data.db) - phai bat TRUOC moi thao tac doc/ghi file (yeu cau #5)
try:
    import file_crypt
    file_crypt.install(migrate=False)
except Exception as _e:
    print('Khong bat duoc ma hoa file:', _e)



# Thoi gian kiem tra may chu (ms) de TU CAP NHAT du lieu may khac vua sua
AUTO_SYNC_MS = 5000
OFFLINE_RETRY_MS = 30000          # dang offline: 30 giay thu ket noi may chu 1 lan
OFFLINE_CACHE_MS = 20 * 60000     # dang online: 20 phut cap nhat ban sao offline 1 lan

# Dem so hop thoai (messagebox / chon file) dang mo -> luc do KHONG tu cap nhat du lieu
_MODAL_DEPTH = [0]


def _track_modal(fn):
    def w(*a, **k):
        _MODAL_DEPTH[0] += 1
        try:
            return fn(*a, **k)
        finally:
            _MODAL_DEPTH[0] -= 1
    w.__wrapped__ = fn
    return w


for _mod, _names in ((messagebox, ('showinfo', 'showwarning', 'showerror', 'askyesno', 'askokcancel',
                                   'askyesnocancel', 'askquestion', 'askretrycancel')),
                     (filedialog, ('askopenfilename', 'askopenfilenames', 'asksaveasfilename',
                                   'askdirectory', 'askopenfile', 'asksaveasfile'))):
    for _n in _names:
        _f = getattr(_mod, _n, None)
        if _f is not None and not hasattr(_f, '__wrapped__'):
            setattr(_mod, _n, _track_modal(_f))
import tkinter.font as tkfont

from engine import ProductionOrder
import license_code as lic
import secure_store
import datastore
import doc_params
from material_rules import (group_levels, is_header_row, NvlCatalog,
                            check_order)
import i18n
i18n.install()          # ngon ngu giao dien (Tieng Viet / English) - doc tu settings.json

# Cac dong bi an tren FORM Lenh san xuat (row 11='Tong don', row 13='Alternative materials.')
HIDDEN_FORM_ROWS = {11, 13}
# Cac cot dieu khien (VBA) an tren FORM: P-T = 16-20 (So Cont/So Kg-Pallet/Loai Bao/CheekNVL/AuTo)
# an them cot 15 (O, trong) de thu gon phan thua ben phai form
HIDDEN_FORM_COLS = {15, 16, 17, 18, 19, 20}

# Tai lieu co KHUNG NGOAI dam bao quanh (giong Excel goc)
FRAMED_DOC_KEYS = {'msds'}

# Sheet hien dang nhom +/- (Outline) giong Excel
GROUPED_SHEETS = {'Đơn giá vật tư'}
# Sheet Don gia vat tu (Unit price): STT tu danh lai + Ma hoa dua vao SX khong duoc trung
UNIT_PRICE_TITLE = 'Đơn giá vật tư'
REPLACE_TITLE = 'NVL thay thế'     # sheet Replace (NVL thay the)
# Sheet hien ten tab bang TEN TIENG VIET (khong lay ten sheet nguon tieng Anh)
VI_TAB_TITLE = {UNIT_PRICE_TITLE}
# Khoa cu cua tab Unit price trong settings (tab_names / tab_order) -> khoa moi
TAB_KEY_ALIAS = {'Unit price': UNIT_PRICE_TITLE, 'Đơn Giá': UNIT_PRICE_TITLE}
UP_COL_STT, UP_COL_CODE, UP_COL_MAVT, UP_COL_TEN = 0, 1, 2, 3


def _cell_str(x):
    return '' if x is None else str(x).strip()


def _cell_disp(x):
    """Gia tri HIEN THI tren bang: so thuc bo duoi sai so (2.5400000000000005 -> 2.54,
    389.5484126984127 -> 389.548412698). Du lieu luu khong doi."""
    if isinstance(x, float):
        try:
            t = '%.12g' % x
            if 'e' in t or 'E' in t:
                return str(x)
            return t if '.' in t else t + '.0'
        except Exception:
            return str(x)
    return _cell_str(x)


def norm_code(x):
    """Chuan hoa Ma hoa de so trung: bo khoang trang thua, khong phan biet hoa/thuong."""
    return ' '.join(str(x if x is not None else '').split()).lower()


def up_is_group_row(r):
    """Dong tieu de nhom cua Unit price: khong co Ma hoa, khong co Ma vat tu, chi co Ten."""
    return (len(r) > UP_COL_TEN and _cell_str(r[UP_COL_CODE]) == ''
            and _cell_str(r[UP_COL_MAVT]) == '' and _cell_str(r[UP_COL_TEN]) != '')


def renumber_unit_price(rows):
    """Danh lai STT 1,2,3... tu tren xuong duoi cho cac dong vat tu (bo qua dong tieu de nhom).
    Tra ve True neu co thay doi."""
    n = 0
    changed = False
    for r in rows:
        if not r:
            continue
        if up_is_group_row(r):
            if _cell_str(r[UP_COL_STT]) != '':
                r[UP_COL_STT] = ''
                changed = True
            continue
        if not any(_cell_str(x) for x in r[1:]):
            continue                     # dong trong hoan toan
        n += 1
        if r[UP_COL_STT] != n:
            r[UP_COL_STT] = n
            changed = True
    return changed


def find_dup_code(rows, code, skip_idx=None):
    """Tim dong khac co cung Ma hoa dua vao SX. Tra ve (idx, row) hoac None."""
    key = norm_code(code)
    if not key:
        return None
    for i, r in enumerate(rows):
        if i == skip_idx or len(r) <= UP_COL_CODE:
            continue
        if norm_code(r[UP_COL_CODE]) == key:
            return i, r
    return None

FONT = ('Segoe UI', 10)
FONT_B = ('Segoe UI', 10, 'bold')
FONT_TITLE = ('Segoe UI', 15, 'bold')


def resource_path(name):
    """Tim file du lieu. Uu tien file dat CANH .exe/.py (de sua khong can build lai),
    neu khong co moi dung ban nhung trong .exe (PyInstaller _MEIPASS)."""
    # 1) thu muc chua file chay (canh .exe hoac .py)
    if getattr(sys, 'frozen', False):
        ext_dir = os.path.dirname(sys.executable)
    else:
        ext_dir = os.path.dirname(os.path.abspath(__file__))
    ext = os.path.join(ext_dir, name)
    if os.path.exists(ext):
        return ext
    # 2) ban nhung trong .exe
    base = getattr(sys, '_MEIPASS', ext_dir)
    return os.path.join(base, name)


_ICON_IMG = []


def _apply_app_icon(win):
    """Bieu tuong HitechPlast cho cua so (thay hinh long chim mac dinh cua Tk).
    iconbitmap(default=...) -> moi cua so con (Toplevel) cua root nay cung dung."""
    try:
        ico = resource_path('hitechplast.ico')
        if os.path.exists(ico) and sys.platform.startswith('win'):
            win.iconbitmap(default=ico)
            return
        png = resource_path('hitechplast_icon_256.png')
        if os.path.exists(png):
            img = tk.PhotoImage(master=win, file=png)
            _ICON_IMG.append(img)
            win.iconphoto(True, img)
    except Exception:
        pass


_TK_INIT = tk.Tk.__init__


def _tk_init_with_icon(self, *a, **k):
    _TK_INIT(self, *a, **k)
    _apply_app_icon(self)


tk.Tk.__init__ = _tk_init_with_icon      # ap dung cho moi cua so goc (dang nhap, phan mem chinh, hop loi...)


# =====================================================================
#  TAB 1: LENH SAN XUAT (co tinh toan)
# =====================================================================
class OrderTab(tk.Frame):
    def __init__(self, master, app):
        super().__init__(master, bg='#f4f6f9')
        self.app = app
        self.data = app.data
        self.current_order = None
        self._filtered = []
        self._build()
        self._reload_orders()

    @property
    def cont_index(self):
        return getattr(self.app, 'cv_index', {})

    def _orders_of_kind(self):
        return self.data['orders'].get(self.kind.get(), [])

    def _build(self):
        bar = tk.Frame(self, bg='#f4f6f9')
        bar.pack(fill='x', padx=12, pady=(12, 6))

        tk.Label(bar, text='Loại đơn:', bg='#f4f6f9', font=FONT_B).grid(row=0, column=0, sticky='w')
        self.kind = tk.StringVar(value='Domestic')
        km = ttk.Combobox(bar, textvariable=self.kind, width=11, state='readonly',
                          values=['Domestic', 'Export', 'Sample'], font=FONT)
        km.grid(row=0, column=1, padx=(4, 14))
        km.bind('<<ComboboxSelected>>', lambda e: self._reload_orders())

        tk.Label(bar, text='Tìm:', bg='#f4f6f9', font=FONT_B).grid(row=0, column=2, sticky='w')
        self.search = tk.StringVar()
        se = tk.Entry(bar, textvariable=self.search, width=18, font=FONT)
        se.grid(row=0, column=3, padx=(4, 8))
        se.bind('<KeyRelease>', lambda e: self._filter_orders())

        tk.Label(bar, text='Đơn hàng:', bg='#f4f6f9', font=FONT_B).grid(row=0, column=4, sticky='w')
        self.order_sel = tk.StringVar()
        self.order_cb = ttk.Combobox(bar, textvariable=self.order_sel, width=44,
                                     state='readonly', font=FONT)
        self.order_cb.grid(row=0, column=5, padx=(4, 14))
        self.order_cb.bind('<<ComboboxSelected>>', lambda e: self._select_order())

        tk.Label(bar, text='Sản lượng (kg):', bg='#f4f6f9', font=FONT_B).grid(row=0, column=6, sticky='w')
        self.san_luong = tk.StringVar(value='0')
        sl = tk.Entry(bar, textvariable=self.san_luong, width=11, font=FONT_B)
        sl.grid(row=0, column=7, padx=(4, 8))
        sl.bind('<KeyRelease>', lambda e: (self._recalc(), self._calc_xep()))

        tk.Label(bar, text='Kg/Pallet:', bg='#f4f6f9', font=FONT).grid(row=0, column=8, sticky='w')
        self.kg_pallet = tk.StringVar(value='1000')
        kp = tk.Entry(bar, textvariable=self.kg_pallet, width=7, font=FONT)
        kp.grid(row=0, column=9, padx=(4, 4))
        kp.bind('<KeyRelease>', lambda e: (self._recalc(), self._calc_xep()))

        # --- Hang 2: Loai bao / Quy cach xep bao / Tong so cont ---
        tk.Label(bar, text='Loại bao:', bg='#f4f6f9', font=FONT_B).grid(row=1, column=0, sticky='w', pady=(6, 0))
        self.loai_bao = tk.StringVar(value='Bao 25Kg')
        lb_cb = ttk.Combobox(bar, textvariable=self.loai_bao, width=12, state='readonly', font=FONT,
                             values=['Bao 25Kg', 'Bao Zumbo'])
        lb_cb.grid(row=1, column=1, sticky='w', padx=(4, 14), pady=(6, 0))
        lb_cb.bind('<<ComboboxSelected>>', lambda e: self._calc_xep())

        tk.Label(bar, text='Quy cách xếp bao:', bg='#f4f6f9', font=FONT_B).grid(row=1, column=2, columnspan=2, sticky='w', pady=(6, 0))
        self.quy_cach = tk.StringVar(value='Hàng Xếp Rời')
        ttk.Combobox(bar, textvariable=self.quy_cach, width=16, state='readonly', font=FONT,
                     values=['Hàng Xếp Rời', 'Hàng Xếp Pallet']).grid(row=1, column=4, sticky='w', padx=(4, 14), pady=(6, 0))

        tk.Label(bar, text='Tổng số Cont:', bg='#f4f6f9', font=FONT_B).grid(row=1, column=6, sticky='w', pady=(6, 0))
        self.so_cont = tk.StringVar(value='')
        ec = tk.Entry(bar, textvariable=self.so_cont, width=11, font=FONT_B)
        ec.grid(row=1, column=7, sticky='w', padx=(4, 8), pady=(6, 0))
        ec.bind('<KeyRelease>', lambda e: (self._recalc(), self._calc_xep()))

        tk.Label(bar, text='% Hao hụt:', bg='#f4f6f9', font=FONT_B, fg='#b8860b').grid(row=1, column=8, sticky='w', pady=(6, 0))
        self.hao_hut = tk.StringVar(value='0.7')
        eh = tk.Entry(bar, textvariable=self.hao_hut, width=7, font=FONT_B)
        eh.grid(row=1, column=9, sticky='w', padx=(4, 4), pady=(6, 0))
        eh.bind('<KeyRelease>', lambda e: self._recalc())

        # --- Hang 3: Cach xep pallet (tu hien khi Quy cach = Hang Xep Pallet) ---
        tk.Label(bar, text='Cách xếp pallet:', bg='#f4f6f9', font=FONT_B, fg='#1d3557').grid(
            row=2, column=0, sticky='w', pady=(6, 0))
        self.cach_xep = tk.StringVar(value='')
        exp = tk.Entry(bar, textvariable=self.cach_xep, width=64, font=FONT)
        exp.grid(row=2, column=1, columnspan=5, sticky='w', padx=(4, 8), pady=(6, 0))
        tk.Button(bar, text='↻ Tự tính', font=('Segoe UI', 9),
                  command=self._calc_xep).grid(row=2, column=6, sticky='w', pady=(6, 0))

        # Day chuyen (o N6 tren form, cac trang sau =N6). VD: TC2,TC5,TC6
        tk.Label(bar, text='Dây chuyền:', bg='#f4f6f9', font=FONT_B, fg='#1d3557').grid(
            row=2, column=7, sticky='e', pady=(6, 0))
        self.day_chuyen = tk.StringVar(value='')
        edc = tk.Entry(bar, textvariable=self.day_chuyen, width=16, font=FONT_B)
        edc.grid(row=2, column=8, columnspan=2, sticky='w', padx=(4, 4), pady=(6, 0))
        edc.bind('<FocusOut>', lambda e: self._save_day_chuyen())
        edc.bind('<Return>', lambda e: self._save_day_chuyen())

        # --- Hang 4: Xep lan loai hang (toi da 3 hang phu) ---
        self.mix_var = tk.BooleanVar(value=False)
        tk.Checkbutton(bar, text='Xếp lẫn loại hàng', variable=self.mix_var, bg='#f4f6f9',
                       font=FONT_B, fg='#8b0000', command=self._calc_xep).grid(
            row=3, column=0, columnspan=2, sticky='w', pady=(4, 0))
        self.ma_phu = []
        self.kg_phu = []
        for k in range(3):
            r = 3 + k
            c0 = 2 if k == 0 else 0
            if k == 0:
                cc = 2
            else:
                tk.Label(bar, text=f'Hàng phụ {k+1}:', bg='#f4f6f9', font=FONT).grid(
                    row=r, column=0, columnspan=2, sticky='e', pady=(2, 0))
                cc = 2
            tk.Label(bar, text=('Mã hàng phụ:' if k == 0 else 'Mã:'), bg='#f4f6f9', font=FONT).grid(
                row=r, column=cc, sticky='e', pady=(2, 0))
            vma = tk.StringVar(value='')
            ema = tk.Entry(bar, textvariable=vma, width=16, font=FONT)
            ema.grid(row=r, column=cc + 1, columnspan=2, sticky='w', padx=(4, 8), pady=(2, 0))
            ema.bind('<KeyRelease>', lambda e: self._calc_xep())
            tk.Label(bar, text='KL (kg):', bg='#f4f6f9', font=FONT).grid(row=r, column=6, sticky='e', pady=(2, 0))
            vkg = tk.StringVar(value='')
            ekg = tk.Entry(bar, textvariable=vkg, width=10, font=FONT)
            ekg.grid(row=r, column=7, sticky='w', padx=(4, 8), pady=(2, 0))
            ekg.bind('<KeyRelease>', lambda e: self._calc_xep())
            self.ma_phu.append(vma)
            self.kg_phu.append(vkg)

        # doi Quy cach -> tu tinh lai
        self.quy_cach.trace_add('write', lambda *a: self._calc_xep())

        info = tk.Frame(self, bg='#eaf0f7')
        info.pack(fill='x', padx=12, pady=(0, 6))
        self.lbl_info = tk.Label(info, text='Chọn một đơn hàng để bắt đầu…',
                                 bg='#eaf0f7', font=FONT, anchor='w', justify='left')
        self.lbl_info.pack(fill='x', padx=10, pady=6)

        mid = tk.Frame(self, bg='#f4f6f9')
        mid.pack(fill='both', expand=True, padx=12)
        cols = ('stt', 'ten', 'ma_nvl', 'pct', 'mass', 'tong_don', 'total')
        heads = ['STT', 'Nguyên vật liệu', 'Mã NVL', 'Tỉ lệ %',
                 'KL/mẻ (kg)', 'Tổng đơn (kg)', 'Total order (kg)']
        widths = [45, 250, 90, 85, 110, 140, 150]
        self.tree = ttk.Treeview(mid, columns=cols, show='headings', height=13)
        for c, h, w in zip(cols, heads, widths):
            self.tree.heading(c, text=h)
            # can GIUA tat ca cac cot -> so lieu thang hang voi tieu de (yeu cau 27/09/2026)
            self.tree.column(c, width=w, anchor='center')
        self.tree.pack(side='left', fill='both', expand=True)
        sb = ttk.Scrollbar(mid, orient='vertical', command=self.tree.yview)
        sb.pack(side='right', fill='y')
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.tag_configure('total', background='#fff3cd')

        sf = tk.Frame(self, bg='#f4f6f9')
        sf.pack(fill='x', padx=12, pady=10)
        self.cards = {}
        specs = [('co_me', 'Cỡ mẻ (kg/mẻ)', '#2d6a4f'), ('so_me', 'Số mẻ', '#1d3557'),
                 ('sltt', 'SL thực tế (kg)', '#6a4c93'), ('cont', 'Số Cont', '#9d4edd'),
                 ('pallet', 'Số Pallet', '#e07a00'), ('tong_pct', 'Tổng tỉ lệ %', '#495057')]
        for i, (k, lbl, color) in enumerate(specs):
            card = tk.Frame(sf, bg=color, height=62)
            card.grid(row=0, column=i, padx=5, sticky='nsew')
            card.pack_propagate(False)
            tk.Label(card, text=lbl, bg=color, fg='white', font=('Segoe UI', 9)).pack(pady=(8, 0))
            v = tk.Label(card, text='–', bg=color, fg='white', font=('Segoe UI', 15, 'bold'))
            v.pack()
            self.cards[k] = v
            sf.grid_columnconfigure(i, weight=1)

        bot = tk.Frame(self, bg='#f4f6f9')
        bot.pack(fill='x', padx=12, pady=(0, 12))
        self.btn_excel = tk.Button(bot, text='⭳  Xuất Lệnh sản xuất (Excel)', font=FONT_B,
                                   bg='#2d6a4f', fg='white', relief='flat', padx=14, pady=7,
                                   command=self._export_excel)
        self.btn_excel.pack(side='right')
        self.btn_pdf = tk.Button(bot, text='📄  Xuất PDF', font=FONT_B, bg='#b0413e', fg='white',
                                 relief='flat', padx=14, pady=7, command=self._export_pdf)
        self.btn_pdf.pack(side='right', padx=8)
        self.btn_print = tk.Button(bot, text='🖨️  In', font=FONT_B, bg='#1d3557', fg='white',
                                   relief='flat', padx=14, pady=7, command=self._print)
        self.btn_print.pack(side='right', padx=(0, 8))
        self.btn_preview = tk.Button(bot, text='🖨  Xem trước', font=FONT, bg='#1d3557', fg='white',
                                     relief='flat', padx=14, pady=7, command=self._preview)
        self.btn_preview.pack(side='right', padx=8)
        tk.Button(bot, text='📁  Đường dẫn ảnh (Tem/Bao bì)', font=FONT, relief='flat',
                  padx=12, pady=7, command=self._open_img_settings).pack(side='left')
        self.btn_stock = tk.Button(bot, text='📦  Kiểm tra tồn kho', font=FONT_B, bg='#6a4c93', fg='white',
                                   relief='flat', padx=14, pady=7, command=self._check_stock)
        self.btn_stock.pack(side='left', padx=8)
        # canh bao khi ton kho qua cu (> 5 ngay) -> cac nut tren bi khoa
        self.lbl_stock_lock = tk.Label(self, text='', bg='#fdecea', fg='#b00020',
                                       font=('Segoe UI', 9, 'bold'), anchor='w', justify='left')
        self._stock_btns = {self.btn_excel: '#2d6a4f', self.btn_pdf: '#b0413e', self.btn_print: '#1d3557',
                            self.btn_preview: '#1d3557', self.btn_stock: '#6a4c93'}
        self._stock_bot = bot
        self.after(300, self._refresh_stock_lock)

    # ---------------- KHOA THEO TON KHO (qua 5 ngay chua nhan dien) ----------------
    def _stock_lock_reason(self):
        try:
            import stock_import
        except Exception:
            return ''
        by = {sh.get('title'): sh for sh in self.app.sheets}
        return stock_import.stock_lock_reason(by.get(UNIT_PRICE_TITLE))

    def _refresh_stock_lock(self, repeat=True):
        """Bat/tat cac nut Kiem tra ton kho / Xem truoc / In / Xuat PDF / Xuat lenh."""
        try:
            reason = self._stock_lock_reason()
            for b, color in self._stock_btns.items():
                if reason:
                    b.config(state='disabled', bg='#b8bec6', cursor='')
                else:
                    b.config(state='normal', bg=color, cursor='hand2')
            if reason:
                self.lbl_stock_lock.config(text='🔒 ĐÃ KHOÁ in / xuất lệnh — ' + reason.replace('\n', ' '))
                if not self.lbl_stock_lock.winfo_ismapped():
                    self.lbl_stock_lock.pack(fill='x', padx=12, pady=(0, 4), before=self._stock_bot)
            elif self.lbl_stock_lock.winfo_ismapped():
                self.lbl_stock_lock.pack_forget()
        except Exception:
            pass
        if repeat:
            try:
                self.after(60000, self._refresh_stock_lock)     # tu kiem tra lai moi phut
            except Exception:
                pass

    def _stock_blocked(self):
        """True = dang bi khoa (da hien thong bao)."""
        reason = self._stock_lock_reason()
        if reason:
            self._refresh_stock_lock(repeat=False)
            messagebox.showwarning('Tồn kho quá cũ', reason)
            return True
        return False

    def _check_stock(self):
        """KiemTraTonKho_ThemThayThe (VBA) - chi HIEN bang ket qua, khong ghi thay the vao lenh."""
        if self._stock_blocked():
            return
        if not self.current_order:
            messagebox.showwarning('Chưa chọn đơn', 'Vui lòng chọn một đơn hàng.')
            return
        try:
            sl = float(self.san_luong.get() or 0)
        except ValueError:
            sl = 0
        self.current_order['hao_hut'] = self._hao()
        try:
            import stock_check
            import stock_import
        except Exception as e:
            messagebox.showerror('Lỗi', f'Không nạp được stock_check.py:\n{e}')
            return
        by = {sh.get('title'): sh for sh in self.app.sheets}
        up = by.get(UNIT_PRICE_TITLE)
        rp = by.get('NVL thay thế')
        if not up:
            messagebox.showerror('Thiếu dữ liệu', 'Không tìm thấy sheet Unit price (Đơn giá vật tư).')
            return
        rows = ProductionOrder(self.current_order, sl).rows()
        res = stock_check.check_order(rows, up, rp)
        # NVL don truoc (sheet Xuat NVL <- Excel 15.XUAT NL): loi o day KHONG lam hong kiem tra ton kho
        prev, on_cmp = None, None
        try:
            import xuat_nvl
            prev = xuat_nvl.for_stock_check(self.app.sheets, self.current_order)
            if prev is not None:
                on_cmp = lambda dlg: xuat_nvl.CompareDialog(dlg, self.current_order, self.app.sheets)
        except Exception as e:
            print('xuat_nvl loi:', e)
        stock_check.StockCheckDialog(
            self, self.current_order, res, stock_import.last_update_text(up),
            on_fix=lambda mode, cols, dlg: self._fix_stock(mode, cols, dlg, res, up, sl),
            on_restore=self._restore_formula,
            admin=bool(self.app) and self.app.user.get('tabs', 'all') == 'all',
            prev=prev, on_compare=on_cmp)

    # ---------------- SUA LENH THEO NVL THAY THE (tu bang Kiem tra ton kho) ----------------
    def _fix_stock(self, mode, cols, dlg, res, up, sl):
        """mode 'thay' = thay han NVL thieu bang NVL thay the;
        mode 'chia' = chia deu ton kho NVL goc cho cac me, phan con lai bu bang NVL thay the."""
        import copy
        import datetime
        import stock_check
        o = self.current_order
        if not o:
            return
        plan = stock_check.plan_fix(o, sl, res, up, mode=mode, only=cols)
        title = 'Sửa lệnh theo NVL thay thế' if mode == 'thay' else 'Chia đều NVL gốc + bù thay thế'
        if plan['error']:
            messagebox.showwarning(title, plan['error'], parent=dlg)
            return
        if not plan['done']:
            messagebox.showinfo(title, 'Không có NVL nào sửa được.\n\n' + '\n'.join(
                '• ' + x for x in plan['skipped']) if plan['skipped'] else
                'Không có NVL nào sửa được (dòng đang chọn đã đủ hàng?).', parent=dlg)
            return
        msg = (f"Đơn hàng {o.get('don_hang') or ''} – {o.get('ma_hang') or ''}\n"
               f"Số mẻ: {plan['so_me_cu']}" + (f" → {plan['so_me_moi']}" if plan['so_me_moi'] != plan['so_me_cu'] else '')
               + '\n\nSẽ sửa công thức:\n' + '\n'.join('• ' + x for x in plan['done']))
        if plan['skipped']:
            msg += '\n\nKhông sửa:\n' + '\n'.join('• ' + x for x in plan['skipped'])
        msg += '\n\nLưu thay đổi vào lệnh? (công thức cũ được giữ lại để Khôi phục)'
        if not messagebox.askyesno(title, msg, parent=dlg):
            return
        old = o.get('materials')
        had_goc = 'materials_goc' in o
        if not had_goc:
            o['materials_goc'] = copy.deepcopy(old or [])
            o['materials_goc_luc'] = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        o['materials'] = plan['materials']
        try:
            save_data_file(self.app.data)
        except Exception as e:
            o['materials'] = old
            if not had_goc:
                o.pop('materials_goc', None)
                o.pop('materials_goc_luc', None)
            messagebox.showerror('Lỗi lưu', f'Không lưu được lệnh sản xuất:\n{e}', parent=dlg)
            return
        self._after_formula_change(dlg)

    def _restore_formula(self, dlg):
        o = self.current_order
        if not o or not o.get('materials_goc'):
            return
        if not messagebox.askyesno('Khôi phục công thức gốc',
                                   'Trả công thức NVL của lệnh về như trước khi sửa theo NVL thay thế'
                                   f" ({o.get('materials_goc_luc') or ''})?\n\n"
                                   'Mọi thay đổi NVL sau thời điểm đó sẽ mất.', parent=dlg):
            return
        old, goc, luc = o.get('materials'), o.get('materials_goc'), o.get('materials_goc_luc')
        o['materials'] = goc
        o.pop('materials_goc', None)
        o.pop('materials_goc_luc', None)
        try:
            save_data_file(self.app.data)
        except Exception as e:
            o['materials'], o['materials_goc'], o['materials_goc_luc'] = old, goc, luc
            messagebox.showerror('Lỗi lưu', f'Không lưu được lệnh sản xuất:\n{e}', parent=dlg)
            return
        self._after_formula_change(dlg)

    def _after_formula_change(self, dlg):
        """Cong thuc vua doi -> tinh lai lenh + mo lai bang kiem tra ton kho voi so lieu moi."""
        try:
            dlg.destroy()
        except Exception:
            pass
        self._recalc()
        self._check_stock()

    def _open_img_settings(self):
        ImagePathDialog(self, self.app)

    def _reload_orders(self):
        self.search.set('')
        self._filter_orders()

    def refresh_live(self):
        """May khac vua sua don hang -> cap nhat danh sach, GIU don dang chon + so lieu dang nhap."""
        cur = self.current_order
        cur_key = str(cur.get('ma_sp') or '') if cur else None
        q = self.search.get().strip().lower()
        self._filtered = []
        for o in self._orders_of_kind():
            blob = ' '.join(str(o.get(k) or '') for k in
                            ('ma_sp', 'don_hang', 'ma_hang', 'khach_hang')).lower()
            if q in blob:
                self._filtered.append(o)
        labels = [self._order_label(o) for o in self._filtered]
        self.order_cb['values'] = labels
        idx = None
        if cur_key is not None:
            for i, o in enumerate(self._filtered):
                if str(o.get('ma_sp') or '') == cur_key:
                    idx = i
                    break
        if idx is not None:
            self.order_cb.current(idx)
            self.current_order = self._filtered[idx]
            self._recalc()
        elif labels:
            self.order_cb.current(0)
            self._select_order()
        else:
            self.order_cb.set('')
            self.current_order = None
            self._clear()

    def _order_label(self, o):
        return f"[{o.get('ma_sp')}] {o.get('don_hang')}  •  {o.get('ma_hang')}  ({o.get('khach_hang')})"

    def _filter_orders(self):
        q = self.search.get().strip().lower()
        self._filtered = []
        for o in self._orders_of_kind():
            blob = ' '.join(str(o.get(k) or '') for k in
                            ('ma_sp', 'don_hang', 'ma_hang', 'khach_hang')).lower()
            if q in blob:
                self._filtered.append(o)
        labels = [self._order_label(o) for o in self._filtered]
        self.order_cb['values'] = labels
        if labels:
            self.order_cb.current(0)
            self._select_order()
        else:
            self.order_cb.set('')
            self.current_order = None
            self._clear()

    def _select_order(self):
        i = self.order_cb.current()
        if i < 0 or i >= len(self._filtered):
            return
        o = self._filtered[i]
        self.current_order = o
        # % hao hut: lay tu don neu da luu, mac dinh 0.7%
        hh = o.get('hao_hut')
        self.hao_hut.set(str(hh) if hh not in (None, '') else '0.7')
        # Day chuyen: lay tu don neu da luu
        self.day_chuyen.set(str(o.get('day_chuyen') or ''))
        key = (str(o.get('don_hang') or '').strip(), str(o.get('ma_hang') or '').strip())
        cv = self.cont_index.get(key) or {}

        # --- Nap tu Container Volume (reset ve mac dinh truoc, roi lay gia tri co san) ---
        # San luong
        sl = cv.get('san_luong')
        if sl not in (None, ''):
            try:
                sl = float(sl)
                self.san_luong.set(str(int(sl) if sl.is_integer() else sl))
            except (TypeError, ValueError):
                self.san_luong.set(str(sl))
        else:
            self.san_luong.set('0')
        # Kg/Pallet
        self.kg_pallet.set(str(cv['kg_pallet']) if cv.get('kg_pallet') not in (None, '') else '1000')
        # Loai bao / Quy cach xep bao (mac dinh neu trong)
        self.loai_bao.set(cv.get('loai_bao') or 'Bao 25Kg')
        self.quy_cach.set(cv.get('quy_cach') or 'Hàng Xếp Rời')
        # Tong so Cont: lay tu CV, neu trong thi tinh = ROUNDUP(SL / 28000)  (1 cont <= 28.000 kg, <= 20 pallet)
        self.pallet_cont = cv.get('pallet_cont') or 20
        so_cont = cv.get('so_cont')
        if so_cont in (None, ''):
            from engine import tinh_so_cont
            so_cont = tinh_so_cont(self.san_luong.get()) or ''
        self.so_cont.set(str(so_cont))
        # Cach xep pallet: lay tu CV neu co, sau do _recalc se tu tinh neu Quy cach = Hang Xep Pallet
        self.cach_xep.set(cv.get('cach_xep') or '')
        self._recalc()

    def _recalc(self):
        o = self.current_order
        if not o:
            return
        try:
            sl = float(self.san_luong.get() or 0)
        except ValueError:
            sl = 0
        try:
            kgp = float(self.kg_pallet.get() or 1000)
        except ValueError:
            kgp = 1000
        o['hao_hut'] = self._hao()
        po = ProductionOrder(o, sl)
        rows = po.rows()
        s = po.summary()
        cont_auto, pallet = po.so_cont(kgp)
        # So Cont lay tu o nhap (nap tu Container Volume); neu trong thi dung tinh tu dong
        cont = self.so_cont.get().strip() or str(cont_auto)
        from engine import chia_pallet
        kq = chia_pallet(sl, kgp, self.loai_bao.get(), so_cont=self.so_cont.get().strip(),
                         pallet_cont=getattr(self, 'pallet_cont', 20))
        if kq:
            pallet = kq['tong_pallet']
        # Xep lan loai hang: tong pallet = pallet day tung ma + pallet lan
        if hasattr(self, 'mix_var') and self.mix_var.get():
            from engine import xep_lan_pallet
            kl = xep_lan_pallet(sl, kgp, self.loai_bao.get(), ma_hang=o.get('ma_hang') or '',
                                phu=[(m.get(), k.get()) for m, k in zip(self.ma_phu, self.kg_phu)])
            if kl:
                pallet = kl['tong_pallet']
        self.lbl_info.config(text=(
            f"Đơn hàng: {o.get('don_hang')}    |    Mã hàng: {o.get('ma_hang')}    |    "
            f"Khách hàng: {o.get('khach_hang')}    |    Loại: {o.get('loai') or '-'}    |    "
            f"Số NVL: {len(rows)}"))
        self.tree.delete(*self.tree.get_children())
        for i, r in enumerate(rows, 1):
            self.tree.insert('', 'end', values=(
                i, r['ten'], r['ma_nvl'] or '', f"{r['percentage']*100:.3f}",
                f"{r['mass_batch']:.3f}", f"{r['tong_don']:,.2f}", f"{r['total_order']:,.2f}"))
        self.tree.insert('', 'end', tags=('total',), values=(
            '', 'TỔNG CỘNG (Total)', '', f"{s['tong_pct']*100:.2f}",
            f"{po.N10():,.2f}", f"{po.N11():,.2f}", f"{po.N12():,.2f}"))
        self.cards['co_me'].config(text=f"{s['co_me']:,.1f}")
        self.cards['so_me'].config(text=f"{s['so_me']}")
        self.cards['sltt'].config(text=f"{s['san_luong_thuc_te']:,.0f}")
        self.cards['cont'].config(text=f"{cont}")
        self.cards['pallet'].config(text=f"{pallet}")
        self.cards['tong_pct'].config(text=f"{s['tong_pct']*100:.1f}%")

    def _clear(self):
        self.tree.delete(*self.tree.get_children())
        for v in self.cards.values():
            v.config(text='–')
        self.lbl_info.config(text='Không có đơn hàng phù hợp.')

    def _hao(self):
        try:
            return float(self.hao_hut.get() or 0.7)
        except ValueError:
            return 0.7

    def _calc_xep(self):
        """Tu tinh 'Cach xep pallet' khi Quy cach = Hang Xep Pallet (van sua duoc)."""
        if not hasattr(self, 'cach_xep'):
            return
        if (self.quy_cach.get() or '').strip().lower() != 'hàng xếp pallet':
            return
        from engine import cach_xep_pallet
        o = self.current_order
        ma = (o.get('ma_hang') if o else '') or ''
        mix = bool(self.mix_var.get())
        phu = [(m.get(), k.get()) for m, k in zip(self.ma_phu, self.kg_phu)]
        s = cach_xep_pallet(self.san_luong.get(), self.kg_pallet.get(), self.loai_bao.get(),
                            ma_hang=ma, mix=mix, phu=phu, so_cont=self.so_cont.get().strip(),
                            pallet_cont=getattr(self, 'pallet_cont', 20))
        if s:
            self.cach_xep.set(s)
        # Xep lan: cap nhat o 'Pallet' = pallet day tung ma + pallet lan
        if mix and hasattr(self, 'cards') and 'pallet' in self.cards:
            from engine import xep_lan_pallet
            kl = xep_lan_pallet(self.san_luong.get(), self.kg_pallet.get(), self.loai_bao.get(),
                                ma_hang=ma, phu=phu)
            if kl:
                self.cards['pallet'].config(text=f"{kl['tong_pallet']}")

    def _save_day_chuyen(self):
        """Ghi Day chuyen vao don dang chon va luu data.json (neu co thay doi)."""
        o = self.current_order
        if o is None:
            return
        v = ' '.join(self.day_chuyen.get().split())
        if (o.get('day_chuyen') or '') == v:
            return
        o['day_chuyen'] = v
        try:
            save_data_file(self.app.data)
        except Exception as e:
            messagebox.showerror('Lỗi lưu', f'Không lưu được Dây chuyền:\n{e}')

    def _current_po(self):
        try:
            sl = float(self.san_luong.get() or 0)
        except ValueError:
            sl = 0
        if self.current_order is not None:
            self.current_order['hao_hut'] = self._hao()
        return ProductionOrder(self.current_order, sl)

    def _opts(self):
        o = self.current_order or {}
        st = getattr(self.app, 'settings', {}) or {}
        kh, ma = o.get('khach_hang'), o.get('ma_hang')
        stamp_img = pkg_img = None
        self._img_warn = []
        try:
            from image_finder import find_image_ex, customer_code, folder_info
            # Thu muc anh dat theo ma KH cu (VD KSP) -> lay tu Don hang (PO/26001/01.KSP);
            # khong co Don hang thi dung cot Khach hang
            _dp = str(o.get('don_hang') or '').split('/')
            if len(_dp) >= 3 and _dp[2].strip():
                kh = _dp[2].strip()
            code = customer_code(kh)
            # Nguon anh: THEO MAY CHU (tai anh cua khach ve bo nho dem) hoac THEO MAY NAY
            import image_sync
            stamp_roots = image_sync.roots(st, 'stamp')
            pkg_roots = image_sync.roots(st, 'pkg')
            if image_sync.mode(st) == 'server' and code:
                for _k in ('stamp', 'pkg'):
                    image_sync.sync_folder(_k, code)

            def _find(roots_):
                kind_ = None
                for r_ in roots_:
                    p_, k_ = find_image_ex(r_, kh, ma)
                    if k_ == 'specific':
                        return p_, k_
                    if p_ and kind_ is None:
                        kind_ = (p_, k_)
                return kind_ or (None, None)
            stamp_img, sk = _find(stamp_roots)
            pkg_img, _pk = _find(pkg_roots)
            # canh bao khi khong co TEM rieng theo ma hang
            if sk == 'generic':
                self._img_warn.append(
                    f"ℹ Không có Tem riêng cho mã hàng '{ma}'.\n"
                    f"    Đang dùng ảnh Tem CHUNG của khách '{code}'.")
            elif sk is None:
                fpath, exists, files = folder_info(stamp_roots[0] if stamp_roots else '', kh)
                if image_sync.mode(st) == 'server':
                    fpath = f"(máy chủ) {code}  →  bộ nhớ đệm: {fpath}"
                if not exists:
                    detail = f"    Không thấy thư mục:\n    {fpath}"
                elif not files:
                    detail = f"    Thư mục rỗng (không có ảnh):\n    {fpath}"
                else:
                    lst = '\n'.join('       • ' + f for f in files[:12])
                    detail = (f"    Thư mục: {fpath}\n"
                              f"    Có {len(files)} ảnh nhưng không khớp mã hàng '{ma}'\n"
                              f"    và không có ảnh chung tên '{code}':\n{lst}")
                self._img_warn.append(
                    f"⚠ Không hiển thị được Tem cho mã hàng '{ma}' (khách '{code}').\n"
                    f"{detail}\n\n"
                    f"    → Đặt tên file: '{code}-{ma}.png' (riêng) hoặc '{code}.png' (chung).")
            # neu tim thay anh nhung thieu Pillow -> anh .jpg khong hien o Xem truoc
            if stamp_img or pkg_img:
                try:
                    import PIL  # noqa
                except Exception:
                    jpg = [p for p in (stamp_img, pkg_img)
                           if p and p.lower().endswith(('.jpg', '.jpeg'))]
                    if jpg:
                        self._img_warn.append(
                            "⚠ Thiếu thư viện Pillow nên ảnh .jpg KHÔNG hiện ở phần Xem trước.\n"
                            "    Chạy: pip install Pillow  (bản .exe đã kèm sẵn Pillow).\n"
                            "    (Xuất Excel/PDF vẫn chèn ảnh bình thường.)")
        except Exception as e:
            self._img_warn.append('Lỗi tìm ảnh Tem: ' + str(e))
        self._save_day_chuyen()
        return {'loai_bao': self.loai_bao.get(),
                'quy_cach': self.quy_cach.get(),
                'day_chuyen': ' '.join(self.day_chuyen.get().split()),
                'so_cont': self.so_cont.get().strip(),
                'cach_xep': self.cach_xep.get().strip(),
                'stamp_img': stamp_img,
                'pkg_img': pkg_img}

    def _warn_images(self):
        w = getattr(self, '_img_warn', None)
        if w:
            messagebox.showinfo('Ảnh Tem', '\n\n'.join(w))

    def _warn_ty_le_b9(self, san_luong):
        """VBA KiemTraTyLeB9: canh bao neu ty le B9 lech so sau P10/P20 cua ma hang.
        Hien hop thoai tu dong dong sau 25 giay (giong TimeOutSeconds:=25)."""
        try:
            import form_rules
            po = ProductionOrder(self.current_order, san_luong)
            msg = form_rules.kiem_tra_ty_le_b9({}, self.current_order, po)
        except Exception:
            return
        if msg:
            timed_msgbox(self, form_rules.TIEU_DE_CANH_BAO_B9, msg,
                         form_rules.TIMEOUT_CANH_BAO_B9)

    def _export_excel(self):
        if self._stock_blocked():
            return
        if not self.current_order:
            messagebox.showwarning('Chưa chọn đơn', 'Vui lòng chọn một đơn hàng.')
            return
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            messagebox.showerror('Thiếu thư viện', 'Cần cài openpyxl:\n\npip install openpyxl')
            return
        try:
            from order_form import fill_order_form
        except Exception as e:
            messagebox.showerror('Lỗi', f'Không nạp được order_form:\n{e}')
            return
        o = self.current_order
        o['hao_hut'] = self._hao()
        kind = self.kind.get()
        try:
            sl = float(self.san_luong.get() or 0)
        except ValueError:
            sl = 0
        default = __import__('doc_tools').lsx_file_name(o, '.xlsx')
        path = filedialog.asksaveasfilename(defaultextension='.xlsx', initialfile=default,
                                            filetypes=[('Excel', '*.xlsx')])
        if not path:
            return
        try:
            opts = self._opts()
            self._warn_images()
            self._warn_ty_le_b9(sl)
            fill_order_form(kind, o, sl, path, opts)
            messagebox.showinfo('Xong', 'Đã xuất Lệnh sản xuất (đúng mẫu form):\n'
                                        f'{path}\n\nMở bằng Excel để xem/in — công thức sẽ '
                                        'tự tính lại khi mở.')
        except FileNotFoundError:
            messagebox.showerror('Thiếu template', 'Không tìm thấy template.xlsx cạnh phần mềm.\n'
                                 'Hãy để template.xlsx cùng thư mục với phần mềm (.exe).')
        except Exception as e:
            messagebox.showerror('Lỗi xuất form', str(e))

    def _preview(self):
        """Ve form Lenh san xuat ngay trong Python (Canvas) - khong can Excel."""
        if self._stock_blocked():
            return
        if not self.current_order:
            messagebox.showwarning('Chưa chọn đơn', 'Vui lòng chọn một đơn hàng.')
            return
        try:
            sl = float(self.san_luong.get() or 0)
        except ValueError:
            sl = 0
        self.current_order['hao_hut'] = self._hao()
        opts = self._opts()
        self._warn_images()
        self._warn_ty_le_b9(sl)
        FormPreview(self, self.kind.get(), self.current_order, sl, opts)

    def _export_pdf(self):
        if self._stock_blocked():
            return
        if not self.current_order:
            messagebox.showwarning('Chưa chọn đơn', 'Vui lòng chọn một đơn hàng.')
            return
        try:
            sl = float(self.san_luong.get() or 0)
        except ValueError:
            sl = 0
        self.current_order['hao_hut'] = self._hao()
        opts = self._opts()
        self._warn_images()
        self._warn_ty_le_b9(sl)
        export_order_pdf(self, self.kind.get(), self.current_order, sl, self._open_file, opts)

    def _print(self):
        if self._stock_blocked():
            return
        if not self.current_order:
            messagebox.showwarning('Chưa chọn đơn', 'Vui lòng chọn một đơn hàng.')
            return
        try:
            sl = float(self.san_luong.get() or 0)
        except ValueError:
            sl = 0
        self.current_order['hao_hut'] = self._hao()
        opts = self._opts()
        self._warn_images()
        self._warn_ty_le_b9(sl)
        PrintDialog(self, self.kind.get(), self.current_order, sl, opts, settings_owner=self.app)

    def _open_file(self, path):
        try:
            if sys.platform.startswith('win'):
                os.startfile(path)                      # noqa
            elif sys.platform == 'darwin':
                import subprocess
                subprocess.Popen(['open', path])
            else:
                import subprocess
                subprocess.Popen(['xdg-open', path])
        except Exception as e:
            messagebox.showinfo('Đã tạo file', f'Đã tạo:\n{path}\n\n(Không tự mở được: {e})')


def timed_msgbox(parent, title, message, timeout=25):
    """MsgBox vbOKOnly tu dong dong sau `timeout` giay (thay cho TimeOutSeconds VBA)."""
    top = tk.Toplevel(parent)
    top.title(title)
    top.configure(bg='white')
    top.resizable(False, False)
    try:
        top.transient(parent)
    except Exception:
        pass
    frm = tk.Frame(top, bg='white', padx=18, pady=14)
    frm.pack(fill='both', expand=True)
    tk.Label(frm, text='⚠', fg='#d97706', bg='white',
             font=('Segoe UI', 26)).grid(row=0, column=0, sticky='n', padx=(0, 12))
    tk.Label(frm, text=message, bg='white', justify='left', wraplength=460,
             font=('Segoe UI', 10)).grid(row=0, column=1, sticky='w')
    left = [int(timeout)]
    btn = ttk.Button(frm, text=f'OK ({left[0]})', command=top.destroy, width=12)
    btn.grid(row=1, column=0, columnspan=2, pady=(14, 0))

    def _tick():
        if not top.winfo_exists():
            return
        left[0] -= 1
        if left[0] <= 0:
            top.destroy()
            return
        btn.configure(text=f'OK ({left[0]})')
        top.after(1000, _tick)

    top.bind('<Return>', lambda e: top.destroy())
    top.bind('<Escape>', lambda e: top.destroy())
    top.update_idletasks()
    try:
        x = parent.winfo_rootx() + (parent.winfo_width() - top.winfo_width()) // 2
        y = parent.winfo_rooty() + (parent.winfo_height() - top.winfo_height()) // 3
        top.geometry(f'+{max(x, 0)}+{max(y, 0)}')
    except Exception:
        pass
    btn.focus_set()
    top.grab_set()
    top.after(1000, _tick)
    parent.wait_window(top)


def export_order_pdf(parent, kind, order, san_luong, open_cb=None, opts=None):
    """Xuat form Lenh san xuat ra PDF (dung chung cho tab va cua so xem truoc)."""
    try:
        from pdf_export import render_pdf
    except ImportError:
        messagebox.showerror('Thiếu thư viện',
                             'Cần cài reportlab để xuất PDF:\n\npip install reportlab',
                             parent=parent)
        return
    except Exception as e:
        messagebox.showerror('Lỗi', str(e), parent=parent)
        return
    # ten file: LSX_<Don hang>_<Ma hang>_<ngay-thang-nam>_<gio> -> khong bao gio trung
    import doc_tools
    default = doc_tools.lsx_file_name(order, '.pdf')
    path = filedialog.asksaveasfilename(defaultextension='.pdf', initialfile=default,
                                        filetypes=[('PDF', '*.pdf')], parent=parent)
    if not path:
        return
    try:
        render_pdf(kind, order, san_luong, path, opts)
    except Exception as e:
        messagebox.showerror('Lỗi xuất PDF', str(e), parent=parent)
        return
    if messagebox.askyesno('Xong', f'Đã xuất PDF:\n{path}\n\nMở file ngay?', parent=parent):
        if open_cb:
            open_cb(path)


# =====================================================================
#  IN LENH SAN XUAT - chon may in, so ban
# =====================================================================
class PrintDialog(tk.Toplevel):
    """Hop thoai In: chon may in + so ban -> tao PDF tam roi gui toi may in."""

    def __init__(self, parent, kind, order, san_luong, opts=None, settings_owner=None,
                 make_pdf=None, landscape=True, heading=None, subtitle=None, pdf_prefix=None):
        """make_pdf(path): ham tu tao PDF (dung cho TDS/COA/MSDS). Khong co -> in Lenh san xuat."""
        super().__init__(parent)
        self.kind, self.order, self.san_luong, self.opts = kind, (order or {}), san_luong, (opts or {})
        self.settings_owner = settings_owner      # doi tuong co .settings (App) de nho may in
        self._custom_make = make_pdf
        self.landscape = landscape
        self.pdf_prefix = pdf_prefix
        heading = heading or 'In Lệnh sản xuất'
        self.title('🖨️ ' + heading)
        self.configure(bg='#f4f6f9')
        self.resizable(False, False)
        self.transient(parent)
        self._pdf = None
        order = self.order

        tk.Label(self, text='🖨️  ' + heading, bg='#f4f6f9', font=FONT_TITLE).grid(
            row=0, column=0, columnspan=3, sticky='w', padx=16, pady=(14, 4))
        if subtitle is None:
            subtitle = (f"Đơn hàng: {order.get('don_hang') or ''}   •   Mã hàng: {order.get('ma_hang') or ''}"
                        f"   •   Loại: {kind}")
        tk.Label(self, text=subtitle, bg='#f4f6f9', fg='#555', font=FONT).grid(
            row=1, column=0, columnspan=3, sticky='w', padx=16, pady=(0, 10))

        tk.Label(self, text='Máy in:', bg='#f4f6f9', font=FONT_B).grid(row=2, column=0, sticky='e', padx=(16, 6))
        self.printer = tk.StringVar()
        self.cb = ttk.Combobox(self, textvariable=self.printer, width=46, state='readonly', font=FONT)
        self.cb.grid(row=2, column=1, sticky='w', pady=4)
        tk.Button(self, text='↻', font=FONT, width=3, command=self._load_printers).grid(
            row=2, column=2, sticky='w', padx=(4, 16))

        tk.Label(self, text='Số bản:', bg='#f4f6f9', font=FONT_B).grid(row=3, column=0, sticky='e', padx=(16, 6))
        self.copies = tk.StringVar(value='1')
        tk.Spinbox(self, from_=1, to=50, textvariable=self.copies, width=6, font=FONT).grid(
            row=3, column=1, sticky='w', pady=4)

        self.msg = tk.Label(self, text='Đang tìm máy in…', bg='#f4f6f9', fg='#666', font=('Segoe UI', 9),
                            justify='left', wraplength=460)
        self.msg.grid(row=4, column=0, columnspan=3, sticky='w', padx=16, pady=(6, 4))

        bf = tk.Frame(self, bg='#f4f6f9')
        bf.grid(row=5, column=0, columnspan=3, sticky='e', padx=16, pady=(6, 14))
        self.btn_print = tk.Button(bf, text='🖨️  In', font=FONT_B, bg='#1d3557', fg='white',
                                   relief='flat', padx=18, pady=6, command=self._do_print)
        self.btn_print.pack(side='right')
        tk.Button(bf, text='📄 Mở PDF để in (Ctrl+P)', font=FONT, relief='flat', padx=10, pady=6,
                  command=self._open_pdf).pack(side='right', padx=6)
        tk.Button(bf, text='Đóng', font=FONT, relief='flat', padx=12, pady=6,
                  command=self.destroy).pack(side='right')
        self.after(50, self._load_printers)

    def _settings(self):
        return getattr(self.settings_owner, 'settings', None)

    def _load_printers(self):
        self.msg.config(text='Đang tìm máy in…')
        self.update_idletasks()
        try:
            from print_util import list_printers, find_sumatra
            names, default = list_printers()
        except Exception as e:
            names, default, find_sumatra = [], '', (lambda: None)
            self.msg.config(text=f'Lỗi đọc danh sách máy in: {e}')
        self.cb['values'] = names
        st = self._settings() or {}
        last = st.get('printer')
        if last in names:
            self.printer.set(last)
        elif default in names:
            self.printer.set(default)
        elif names:
            self.printer.set(names[0])
        if not names:
            self.msg.config(text='⚠ Không tìm thấy máy in nào. Bạn vẫn có thể bấm "Mở PDF để in".',
                            fg='#c0392b')
        else:
            how = (f"Sẽ in bằng SumatraPDF (khổ {'ngang' if self.landscape else 'dọc'}, vừa trang)." if find_sumatra()
                   else 'Sẽ in qua trình đọc PDF mặc định của Windows.\n'
                        'Mẹo: đặt SumatraPDF.exe cạnh phần mềm để in ổn định & đúng số bản hơn.')
            self.msg.config(text=f'Tìm thấy {len(names)} máy in' +
                                 (f' (mặc định: {default})' if default else '') + '.\n' + how, fg='#555')

    def _make_pdf(self):
        try:
            from pdf_export import render_pdf
            from print_util import temp_pdf_path
        except ImportError:
            messagebox.showerror('Thiếu thư viện', 'Cần cài reportlab để in:\n\npip install reportlab',
                                 parent=self)
            return None
        prefix = self.pdf_prefix or f"LenhSX_{str(self.order.get('don_hang') or '').replace('/', '-')}_"
        path = temp_pdf_path(prefix)
        try:
            if self._custom_make is not None:
                self._custom_make(path)
            else:
                render_pdf(self.kind, self.order, self.san_luong, path, self.opts)
        except Exception as e:
            messagebox.showerror('Lỗi tạo bản in', str(e), parent=self)
            return None
        self._pdf = path
        return path

    def _open_pdf(self):
        path = self._make_pdf()
        if not path:
            return
        try:
            if sys.platform.startswith('win'):
                os.startfile(path)                      # noqa
            else:
                import subprocess
                subprocess.Popen(['xdg-open', path])
        except Exception as e:
            messagebox.showinfo('Đã tạo PDF', f'{path}\n\n(Không tự mở được: {e})', parent=self)

    def _do_print(self):
        printer = self.printer.get().strip()
        try:
            copies = max(1, int(self.copies.get() or 1))
        except ValueError:
            copies = 1
        path = self._make_pdf()
        if not path:
            return
        st = self._settings()
        if st is not None and printer:
            st['printer'] = printer
            save_settings(st)
        self.btn_print.config(state='disabled', text='Đang in…')
        self.msg.config(text=f'Đang gửi tới máy in "{printer or "mặc định"}"…', fg='#1d3557')
        import threading

        def work():
            try:
                from print_util import print_pdf
                ok, info = print_pdf(path, printer, copies, landscape=self.landscape)
            except Exception as e:
                ok, info = False, str(e)
            self.after(0, lambda: self._done(ok, info))
        threading.Thread(target=work, daemon=True).start()

    def _done(self, ok, info):
        try:
            self.btn_print.config(state='normal', text='🖨️  In')
        except Exception:
            return
        if ok:
            messagebox.showinfo('Đã gửi lệnh in', info, parent=self)
            self.destroy()
        else:
            self.msg.config(text=info, fg='#c0392b')
            messagebox.showwarning('Chưa in được', info, parent=self)


# =====================================================================
#  XEM TRUOC LENH SAN XUAT - ve bang Canvas (khong can Excel)
# =====================================================================
SIGN_BY_KIND = {
    'Domestic': [('QL. sản xuất', 'Nguyễn Thị Lan'), ('Trưởng BP. QC', 'Nguyễn Hoài Nam'),
                 ('Trợ lý sản xuất', 'Nguyễn Xuân Duy'), ('Giám đốc nhà máy', 'Vishal Soan')],
    'Export':   [('QL. sản xuất', 'Nguyễn Thị Lan'), ('Trưởng BP. QC', 'Nguyễn Hoài Nam'),
                 ('Trợ lý sản xuất', 'Nguyễn Xuân Duy'), ('Giám đốc nhà máy', 'Vishal Soan')],
    'Sample':   [('QL. sản xuất', 'Nguyễn Thị Lan'), ('Trưởng BP. QC', 'Nguyễn Hoài Nam'),
                 ('Trợ lý sản xuất', 'Nguyễn Xuân Duy'), ('Giám đốc nhà máy', 'Vishal Soan')],
}


class LayoutRenderMixin:
    """Ve 1 layout (o, vung gop, vien, chu) len self.cv. Dung chung cho form & tai lieu."""
    @staticmethod
    def _fmt(v, is_pct):
        if isinstance(v, (int, float)):
            if is_pct:
                return f"{v*100:.2f}%"
            if float(v).is_integer():
                return f"{int(v)}"
            return f"{v:,.2f}"
        return str(v)

    def _cellref(self, s):
        import re
        m = re.fullmatch(r"=\$?([A-Z]{1,3})\$?(\d+)", str(s).strip())
        if not m:
            return None
        from openpyxl.utils import column_index_from_string
        return (int(m.group(2)), column_index_from_string(m.group(1)))

    _zoom = 1.0                      # ti le phong to ban xem truoc (FormPreview co nut +/-)

    def _font(self, size, bold):
        import tkinter.font as tkfont
        if not hasattr(self, '_fontcache'):
            self._fontcache = {}
        size = max(5, int(round(size * self._zoom)))
        key = (size, bold)
        if key not in self._fontcache:
            self._fontcache[key] = tkfont.Font(
                family='Segoe UI', size=size, weight='bold' if bold else 'normal')
        return self._fontcache[key]

    @staticmethod
    def _wrap_lines(text, font, avail_w):
        import math
        total = 0
        for seg in str(text).split('\n'):
            if seg == '':
                total += 1
                continue
            w = font.measure(seg)
            total += max(1, math.ceil(w / max(20, avail_w)))
        return total

    def _cell_meta(self, c, dyn):
        r, col = c['r'], c['c']
        if (r, col) in dyn:
            text = dyn[(r, col)]
        elif c['f']:
            ref = self._cellref(c['tv'])
            text = dyn[ref] if (ref and ref in dyn) else self._fmt(c['cv'], c['p'])
        else:
            text = self._fmt(c['tv'] if c['tv'] != '' else c['cv'], c['p'])
        from pdf_export import _render_size
        size = _render_size(c['sz'])
        return ('' if text in ('', 'None') else str(text)), size, bool(c['b'])

    def _font_i(self, size, bold):
        """Font NGHIENG (o co dinh dang 'i' trong layout)."""
        import tkinter.font as tkfont
        if not hasattr(self, '_fontcache_i'):
            self._fontcache_i = {}
        size = max(5, int(round(size * self._zoom)))
        key = (size, bold)
        if key not in self._fontcache_i:
            self._fontcache_i[key] = tkfont.Font(family='Segoe UI', size=size, slant='italic',
                                                 weight='bold' if bold else 'normal')
        return self._fontcache_i[key]

    def _place_logo(self, x_of, y_of, header_rows=(1,)):
        try:
            import math
            p = resource_path('logo.png')
            if not os.path.exists(p):
                return
            self._logo_imgs = []          # giu tham chieu de khong bi thu gom
            cx = (x_of(1) + x_of(4)) / 2
            for h in header_rows:
                base = tk.PhotoImage(file=p)
                # logo nam duoi dong ten cong ty: vung hang h+1 .. h+2
                rw = max(40, x_of(4) - x_of(1) - 14)
                rh = max(24, y_of(h + 3) - y_of(h + 1) - 4)
                k = max(1, math.ceil(max(base.width() / rw, base.height() / rh)))
                img = base.subsample(k)
                self._logo_imgs.append(img)
                cy = (y_of(h + 1) + y_of(h + 3)) / 2
                self.cv.create_image(cx, cy, image=img, anchor='center')
        except Exception:
            pass

    def _place_form_images(self, x_of, y_of):
        """Mac dinh khong lam gi (chi FormPreview ve anh Tem/Bao bi)."""
        pass

    # vung o chua anh tren FORM (rows 114-121)
    IMG_STAMP_BOX = (6, 114, 10, 121)   # Tem: cols 6-10
    IMG_PKG_BOX = (11, 114, 13, 121)    # Hinh anh (Picture): cols 11-13

    def _draw_layout(self, lay, dyn, place_logo=True):
        cols = lay['cols']
        rowh = lay['rowh']
        ncol = lay['ncol']
        nrow = len(rowh)
        cv = self.cv
        hidden = set(HIDDEN_FORM_ROWS) if place_logo else set()   # an dong 'Tong don' & 'Alternative'
        if place_logo and str(getattr(self, 'kind', '')).strip().lower() == 'sample':
            try:
                from pdf_export import SAMPLE_HIDDEN_ROWS
                hidden |= SAMPLE_HIDDEN_ROWS          # Sample: chi hien 1 trang gon (giong Excel)
            except Exception:
                pass
        hcols = HIDDEN_FORM_COLS if place_logo else set()    # an cot dieu khien P-T

        z = self._zoom
        xs = [8]
        for i, w in enumerate(cols):
            ww = 0 if (i + 1) in hcols else w   # cot an -> rong 0
            xs.append(xs[-1] + int((ww * 7 + 5) * z) if ww else xs[-1])

        def x_of(cc):
            cc = min(max(cc, 1), ncol + 1)
            return xs[cc - 1]

        covered = set()
        anchor_span = {}
        for (r0, c0, r1, c1) in lay['merges']:
            anchor_span[(r0, c0)] = (r1, c1)
            for rr in range(r0, r1 + 1):
                for cc in range(c0, c1 + 1):
                    if (rr, cc) != (r0, c0):
                        covered.add((rr, cc))

        info = {}
        occ = {}
        for c in lay['cells']:
            r, col = c['r'], c['c']
            if r in hidden or col in hcols:    # bo qua dong/cot bi an
                continue
            if (r, col) in covered:
                occ.setdefault(r, set()).add(col)
                continue
            text, size, bold = self._cell_meta(c, dyn)
            r1, c1 = anchor_span.get((r, col), (r, col))
            fg = c['color'] or 'black'
            if str(fg).lower() in ('#ffffff', 'white'):
                fg = 'black'
            info[(r, col)] = (text, size, bold, c['al'], c['fill'], fg, r1, c1)
            if text or c['fill'] or (r1 != r or c1 != col):
                for cc in range(col, c1 + 1):
                    occ.setdefault(r, set()).add(cc)

        def spill_c1(r, col, c1, al, text):
            if al in ('center', 'right') or not text or c1 != col:
                return c1
            row_occ = occ.get(r, set())
            for cc in range(col + 1, ncol + 1):
                if cc in row_occ:
                    return cc - 1
            return ncol

        def text_span(r, col, c1, al):
            """Mo rong vung chu sang o TRONG ben canh (giong Excel tran chu)."""
            row_occ = occ.get(r, set())
            lc, rc = col, c1
            cc = c1 + 1
            while cc <= ncol and cc not in row_occ:
                rc = cc; cc += 1
            cc = col - 1
            while cc >= 1 and cc not in row_occ:
                lc = cc; cc -= 1
            if al == 'center':
                k = min(col - lc, rc - c1)
                return col - k, c1 + k
            if al == 'right':
                return lc, c1
            return col, rc

        LINE_PAD = 4
        base_h = [max(int(15 * z), int(rowh[i] * 1.33 * z)) if i < nrow else int(15 * z) for i in range(nrow)]
        for hr in hidden:              # thu gon chieu cao dong an ve 0
            if 1 <= hr <= nrow:
                base_h[hr - 1] = 0
        if place_logo and nrow >= 1:
            base_h[0] = max(base_h[0], int(22 * z))
        if place_logo and nrow >= 2:
            base_h[1] = max(base_h[1], int(58 * z))
        multi = []
        for (r, col), (text, size, bold, al, fill, fg, r1, c1) in info.items():
            if not text:
                continue
            _lc, _rc = text_span(r, col, c1, al)
            avail_w = x_of(_rc + 1) - x_of(_lc) - 6
            f = self._font(size, bold)
            lh = f.metrics('linespace') + 2
            needed = self._wrap_lines(text, f, avail_w) * lh + LINE_PAD
            if r1 == r:
                i = r - 1
                if 0 <= i < nrow:
                    base_h[i] = max(base_h[i], needed)
            else:
                multi.append((r, r1, needed))

        def build_ys():
            ys = [8]
            for i in range(nrow):
                ys.append(ys[-1] + base_h[i])
            return ys

        ys = build_ys()
        for (r0, r1, needed) in multi:
            i0, i1 = r0 - 1, r1 - 1
            if i0 < 0 or i1 >= nrow:
                continue
            have = ys[i1 + 1] - ys[i0]
            if needed > have:
                base_h[i1] += (needed - have)
        ys = build_ys()

        def y_of(rr):
            rr = min(max(rr, 1), nrow + 1)
            return ys[rr - 1]

        # --- cac duong vien doc bi CHU tran DE len -> se bo di (khong ve) ---
        covered_vlines = set()   # (row, cot ranh gioi)
        for (r, col), (text, size, bold, al, fill, fg, r1, c1) in info.items():
            if not text:
                continue
            lc, rc = text_span(r, col, c1, al)
            if lc == col and rc == c1:
                continue
            f = self._font(size, bold)
            tw = max((f.measure(seg) for seg in str(text).split('\n')), default=0)
            if al == 'center':
                cx = (x_of(col) + x_of(c1 + 1)) / 2
                ex0, ex1 = cx - tw / 2, cx + tw / 2
            elif al == 'right':
                ex1 = x_of(c1 + 1) - 3
                ex0 = ex1 - tw
            else:
                ex0 = x_of(col) + 3
                ex1 = ex0 + tw
            for k in range(lc + 1, rc + 1):
                if ex0 + 1 < x_of(k) < ex1 - 1:
                    covered_vlines.add((r, k))

        for (r, col), (text, size, bold, al, fill, fg, r1, c1) in info.items():
            if fill:
                cv.create_rectangle(x_of(col), y_of(r), x_of(c1 + 1), y_of(r1 + 1), outline='', fill=fill)
        italic_cells = {(c['r'], c['c']) for c in lay['cells'] if c.get('i')}
        seg_box = {}          # bbox cac duong vien theo hang -> ve KHUNG NGOAI tung trang
        for c in lay['cells']:
            bd = c.get('bd')
            if not bd or not any(bd):
                continue
            r, col = c['r'], c['c']
            if r in hidden or col in hcols or col > ncol:
                continue
            x0, y0, x1, y1 = x_of(col), y_of(r), x_of(col + 1), y_of(r + 1)
            lines = []
            if bd[0] and (r, col) not in covered_vlines:
                lines.append((x0, y0, x0, y1))
            if bd[1]:
                lines.append((x0, y0, x1, y0))
            if bd[2] and (r, col + 1) not in covered_vlines:
                lines.append((x1, y0, x1, y1))
            if bd[3]:
                lines.append((x0, y1, x1, y1))
            for (lx0, ly0, lx1, ly1) in lines:
                cv.create_line(lx0, ly0, lx1, ly1, fill='#555')
                b = seg_box.get(r)
                seg_box[r] = ((min(b[0], lx0), min(b[1], ly0), max(b[2], lx1), max(b[3], ly1))
                              if b else (lx0, ly0, lx1, ly1))
        # ghi lai o de Admin bam chon sua (toa do truoc khi dich chuyen)
        self._cell_rects = [(x_of(col), y_of(r), x_of(c1 + 1), y_of(r1 + 1), r, col)
                            for (r, col), (_t, _s, _b, _a, _f, _g, r1, c1) in info.items()]
        for (r, col), (text, size, bold, al, fill, fg, r1, c1) in info.items():
            if not text:
                continue
            _lc, _rc = text_span(r, col, c1, al)
            x0, y0 = x_of(col), y_of(r)
            y1 = y_of(r1 + 1)
            wrap_w = x_of(_rc + 1) - x_of(_lc) - 5
            if al == 'center':
                ax, anchor = (x_of(col) + x_of(c1 + 1)) / 2, 'center'
            elif al == 'right':
                ax, anchor = x_of(c1 + 1) - 3, 'e'
            else:
                ax, anchor = x0 + 3, 'w'
            fnt = self._font(size, bold)
            if (r, col) in italic_cells:
                fnt = self._font_i(size, bold)
            cv.create_text(ax, (y0 + y1) / 2, text=text, font=fnt,
                           fill=fg, anchor=anchor, width=max(20, wrap_w))
        # --- cac dong tieu de 'CÔNG TY CPQT' (dau moi trang) ---
        starts = sorted({c['r'] for c in lay['cells']
                         if ((isinstance(c.get('tv'), str) and 'CÔNG TY CPQT' in c['tv'])
                             or (isinstance(c.get('cv'), str) and 'CÔNG TY CPQT' in c['cv']))
                         and c['r'] not in hidden})
        # --- KHUNG NGOAI dam quanh bang cua tung trang (giong Excel / PDF) ---
        if place_logo or getattr(self, 'layout_key', '') in FRAMED_DOC_KEYS:
            bounds = (starts or [1]) + [nrow + 1]
            for a, b in zip(bounds, bounds[1:]):
                boxes = [v for rr, v in seg_box.items() if a <= rr < b]
                if boxes:
                    cv.create_rectangle(min(v[0] for v in boxes), min(v[1] for v in boxes),
                                        max(v[2] for v in boxes), max(v[3] for v in boxes),
                                        outline='black', width=2)
        if place_logo:
            self._place_logo(x_of, y_of, starts or (1,))
            self._place_form_images(x_of, y_of)
        # --- vach ngan trang de nhin ro 4 trang khi in ---
        if len(starts) > 1:
            xr = x_of(ncol + 1)
            for pg, r in enumerate(starts[1:], 2):
                yy = y_of(r)
                cv.create_line(6, yy, xr, yy, fill='#c0392b', dash=(6, 4), width=2)
                cv.create_text(10, yy - 8, text=f'— Trang {pg} —', anchor='w',
                               fill='#c0392b', font=('Segoe UI', 8, 'bold'))
        cv.configure(scrollregion=(0, 0, xs[-1] + 20, ys[-1] + 20))


class FormPreview(LayoutRenderMixin, tk.Toplevel):
    """Ve lai form Lenh san xuat (du 4 trang) tu layout goc - khong can Excel."""
    _LAYOUTS = {}

    def __init__(self, master, kind, order, san_luong, opts=None):
        super().__init__(master)
        self.kind = kind
        self.order = order
        self.opts = opts or {}
        self.po = ProductionOrder(order, san_luong)
        self.title(f"Xem trước Lệnh sản xuất ({kind}) - {order.get('don_hang')}")
        self.geometry('1180x820')
        self.configure(bg='#e9edf2')

        bar = tk.Frame(self, bg='#e9edf2')
        bar.pack(fill='x')
        tk.Button(bar, text='⭳  Xuất Excel form gốc', font=FONT, bg='#2d6a4f', fg='white',
                  relief='flat', padx=12, pady=5, command=self._export).pack(side='right', padx=6, pady=6)
        tk.Button(bar, text='📄  Xuất PDF', font=FONT_B, bg='#b0413e', fg='white',
                  relief='flat', padx=12, pady=5, command=self._export_pdf).pack(side='right', padx=6, pady=6)
        tk.Button(bar, text='🖨️  In', font=FONT_B, bg='#1d3557', fg='white',
                  relief='flat', padx=14, pady=5, command=self._print).pack(side='right', padx=6, pady=6)
        tk.Label(bar, text='Bản xem trước mô phỏng form Excel (cuộn để xem đủ các trang)',
                 bg='#e9edf2', fg='#555', font=FONT).pack(side='left', padx=12)
        # --- phong to / thu nho (Ctrl + lan chuot) ---
        self._zoom = 1.0
        zf = tk.Frame(bar, bg='#e9edf2')
        zf.pack(side='left', padx=6)
        tk.Button(zf, text='－', font=FONT_B, width=3, relief='groove',
                  command=lambda: self._set_zoom(self._zoom - 0.1)).pack(side='left')
        self._zoom_lbl = tk.Label(zf, text='100%', width=6, bg='#e9edf2', font=FONT_B)
        self._zoom_lbl.pack(side='left')
        tk.Button(zf, text='＋', font=FONT_B, width=3, relief='groove',
                  command=lambda: self._set_zoom(self._zoom + 0.1)).pack(side='left')
        tk.Button(zf, text='Vừa khung', font=FONT, relief='groove',
                  command=self._zoom_fit).pack(side='left', padx=(6, 0))
        try:
            self.state('zoomed')                      # mo toan man hinh cho de doc
        except tk.TclError:
            pass

        wrap = tk.Frame(self, bg='#e9edf2')
        wrap.pack(fill='both', expand=True)
        self.cv = tk.Canvas(wrap, bg='white', highlightthickness=0)
        vs = ttk.Scrollbar(wrap, orient='vertical', command=self.cv.yview)
        hs = ttk.Scrollbar(self, orient='horizontal', command=self.cv.xview)
        self.cv.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        vs.pack(side='right', fill='y')
        self.cv.pack(side='left', fill='both', expand=True)
        hs.pack(fill='x')
        self.cv.bind_all('<MouseWheel>', self._on_wheel)

        lay = self._load_layout(kind)
        self._lay = lay
        if not lay:
            tk.Label(self.cv, text='Không tìm thấy layout_%s.json' % kind.lower(),
                     fg='red', bg='white', font=FONT_TITLE).pack(pady=40)
            return
        self._render(lay)
        self.after(150, self._zoom_fit)               # lan dau: vua chieu ngang cua so

    def _on_wheel(self, e):
        try:
            if e.state & 0x0004:                      # Ctrl + lan chuot = zoom
                self._set_zoom(self._zoom + (0.1 if e.delta > 0 else -0.1))
            else:
                self.cv.yview_scroll(int(-e.delta / 120), 'units')
        except tk.TclError:
            pass

    def _set_zoom(self, z):
        z = round(min(2.5, max(0.5, z)), 2)
        if abs(z - self._zoom) < 1e-6 or not getattr(self, '_lay', None):
            return
        self._zoom = z
        self._fontcache = {}
        self._fontcache_i = {}
        self._form_imgs = []
        self.cv.delete('all')
        self._render(self._lay)
        self._zoom_lbl.config(text=f'{int(round(z * 100))}%')

    def _zoom_fit(self):
        """Phong to cho bang vua chieu ngang khung xem (khong nho hon 100%)."""
        try:
            self.update_idletasks()
            bb = self.cv.bbox('all')
            if not bb:
                return
            w_now = (bb[2] - bb[0]) / self._zoom
            avail = self.cv.winfo_width() - 30
            if w_now > 0 and avail > 200:
                self._set_zoom(max(1.0, avail / w_now))
        except tk.TclError:
            pass

    @classmethod
    def _load_layout(cls, kind):
        if kind in cls._LAYOUTS:
            return cls._LAYOUTS[kind]
        try:
            p = resource_path('layout_%s.json' % kind.lower())
            lay = json.load(open(p, encoding='utf-8'))
            try:
                import doc_tools
                lay = doc_tools.patch_lsx_layout(lay)   # can giua khoi tron (anh 3) + tieu de III
            except Exception:
                pass
        except Exception:
            lay = None
        cls._LAYOUTS[kind] = lay
        return lay

    # ---- gia tri dong theo don duoc chon ----
    def _dynamic(self):
        po = self.po
        o = self.order
        rows = po.rows()
        d = {}
        d[(4, 5)] = o.get('don_hang') or ''      # E4 don hang
        d[(4, 11)] = o.get('ma_hang') or ''      # K4 ma hang
        d[(5, 5)] = f"{po.san_luong:,.0f}"       # E5 san luong
        d[(5, 11)] = f"{po.K5()}"                # K5 so me
        # K6 "Ma Khach Hang" = MA SAN PHAM (theo cong thuc goc XLOOKUP -> cot Ma san pham)
        _msp = o.get('ma_san_pham')
        if _msp in (None, ''):
            for _m in o.get('materials', []):
                if _m.get('nhom') not in (None, ''):
                    _msp = _m['nhom']
                    break
        d[(6, 11)] = str(_msp) if _msp not in (None, '') else ''
        # So series (I3 va cac trang =I3): lay so GIUA cua Don hang, 3 so cuoi
        seg = (o.get('don_hang') or '').split('/')
        ser = seg[1][-3:] if len(seg) >= 2 else ''
        for rr in (3, 39, 65, 100):
            d[(rr, 9)] = ser
        # Text pha tron theo loai P10/P20 (=LayLoaiP10P20(K4)): P10->HDPE, khac->T3034
        import re as _re
        _m = _re.search(r'P(\d{2})', str(o.get('ma_hang') or '').upper())
        _res = 'HDPE' if (_m and _m.group(1) == '10') else 'T3034'
        _pct = '70% HDPE + 30%' if _res == 'HDPE' else '50% T3034 + 50%'
        d[(104, 4)] = f'Pha {_pct} Filler phân tán tốt , không chấm trắng, đen'
        d[(105, 4)] = f'Pha {_pct} Filler không hoa màng,Không có bóng nước chân kim'
        # Bang 'Thong so van hanh don truoc' (TC1-6) + III van hanh may - tra theo Ma hang
        try:
            from machine_lookup import fill_dynamic as _mfill
            _mfill(d, o.get('ma_hang'))
        except Exception:
            pass
        for i in range(12):
            c = 2 + i
            if i < len(rows):
                r = rows[i]
                d[(8, c)] = r['ten'] or ''
                d[(9, c)] = f"{r['percentage']*100:.4f}%"
                d[(10, c)] = f"{r['mass_batch']:.2f}"
                d[(11, c)] = f"{r['tong_don']:,.2f}"
                d[(12, c)] = f"{r['total_order']:,.2f}"
            else:
                d[(8, c)] = ''
                d[(9, c)] = ''
                d[(10, c)] = '0.00'
                d[(11, c)] = '0.00'
                d[(12, c)] = '0.00'
        d[(9, 14)] = f"{po.N9()*100:.2f}%"       # N9 tong %
        d[(10, 14)] = f"{po.N10():,.2f}"         # N10 co me
        d[(11, 14)] = f"{po.N11():,.2f}"         # N11
        d[(12, 14)] = f"{po.N12():,.2f}"         # N12
        # tuy chon: Loai bao (S4=4,19), So cont (Q4=4,17), Quy cach xep bao (B117=117,2)
        op = getattr(self, 'opts', {}) or {}
        if op.get('loai_bao'):
            d[(4, 19)] = op['loai_bao']
        if op.get('so_cont'):
            d[(4, 17)] = op['so_cont']
        # Day chuyen: N6 va cac o =N6 o cac trang sau (N42, N68, N103)
        if op.get('day_chuyen'):
            for rr in (6, 42, 68, 103):
                d[(rr, 14)] = op['day_chuyen']
        # B117 (Quy cach xep bao): neu Hang Xep Pallet + co chuoi cach xep -> hien chuoi do
        qc = op.get('quy_cach') or ''
        if qc.strip().lower() == 'hàng xếp pallet' and op.get('cach_xep'):
            d[(117, 2)] = op['cach_xep']
        elif qc:
            d[(117, 2)] = qc
        # === Cac SUB VBA: khoi tron (ANH 3), nhiet do may tao hat (ANH 4), dong goi ... ===
        try:
            import form_rules
            self.rule_warnings = []
            form_rules.apply_all(d, o, po, op, warnings=self.rule_warnings,
                                 online_weather=bool(op.get('online_weather')),
                                 kind=getattr(self, 'kind', None))
        except Exception as _e:
            pass
        return d

    @staticmethod
    def _fmt(v, is_pct):
        if isinstance(v, (int, float)):
            if is_pct:
                return f"{v*100:.2f}%"
            if float(v).is_integer():
                return f"{int(v)}"
            return f"{v:,.2f}"
        return str(v)

    def _cellref(self, s):
        # '=E4' -> (4,5); tra ve None neu khong phai tham chieu don gian
        import re
        m = re.fullmatch(r"=\$?([A-Z]{1,3})\$?(\d+)", str(s).strip())
        if not m:
            return None
        from openpyxl.utils import column_index_from_string
        return (int(m.group(2)), column_index_from_string(m.group(1)))

    def _font(self, size, bold):
        import tkinter.font as tkfont
        key = (size, bold)
        if not hasattr(self, '_fontcache'):
            self._fontcache = {}
        if key not in self._fontcache:
            self._fontcache[key] = tkfont.Font(
                family='Segoe UI', size=size, weight='bold' if bold else 'normal')
        return self._fontcache[key]

    @staticmethod
    def _wrap_lines(text, font, avail_w):
        """So dong sau khi bo chu vao be rong avail_w (co tinh \\n va tu dai)."""
        import math
        total = 0
        for seg in str(text).split('\n'):
            if seg == '':
                total += 1
                continue
            w = font.measure(seg)
            total += max(1, math.ceil(w / max(20, avail_w)))
        return total

    def _cell_meta(self, c, dyn):
        """Tra ve (text, size, bold) da phan giai gia tri hien thi."""
        r, col = c['r'], c['c']
        if (r, col) in dyn:
            text = dyn[(r, col)]
        elif c['f']:
            ref = self._cellref(c['tv'])
            text = dyn[ref] if (ref and ref in dyn) else self._fmt(c['cv'], c['p'])
        else:
            text = self._fmt(c['tv'] if c['tv'] != '' else c['cv'], c['p'])
        from pdf_export import _form_pt
        size = _form_pt(c['sz'])             # chu to nhu ban Excel in ra
        return ('' if text in ('', 'None') else str(text)), size, bool(c['b'])

    def _render(self, lay):
        # gian cot / gian hang giong het ban PDF (phu kin to A4 ngang)
        try:
            from pdf_export import tune_form_layout
            key = id(lay)
            if getattr(self, '_tuned_key', None) != key:
                self._tuned_lay = tune_form_layout(self.kind, self.order, self.po.san_luong, lay, self.opts)
                self._tuned_key = key
            lay = self._tuned_lay
        except Exception:
            pass
        self._draw_layout(lay, self._dynamic(), place_logo=True)

    def _place_form_images(self, x_of, y_of):
        """Ve anh Tem (cols 6-10) va Hinh anh bao bi (cols 11-13) vao vung rows 114-121."""
        op = getattr(self, 'opts', {}) or {}
        self._form_imgs = getattr(self, '_form_imgs', [])
        # (path, box, fill?) - Tem gian day khung, Bao bi giu ti le
        boxes = [(op.get('stamp_img'), self.IMG_STAMP_BOX, True),
                 (op.get('pkg_img'), self.IMG_PKG_BOX, False)]
        for path, (c0, r0, c1, r1), fill in boxes:
            if not path or not os.path.exists(path):
                continue
            x0, y0 = x_of(c0), y_of(r0)
            x1, y1 = x_of(c1 + 1), y_of(r1 + 1)
            bw, bh = int(x1 - x0 - 6), int(y1 - y0 - 6)
            if bw < 8 or bh < 8:
                continue
            img = self._load_fit_image(path, bw, bh, fill)
            if img is None:
                continue
            self._form_imgs.append(img)
            self.cv.create_image((x0 + x1) / 2, (y0 + y1) / 2, image=img, anchor='center')

    @staticmethod
    def _load_fit_image(path, bw, bh, fill=False):
        """Doc anh (png/jpg) vua khung bw x bh. fill=True: gian day khung; nguoc lai giu ti le."""
        try:
            from PIL import Image, ImageTk
            im = Image.open(path)
            if fill:
                im = im.resize((max(1, bw), max(1, bh)), Image.LANCZOS)
            else:
                im.thumbnail((bw, bh), Image.LANCZOS)
            return ImageTk.PhotoImage(im)
        except Exception:
            pass
        # fallback: chi PNG/GIF, thu nho bang subsample
        try:
            import math
            base = tk.PhotoImage(file=path)
            k = max(1, math.ceil(max(base.width() / bw, base.height() / bh)))
            return base.subsample(k)
        except Exception:
            return None


    def _export(self):
        try:
            from order_form import fill_order_form
        except Exception as e:
            messagebox.showerror('Lỗi', str(e))
            return
        o = self.order
        default = __import__('doc_tools').lsx_file_name(o, '.xlsx')
        path = filedialog.asksaveasfilename(defaultextension='.xlsx', initialfile=default,
                                            filetypes=[('Excel', '*.xlsx')])
        if not path:
            return
        try:
            fill_order_form(self.kind, o, self.po.san_luong, path, self.opts)
            messagebox.showinfo('Xong', f'Đã xuất form gốc:\n{path}')
        except Exception as e:
            messagebox.showerror('Lỗi xuất form', str(e))

    def _export_pdf(self):
        export_order_pdf(self, self.kind, self.order, self.po.san_luong, self._open_file, self.opts)

    def _print(self):
        app = self.master.winfo_toplevel() if self.master else None
        PrintDialog(self, self.kind, self.order, self.po.san_luong, self.opts,
                    settings_owner=app)

    def _open_file(self, path):
        try:
            if sys.platform.startswith('win'):
                os.startfile(path)             # noqa
            elif sys.platform == 'darwin':
                import subprocess
                subprocess.Popen(['open', path])
            else:
                import subprocess
                subprocess.Popen(['xdg-open', path])
        except Exception as e:
            messagebox.showinfo('Đã tạo', f'{path}\n({e})', parent=self)


# =====================================================================
#  TAB CHUNG: hien thi mot sheet dang bang (co tim kiem)
# =====================================================================
def data_writable_path():
    """data.json ben canh app (.py/.exe) - noi ghi duoc de luu don moi."""
    if getattr(sys, 'frozen', False):
        base = os.path.dirname(sys.executable)
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, 'data.json')


def sheets_writable_path():
    if getattr(sys, 'frozen', False):
        base = os.path.dirname(sys.executable)
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, 'sheets.json')


def reset_sheet_caches():
    """Xoa cache sheets.json trong cac module tra cuu (Tieu chuan dong goi, Don gia,
    Thong so van hanh may) -> lan Xem truoc / Xuat Excel / PDF sau se doc du lieu MOI."""
    for mod_name in ('form_rules', 'machine_lookup'):
        try:
            mod = __import__(mod_name)
            mod.reset_cache()
        except Exception:
            pass


def save_sheets_file(sheets, force_titles=()):
    """Luu cac sheet (file sheets.json hoac MAY CHU) va xoa cache de cac noi khac cap nhat luon.
    Che do may chu: chi ghi sheet thay doi; may khac vua sua -> tu gop theo dong."""
    datastore.save_sheets(sheets, force_titles=force_titles)
    reset_sheet_caches()


def save_data_file(data, force=False):
    """Luu don hang (file data.json hoac MAY CHU - chi ghi lenh thay doi)."""
    datastore.save_data(data, force=force)


def users_writable_path():
    if getattr(sys, 'frozen', False):
        base = os.path.dirname(sys.executable)
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, 'users.json')


def settings_writable_path():
    if getattr(sys, 'frozen', False):
        base = os.path.dirname(sys.executable)
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, 'settings.json')


# Duong dan mac dinh (may nguoi dung)
DEFAULT_STAMP_DIR = r'C:\Tool-QC-QA-RD-HitechPlast\Data\HitechPlast order creation software\Year 2026\Stamp images'
DEFAULT_PKG_DIR = r'C:\Tool-QC-QA-RD-HitechPlast\Data\HitechPlast order creation software\Year 2026\Packaging Picture'


def load_settings():
    p = settings_writable_path()
    s = {}
    if os.path.exists(p):
        try:
            s = json.load(open(p, encoding='utf-8'))
        except Exception:
            s = {}
    s.setdefault('stamp_dir', DEFAULT_STAMP_DIR)
    s.setdefault('pkg_dir', DEFAULT_PKG_DIR)
    return s


def save_settings(s):
    try:
        json.dump(s, open(settings_writable_path(), 'w', encoding='utf-8'),
                  ensure_ascii=False, indent=1)
    except Exception:
        pass


POLICY_KEY = 'app_policy'          # cai dat chung moi may: {'instance': 'single' | 'multi'}


def load_app_policy():
    """Cai dat chung (admin dat trong Cap quyen). Che do may chu: doc tu may chu (htp_kv) va ghi nho
    vao settings.json (dung khi mat ket noi). Che do file: settings.json."""
    pol = None
    if datastore.is_server():
        try:
            pol = datastore.server().kv_get(POLICY_KEY, track=False)
        except Exception:
            pol = None
    st = load_settings()
    if isinstance(pol, dict):
        if st.get(POLICY_KEY) != pol:
            st[POLICY_KEY] = pol
            save_settings(st)
    else:
        pol = st.get(POLICY_KEY) if isinstance(st.get(POLICY_KEY), dict) else {}
    out = {'instance': 'single', 'offline_main': True, 'offline_qc': True, 'show_server_setup': True}
    out.update(pol or {})
    return out


def offline_policy():
    """(chep_du_lieu_chinh, chep_du_lieu_QC) - doc tu ban ghi nho settings.json (KHONG hoi may chu).
    Admin dat trong 🔑 Cap quyen -> 'Ban sao offline'."""
    try:
        pol = load_settings().get(POLICY_KEY) or {}
        if not isinstance(pol, dict):
            pol = {}
    except Exception:
        pol = {}
    return bool(pol.get('offline_main', True)), bool(pol.get('offline_qc', True))


def cached_policy():
    """Cai dat chung ghi nho trong settings.json (doc duoc ca khi KHONG ket noi may chu)."""
    try:
        pol = load_settings().get(POLICY_KEY) or {}
        return pol if isinstance(pol, dict) else {}
    except Exception:
        return {}


def save_app_policy(pol):
    if datastore.is_server():
        datastore.server().kv_put(POLICY_KEY, pol, force=True)
    st = load_settings()
    st[POLICY_KEY] = pol
    save_settings(st)


# ---- THU MUC NGUON ANH TREN MAY CHU (chi may chu ghi; may khac xem) ----
#  Luu y: 'stamp_dir'/'pkg_dir' trong settings.json la THU MUC TREN MAY NAY (cuc bo, moi may rieng).
#  Rieng THU MUC NGUON cua CHINH MAY CHU duoc luu vao kho chung (kv) de moi may XEM duoc.
IMG_CFG_KEY = 'img_cfg'


def _is_server_host():
    """May DANG CHAY co phai chinh MAY CHU CSDL dang hoat dong khong.
    Ho tro 2 may chu (chinh + du phong): may nao dang la server (dia chi cau hinh
    trung LAN/Tailscale/Radmin/ten may cua may nay, hoac cau hinh 'localhost') -> True.
    => Khi failover sang may du phong, chinh may du phong se la host va ghi folder cua no."""
    if not datastore.is_server():
        return False
    try:
        import presence
        import socket
        hosts = set(presence.server_hosts() or [])
        if not hosts:
            return False
        if any(h in ('localhost', '127.0.0.1', '::1') for h in hosts):
            return True
        mine = set()
        for fn in (presence.lan_ip, presence.ts_ip, presence.rm_ip):
            try:
                v = fn()
                if v:
                    mine.add(str(v).strip().lower())
            except Exception:
                pass
        try:
            mine.add(socket.gethostname().strip().lower())
        except Exception:
            pass
        return bool(mine & hosts)
    except Exception:
        return False


def server_img_dirs():
    """Thu muc nguon anh cau hinh TREN MAY CHU (de hien thi cho moi may). {} neu chua co."""
    if not datastore.is_server():
        return {}
    try:
        cfg = datastore.server().kv_get(IMG_CFG_KEY, track=False)
        return cfg if isinstance(cfg, dict) else {}
    except Exception:
        return {}


def save_img_cfg(st):
    """Luu cai dat anh: LUON ghi may nay (settings.json cuc bo). CHI khi may nay LA MAY CHU
    moi ghi thu muc nguon len kho chung (kv) de cac may khac xem duoc."""
    save_settings(st)
    if _is_server_host():
        try:
            datastore.server().kv_put(IMG_CFG_KEY, {
                'stamp_dir': st.get('stamp_dir', ''),
                'pkg_dir': st.get('pkg_dir', ''),
                'host': os.environ.get('COMPUTERNAME') or '',
            }, force=True)
        except Exception:
            pass


def load_store():
    """Doc kho tai khoan ma hoa (users.dat). Lan dau -> nhap tu users.json."""
    d = secure_store.load()
    if d is None:
        d = first_run_init()
    return d


def first_run_init():
    """Lan dau chay: nhap users.json (neu co) vao users.dat roi xoa users.json."""
    ext = os.path.join(secure_store.app_dir(), 'users.json')
    src = ext if os.path.exists(ext) else resource_path('users.json')
    data = None
    if os.path.exists(src):
        try:
            data = secure_store.import_plain_json(src)
        except Exception:
            data = None
    if not data or not data.get('users'):
        data = {'users': [{'username': 'admin', 'password': 'admin123',
                           'hoten': 'Quản trị viên', 'tabs': 'all'}], 'license_log': []}
    data = secure_store.save(data)
    if os.path.exists(ext):
        try:
            os.remove(ext)                  # xoa ban chu thuong (tranh lo mat khau / sua tay)
        except Exception:
            pass
    return data


DEFAULT_APP_YEAR = 2026
APP_SUBTITLE = 'Quản Lý & Lệnh Sản Xuất By Nguyễn Hoài Nam'


def get_app_year():
    """Nam hien thi tren phan mem (admin doi trong cua so Cap quyen, luu trong users.dat)."""
    try:
        y = int(load_store().get('app_year') or DEFAULT_APP_YEAR)
        return y if 2000 <= y <= 2100 else DEFAULT_APP_YEAR
    except Exception:
        return DEFAULT_APP_YEAR


def set_app_year(year):
    d = load_store()
    d['app_year'] = int(year)
    secure_store.save(d)


# list users da tai -> rev cua ban goc (de GOP khi may khac vua sua - che do may chu)
_USERS_SRC = {}


def save_users_list(users):
    src = _USERS_SRC.get(id(users))
    if datastore.is_server() and src is not None and src[0] is users:
        base = secure_store._LOADED.get(src[1])
        if base is not None:
            import copy
            mine = copy.deepcopy(base)
            mine['users'] = users
            secure_store.save(mine)       # may chu: gop voi thay doi cua may khac (theo tung tai khoan/truong)
            return
    try:
        d = load_store()
    except secure_store.StoreError:
        d = secure_store.best_available() or {}
    d['users'] = users
    secure_store.save(d)


def update_user_in_file(username, **changes):
    """Cap nhat 1 ban ghi user trong users.json."""
    try:
        users = load_users()
        for u in users:
            if u.get('username') == username:
                u.update(changes)
        save_users_list(users)
        return True
    except Exception:
        return False


class ChangePassword(tk.Toplevel):
    """Doi mat khau: cu / moi / xac nhan. Nut 'Doi mat khau' chi hien khi moi==xac nhan."""
    def __init__(self, app, user):
        super().__init__(app)
        self.app = app
        self.user = user
        self.title('Đổi mật khẩu')
        self.geometry('420x260')
        self.configure(bg='#f4f6f9')
        self.resizable(False, False)
        try:
            self.grab_set()
        except Exception:
            pass

        tk.Label(self, text='🔒 Đổi mật khẩu', bg='#f4f6f9', font=FONT_TITLE).pack(anchor='w', padx=16, pady=(14, 8))
        box = tk.Frame(self, bg='#f4f6f9')
        box.pack(fill='x', padx=16)
        self.v_old = tk.StringVar()
        self.v_new = tk.StringVar()
        self.v_cf = tk.StringVar()
        rows = [('Mật khẩu cũ:', self.v_old), ('Mật khẩu mới:', self.v_new),
                ('Xác nhận mật khẩu mới:', self.v_cf)]
        for lb, var in rows:
            fr = tk.Frame(box, bg='#f4f6f9')
            fr.pack(fill='x', pady=4)
            tk.Label(fr, text=lb, bg='#f4f6f9', font=FONT, width=20, anchor='e').pack(side='left')
            e = tk.Entry(fr, textvariable=var, show='●', width=22, font=FONT)
            e.pack(side='left')
            e.bind('<KeyRelease>', lambda ev: self._check())

        self.msg = tk.Label(self, text='', bg='#f4f6f9', fg='#c0392b', font=FONT)
        self.msg.pack(pady=(6, 0))

        self.btnfr = tk.Frame(self, bg='#f4f6f9')
        self.btnfr.pack(fill='x', padx=16, pady=8)
        tk.Button(self.btnfr, text='Đóng', font=FONT, relief='flat', padx=12, pady=6,
                  command=self.destroy).pack(side='right', padx=6)
        # nut doi mat khau - an ban dau, chi hien khi 2 mat khau moi trung
        self.btn_change = tk.Button(self.btnfr, text='✔ Đổi mật khẩu', font=FONT_B,
                                    bg='#2d6a4f', fg='white', relief='flat', padx=16, pady=6,
                                    command=self._do)

    def _check(self):
        n, c = self.v_new.get(), self.v_cf.get()
        self.btn_change.pack_forget()
        if not n:
            self.msg.config(text='')
        elif n == c:
            self.msg.config(text='✓ Hai mật khẩu mới trùng khớp', fg='#2d6a4f')
            self.btn_change.pack(side='right', padx=6)   # HIEN nut khi trung
        else:
            self.msg.config(text='✗ Mật khẩu xác nhận chưa khớp', fg='#c0392b')

    def _do(self):
        if str(self.v_old.get()) != str(self.user.get('password', '')):
            messagebox.showwarning('Sai', 'Mật khẩu cũ không đúng.', parent=self)
            return
        if not self.v_new.get():
            messagebox.showwarning('Thiếu', 'Nhập mật khẩu mới.', parent=self)
            return
        new = self.v_new.get()
        self.user['password'] = new
        self.user['pw_changed'] = True
        update_user_in_file(self.user.get('username'), password=new, pw_changed=True)
        messagebox.showinfo('Xong', 'Đã đổi mật khẩu thành công.', parent=self)
        self.destroy()


class UserAdmin(tk.Toplevel):
    """Quan tri: cap quyen cho cac nick thanh vien (sua users.json)."""
    def __init__(self, master, perms):
        super().__init__(master)
        self.perms = perms                 # [(key, label), ...]
        self.title('Cấp quyền người dùng - Quản trị')
        sh = self.winfo_screenheight()
        self.geometry('1060x%d' % max(600, min(780, sh - 80)))
        self.minsize(760, 560)
        self.configure(bg='#f4f6f9')
        self.users = load_users()          # list dict
        self.idx = None

        hdr = tk.Frame(self, bg='#f4f6f9')
        hdr.pack(fill='x', padx=14, pady=(12, 6))
        tk.Label(hdr, text='🔑 Cấp quyền cho các tài khoản', bg='#f4f6f9',
                 font=FONT_TITLE).pack(side='left')
        yf = tk.Frame(hdr, bg='#e3f2fd', padx=8, pady=4)
        yf.pack(side='right')
        tk.Label(yf, text='📅 Năm phần mềm:', bg='#e3f2fd', font=FONT_B).pack(side='left')
        self.v_year = tk.StringVar(value=str(get_app_year()))
        try:
            import year_data
            ychoices = [str(y) for y in year_data.year_choices(get_app_year())]
        except Exception:
            ychoices = [str(get_app_year())]
        self.cb_year = ttk.Combobox(yf, textvariable=self.v_year, width=6, font=FONT_B,
                                    justify='center', values=ychoices)
        self.cb_year.pack(side='left', padx=4)
        self.cb_year.bind('<Return>', lambda e: self._save_year())
        tk.Button(yf, text='📅 Chuyển năm', font=FONT, bg='#1d3557', fg='white', relief='flat',
                  padx=8, command=self._save_year).pack(side='left')
        self.lbl_years = tk.Label(yf, text='', bg='#e3f2fd', fg='#555', font=('Segoe UI', 8))
        self.lbl_years.pack(side='left', padx=(6, 0))
        self._show_saved_years()

        bot = tk.Frame(self, bg='#f4f6f9')
        bot.pack(side='bottom', fill='x', padx=14, pady=10)
        tk.Button(bot, text='💾 Lưu ra file', font=FONT_B, bg='#2d6a4f', fg='white',
                  relief='flat', padx=16, pady=7, command=self._save_file).pack(side='right')
        tk.Button(bot, text='Đóng', font=FONT, relief='flat', padx=14, pady=7,
                  command=self.destroy).pack(side='right', padx=8)
        tk.Button(bot, text='🖧 Máy chủ dữ liệu (dùng chung)', font=FONT, bg='#dfe6ee', relief='flat',
                  padx=12, pady=7, command=self._server_setup).pack(side='left')
        tk.Button(bot, text='📋 Nhật ký bất thường', font=FONT_B, bg='#c62828', fg='white', relief='flat',
                  padx=12, pady=7, command=self._audit_log).pack(side='left', padx=(8, 0))
        tk.Button(bot, text='📥 Cập nhật dữ liệu từ bản cũ', font=FONT_B, bg='#0b7285', fg='white',
                  relief='flat', padx=12, pady=7, command=self._migrate_old).pack(side='left', padx=8)

        body = tk.Frame(self, bg='#f4f6f9')
        body.pack(fill='both', expand=True, padx=14)

        # trai: danh sach user
        left = tk.Frame(body, bg='#f4f6f9')
        left.pack(side='left', fill='y')
        tk.Label(left, text='Tài khoản:', bg='#f4f6f9', font=FONT_B).pack(anchor='w')
        self.lb = tk.Listbox(left, width=24, height=20, font=FONT, exportselection=False)
        self.lb.pack(fill='y', expand=True)
        self.lb.bind('<<ListboxSelect>>', lambda e: self._select())
        bl = tk.Frame(left, bg='#f4f6f9')
        bl.pack(fill='x', pady=4)
        tk.Button(bl, text='➕ Thêm', font=FONT, command=self._add).pack(side='left')
        tk.Button(bl, text='🗑️ Xóa', font=FONT, command=self._del).pack(side='left', padx=4)

        # phai: chi tiet
        right = tk.Frame(body, bg='#f4f6f9')
        right.pack(side='left', fill='both', expand=True, padx=(14, 0))
        self.v_user = tk.StringVar()
        self.v_pass = tk.StringVar()
        self.v_hoten = tk.StringVar()
        self.v_expire = tk.StringVar()
        for lb_txt, var in [('Tên đăng nhập:', self.v_user), ('Mật khẩu:', self.v_pass),
                            ('Họ tên:', self.v_hoten), ('Hạn sử dụng:', self.v_expire)]:
            fr = tk.Frame(right, bg='#f4f6f9')
            fr.pack(fill='x', pady=3)
            tk.Label(fr, text=lb_txt, bg='#f4f6f9', font=FONT, width=14, anchor='e').pack(side='left')
            tk.Entry(fr, textvariable=var, width=30, font=FONT).pack(side='left')
        tk.Label(right, text='Hạn sử dụng nhập dạng ddmmyyyy (vd 21122026). Đến đúng ngày đó '
                             'tài khoản bị khóa. Để trống = không giới hạn.',
                 bg='#f4f6f9', fg='#666', font=('Segoe UI', 8), wraplength=420,
                 justify='left').pack(anchor='w', padx=(4, 0))
        rc = tk.Frame(self, bg='#fff3cd')
        rc.pack(side='bottom', fill='x', padx=14)
        tk.Label(rc, text='🗝 Mã khôi phục (admin ghi lại & giữ bí mật, dùng khi dữ liệu tài khoản bị xóa/sửa):',
                 bg='#fff3cd', font=('Segoe UI', 8)).pack(side='left', padx=6, pady=3)
        rv = tk.StringVar(value=lic.recovery_code())
        tk.Entry(rc, textvariable=rv, state='readonly', font=('Consolas', 10, 'bold'), width=22,
                 readonlybackground='#fff3cd', relief='flat').pack(side='left')
        # che do dang nhap nhieu may / dang xuat (dung nhip tim presence.py - che do may chu)
        sess = tk.Frame(right, bg='#eef3f8', padx=8, pady=4, highlightthickness=1, highlightbackground='#d0d7e1')
        sess.pack(fill='x', pady=(6, 0))
        self.v_login_mode = tk.StringVar(value='multi')
        self.v_logout_mode = tk.StringVar(value='one')
        r1 = tk.Frame(sess, bg='#eef3f8')
        r1.pack(fill='x')
        tk.Label(r1, text='Đăng nhập:', bg='#eef3f8', font=FONT_B, width=11, anchor='w').pack(side='left')
        tk.Radiobutton(r1, text='Nhiều máy cùng lúc', variable=self.v_login_mode, value='multi',
                       bg='#eef3f8', font=FONT).pack(side='left')
        tk.Radiobutton(r1, text='Chỉ 1 máy (đăng nhập máy mới → máy cũ tự đăng xuất)',
                       variable=self.v_login_mode, value='single', bg='#eef3f8', font=FONT).pack(side='left', padx=8)
        r2 = tk.Frame(sess, bg='#eef3f8')
        r2.pack(fill='x')
        tk.Label(r2, text='Đăng xuất:', bg='#eef3f8', font=FONT_B, width=11, anchor='w').pack(side='left')
        tk.Radiobutton(r2, text='Chỉ máy đang dùng', variable=self.v_logout_mode, value='one',
                       bg='#eef3f8', font=FONT).pack(side='left')
        tk.Radiobutton(r2, text='Tất cả các máy đang đăng nhập tài khoản này',
                       variable=self.v_logout_mode, value='all', bg='#eef3f8', font=FONT).pack(side='left', padx=8)
        tk.Button(r2, text='⏏ Đăng xuất ngay trên mọi máy', font=FONT_B, bg='#6c757d', fg='white',
                  relief='flat', padx=8, command=self._logout_everywhere).pack(side='right')
        r3 = tk.Frame(sess, bg='#eef3f8')
        r3.pack(fill='x', pady=(2, 0))
        tk.Label(r3, text='Mỗi máy:', bg='#eef3f8', font=FONT_B, width=11, anchor='w').pack(side='left')
        try:
            inst = load_app_policy().get('instance', 'single')
        except Exception:
            inst = 'single'
        self.v_instance = tk.StringVar(value=inst)
        tk.Radiobutton(r3, text='Chỉ mở 1 phần mềm', variable=self.v_instance, value='single',
                       bg='#eef3f8', font=FONT, command=self._save_instance).pack(side='left')
        tk.Radiobutton(r3, text='Cho phép mở nhiều phần mềm', variable=self.v_instance, value='multi',
                       bg='#eef3f8', font=FONT, command=self._save_instance).pack(side='left', padx=8)
        self.lbl_instance = tk.Label(r3, text='(áp dụng chung cho tất cả các máy)', bg='#eef3f8', fg='#666',
                                     font=('Segoe UI', 8))
        self.lbl_instance.pack(side='left')
        # --- Ban sao offline: chep du lieu tu may chu ve tung may (nang -> co the gay giat) ---
        r4 = tk.Frame(sess, bg='#eef3f8')
        r4.pack(fill='x', pady=(2, 0))
        tk.Label(r4, text='Bản offline:', bg='#eef3f8', font=FONT_B, width=11, anchor='w').pack(side='left')
        try:
            _pol = load_app_policy()
        except Exception:
            _pol = {}
        self.v_off_main = tk.BooleanVar(value=bool(_pol.get('offline_main', True)))
        self.v_off_qc = tk.BooleanVar(value=bool(_pol.get('offline_qc', True)))
        tk.Checkbutton(r4, text='Chép sheet + lệnh SX về máy', variable=self.v_off_main, bg='#eef3f8',
                       font=FONT, command=self._save_offline_pol).pack(side='left')
        tk.Checkbutton(r4, text='Chép dữ liệu QC về máy (nặng, dễ gây giật)', variable=self.v_off_qc,
                       bg='#eef3f8', font=FONT, command=self._save_offline_pol).pack(side='left', padx=8)
        self.lbl_offline = tk.Label(sess, text='Tắt = không giật khi đang làm việc, nhưng khi MÁY CHỦ TẮT thì máy '
                                               'con không mở được phần mềm bằng bản offline (hoặc dùng bản cũ).',
                                    bg='#eef3f8', fg='#666', font=('Segoe UI', 8), justify='left', anchor='w')
        self.lbl_offline.pack(fill='x', padx=(92, 0))
        # --- Cho may con thay nut "🖧 Cai dat may chu" khi khong ket noi duoc may chu (yeu cau #1) ---
        r5 = tk.Frame(sess, bg='#eef3f8')
        r5.pack(fill='x', pady=(2, 0))
        tk.Label(r5, text='Máy chủ:', bg='#eef3f8', font=FONT_B, width=11, anchor='w').pack(side='left')
        self.v_show_setup = tk.BooleanVar(value=bool(_pol.get('show_server_setup', True)))
        tk.Checkbutton(r5, text='Hiện nút "🖧 Cài đặt máy chủ" khi không kết nối được máy chủ',
                       variable=self.v_show_setup, bg='#eef3f8', font=FONT,
                       command=self._save_setup_pol).pack(side='left')
        self.lbl_setup = tk.Label(sess, text='Tắt = máy con chỉ thấy "🔁 Thử lại" và "📂 Dùng offline" (đường dẫn '
                                             'bản sao offline); quản trị viên vẫn vào được bằng nút 🔑.',
                                  bg='#eef3f8', fg='#666', font=('Segoe UI', 8), justify='left', anchor='w')
        self.lbl_setup.pack(fill='x', padx=(92, 0))
        self.lbl_status = tk.Label(right, text='', bg='#f4f6f9', font=FONT_B)
        self.lbl_status.pack(anchor='w', pady=(2, 0))
        act = tk.Frame(right, bg='#f4f6f9')
        act.pack(fill='x', pady=(6, 4))
        tk.Button(act, text='✔ Áp dụng vào tài khoản', font=FONT_B, bg='#1d3557', fg='white',
                  relief='flat', padx=12, pady=6, command=self._apply).pack(side='left')
        tk.Button(act, text='🎟 Tạo mã gia hạn', font=FONT_B, bg='#7b2cbf', fg='white',
                  relief='flat', padx=12, pady=6, command=self._make_code).pack(side='left', padx=8)
        tk.Button(act, text='🔓 Mở khóa', font=FONT_B, bg='#e07a00', fg='white',
                  relief='flat', padx=12, pady=6, command=self._unlock).pack(side='left')
        tk.Button(act, text='🔑 Tạo mã mở sheet', font=FONT_B, bg='#0b7285', fg='white',
                  relief='flat', padx=12, pady=6, command=self._make_tab_code).pack(side='left', padx=8)

        self.all_var = tk.BooleanVar()
        tk.Checkbutton(right, text='Toàn quyền (xem tất cả tab)', variable=self.all_var,
                       bg='#f4f6f9', font=FONT_B, command=self._toggle_all).pack(anchor='w', pady=(8, 2))
        tk.Label(right, text='Hoặc chọn từng tab được phép:', bg='#f4f6f9', font=FONT).pack(anchor='w')
        self.tabvars = {}
        # Vung chon tab CO THANH CUON: nhieu sheet van hien du, khong bi cat o cuoi cua so
        box = tk.Frame(right, bg='#f4f6f9', highlightthickness=1, highlightbackground='#d0d7e1')
        box.pack(fill='both', expand=True, pady=(2, 4))
        cv = tk.Canvas(box, bg='#f4f6f9', highlightthickness=0)
        sb = ttk.Scrollbar(box, orient='vertical', command=cv.yview)
        cv.configure(yscrollcommand=sb.set)
        sb.pack(side='right', fill='y')
        cv.pack(side='left', fill='both', expand=True)
        pf = tk.Frame(cv, bg='#f4f6f9')
        win = cv.create_window((0, 0), window=pf, anchor='nw')

        def _sync(_e=None):
            cv.configure(scrollregion=cv.bbox('all'))
            # het can cuon -> an thanh cuon
            if pf.winfo_reqheight() <= cv.winfo_height():
                sb.pack_forget()
                cv.yview_moveto(0)
            elif not sb.winfo_ismapped():
                sb.pack(side='right', fill='y', before=cv)
        pf.bind('<Configure>', _sync)
        cv.bind('<Configure>', lambda e: (cv.itemconfigure(win, width=e.width), _sync()))

        def _wheel(e):
            if pf.winfo_reqheight() > cv.winfo_height():
                cv.yview_scroll(-1 if (getattr(e, 'delta', 0) > 0 or e.num == 4) else 1, 'units')
        for w in (cv, pf):
            w.bind('<Enter>', lambda e: (cv.bind_all('<MouseWheel>', _wheel),
                                         cv.bind_all('<Button-4>', _wheel), cv.bind_all('<Button-5>', _wheel)))
            w.bind('<Leave>', lambda e: (cv.unbind_all('<MouseWheel>'),
                                         cv.unbind_all('<Button-4>'), cv.unbind_all('<Button-5>')))

        ncol = 3
        nrow = (len(self.perms) + ncol - 1) // ncol
        for c in range(ncol):
            pf.grid_columnconfigure(c, weight=1, uniform='perm')
        for i, (key, label) in enumerate(self.perms):
            var = tk.BooleanVar()
            self.tabvars[key] = var
            # xep theo COT (doc tu tren xuong, het cot sang cot ke)
            tk.Checkbutton(pf, text=label, variable=var, bg='#f4f6f9', font=FONT, anchor='w').grid(
                row=i % nrow, column=i // nrow, sticky='w', padx=(4, 12), pady=0)
        tk.Label(right, text=f'Tổng cộng {len(self.perms)} tab – lăn chuột trong khung để xem hết.',
                 bg='#f4f6f9', fg='#666', font=('Segoe UI', 8)).pack(anchor='w')


        self._refresh_list()

    def _server_setup(self):
        import server_setup
        server_setup.open_dialog(self)

    def _audit_log(self):
        import audit_ui
        audit_ui.open_window(self)

    def _migrate_old(self):
        """ADMIN: chep du lieu tu thu muc phan mem BAN CU sang ban nay (migrate_data.py)."""
        try:
            import migrate_data
        except Exception as e:
            messagebox.showerror('Lỗi', f'Không nạp được chức năng cập nhật dữ liệu:\n{e}', parent=self)
            return
        migrate_data.open_dialog(self, app=self.master)

    def _show_saved_years(self):
        try:
            import year_data
            ys = sorted(set(year_data.archived_years()) | {get_app_year()})
            self.lbl_years.config(text='Có dữ liệu: ' + ', '.join(str(y) for y in ys))
        except Exception:
            pass

    def _save_year(self):
        """Chon nam: luu du lieu nam dang mo, nap du lieu nam duoc chon (hoac lam moi)."""
        t = self.v_year.get().strip()
        if not (t.isdigit() and 2000 <= int(t) <= 2100):
            messagebox.showwarning('Sai năm', 'Nhập năm dạng 4 số, vd 2027.', parent=self)
            return
        new_y, old_y = int(t), get_app_year()
        if datastore.is_offline():
            messagebox.showwarning('Đang OFFLINE', 'Máy chủ chưa bật – phần mềm đang chạy OFFLINE.\n'
                                   'Không chuyển năm được lúc này (để tránh lẫn dữ liệu giữa các máy).\n'
                                   'Đợi máy chủ bật (phần mềm tự kết nối) rồi chuyển năm.', parent=self)
            return
        if new_y == old_y:
            messagebox.showinfo('Năm phần mềm', f'Phần mềm đang ở năm {old_y}.', parent=self)
            return
        app = self.master
        if not hasattr(app, 'sheets') or not hasattr(app, 'data'):
            messagebox.showerror('Lỗi', 'Không truy cập được dữ liệu phần mềm.', parent=self)
            return
        import year_data
        where = ('trên MÁY CHỦ (mọi máy sẽ tự chuyển theo)' if datastore.is_server()
                 else f'trong thư mục "{year_data.ARCHIVE_DIR_NAME}"')
        self.config(cursor='watch')
        self.update_idletasks()
        try:                        # chi dem dong (khong tai ca nam ve) -> hop xac nhan hien nhanh
            saved = year_data.year_summary(new_y)
        finally:
            self.config(cursor='')
        sheets_txt = ('Data Export, Data Domestic, Data Sample, Container Volume, '
                      'Bảng tổng thành phẩm')
        if saved:
            what = (f'Năm {new_y} ĐÃ CÓ dữ liệu lưu (lưu lúc {saved[0]})\n'
                    f'→ sẽ NẠP LẠI để xem/sửa:\n{saved[1]}')
        else:
            what = (f'Năm {new_y} CHƯA có dữ liệu\n→ các sheet {sheets_txt} sẽ được '
                    f'LÀM MỚI (để trống, giữ tiêu đề cột).')
        if not messagebox.askyesno(
                'Chuyển năm phần mềm',
                f'Chuyển từ năm {old_y} sang năm {new_y}?\n\n'
                f'• Dữ liệu năm {old_y} ({sheets_txt}) sẽ được LƯU LẠI {where}'
                f' — chọn lại {old_y} bất cứ lúc nào để xem.\n'
                f'• {what}\n\n'
                'Các sheet khác (Đơn giá, NVL thay thế, Thông số máy, Tiêu chuẩn đóng gói, '
                'CLAIM/TDS/COA/MSDS) giữ nguyên.\nDữ liệu sẽ cập nhật ngay, không cần mở lại phần mềm.',
                parent=self):
            return

        def save_data(d):
            save_data_file(d)

        def finish(res):
            app._year_switching = False
            try:
                win.grab_release()
                win.destroy()
            except Exception:
                pass
            if isinstance(res, Exception):
                messagebox.showerror('Lỗi chuyển năm',
                                     f'Không chuyển được năm:\n{res}\n\n'
                                     f'Dữ liệu năm {old_y} vẫn còn {where} (nếu đã lưu kịp).', parent=self)
                return
            had, _p_old = res
            set_app_year(new_y)
            self.destroy()
            # cap nhat NGAY tren cua so dang mo (khong dong / mo lai phan mem)
            if hasattr(app, 'switch_year_live'):
                ok = app.switch_year_live(new_y)
            elif hasattr(app, 'apply_year'):
                app.apply_year(new_y)
                ok = True
            else:
                ok = False
            if ok:
                messagebox.showinfo('Đã chuyển năm',
                                    f'Đã lưu dữ liệu năm {old_y} và chuyển sang năm {new_y}'
                                    + (' (đã nạp lại dữ liệu cũ).' if had else ' (dữ liệu mới, trống).')
                                    + '\nĐã cập nhật ngay – không cần mở lại phần mềm.', parent=app)

        # --- cua so tien trinh (khong treo "Not Responding")
        win = tk.Toplevel(self)
        win.title('Đang chuyển năm…')
        win.transient(self)
        win.resizable(False, False)
        win.protocol('WM_DELETE_WINDOW', lambda: None)       # dang ghi du lieu -> khong cho dong
        tk.Label(win, text=f'⏳ Đang chuyển từ năm {old_y} sang năm {new_y}…', font=('Segoe UI', 11, 'bold'),
                 padx=20, pady=(12)).pack(anchor='w')
        lbl = tk.Label(win, text='Chuẩn bị…', font=('Segoe UI', 9), fg='#1d3557', padx=20, anchor='w', width=52)
        lbl.pack(anchor='w')
        pb = ttk.Progressbar(win, mode='indeterminate', length=380)
        pb.pack(padx=20, pady=(8, 6))
        tk.Label(win, text='Vui lòng đợi, đừng tắt phần mềm.', font=('Segoe UI', 8), fg='#888',
                 padx=20, pady=(0)).pack(anchor='w', pady=(0, 12))
        win.update_idletasks()
        try:
            x = self.winfo_rootx() + (self.winfo_width() - win.winfo_width()) // 2
            y = self.winfo_rooty() + (self.winfo_height() - win.winfo_height()) // 3
            win.geometry(f'+{max(0, x)}+{max(0, y)}')
        except Exception:
            pass
        win.grab_set()
        pb.start(12)
        app._year_switching = True              # tam dung kiem tra may chu (_poll_server) trong luc ghi

        if not datastore.is_server():
            # du lieu tren may nay (sqlite chi dung o luong giao dien) -> chay tai cho, van cap nhat chu
            def say(t):
                try:
                    lbl.config(text=t)
                    win.update()
                except Exception:
                    pass
            say(f'Lưu dữ liệu năm {old_y} và nạp năm {new_y}…')
            try:
                res = year_data.switch_year(old_y, new_y, app.sheets, app.data, save_sheets_file, save_data)
            except Exception as e:
                res = e
            finish(res)
            return

        # --- may chu: chay o LUONG NEN, giao dien van muot; bao tien trinh qua hang doi
        import queue as _queue
        q = _queue.Queue()

        def work():
            try:
                r = year_data.switch_year(old_y, new_y, app.sheets, app.data, save_sheets_file, save_data,
                                          progress=lambda t: q.put(('msg', t)))
                q.put(('done', r))
            except Exception as e:
                q.put(('done', e))
        threading.Thread(target=work, daemon=True).start()

        def poll():
            try:
                while True:
                    kind, val = q.get_nowait()
                    if kind == 'msg':
                        lbl.config(text=val)
                    else:
                        finish(val)
                        return
            except _queue.Empty:
                pass
            except Exception:
                pass
            self.after(120, poll)
        poll()

    def _refresh_list(self):
        self.lb.delete(0, 'end')
        for u in self.users:
            ok, why, _ = lic.check_account(u)
            flag = '' if ok else ('  ⛔' if why == 'expired' else '  ⚠')
            lk, lk_until = lic.lock_info(u)
            if lk:
                flag += '  ⛔🔒' if lk_until is None else '  🔒'
            self.lb.insert('end', f"{u.get('username')}  ({u.get('hoten') or ''}){flag}")

    def _toggle_all(self):
        state = 'disabled' if self.all_var.get() else 'normal'
        for key in self.tabvars:
            pass  # checkbox van bam duoc; chi bo qua khi luu neu all

    def _select(self):
        sel = self.lb.curselection()
        if not sel:
            return
        self.idx = sel[0]
        u = self.users[self.idx]
        self.v_user.set(u.get('username') or '')
        self.v_pass.set(u.get('password') or '')
        self.v_hoten.set(u.get('hoten') or '')
        self.v_expire.set(u.get('expire') or '')
        self.v_login_mode.set(u.get('login_mode') or 'multi')
        self.v_logout_mode.set(u.get('logout_mode') or 'one')
        self._show_status(u)
        tabs = u.get('tabs', [])
        is_all = (tabs == 'all')
        self.all_var.set(is_all)
        allowed = set() if is_all else set(tabs or [])
        for key, var in self.tabvars.items():
            var.set(key in allowed)

    def _show_status(self, u):
        locked, until = lic.lock_info(u)
        if locked:
            self.lbl_status.config(text=('⛔ Bị khóa VĨNH VIỄN (sai mật khẩu)' if until is None else
                                         f'🔒 Đang bị khóa (sai mật khẩu) đến {lic.lock_text(until)}')
                                   + ' → bấm "Mở khóa"', fg='#c0392b')
            return
        ok, why, exp_d = lic.check_account(u)
        if u.get('tabs') == 'all':
            txt, col = 'Tài khoản toàn quyền: không áp dụng hạn sử dụng.', '#1d3557'
        elif why == 'clock':
            txt, col = (f'⚠ Bị chặn "Sai ngày giờ" (lần đăng nhập ghi nhận: {lic.show_date(u.get("last_login"))})'
                        ' → bấm "Mở khóa" để đặt lại', '#c0392b')
        elif why == 'tampered':
            txt, col = '⚠ Chưa cấp hạn / hạn bị sửa tay → bấm Áp dụng + Lưu để cấp lại.', '#c0392b'
        elif why == 'expired':
            txt, col = f'⛔ Đã hết hạn từ {exp_d.strftime("%d/%m/%Y")}', '#c0392b'
        elif exp_d:
            txt, col = f'✔ Còn hạn đến {exp_d.strftime("%d/%m/%Y")} ({lic.days_left(u)} ngày)', '#2d6a4f'
        else:
            txt, col = '✔ Không giới hạn thời gian', '#2d6a4f'
        self.lbl_status.config(text=txt, fg=col)

    def _unlock(self):
        if self.idx is None or self.idx >= len(self.users):
            messagebox.showwarning('Chưa chọn', 'Chọn tài khoản cần mở khóa.', parent=self)
            return
        u = self.users[self.idx]
        lic.unlock(u)
        # dat lai ngay dang nhap ghi nhan = hom nay (go loi "Sai ngay gio" khi may khac tung de vuot ngay)
        from datetime import date as _date
        u['last_login'] = lic.fmt_date(_date.today())
        update_user_in_file(u.get('username'), last_login=u['last_login'],
                            **{k: u.get(k) for k in lic.LOCK_FIELDS})
        self._refresh_list()
        self.lb.selection_set(self.idx)
        self._show_status(u)
        messagebox.showinfo('Đã mở khóa', f'Đã mở khóa tài khoản {u.get("username")}.\n'
                                          'Đã đặt lại ngày đăng nhập (gỡ lỗi "Sai ngày giờ").', parent=self)

    def _make_tab_code(self):
        un = self.v_user.get().strip()
        if not un:
            messagebox.showwarning('Chưa chọn', 'Chọn tài khoản cần tạo mã mở sheet.', parent=self)
            return
        if self.all_var.get():
            messagebox.showinfo('Toàn quyền', 'Tài khoản này đang Toàn quyền - không cần mã mở sheet.', parent=self)
            return
        import sheet_code
        pre = [k for k, v in self.tabvars.items() if v.get()]
        sheet_code.open_generator(self, un, self.perms, preselect=pre, on_log=append_license_log)

    def _make_code(self):
        un = self.v_user.get().strip()
        if not un:
            messagebox.showwarning('Chưa chọn', 'Chọn tài khoản cần tạo mã gia hạn.', parent=self)
            return
        CodeGenerator(self, un)

    def _collect_tabs(self):
        if self.all_var.get():
            return 'all'
        return [key for key, var in self.tabvars.items() if var.get()]

    def _apply(self):
        un = self.v_user.get().strip()
        if not un:
            messagebox.showwarning('Thiếu', 'Nhập tên đăng nhập.', parent=self)
            return
        try:
            exp_d = lic.parse_date(self.v_expire.get())
        except ValueError as e:
            messagebox.showwarning('Sai ngày', str(e), parent=self)
            return
        old = {}
        if self.idx is not None and self.idx < len(self.users):
            old = dict(self.users[self.idx])
        rec = dict(old)                     # giu lai pw_changed, used_codes, last_login...
        rec.update({'username': un, 'password': self.v_pass.get(),
                    'hoten': self.v_hoten.get().strip(), 'tabs': self._collect_tabs(),
                    'login_mode': self.v_login_mode.get(), 'logout_mode': self.v_logout_mode.get()})
        lic.set_expire(rec, lic.fmt_date(exp_d))
        self.v_expire.set(rec['expire'])
        if self.idx is None or self.idx >= len(self.users):
            # trung ten?
            if any(u.get('username') == un for u in self.users):
                messagebox.showwarning('Trùng', 'Tên đăng nhập đã tồn tại.', parent=self)
                return
            rec['pw_changed'] = False       # nick moi -> se duoc hoi doi MK cho den khi doi that
            self.users.append(rec)
            self.idx = len(self.users) - 1
        else:
            self.users[self.idx] = rec
        self._refresh_list()
        self.lb.selection_clear(0, 'end')
        self.lb.selection_set(self.idx)
        self._show_status(rec)
        messagebox.showinfo('OK', 'Đã áp dụng (nhớ bấm "Lưu ra file").', parent=self)

    def _save_offline_pol(self):
        try:
            pol = load_app_policy()
            pol['offline_main'] = bool(self.v_off_main.get())
            pol['offline_qc'] = bool(self.v_off_qc.get())
            save_app_policy(pol)
            on = [t for t, v in (('sheet + lệnh SX', pol['offline_main']), ('dữ liệu QC', pol['offline_qc'])) if v]
            self.lbl_offline.config(text='✔ Đã lưu – ' + ('chép về máy: ' + ', '.join(on) if on else
                                                          'KHÔNG chép dữ liệu về máy') +
                                         '. Máy này áp dụng ngay, máy khác từ lần mở phần mềm tiếp theo.',
                                    fg='#2d6a4f')
        except Exception as e:
            self.lbl_offline.config(text=f'✖ Không lưu được: {e}', fg='#c0392b')

    def _save_setup_pol(self):
        try:
            pol = load_app_policy()
            pol['show_server_setup'] = bool(self.v_show_setup.get())
            save_app_policy(pol)
            self.lbl_setup.config(text=('✔ Đã lưu – máy con ' + ('ĐƯỢC' if pol['show_server_setup'] else 'KHÔNG')
                                        + ' thấy nút Cài đặt máy chủ (áp dụng từ lần mở phần mềm tiếp theo).'),
                                  fg='#2d6a4f')
        except Exception as e:
            self.lbl_setup.config(text=f'✖ Không lưu được: {e}', fg='#c0392b')

    def _save_instance(self):
        try:
            pol = load_app_policy()
            pol['instance'] = self.v_instance.get()
            save_app_policy(pol)
            self.lbl_instance.config(text='✔ Đã lưu – áp dụng từ lần mở phần mềm tiếp theo trên mọi máy',
                                     fg='#2d6a4f')
        except Exception as e:
            self.lbl_instance.config(text=f'✖ Không lưu được: {e}', fg='#c0392b')

    def _logout_everywhere(self):
        un = self.v_user.get().strip()
        if not un:
            messagebox.showwarning('Chưa chọn', 'Chọn tài khoản trước.', parent=self)
            return
        if not datastore.is_server():
            messagebox.showinfo('Đăng xuất', 'Chỉ dùng được khi phần mềm chạy chế độ MÁY CHỦ (dữ liệu chung).',
                                parent=self)
            return
        me = str((getattr(self.master, 'user', None) or {}).get('username') or '')
        if not messagebox.askyesno('Đăng xuất trên mọi máy',
                                   f'Đăng xuất tài khoản "{un}" trên TẤT CẢ các máy đang đăng nhập?\n'
                                   '(Các máy đó tự về màn hình đăng nhập trong ~20 giây.)', parent=self):
            return
        try:
            import presence
            n = presence.logout_user(un, by='Quản trị viên', msg='Quản trị viên đăng xuất tài khoản trên mọi máy.',
                                     exclude_me=(un.lower() == me.lower()))
        except Exception as e:
            messagebox.showerror('Lỗi', f'Không gửi được lệnh:\n{e}', parent=self)
            return
        messagebox.showinfo('Đã gửi', f'Đã gửi lệnh đăng xuất tới {n} máy.' if n else
                            f'Tài khoản "{un}" không đang đăng nhập ở máy nào khác.', parent=self)

    def _add(self):
        self.idx = None
        self.v_user.set('')
        self.v_pass.set('')
        self.v_hoten.set('')
        self.v_expire.set('')
        self.v_login_mode.set('multi')
        self.v_logout_mode.set('one')
        self.lbl_status.config(text='')
        self.all_var.set(False)
        for var in self.tabvars.values():
            var.set(False)
        self.lb.selection_clear(0, 'end')

    def _del(self):
        if self.idx is None or self.idx >= len(self.users):
            messagebox.showwarning('Chưa chọn', 'Chọn tài khoản để xóa.', parent=self)
            return
        u = self.users[self.idx]
        if u.get('tabs') == 'all' and sum(1 for x in self.users if x.get('tabs') == 'all') <= 1:
            messagebox.showwarning('Không thể xóa', 'Phải còn ít nhất 1 tài khoản Toàn quyền.', parent=self)
            return
        if not messagebox.askyesno('Xóa', f"Xóa tài khoản {u.get('username')}?", parent=self):
            return
        del self.users[self.idx]
        self.idx = None
        self._refresh_list()
        self._add()

    def _save_file(self):
        # giu trang thai khoa / ma da dung / lan dang nhap moi nhat trong file (co the vua thay doi)
        fresh = {x.get('username'): x for x in load_users()}
        for u in self.users:
            f = fresh.get(u.get('username'))
            if f:
                for k in lic.LOCK_FIELDS + ('used_codes', 'used_tab_codes', 'last_login'):
                    if k in f:
                        u[k] = f[k]
        data = {'_huong_dan': 'Sua bang chuong trinh (nut Cap quyen) hoac sua tay. '
                              "tabs='all' = toan quyen; hoac liet ke ten tab.",
                'users': self.users}
        try:
            save_users_list(self.users)
            messagebox.showinfo('Đã lưu', 'Đã lưu (file mã hóa users.dat).\nThay đổi có hiệu lực khi '
                                          'tài khoản đăng nhập lại.', parent=self)
        except Exception as e:
            messagebox.showerror('Lỗi lưu', str(e), parent=self)


def append_license_log(entry):
    """Luu lich su cac ma gia han admin da tao (nam trong users.dat ma hoa)."""
    try:
        d = load_store()
        d.setdefault('license_log', []).append(entry)
        secure_store.save(d)
    except Exception:
        pass


class CodeGenerator(tk.Toplevel):
    """Admin: tao ma gia han cho 1 tai khoan."""
    def __init__(self, master, username):
        super().__init__(master)
        self.username = username
        self.title('Tạo mã gia hạn')
        self.geometry('500x300')
        self.configure(bg='#f4f6f9')
        self.resizable(False, False)
        try:
            self.transient(master)
            self.grab_set()
        except Exception:
            pass
        tk.Label(self, text='🎟 Tạo mã gia hạn', bg='#f4f6f9', font=FONT_TITLE).pack(anchor='w', padx=16, pady=(14, 6))
        fr = tk.Frame(self, bg='#f4f6f9')
        fr.pack(fill='x', padx=16)
        tk.Label(fr, text='Cấp cho tài khoản:', bg='#f4f6f9', font=FONT, width=18, anchor='e').grid(row=0, column=0, pady=4)
        tk.Label(fr, text=username, bg='#f4f6f9', font=FONT_B, fg='#1d3557').grid(row=0, column=1, sticky='w')
        tk.Label(fr, text='Gia hạn đến (ddmmyyyy):', bg='#f4f6f9', font=FONT, width=18, anchor='e').grid(row=1, column=0, pady=4)
        self.v_date = tk.StringVar()
        e = tk.Entry(fr, textvariable=self.v_date, width=16, font=FONT)
        e.grid(row=1, column=1, sticky='w')
        e.bind('<Return>', lambda ev: self._gen())
        e.focus_set()
        tk.Button(fr, text='Tạo mã', font=FONT_B, bg='#7b2cbf', fg='white', relief='flat',
                  padx=12, command=self._gen).grid(row=1, column=2, padx=8)

        tk.Label(self, text='Mã gia hạn (gửi cho thành viên, mỗi mã dùng 1 lần):',
                 bg='#f4f6f9', font=FONT).pack(anchor='w', padx=16, pady=(14, 2))
        self.v_code = tk.StringVar()
        tk.Entry(self, textvariable=self.v_code, font=('Consolas', 14, 'bold'), width=36,
                 state='readonly', readonlybackground='white', fg='#7b2cbf',
                 justify='center').pack(padx=16, ipady=4)
        self.msg = tk.Label(self, text='', bg='#f4f6f9', fg='#2d6a4f', font=FONT)
        self.msg.pack(pady=4)
        bt = tk.Frame(self, bg='#f4f6f9')
        bt.pack(fill='x', padx=16, pady=6)
        tk.Button(bt, text='Đóng', font=FONT, relief='flat', padx=12, pady=5,
                  command=self.destroy).pack(side='right')
        tk.Button(bt, text='📋 Copy mã', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat',
                  padx=12, pady=5, command=self._copy).pack(side='right', padx=8)

    def _gen(self):
        try:
            d = lic.parse_date(self.v_date.get())
            if not d:
                raise ValueError('Nhập ngày gia hạn, vd 21122027')
            from datetime import date as _date
            if d <= _date.today():
                raise ValueError('Ngày gia hạn phải sau hôm nay.')
            code = lic.make_code(self.username, d)
        except ValueError as e:
            messagebox.showwarning('Lỗi', str(e), parent=self)
            return
        self.v_code.set(code)
        self.msg.config(text=f'Mã cho "{self.username}" → gia hạn đến {d.strftime("%d/%m/%Y")}')
        from datetime import datetime as _dt
        append_license_log({'tao_luc': _dt.now().strftime('%d/%m/%Y %H:%M'),
                            'username': self.username, 'gia_han_den': lic.fmt_date(d),
                            'ma': code, 'code_id': lic.code_id(code)})
        self._copy(silent=True)

    def _copy(self, silent=False):
        c = self.v_code.get()
        if not c:
            return
        self.clipboard_clear()
        self.clipboard_append(c)
        if not silent:
            self.msg.config(text='Đã copy mã vào clipboard.')


class RenewDialog(tk.Toplevel):
    """Thanh vien nhap ma gia han. self.ok = True neu gia han thanh cong."""
    def __init__(self, master, user, reason_text):
        super().__init__(master)
        self.user = user
        self.ok = False
        self.title('Gia hạn tài khoản')
        self.geometry('480x270')
        self.configure(bg='#f4f6f9')
        self.resizable(False, False)
        try:
            self.transient(master)
            self.grab_set()
        except Exception:
            pass
        tk.Label(self, text='⏳ Gia hạn tài khoản', bg='#f4f6f9', font=FONT_TITLE).pack(anchor='w', padx=16, pady=(14, 4))
        tk.Label(self, text=reason_text, bg='#f4f6f9', fg='#c0392b', font=FONT,
                 wraplength=440, justify='left').pack(anchor='w', padx=16)
        tk.Label(self, text='Nhập mã gia hạn do quản trị viên cấp:', bg='#f4f6f9',
                 font=FONT_B).pack(anchor='w', padx=16, pady=(12, 2))
        self.v_code = tk.StringVar()
        e = tk.Entry(self, textvariable=self.v_code, font=('Consolas', 14, 'bold'), width=36,
                     justify='center')
        e.pack(padx=16, ipady=4)
        e.bind('<Return>', lambda ev: self._do())
        e.focus_set()
        self.msg = tk.Label(self, text='', bg='#f4f6f9', fg='#c0392b', font=FONT, wraplength=440)
        self.msg.pack(pady=4)
        bt = tk.Frame(self, bg='#f4f6f9')
        bt.pack(fill='x', padx=16, pady=6)
        tk.Button(bt, text='Đóng', font=FONT, relief='flat', padx=12, pady=5,
                  command=self.destroy).pack(side='right')
        tk.Button(bt, text='✔ Gia hạn', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat',
                  padx=14, pady=5, command=self._do).pack(side='right', padx=8)

    def _do(self):
        # doc lai ban moi nhat trong file (de biet ma nao da dung)
        fresh = next((u for u in load_users() if u.get('username') == self.user.get('username')), None)
        rec = dict(fresh or self.user)
        rec['used_codes'] = list(rec.get('used_codes') or [])
        try:
            new_d = lic.apply_code(rec, self.v_code.get())
        except ValueError as e:
            self.msg.config(text='✗ ' + str(e))
            return
        changes = {'expire': rec['expire'], 'expire_sig': rec['expire_sig'],
                   'used_codes': rec['used_codes']}
        if not update_user_in_file(rec.get('username'), **changes):
            self.msg.config(text='✗ Không ghi được dữ liệu tài khoản (kiểm tra quyền ghi thư mục).')
            return
        self.user.update(changes)
        self.ok = True
        messagebox.showinfo('Thành công', f'Đã gia hạn tài khoản đến {new_d.strftime("%d/%m/%Y")}.',
                            parent=self)
        self.destroy()


class RowEditor(tk.Toplevel):
    """Sua/them 1 dong bang chung: moi cot 1 o nhap.
    Khi THEM (values=None) va co source_rows -> cho phep nap tu 1 dong cu roi sua."""
    def __init__(self, master, title, headers, values, on_save, source_rows=None,
                 readonly_cols=None, note=None, defaults=None, keep_on_load=None,
                 code_lookup=None, code_pairs=None):
        """defaults: {cot: gia tri} dien san khi THEM moi (VD STT = so lon nhat + 1).
        keep_on_load: cac cot KHONG bi ghi de khi 'Nap tu dong cu' (VD STT).
        code_lookup(ten) -> [ma,...]; code_pairs: [(cot ten, cot ma), ...] -> go ten la tu hien ma."""
        super().__init__(master)
        self.on_save = on_save
        self.headers = headers
        self.readonly_cols = dict(readonly_cols or {})   # {chi so cot: chu thich}
        self.source_rows = source_rows or []
        self.keep_on_load = set(keep_on_load or ())
        self.code_lookup = code_lookup
        self.code_pairs = list(code_pairs or []) if code_lookup else []
        self._loading = False
        self.title(title)
        self.configure(bg='#f4f6f9')
        self.resizable(False, True)
        tk.Label(self, text=title, bg='#f4f6f9', font=FONT_B).pack(anchor='w', padx=14, pady=(12, 6))

        # thanh nap tu dong cu (chi khi THEM moi)
        if values is None and self.source_rows:
            nb = tk.Frame(self, bg='#eaf0f7')
            nb.pack(fill='x', padx=14, pady=(0, 6))
            tk.Label(nb, text='Nạp từ dòng cũ:', bg='#eaf0f7', font=FONT_B).pack(side='left', padx=(6, 4), pady=6)
            self._src = tk.StringVar()
            self._src_cb = ttk.Combobox(nb, textvariable=self._src, width=52, state='readonly',
                                        font=FONT, values=[self._row_label(r) for r in self.source_rows])
            self._src_cb.pack(side='left', pady=6)
            tk.Button(nb, text='Nạp ↧', font=FONT, command=self._load_row).pack(side='left', padx=6)
            tk.Label(nb, text='(hoặc để trống rồi đánh mới)', bg='#eaf0f7', fg='#666',
                     font=('Segoe UI', 9)).pack(side='left', padx=6)

        if note:
            tk.Label(self, text=note, bg='#fff8e1', fg='#6d4c00', font=('Segoe UI', 9),
                     justify='left', anchor='w', wraplength=520).pack(fill='x', padx=14, pady=(0, 6))
        box = tk.Frame(self, bg='#f4f6f9')
        box.pack(fill='both', expand=True, padx=14)
        self.vars = []
        self._code_cb = {}            # cot ma -> Combobox (chon 1 trong cac ma cua ten)
        self._code_hint = {}          # cot ma -> Label goi y
        code_cols = {c for _n, c in self.code_pairs}
        dflt = dict(defaults or {}) if values is None else {}
        for i, h in enumerate(headers):
            tk.Label(box, text=h + ':', bg='#f4f6f9', font=FONT, width=20, anchor='e').grid(
                row=i, column=0, sticky='e', padx=4, pady=2)
            if values and i < len(values) and values[i] is not None:
                v0 = str(values[i])
            elif i in dflt:
                v0 = str(dflt[i])
            else:
                v0 = ''
            var = tk.StringVar(value=v0)
            if i in code_cols and i not in self.readonly_cols:
                ent = ttk.Combobox(box, textvariable=var, width=40, font=FONT, values=[])
                self._code_cb[i] = ent
                hint = tk.Label(box, text='', bg='#f4f6f9', fg='#6a4c93', font=('Segoe UI', 9))
                hint.grid(row=i, column=2, sticky='w')
                self._code_hint[i] = hint
            else:
                ent = tk.Entry(box, textvariable=var, width=42, font=FONT)
            ent.grid(row=i, column=1, sticky='w', padx=4, pady=2)
            if i in self.readonly_cols:
                ent.config(state='readonly')
                tk.Label(box, text=self.readonly_cols[i], bg='#f4f6f9', fg='#666',
                         font=('Segoe UI', 9)).grid(row=i, column=2, sticky='w')
            self.vars.append(var)
        # go Ten (Ma hoa) -> tu dien CODE theo sheet Don gia vat tu
        for n_i, c_i in self.code_pairs:
            if n_i < len(self.vars) and c_i < len(self.vars):
                self.vars[n_i].trace_add('write', lambda *_a, n=n_i, c=c_i: self._fill_code(n, c))
                self._fill_code(n_i, c_i, overwrite=False)
        bot = tk.Frame(self, bg='#f4f6f9')
        bot.pack(fill='x', padx=14, pady=12)
        tk.Button(bot, text='💾 Lưu', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat',
                  padx=16, pady=6, command=self._ok).pack(side='right')
        tk.Button(bot, text='Đóng', font=FONT, relief='flat', padx=14, pady=6,
                  command=self.destroy).pack(side='right', padx=8)

    def _row_label(self, r):
        parts = [str(x) for x in r[:5] if str(x).strip()]
        return '  |  '.join(parts[:4]) if parts else '(dòng trống)'

    def _load_row(self):
        i = self._src_cb.current()
        if i < 0 or i >= len(self.source_rows):
            return
        r = self.source_rows[i]
        self._loading = True             # nap dong cu: giu nguyen CODE cua dong cu, khong tra lai
        try:
            for j, var in enumerate(self.vars):
                if j in self.readonly_cols or j in self.keep_on_load:
                    continue             # cot tu dong (VD: STT) khong nap tu dong cu
                var.set(str(r[j]) if j < len(r) and r[j] is not None else '')
        finally:
            self._loading = False
        for n_i, c_i in self.code_pairs:
            self._fill_code(n_i, c_i, overwrite=False)

    def _fill_code(self, n_i, c_i, overwrite=True):
        """Ten NVL (Ma hoa) -> CODE theo sheet Don gia vat tu. Nhieu ma -> chon trong o so xuong;
        cung ten voi ben kia (VD PET RF -> PET RF) thi uu tien ma KHAC ma ben kia."""
        if not self.code_lookup:
            return
        ten = self.vars[n_i].get().strip()
        codes = []
        if ten:
            try:
                codes = [str(c) for c in (self.code_lookup(ten) or [])]
            except Exception:
                codes = []
        cb = self._code_cb.get(c_i)
        if cb is not None:
            cb.configure(values=codes)
        hint = self._code_hint.get(c_i)
        if hint is not None:
            if not ten:
                hint.config(text='', fg='#6a4c93')
            elif not codes:
                hint.config(text='(tên không có trong Đơn giá vật tư)', fg='#c0392b')
            elif len(codes) > 1:
                hint.config(text=f'({len(codes)} mã — bấm ▼ để chọn)', fg='#6a4c93')
            else:
                hint.config(text='(lấy theo Đơn giá vật tư)', fg='#2d6a4f')
        if self._loading or not codes:
            return
        cur = self.vars[c_i].get().strip()
        # code ben doi dien (cap con lai) khi 2 ben cung ten - tranh chon trung ma
        other = None
        for n2, c2 in self.code_pairs:
            if c2 != c_i and self.vars[n2].get().strip().lower() == ten.lower():
                other = self.vars[c2].get().strip()
        if cur in codes and (len(codes) == 1 or not other or cur != other):
            return                       # ma dang co da dung voi ten -> giu nguyen
        if cur and not overwrite:
            return                       # luc mo cua so / nap dong cu: khong doi ma cu
        pick = codes[0]
        if other and len(codes) > 1:
            for c in codes:
                if c != other:
                    pick = c
                    break
        self.vars[c_i].set(pick)

    def _ok(self):
        # on_save tra ve False -> du lieu chua hop le, giu cua so de sua tiep
        if self.on_save([v.get() for v in self.vars]) is False:
            return
        self.destroy()


def _pct_str(p):
    """Ti le (0..1) -> chuoi % du chinh xac (12 so le, bo so 0 thua). VD 0.719683 -> '71.9683'."""
    try:
        v = round(float(p or 0) * 100, 12)
    except (TypeError, ValueError):
        return '0'
    t = ('%.12f' % v).rstrip('0').rstrip('.')
    return t if t not in ('', '-0') else '0'


class OrderEditor(tk.Toplevel):
    """Them lenh san xuat moi: copy tu 1 don co san roi sua."""
    NMAT = 12

    def __init__(self, master, app, kind, dl_tab, mode='add'):
        super().__init__(master)
        self.app = app
        self.kind = kind
        self.dl_tab = dl_tab
        self.mode = mode                 # 'add' = them moi, 'edit' = sua don cu
        self.edit_target = None          # don dang sua (mode edit)
        self._copied_loai = ''
        self.title(('Chỉnh sửa' if mode == 'edit' else 'Thêm') + f' lệnh sản xuất - {kind}')
        self.geometry('1000x620')
        self.configure(bg='#f4f6f9')
        self.orders = app.data['orders'].get(kind, [])
        self.nvl_list = self._build_nvl_list()   # danh sach NVL tu Don gia
        self.catalog = NvlCatalog.from_sheets(app.sheets)   # tra Ten -> Ma + loai PP/PE/bot da
        self._ac_pop = None                        # popup goi y
        self._ac_ten = None                        # o Ten dang go
        self._ac_ma = None                         # o Ma tuong ung
        self.minsize(760, 420)

        # --- Nut Luu/Dong GHIM day + vung noi dung CUON DUOC (co thanh keo) ---
        self._botbar = tk.Frame(self, bg='#e9edf2')
        self._botbar.pack(side='bottom', fill='x')
        _outer = tk.Frame(self, bg='#f4f6f9')
        _outer.pack(side='top', fill='both', expand=True)
        self._sc = tk.Canvas(_outer, bg='#f4f6f9', highlightthickness=0)
        _vs = ttk.Scrollbar(_outer, orient='vertical', command=self._sc.yview)
        self._sc.configure(yscrollcommand=_vs.set)
        _vs.pack(side='right', fill='y')
        self._sc.pack(side='left', fill='both', expand=True)
        self.body = tk.Frame(self._sc, bg='#f4f6f9')
        self._sc_win = self._sc.create_window((0, 0), window=self.body, anchor='nw')
        self.body.bind('<Configure>', lambda e: self._sc.configure(scrollregion=self._sc.bbox('all')))
        self._sc.bind('<Configure>', lambda e: self._sc.itemconfigure(self._sc_win, width=e.width))
        self._sc.bind('<Enter>', lambda e: self._sc.bind_all(
            '<MouseWheel>', lambda ev: self._sc.yview_scroll(int(-ev.delta / 120), 'units')))
        self._sc.bind('<Leave>', lambda e: self._sc.unbind_all('<MouseWheel>'))

        top = tk.Frame(self.body, bg='#f4f6f9')
        top.pack(fill='x', padx=14, pady=10)
        lbl = 'Chọn lệnh cần sửa:' if mode == 'edit' else 'Copy từ đơn có sẵn:'
        tk.Label(top, text=lbl, bg='#f4f6f9', font=FONT_B).pack(side='left')
        self.src = tk.StringVar()
        self.src_cb = ttk.Combobox(top, textvariable=self.src, width=48, state='readonly',
                                   font=FONT, values=[self._olabel(o) for o in self.orders])
        self.src_cb.pack(side='left', padx=6)
        tk.Button(top, text=('Nạp để sửa ↧' if mode == 'edit' else 'Nạp ↧'),
                  font=FONT, command=self._load_src).pack(side='left')

        # --- thong tin don (da bo Loai) ---
        info = tk.Frame(self.body, bg='#f4f6f9')
        info.pack(fill='x', padx=14)
        self.v = {}
        defs = [('ma_sp', 'Mã SP'), ('don_hang', 'Đơn hàng số'),
                ('ma_hang', 'Mã hàng'), ('khach_hang', 'Khách hàng'),
                ('ma_san_pham', 'Mã sản phẩm'), ('hao_hut', '% Hao hụt')]
        hints = self._field_hints()
        for i, (k, lb) in enumerate(defs):
            r0, c0 = (i // 2) * 2, (i % 2) * 2
            tk.Label(info, text=lb + ':', bg='#f4f6f9', font=FONT).grid(row=r0, column=c0,
                                                                        sticky='e', padx=4, pady=(3, 0))
            var = tk.StringVar()
            st = 'readonly' if (k == 'ma_sp' and mode == 'edit') else 'normal'
            tk.Entry(info, textvariable=var, width=30, font=FONT, state=st).grid(
                row=r0, column=c0 + 1, sticky='w', padx=4, pady=(3, 0))
            # dong goi y VI DU ngay duoi o nhap (tranh nhap sai dinh dang)
            if hints.get(k):
                tk.Label(info, text=hints[k], bg='#f4f6f9', fg='#8a8a8a',
                         font=('Segoe UI', 8, 'italic')).grid(row=r0 + 1, column=c0 + 1,
                                                              sticky='w', padx=6, pady=(0, 2))
            self.v[k] = var
        self.v['hao_hut'].set('0.7')          # mac dinh 0.7%
        if mode == 'add':
            self.v['ma_sp'].set(self._next_masp())

        # --- Thong so lenh SX (chi Export/Domestic) - luu vao Container Volume ---
        self.cv_vars = {}
        if kind in ('Export', 'Domestic'):
            cvf = tk.LabelFrame(self.body, text='Thông số lệnh SX (lưu vào Container Volume)',
                                bg='#f4f6f9', font=FONT_B, fg='#1d3557')
            cvf.pack(fill='x', padx=14, pady=(8, 0))
            self.cv_vars['san_luong'] = tk.StringVar()
            self.cv_vars['so_cont'] = tk.StringVar()
            self.cv_vars['kg_pallet'] = tk.StringVar(value='1000')
            self.cv_vars['loai_bao'] = tk.StringVar(value='Bao 25Kg')
            self.cv_vars['quy_cach'] = tk.StringVar(value='Hàng Xếp Rời')
            self.cv_vars['so_pallet'] = tk.StringVar()
            self.cv_vars['pallet_cont'] = tk.StringVar(value='20')
            self.cv_vars['cach_xep'] = tk.StringVar()
            self.cv_vars['mix'] = tk.BooleanVar(value=False)
            self.cv_vars['ma_phu'] = [tk.StringVar(), tk.StringVar(), tk.StringVar()]
            self.cv_vars['kg_phu'] = [tk.StringVar(), tk.StringVar(), tk.StringVar()]
            e_sl = tk.Entry(cvf, textvariable=self.cv_vars['san_luong'], width=14, font=FONT)
            e_kp = tk.Entry(cvf, textvariable=self.cv_vars['kg_pallet'], width=10, font=FONT)
            e_pc = tk.Entry(cvf, textvariable=self.cv_vars['pallet_cont'], width=10, font=FONT)
            tk.Label(cvf, text='Sản lượng đơn (Kg):', bg='#f4f6f9', font=FONT).grid(row=0, column=0, sticky='e', padx=4, pady=3)
            e_sl.grid(row=0, column=1, sticky='w', padx=4)
            tk.Label(cvf, text='Kg/Pallet:', bg='#f4f6f9', font=FONT).grid(row=0, column=2, sticky='e', padx=4)
            e_kp.grid(row=0, column=3, sticky='w', padx=4)
            tk.Label(cvf, text='Tổng số Cont:', bg='#f4f6f9', font=FONT).grid(row=0, column=4, sticky='e', padx=4)
            e_scont = tk.Entry(cvf, textvariable=self.cv_vars['so_cont'], width=10, font=FONT)
            e_scont.grid(row=0, column=5, sticky='w', padx=4)
            e_scont.bind('<KeyRelease>', lambda ev: (self._cv_calc_pallet(), self._cv_calc_xep()))
            tk.Label(cvf, text='Loại bao:', bg='#f4f6f9', font=FONT).grid(row=1, column=0, sticky='e', padx=4, pady=3)
            ttk.Combobox(cvf, textvariable=self.cv_vars['loai_bao'], width=12, state='readonly',
                         font=FONT, values=['Bao 25Kg', 'Bao Zumbo']).grid(row=1, column=1, sticky='w', padx=4)
            tk.Label(cvf, text='Quy cách xếp bao:', bg='#f4f6f9', font=FONT).grid(row=1, column=2, columnspan=2, sticky='e', padx=4)
            qc_cb = ttk.Combobox(cvf, textvariable=self.cv_vars['quy_cach'], width=16, state='readonly',
                                 font=FONT, values=['Hàng Xếp Rời', 'Hàng Xếp Pallet'])
            qc_cb.grid(row=1, column=4, columnspan=2, sticky='w', padx=4)
            qc_cb.bind('<<ComboboxSelected>>', lambda ev: self._cv_calc_xep())
            tk.Label(cvf, text='Tổng số Pallet:', bg='#f4f6f9', font=FONT).grid(row=2, column=0, sticky='e', padx=4, pady=3)
            tk.Entry(cvf, textvariable=self.cv_vars['so_pallet'], width=10, font=FONT).grid(row=2, column=1, sticky='w', padx=4)
            tk.Label(cvf, text='Pallet/Cont:', bg='#f4f6f9', font=FONT).grid(row=2, column=2, sticky='e', padx=4)
            e_pc.grid(row=2, column=3, sticky='w', padx=4)
            tk.Label(cvf, text='(Cont = SL÷28.000;  Pallet = SL÷Kg/Pallet, tối đa Pallet/Cont × Cont — có thể sửa)',
                     bg='#f4f6f9', fg='#666', font=('Segoe UI', 9)).grid(row=3, column=0, columnspan=6, sticky='w', pady=(2, 4))
            # --- Cach xep pallet (tu tinh khi Quy cach = Hang Xep Pallet, van sua duoc) ---
            tk.Label(cvf, text='Cách xếp pallet:', bg='#f4f6f9', font=FONT_B, fg='#1d3557').grid(
                row=4, column=0, sticky='e', padx=4, pady=(2, 4))
            tk.Entry(cvf, textvariable=self.cv_vars['cach_xep'], width=52, font=FONT).grid(
                row=4, column=1, columnspan=4, sticky='w', padx=4, pady=(2, 4))
            tk.Button(cvf, text='↻ Tự tính', font=('Segoe UI', 9),
                      command=self._cv_calc_xep).grid(row=4, column=5, sticky='w', padx=4, pady=(2, 4))
            # --- Xep lan loai hang (toi da 3 hang phu) ---
            tk.Checkbutton(cvf, text='Xếp lẫn loại hàng', variable=self.cv_vars['mix'], bg='#f4f6f9',
                           font=FONT_B, fg='#8b0000', command=self._cv_calc_xep).grid(
                row=5, column=0, sticky='e', padx=4, pady=(0, 2))
            mix_ws = []          # cac o Hang phu 1-3: chi HIEN khi tich 'Xep lan loai hang'
            for k in range(3):
                lb1 = tk.Label(cvf, text=f'Hàng phụ {k+1}  —  Mã:', bg='#f4f6f9', font=FONT)
                lb1.grid(row=6 + k, column=0, sticky='e', padx=4, pady=1)
                emp = tk.Entry(cvf, textvariable=self.cv_vars['ma_phu'][k], width=14, font=FONT)
                emp.grid(row=6 + k, column=1, sticky='w', padx=4)
                lb2 = tk.Label(cvf, text='KL (kg):', bg='#f4f6f9', font=FONT)
                lb2.grid(row=6 + k, column=2, sticky='e', padx=4)
                ekpu = tk.Entry(cvf, textvariable=self.cv_vars['kg_phu'][k], width=12, font=FONT)
                ekpu.grid(row=6 + k, column=3, sticky='w', padx=4)
                for e in (emp, ekpu):
                    e.bind('<KeyRelease>', lambda ev: self._cv_calc_xep())
                mix_ws += [lb1, emp, lb2, ekpu]

            def _toggle_mix(*_a):
                try:
                    on = bool(self.cv_vars['mix'].get())
                except Exception:
                    on = False
                for w in mix_ws:
                    if on:
                        w.grid()
                    else:
                        w.grid_remove()
            self.cv_vars['mix'].trace_add('write', _toggle_mix)
            _toggle_mix()
            # tu tinh khi doi san luong / kg-pallet / pallet-cont
            for e in (e_sl, e_kp, e_pc):
                e.bind('<KeyRelease>', lambda ev: self._cv_recalc())

        # --- NHAN DIEN CONG THUC TU ANH / BANG EXCEL ---
        self._build_recipe_zone()

        # --- bang NVL ---
        tk.Label(self.body, text='Nguyên vật liệu (để trống dòng không dùng):', bg='#f4f6f9',
                 font=FONT_B).pack(anchor='w', padx=14, pady=(10, 2))
        grid = tk.Frame(self.body, bg='#f4f6f9')
        grid.pack(fill='x', padx=14)
        # cot: 0 Ten,1 Ma,2 KL,3 %, 4 dem, 5 Kg ao, 6 % ao  (px width)
        COLW = [210, 100, 100, 90, 26, 100, 90]
        for ci, w in enumerate(COLW):
            grid.grid_columnconfigure(ci, minsize=w, weight=0)
        heads = [(0, 'Tên NVL', '#dbe5f1'), (1, 'Mã NVL', '#dbe5f1'),
                 (2, 'KL/mẻ (kg)', '#dbe5f1'), (3, 'Tỉ lệ %', '#dbe5f1'),
                 (5, 'Kg/mẻ ảo', '#fde9c8'), (6, '% ảo', '#fde9c8')]
        for ci, t, bgc in heads:
            tk.Label(grid, text=t, bg=bgc, font=FONT_B, relief='groove',
                     bd=1).grid(row=0, column=ci, sticky='nsew', ipady=2)
        self.rows = []
        for ri in range(1, self.NMAT + 1):
            e_ten = tk.Entry(grid, font=FONT, width=1)
            e_ten.grid(row=ri, column=0, sticky='nsew')
            e_ma = tk.Entry(grid, font=FONT, width=1)
            e_ma.grid(row=ri, column=1, sticky='nsew')
            # goi y ten NVL tu Don gia
            e_ten.bind('<KeyRelease>', lambda ev, t=e_ten, m=e_ma: self._ac_key(ev, t, m))
            e_ten.bind('<Down>', lambda ev: self._ac_focus_list())
            e_ten.bind('<FocusOut>', lambda ev: self.after(200, self._ac_hide))
            # doi Ten NVL -> tu cap nhat lai Ma NVL (ten khong co trong list -> xoa ma)
            e_ten.bind('<FocusOut>', lambda ev, t=e_ten, m=e_ma: self.after(
                250, lambda: self._sync_ma(t, m)), add='+')
            e_ma.bind('<KeyRelease>', lambda ev: self._show_rule(), add='+')
            e_kl = tk.Entry(grid, font=FONT, width=1)
            e_kl.grid(row=ri, column=2, sticky='nsew')
            e_pct = tk.Entry(grid, font=FONT, width=1)
            e_pct.grid(row=ri, column=3, sticky='nsew')
            e_pct.bind('<FocusOut>', lambda ev: self._auto_recalc())
            e_pct.bind('<Return>', lambda ev: self._auto_recalc())
            e_pct.bind('<KeyRelease>', lambda ev: self._show_tong())
            # cot dem (4) de trong -> tach phan ao
            e_klv = tk.Entry(grid, font=FONT, width=1, bg='#fffdf5')
            e_klv.grid(row=ri, column=5, sticky='nsew')
            e_pctv = tk.Entry(grid, font=FONT, width=1, bg='#fffdf5')
            e_pctv.grid(row=ri, column=6, sticky='nsew')
            e_klv.bind('<FocusOut>', lambda ev: self._recalc_virtual())
            e_klv.bind('<Return>', lambda ev: self._recalc_virtual())
            self.rows.append((e_ten, e_ma, e_kl, e_pct, e_klv, e_pctv))

        opt = tk.Frame(self.body, bg='#f4f6f9')
        opt.pack(fill='x', padx=14, pady=(6, 0))
        self.auto_kg = tk.BooleanVar(value=True)
        tk.Checkbutton(opt, text='Tự tính KL/mẻ theo % (dựa trên dòng nền đầu tiên)',
                       variable=self.auto_kg, bg='#f4f6f9', font=FONT).pack(side='left')
        tk.Button(opt, text='🔄 Tính lại KL/mẻ', font=FONT, command=self._recalc_kg).pack(side='left', padx=10)
        # Lam tron KL/me (VD Acid 2.28997 -> 2.300; LLDPE 12.0159 -> 12.000) bang cach chinh lai %;
        # chi de 1 NVL nhua o cuoi cung so le de tong = 100%
        tk.Button(opt, text='🎯 Làm tròn KL/mẻ', font=FONT_B, bg='#1d6fa5', fg='white', relief='flat',
                  padx=10, command=self._round_dialog).pack(side='left', padx=(0, 6))
        self._round_undo = None
        self.btn_round_undo = tk.Button(opt, text='↩ Bỏ làm tròn', font=FONT, relief='groove',
                                        state='disabled', command=self._round_undo_apply)
        self.btn_round_undo.pack(side='left', padx=(0, 6))
        self.tong_lbl = tk.Label(opt, text='', bg='#f4f6f9', fg='#1d3557', font=FONT_B)
        self.tong_lbl.pack(side='left', padx=10)
        # kiem tra ma hang (bot da % + PP/PE) hien ngay khi go
        self.rule_lbl = tk.Label(self.body, text='', bg='#f4f6f9', fg='#1d3557', font=FONT_B,
                                 justify='left', anchor='w', wraplength=900)
        self.rule_lbl.pack(anchor='w', fill='x', padx=14, pady=(4, 0))
        self.v['ma_hang'].trace_add('write', lambda *a: self._show_rule())

        note = tk.Label(self.body, text='Ghi chú: KL/mẻ = (%/100) × tổng mẻ. Tổng mẻ lấy theo dòng nền '
                                   '(dòng đầu tiên có KL & %). Đổi % dòng nào thì KL dòng đó tự đổi.',
                        bg='#f4f6f9', fg='#666', font=('Segoe UI', 9), justify='left', wraplength=720)
        note.pack(anchor='w', padx=14, pady=4)
        note2 = tk.Label(
            self.body,
            text='Quy tắc mã hàng (không đúng sẽ KHÔNG cho lưu):\n'
                 '  • Tên NVL phải có trong danh sách Unit price — đổi tên thì Mã NVL tự cập nhật; '
                 'tên không có trong list thì Mã bị xoá.\n'
                 '  • 2 số sau P10/P20 = phần nhựa + phụ gia ⇒ % bột đá = 100 − số đó.\n'
                 '      VD: lệnh PO/26091/17.TBC, mã hàng HTP-P2015IM-3 ⇒ bột đá phải 85% '
                 '(85 → dưới 86%). Nhập VMI15TL = 86% ⇒ báo lỗi, không lưu.\n'
                 '  • Tổng % PP > PE ⇒ mã hàng phải có "P20";  PE > PP ⇒ phải có "P10".',
            bg='#fff8e1', fg='#6d4c00', font=('Segoe UI', 9), justify='left', anchor='w')
        note2.pack(anchor='w', fill='x', padx=14, pady=(0, 6))

        bot = tk.Frame(self._botbar, bg='#e9edf2')
        bot.pack(fill='x', padx=14, pady=8)
        save_txt = '💾 Lưu thay đổi' if self.mode == 'edit' else '💾 Lưu (thêm mới)'
        tk.Button(bot, text=save_txt, font=FONT_B, bg='#2d6a4f', fg='white',
                  relief='flat', padx=16, pady=7, command=self._save).pack(side='right')
        tk.Button(bot, text='Đóng', font=FONT, relief='flat', padx=14, pady=7,
                  command=self.destroy).pack(side='right', padx=8)

    # ---------------- Nhan dien cong thuc tu ANH / bang Excel ----------------
    def _build_recipe_zone(self):
        box = tk.LabelFrame(self.body, text='Nhận diện công thức từ ảnh / bảng Excel  '
                                            '(tự điền Tên NVL + Tỉ lệ %)',
                            bg='#f4f6f9', font=FONT_B, fg='#1d3557')
        box.pack(fill='x', padx=14, pady=(10, 0))
        try:
            from recipe_ui import RecipeDropZone
            RecipeDropZone(box, self._recipe_result).pack(fill='x', padx=6, pady=6)
        except Exception as e:
            tk.Label(box, text=f'(Không tải được chức năng nhận diện ảnh: {e})', bg='#f4f6f9',
                     fg='#c0392b', font=FONT).pack(anchor='w', padx=6, pady=6)

    def _recipe_matcher(self):
        if getattr(self, '_rmatcher', None) is None:
            from image_recipe import CatalogMatcher
            hist = []
            for lst in self.app.data.get('orders', {}).values():
                for o in lst:
                    for m in o.get('materials') or []:
                        if (m.get('ten') or '').strip():
                            hist.append((m.get('ten'), m.get('ma_nvl')))
            self._rmatcher = CatalogMatcher(self.catalog.items, hist)
        return self._rmatcher

    def _recipe_result(self, result, image):
        from recipe_ui import RecipeReviewDialog
        RecipeReviewDialog(self, result, self._recipe_matcher(), self.catalog,
                           self._apply_recipe, image=image, max_rows=self.NMAT)

    def _apply_recipe(self, vals):
        """Dien ket qua nhan dien vao bang NVL (xoa cac dong cu)."""
        filled = [r for r in self.rows if r[0].get().strip()]
        if filled and not messagebox.askyesno(
                'Ghi đè bảng NVL', 'Bảng NVL đang có dữ liệu.\nXoá và thay bằng công thức vừa nhận diện?',
                parent=self):
            return
        has_mass = any(v.get('mass') for v in vals)
        for e_row in self.rows:
            for e in e_row:
                e.delete(0, 'end')
        for j, v in enumerate(vals[:self.NMAT]):
            e_ten, e_ma, e_kl, e_pct, e_klv, e_pctv = self.rows[j]
            e_ten.insert(0, v['ten'])
            if v['ma']:
                e_ma.insert(0, v['ma'])
            if v.get('pct') is not None:
                e_pct.insert(0, ('%.9f' % v['pct']).rstrip('0').rstrip('.'))
            if v.get('mass'):
                e_kl.insert(0, f"{v['mass']:g}")
            e_ten.configure(bg=('white' if v['ma'] else '#ffd6d6'))
        # anh khong co KL -> dong dau (bot da) mac dinh 275 kg/me lam goc de tinh cac dong con lai
        if not has_mass and vals:
            for j, v in enumerate(vals[:self.NMAT]):
                if v.get('pct') and self.catalog.kind_of(v['ten'], v['ma']) == 'FILLER':
                    self.rows[j][2].insert(0, '275')
                    break
        if self.auto_kg.get():
            self._recalc_kg()
        for (e_ten, e_ma, e_kl, e_pct, e_klv, e_pctv) in self.rows:
            if e_ten.get().strip() and e_kl.get().strip():
                e_klv.insert(0, e_kl.get().strip())
        self._recalc_virtual()
        self._show_tong()
        b = self._base_batch()
        if b:
            self.tong_lbl.config(text=self.tong_lbl.cget('text') + f'  |  Cỡ mẻ: {b:,.1f} kg')

    def _tong_pct(self):
        """Tong % (Decimal) cua cac dong co Ten NVL; None neu co o % khong phai so."""
        from decimal import Decimal, InvalidOperation
        t = Decimal(0)
        for (e_ten, e_ma, e_kl, e_pct, e_klv, e_pctv) in self.rows:
            if not e_ten.get().strip():
                continue
            try:
                t += Decimal((e_pct.get() or '0').strip().replace(',', '.'))
            except InvalidOperation:
                return None
        return t

    def _show_tong(self):
        """Hien Tong % (12 so le) ngay khi go: xanh = dung 100%, do = thieu/vuot."""
        from decimal import Decimal
        t = self._tong_pct()
        if t is None:
            self.tong_lbl.config(text='Tổng %: (có ô % không hợp lệ)', fg='#c0392b')
            return
        ok = abs(t - Decimal(100)) <= Decimal('1e-11')
        self.tong_lbl.config(text=f'Tổng %: {t:.12f}%' + ('  ✔' if ok else '  ✖ (phải = 100%)'),
                             fg=('#2d6a4f' if ok else '#c0392b'))
        self._show_rule()

    # ---------------- Kiem tra Ten/Ma NVL + ma hang ----------------
    def _sync_ma(self, e_ten, e_ma):
        """Ten NVL doi -> cap nhat Ma NVL theo danh sach Unit price.
        Ten khong co trong list -> xoa Ma + to do o Ten (se khong cho luu)."""
        try:
            if not e_ten.winfo_exists():
                return
        except Exception:
            return
        ten = e_ten.get().strip()
        cur = e_ma.get().strip()
        if not ten:
            e_ten.configure(bg='white')
            return
        codes = self.catalog.codes_for(ten)
        if codes:
            if cur not in codes:
                e_ma.delete(0, 'end')
                e_ma.insert(0, codes[0])
            e_ten.configure(bg='white')
        else:
            if cur:
                e_ma.delete(0, 'end')
            e_ten.configure(bg='#ffd6d6')
        self._show_rule()

    def _rows_for_check(self):
        out = []
        for (e_ten, e_ma, e_kl, e_pct, e_klv, e_pctv) in self.rows:
            ten = e_ten.get().strip()
            if not ten:
                continue
            try:
                p = float((e_pct.get() or '0').strip().replace(',', '.'))
            except ValueError:
                p = 0.0
            out.append((ten, e_ma.get().strip(), p))
        return out

    def _show_rule(self):
        lbl = getattr(self, 'rule_lbl', None)
        if lbl is None:
            return
        mats = self._rows_for_check()
        if not mats:
            lbl.config(text='')
            return
        errs, info = check_order(self.v['ma_hang'].get(), mats, self.catalog)
        head = (f"Bột đá: {info['filler']:.4g}%   |   PP: {info['pp']:.4g}%   |   "
                f"PE: {info['pe']:.4g}%")
        if errs:
            lbl.config(text=head + '   ✖ ' + '  /  '.join(e.split('\n')[0] for e in errs),
                       fg='#c0392b')
        else:
            lbl.config(text=head + '   ✔ mã hàng hợp lệ', fg='#2d6a4f')

    def _auto_recalc(self):
        if getattr(self, 'auto_kg', None) and self.auto_kg.get():
            self._recalc_kg()

    def _base_batch(self):
        """Tong me = KL_nen / (%_nen/100) tu dong dau tien co du KL va %."""
        for (e_ten, e_ma, e_kl, e_pct, e_klv, e_pctv) in self.rows:
            if not e_ten.get().strip():
                continue
            try:
                p = float(e_pct.get() or 0)
                k = float(e_kl.get() or 0)
            except ValueError:
                continue
            if p > 0 and k > 0:
                return k / (p / 100.0)
        return None

    def _recalc_kg(self):
        """Tinh lai KL/me cho moi dong = (%/100) x tong me."""
        batch = self._base_batch()
        if not batch:
            return
        for (e_ten, e_ma, e_kl, e_pct, e_klv, e_pctv) in self.rows:
            if not e_ten.get().strip():
                continue
            try:
                p = float(e_pct.get() or 0)
            except ValueError:
                continue
            kg = p / 100.0 * batch
            e_kl.delete(0, 'end')
            e_kl.insert(0, f"{kg:.4f}".rstrip('0').rstrip('.'))
        # hien thi tong
        tp = 0.0
        for (e_ten, e_ma, e_kl, e_pct, e_klv, e_pctv) in self.rows:
            if e_ten.get().strip():
                try:
                    tp += float(e_pct.get() or 0)
                except ValueError:
                    pass
        self._show_tong()
        self.tong_lbl.config(text=self.tong_lbl.cget('text') + f'  |  Cỡ mẻ: {batch:,.1f} kg')

    # ---------------- LAM TRON KL/me bang cach chinh % ----------------
    @staticmethod
    def _dec(t):
        from decimal import Decimal, InvalidOperation
        try:
            return Decimal((t or '').strip().replace(',', '.') or '0')
        except InvalidOperation:
            return None

    @staticmethod
    def _dstr(d, places):
        """Decimal -> chuoi, toi da 'places' so le, bo so 0 thua (12.000 -> '12')."""
        from decimal import Decimal, ROUND_HALF_UP
        q = d.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
        t = format(q, 'f')
        if '.' in t:
            t = t.rstrip('0').rstrip('.')
        return '0' if t in ('', '-0') else t

    def _round_plan(self, step, bal_idx, pct_places=None):
        """Tinh phuong an lam tron. Tra ve (batch, [(idx, ten, kg_cu, kg_moi, pct_cu, pct_moi)], loi).
        Co me giu nguyen = KL_nen / (%_nen/100). Moi dong (tru dong can bang) KL lam tron theo 'step';
        dong can bang (NVL nhua cuoi) = phan con lai -> tong % dung 100.
        pct_places = 4: % chi giu 4 so le -> KL/me chi le tu so thu 4 sau dau phay
        (VD Acid 2.28997 -> 0.7235% -> 2.3001 kg), tong % van dung 100.0000%."""
        from decimal import Decimal, ROUND_HALF_UP
        used = [(i, r) for i, r in enumerate(self.rows) if r[0].get().strip()]
        if len(used) < 2:
            return None, [], 'Cần ít nhất 2 dòng NVL.'
        batch = None
        for i, r in used:
            k, p = self._dec(r[2].get()), self._dec(r[3].get())
            if k is None or p is None:
                return None, [], f'Dòng "{r[0].get().strip()}": KL hoặc % không phải số.'
            if batch is None and k > 0 and p > 0:
                batch = k / (p / Decimal(100))
        if not batch:
            return None, [], 'Chưa có dòng nền (dòng đầu tiên có cả KL/mẻ và %) để tính cỡ mẻ.'
        stepd = Decimal(str(step))
        plan, sum_kg, sum_pct = [], Decimal(0), Decimal(0)
        for i, r in used:
            if i == bal_idx:
                continue
            p_old = self._dec(r[3].get())
            kg_old = p_old / Decimal(100) * batch          # KL dung theo % hien tai
            sd = stepd
            kg_new = (kg_old / sd).quantize(Decimal(1), rounding=ROUND_HALF_UP) * sd
            while kg_new <= 0 < kg_old and sd > Decimal('0.001'):
                sd = sd / 10                               # NVL qua nho (VD 0.04) -> lam tron min hon
                kg_new = (kg_old / sd).quantize(Decimal(1), rounding=ROUND_HALF_UP) * sd
            if pct_places:
                # % lam tron 4 so le -> KL tinh lai tu % (chi lech o so le thu 4)
                pq = Decimal(1).scaleb(-int(pct_places))
                pct_new = (kg_new / batch * Decimal(100)).quantize(pq, rounding=ROUND_HALF_UP)
                if pct_new <= 0 < kg_old:
                    pct_new = pq
                kg_new = (batch * pct_new / Decimal(100)).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP)
            else:
                pct_new = Decimal(self._dstr(kg_new / batch * Decimal(100), 12))
            sum_kg += kg_new
            sum_pct += pct_new
            plan.append((i, r[0].get().strip(), kg_old, kg_new, p_old, pct_new))
        rb = self.rows[bal_idx]
        p_old_b = self._dec(rb[3].get())
        pct_b = Decimal(100) - sum_pct
        if pct_places:
            kg_b = (batch * pct_b / Decimal(100)).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP)
        else:
            kg_b = batch - sum_kg
        if pct_b <= 0 or kg_b <= 0:
            return batch, [], (f'Dòng cân bằng "{rb[0].get().strip()}" bị âm/bằng 0 sau khi làm tròn — '
                               'hãy chọn bước làm tròn nhỏ hơn hoặc dòng cân bằng khác.')
        plan.append((bal_idx, rb[0].get().strip(), p_old_b / Decimal(100) * batch, kg_b, p_old_b, pct_b))
        plan.sort(key=lambda x: x[0])
        return batch, plan, None

    def _round_dialog(self):
        used = [(i, r) for i, r in enumerate(self.rows) if r[0].get().strip()]
        if len(used) < 2:
            messagebox.showwarning('Làm tròn', 'Cần ít nhất 2 dòng NVL.', parent=self)
            return
        # dong can bang mac dinh = NVL NHUA (PP/PE) cuoi cung; khong co thi lay dong cuoi
        bal_default = used[-1][0]
        for i, r in reversed(used):
            if self.catalog.kind_of(r[0].get().strip(), r[1].get().strip()) in ('PP', 'PE'):
                bal_default = i
                break
        dlg = tk.Toplevel(self)
        dlg.title('🎯 Làm tròn KL/mẻ')
        dlg.configure(bg='#f4f6f9')
        dlg.transient(self)
        dlg.grab_set()
        tk.Label(dlg, text='Làm tròn KL/mẻ bằng cách chỉnh lại Tỉ lệ % (cỡ mẻ giữ nguyên)', bg='#f4f6f9',
                 font=FONT_B, fg='#1f3a5f').pack(anchor='w', padx=12, pady=(10, 2))
        tk.Label(dlg, text='Các NVL được làm tròn KL theo bước chọn (VD Acid 2.28997 → 2.3001, '
                           'LLDPE 12.0159 → 12.0001).\nTỉ lệ % giữ 4 số lẻ nên KL/mẻ chỉ lẻ từ số thứ 4 sau dấu phẩy. '
                           '1 dòng "cân bằng" (mặc định NVL nhựa cuối cùng) nhận phần còn lại để tổng % = 100%.',
                 bg='#f4f6f9', fg='#555', font=('Segoe UI', 9), justify='left').pack(anchor='w', padx=12)
        f = tk.Frame(dlg, bg='#f4f6f9')
        f.pack(fill='x', padx=12, pady=6)
        tk.Label(f, text='Bước làm tròn (kg):', bg='#f4f6f9', font=FONT).grid(row=0, column=0, sticky='e')
        steps = ['1', '0.5', '0.1', '0.05', '0.01']
        v_step = tk.StringVar(value='0.1')
        cb_step = ttk.Combobox(f, textvariable=v_step, values=steps, width=8, state='readonly', font=FONT)
        cb_step.grid(row=0, column=1, sticky='w', padx=6)
        tk.Label(f, text='Dòng cân bằng (để lẻ):', bg='#f4f6f9', font=FONT).grid(row=0, column=2, sticky='e',
                                                                              padx=(16, 0))
        names = [f'{r[0].get().strip()}' for i, r in used]
        idxs = [i for i, r in used]
        v_bal = tk.StringVar(value=names[idxs.index(bal_default)])
        cb_bal = ttk.Combobox(f, textvariable=v_bal, values=names, width=22, state='readonly', font=FONT)
        cb_bal.current(idxs.index(bal_default))
        cb_bal.grid(row=0, column=3, sticky='w', padx=6)
        v_p4 = tk.BooleanVar(value=True)
        tk.Checkbutton(f, text='Tỉ lệ % giữ 4 số lẻ (KL chỉ lẻ từ số thứ 4 — sát số hơn)', variable=v_p4,
                       bg='#f4f6f9', font=FONT, command=lambda: refresh()).grid(
            row=1, column=0, columnspan=4, sticky='w', pady=(6, 0))

        cols = ('ten', 'kg_old', 'kg_new', 'p_old', 'p_new')
        tv = ttk.Treeview(dlg, columns=cols, show='headings', height=min(12, len(used)))
        for c, h, w in (('ten', 'Tên NVL', 170), ('kg_old', 'KL/mẻ cũ', 110), ('kg_new', 'KL/mẻ mới', 110),
                        ('p_old', '% cũ', 130), ('p_new', '% mới', 150)):
            tv.heading(c, text=h)
            tv.column(c, width=w, anchor='w' if c == 'ten' else 'e')
        tv.tag_configure('bal', background='#fff3e0')
        tv.pack(fill='both', expand=True, padx=12, pady=4)
        info = tk.Label(dlg, text='', bg='#f4f6f9', font=FONT_B, justify='left', anchor='w')
        info.pack(fill='x', padx=12)
        state = {'plan': None}

        def refresh(*_a):
            tv.delete(*tv.get_children())
            try:
                bal = idxs[cb_bal.current()]
            except Exception:
                bal = bal_default
            p4 = 4 if v_p4.get() else None
            batch, plan, err = self._round_plan(float(v_step.get()), bal, p4)
            state['plan'] = None if err else (batch, plan, p4)
            for (i, ten, kg_o, kg_n, p_o, p_n) in plan:
                if p4:
                    kg_txt, p_txt = f'{kg_n:.4f}', f'{p_n:.4f}'
                else:
                    kg_txt = f'{kg_n:.3f}' if i != bal else self._dstr(kg_n, 5)
                    p_txt = self._dstr(p_n, 12)
                tv.insert('', 'end', tags=('bal',) if i == bal else (),
                          values=(ten + ('  ⚖' if i == bal else ''), self._dstr(kg_o, 5),
                                  kg_txt, self._dstr(p_o, 6), p_txt))
            if err:
                info.config(text='✖ ' + err, fg='#c0392b')
                btn_ok.config(state='disabled')
            else:
                from decimal import Decimal
                tot = sum((p[5] for p in plan), Decimal(0))
                info.config(text=f'Cỡ mẻ: {self._dstr(batch, 4)} kg   |   Tổng % mới: {self._dstr(tot, 12)}%  ✔',
                            fg='#2d6a4f')
                btn_ok.config(state='normal')

        def apply():
            if not state['plan']:
                return
            batch, plan, p4 = state['plan']
            self._round_undo = [(r[2].get(), r[3].get()) for r in self.rows]
            bal = idxs[cb_bal.current()]
            for (i, ten, kg_o, kg_n, p_o, p_n) in plan:
                e_kl, e_pct = self.rows[i][2], self.rows[i][3]
                e_kl.delete(0, 'end')
                if p4:
                    e_kl.insert(0, self._dstr(kg_n, 4))
                else:
                    e_kl.insert(0, f'{kg_n:.3f}' if i != bal else self._dstr(kg_n, 6))
                e_pct.delete(0, 'end')
                e_pct.insert(0, self._dstr(p_n, p4 or 12))
            self.btn_round_undo.config(state='normal')
            self._show_tong()
            self.tong_lbl.config(text=self.tong_lbl.cget('text') + f'  |  Cỡ mẻ: {float(batch):,.1f} kg')
            dlg.destroy()

        bf = tk.Frame(dlg, bg='#f4f6f9')
        bf.pack(fill='x', padx=12, pady=10)
        btn_ok = tk.Button(bf, text='✔ Áp dụng làm tròn', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat',
                           padx=14, pady=5, command=apply)
        btn_ok.pack(side='right')
        tk.Button(bf, text='Đóng', font=FONT, relief='flat', padx=12, pady=5,
                  command=dlg.destroy).pack(side='right', padx=8)
        cb_step.bind('<<ComboboxSelected>>', refresh)
        cb_bal.bind('<<ComboboxSelected>>', refresh)
        refresh()

    def _round_undo_apply(self):
        """Tra lai KL/me + % nhu truoc khi bam Lam tron."""
        if not self._round_undo:
            return
        for (kl, pct), r in zip(self._round_undo, self.rows):
            r[2].delete(0, 'end')
            r[2].insert(0, kl)
            r[3].delete(0, 'end')
            r[3].insert(0, pct)
        self._round_undo = None
        self.btn_round_undo.config(state='disabled')
        self._show_tong()

    def _recalc_virtual(self):
        """Cot AO: doi Kg/me ao -> % ao thay doi.
        Giu nguyen % ao cua DONG NEN (dong dau tien co ten) lam moc,
        cac dong khac tinh theo cung co me = KL_nen / (%_nen/100)."""
        # tim dong nen = dong dau tien co ten
        base_idx = None
        for i, r in enumerate(self.rows):
            if r[0].get().strip():
                base_idx = i
                break

        def num(e):
            try:
                return float(e.get() or 0)
            except ValueError:
                return 0.0

        batch = None
        if base_idx is not None:
            bp = num(self.rows[base_idx][5])   # % ao dong nen (moc, giu nguyen)
            bk = num(self.rows[base_idx][4])   # Kg ao dong nen
            if bp > 0 and bk > 0:
                batch = bk / (bp / 100.0)

        if batch is None:
            # chua co moc -> chia theo tong
            total = sum(num(r[4]) for r in self.rows)
            for r in self.rows:
                r[5].delete(0, 'end')
                k = num(r[4])
                if total > 0 and k > 0:
                    r[5].insert(0, f"{k / total * 100:.4f}".rstrip('0').rstrip('.'))
            return

        # co moc: giu % ao dong nen, tinh lai cac dong con lai
        for i, r in enumerate(self.rows):
            if i == base_idx:
                continue                      # giu nguyen % ao dong nen
            r[5].delete(0, 'end')
            k = num(r[4])
            if k > 0:
                r[5].insert(0, f"{k / batch * 100:.4f}".rstrip('0').rstrip('.'))

    def _next_don_khach(self, src_don, src_khach):
        """Tao Don hang & Khach hang cho don MOI:
        - so don (o giua) = so lon nhat trong loai + 1
        - so khach (phan cuoi Don hang) = so lon nhat cua cung ten khach + 1
        VD: PO/26090/17.TBC -> PO/26092/18.TBC
        Khach hang dang ma moi (VD 'CN10', khong co so dau) -> giu nguyen;
        Khach hang dang cu 'NN.XXX' -> tang so nhu truoc (18.XXX).
        """
        import re
        pat = re.compile(r'^\s*(\d+)\s*\.?\s*(.*)$')
        parts = str(src_don or '').split('/')
        prefix = parts[0] if parts and parts[0] else 'PO'
        # so thu tu khach lay tu phan cuoi cua Don hang (VD 01.KSP);
        # khong co thi lay tu cot Khach hang (dinh dang cu)
        tail = parts[2].strip() if len(parts) >= 3 else str(src_khach or '').strip()
        mk = pat.match(tail)
        cust_num = int(mk.group(1)) if mk else 0
        cust_name = (mk.group(2) if mk else tail).strip()
        max_mid = 0
        max_cust = 0
        mid_width = 5
        cust_width = max(2, len(mk.group(1)) if mk else 2)
        for o in self.orders:
            p = str(o.get('don_hang') or '').split('/')
            if len(p) >= 2:
                try:
                    max_mid = max(max_mid, int(p[1]))
                    mid_width = max(mid_width, len(p[1]))
                except ValueError:
                    pass
            for v in ((p[2] if len(p) >= 3 else ''), o.get('khach_hang')):
                mkk = pat.match(str(v or '').strip())
                if mkk and mkk.group(2).strip().lower() == cust_name.lower():
                    try:
                        max_cust = max(max_cust, int(mkk.group(1)))
                        cust_width = max(cust_width, len(mkk.group(1)))
                    except ValueError:
                        pass
        new_mid = str(max(max_mid, 0) + 1).zfill(mid_width)
        new_cust = str(max(max_cust, cust_num) + 1).zfill(cust_width)
        new_tail = f"{new_cust}.{cust_name}" if cust_name else new_cust
        new_don = f"{prefix}/{new_mid}/{new_tail}"
        kh = str(src_khach or '').strip()
        new_khach = new_tail if (not kh or pat.match(kh)) else kh
        return new_don, new_khach

    def _olabel(self, o):
        return f"[{o.get('ma_sp')}] {o.get('don_hang')} • {o.get('ma_hang')}"

    # ---------------- Goi y ten NVL tu Don gia ----------------
    def _build_nvl_list(self):
        out = []
        try:
            for sh in self.app.sheets:
                if sh.get('title') == 'Đơn giá vật tư':
                    for r in sh['rows']:
                        ma_hoa = str(r[1] or '').strip()
                        ma_vt = str(r[2] or '').strip()
                        ten_vt = str(r[3] or '').strip()
                        if ma_hoa and ma_vt:
                            out.append((ma_hoa, ma_vt, ten_vt))
                    break
        except Exception:
            pass
        return out

    def _ac_matches(self, q):
        q = q.strip().lower()
        if not q:
            return []
        starts, contains = [], []
        for (ma_hoa, ma_vt, ten_vt) in self.nvl_list:
            hay = f"{ma_hoa} {ten_vt}".lower()
            if ma_hoa.lower().startswith(q):
                starts.append((ma_hoa, ma_vt, ten_vt))
            elif q in hay:
                contains.append((ma_hoa, ma_vt, ten_vt))
        return (starts + contains)[:20]

    def _ac_key(self, ev, ten_entry, ma_entry):
        if ev.keysym in ('Up', 'Down', 'Return', 'Escape'):
            return
        self._sync_ma(ten_entry, ma_entry)
        self._ac_ten = ten_entry
        self._ac_ma = ma_entry
        items = self._ac_matches(ten_entry.get())
        if not items:
            self._ac_hide()
            return
        self._ac_items = items
        if self._ac_pop is None or not tk.Toplevel.winfo_exists(self._ac_pop):
            self._ac_pop = tk.Toplevel(self)
            self._ac_pop.wm_overrideredirect(True)
            self._ac_pop.attributes('-topmost', True)
            self._ac_lb = tk.Listbox(self._ac_pop, font=FONT, height=8, activestyle='dotbox')
            self._ac_lb.pack(fill='both', expand=True)
            self._ac_lb.bind('<ButtonRelease-1>', lambda e: self._ac_pick())
            self._ac_lb.bind('<Return>', lambda e: self._ac_pick())
            self._ac_lb.bind('<Escape>', lambda e: self._ac_hide())
        self._ac_lb.delete(0, 'end')
        for (ma_hoa, ma_vt, ten_vt) in items:
            self._ac_lb.insert('end', f"{ma_hoa}   —   {ten_vt}   ({ma_vt})")
        # dat popup ngay duoi o Ten
        x = ten_entry.winfo_rootx()
        y = ten_entry.winfo_rooty() + ten_entry.winfo_height()
        w = max(360, ten_entry.winfo_width())
        self._ac_pop.geometry(f"{w}x160+{x}+{y}")
        self._ac_pop.deiconify()

    def _ac_focus_list(self):
        if self._ac_pop and tk.Toplevel.winfo_exists(self._ac_pop):
            self._ac_lb.focus_set()
            self._ac_lb.selection_clear(0, 'end')
            self._ac_lb.selection_set(0)
            self._ac_lb.activate(0)

    def _ac_pick(self):
        try:
            sel = self._ac_lb.curselection()
            if not sel:
                return
            ma_hoa, ma_vt, ten_vt = self._ac_items[sel[0]]
            if self._ac_ten:
                self._ac_ten.delete(0, 'end')
                self._ac_ten.insert(0, ma_hoa)
                self._ac_ten.configure(bg='white')
            if self._ac_ma:
                self._ac_ma.delete(0, 'end')
                self._ac_ma.insert(0, ma_vt)
                self._ac_ma.configure(state='normal')
        except Exception:
            pass
        self._ac_hide()
        self._show_rule()

    def _ac_hide(self):
        if self._ac_pop is not None:
            try:
                self._ac_pop.withdraw()
            except Exception:
                pass

    def _next_masp(self):
        mx = 0
        for o in self.orders:
            try:
                mx = max(mx, int(str(o.get('ma_sp'))))
            except Exception:
                pass
        return str(mx + 1)

    # Vi du mac dinh theo loai don (dung khi chua co don nao de lay mau)
    _HINT_DEFAULT = {
        'Export':   {'pre': 'PO', 'tail': '17.TBC', 'khach': 'CN10.TBC', 'msp': '807'},
        'Domestic': {'pre': 'DH', 'tail': '43.PGV', 'khach': 'VN12.PGV', 'msp': 'PP-01'},
        'Sample':   {'pre': 'RD', 'tail': '80.PTM-RD', 'khach': '80.PTM-RD', 'msp': 'PP-01'},
    }

    def _field_hints(self):
        """Dong vi du hien duoi moi o nhap. Lay theo don GAN NHAT cung loai (dung dinh dang that)."""
        dft = self._HINT_DEFAULT.get(self.kind, self._HINT_DEFAULT['Export'])
        yy = str(getattr(self.app, 'app_year', 2026))[-2:]
        last = self.orders[-1] if self.orders else {}
        nxt = self._next_masp() if self.orders else '1'
        try:
            n_mid = f'{yy}{int(nxt):03d}'
        except ValueError:
            n_mid = f'{yy}001'
        don_ex = f"{dft['pre']}/{n_mid}/{dft['tail']}"
        if last.get('don_hang') and str(last.get('don_hang')).count('/') >= 2:
            try:
                don_ex, _k = self._next_don_khach(last.get('don_hang'), last.get('khach_hang'))
            except Exception:
                pass
        khach_ex = (last.get('khach_hang') or dft['khach']) if last else dft['khach']
        ma_hang_ex = (last.get('ma_hang') or 'HTP-P1012BF-7') if last else 'HTP-P1012BF-7'
        return {
            'ma_sp': f'VD: {nxt}  (số thứ tự lệnh, tự tăng)',
            'don_hang': f'VD: {don_ex}  (không nhập mỗi "01")',
            'ma_hang': f'VD: {ma_hang_ex}',
            'khach_hang': f'VD: {khach_ex}',
            'ma_san_pham': f"VD: {dft['msp']}",
            'hao_hut': 'VD: 0.7  (= 0,7%, dùng dấu chấm)',
        }

    def _check_don_format(self):
        """Canh bao neu Don hang so khong dung dang XX/26xxx/NN.KHACH. Tra ve True = cho luu."""
        import re
        don = self.v['don_hang'].get().strip()
        if re.match(r'^[A-Za-z]{2}\s*/\s*\d{4,}\s*/\s*\S+', don):
            return True
        ex = self._field_hints().get('don_hang', '').replace('VD: ', '').split('  (')[0]
        return messagebox.askyesno(
            'Kiểm tra Đơn hàng số',
            f'Đơn hàng số "{don or "(trống)"}" có vẻ KHÔNG đúng định dạng.\n\n'
            f'Định dạng đúng:  {ex}\n'
            '   (Loại đơn / năm + số lệnh / số lần.mã khách)\n\n'
            'Vẫn lưu với giá trị này?', icon='warning', parent=self)

    def _load_src(self):
        i = self.src_cb.current()
        if i < 0:
            return
        o = self.orders[i]
        self.v['ma_hang'].set(o.get('ma_hang') or '')
        # Them moi: don hang & khach hang tu tang so; Sua: giu nguyen
        if self.mode == 'add':
            new_don, new_khach = self._next_don_khach(o.get('don_hang'), o.get('khach_hang'))
            self.v['don_hang'].set(new_don)
            self.v['khach_hang'].set(new_khach)
        else:
            self.v['don_hang'].set(o.get('don_hang') or '')
            self.v['khach_hang'].set(o.get('khach_hang') or '')
        self._copied_loai = o.get('loai') or ''
        # Ma san pham: lay tu order, neu chua co thi lay tu 'nhom' cua NVL dau
        msp = o.get('ma_san_pham')
        if not msp:
            for m in o.get('materials', []):
                if m.get('nhom'):
                    msp = m.get('nhom')
                    break
        self.v['ma_san_pham'].set(msp or '')
        hh = o.get('hao_hut')
        self.v['hao_hut'].set(str(hh) if hh not in (None, '') else '0.7')
        if self.mode == 'edit':
            self.edit_target = o
            self.v['ma_sp'].set(str(o.get('ma_sp')))
        mats = [m for m in o.get('materials', []) if (m.get('ten') or '').strip()]
        for j, (e_ten, e_ma, e_kl, e_pct, e_klv, e_pctv) in enumerate(self.rows):
            for e in (e_ten, e_ma, e_kl, e_pct, e_klv, e_pctv):
                e.delete(0, 'end')
            if j < len(mats):
                m = mats[j]
                e_ten.insert(0, m.get('ten') or '')
                e_ma.insert(0, m.get('ma_nvl') or '')
                e_kl.insert(0, f"{float(m.get('mass_batch') or 0):g}")
                e_pct.insert(0, _pct_str(m.get('percentage')))
                # cot AO cung nhan cong thuc (KL ao = KL, roi tinh % ao)
                e_klv.insert(0, f"{float(m.get('mass_batch') or 0):g}")
        for (e_ten, e_ma, *_r) in self.rows:
            t = e_ten.get().strip()
            e_ten.configure(bg=('#ffd6d6' if t and not self.catalog.codes_for(t) else 'white'))
        self._recalc_virtual()
        # nap thong so Container Volume theo don
        self._cv_load(o.get('don_hang'), o.get('ma_hang'))
        # hien thi tong % (12 so le, xanh/do) / co me
        self._show_tong()
        b = self._base_batch()
        if b:
            self.tong_lbl.config(text=self.tong_lbl.cget('text') + f'  |  Cỡ mẻ: {b:,.1f} kg')

    # ---------------- Container Volume ----------------
    def _cv_sheet(self):
        for sh in self.app.sheets:
            if sh.get('title') == 'Container Volume':
                return sh
        return None

    def _cv_find(self, don, ma):
        sh = self._cv_sheet()
        if not sh:
            return None, None
        for i, r in enumerate(sh['rows']):
            if str(r[1] if len(r) > 1 else '').strip() == str(don or '').strip() and \
               str(r[2] if len(r) > 2 else '').strip() == str(ma or '').strip():
                return sh, i
        return sh, None

    def _cv_calc_pallet(self):
        """so_pallet = ROUNDUP(SL / Kg-Pallet), toi da Pallet/Cont x So cont, uu tien so chan."""
        if not self.cv_vars:
            return
        from engine import chia_pallet
        pc = self.cv_vars['pallet_cont'].get() if self.cv_vars.get('pallet_cont') else 20
        kq = chia_pallet(self.cv_vars['san_luong'].get(), self.cv_vars['kg_pallet'].get(),
                         self.cv_vars['loai_bao'].get(), so_cont=self.cv_vars['so_cont'].get(),
                         pallet_cont=pc or 20)
        if kq:
            self.cv_vars['so_pallet'].set(str(kq['tong_pallet']))

    def _cv_calc_cont(self):
        """so_cont = ROUNDUP( san_luong / 28000 )  (1 cont <= 28.000 kg va <= 20 pallet)."""
        if not self.cv_vars:
            return
        from engine import tinh_so_cont
        n = tinh_so_cont(self.cv_vars['san_luong'].get())
        if n > 0:
            self.cv_vars['so_cont'].set(str(n))

    def _cv_calc_xep(self):
        """Tu tinh chuoi 'Cach xep pallet' khi Quy cach = Hang Xep Pallet."""
        if not self.cv_vars or not self.cv_vars.get('cach_xep'):
            return
        from engine import cach_xep_pallet
        if (self.cv_vars['quy_cach'].get() or '').strip().lower() != 'hàng xếp pallet':
            return
        ma = self.v['ma_hang'].get().strip() if self.v.get('ma_hang') else ''
        mix = bool(self.cv_vars['mix'].get()) if self.cv_vars.get('mix') else False
        phu = [(m.get(), k.get()) for m, k in zip(self.cv_vars['ma_phu'], self.cv_vars['kg_phu'])]
        pc = self.cv_vars['pallet_cont'].get() if self.cv_vars.get('pallet_cont') else 20
        s = cach_xep_pallet(self.cv_vars['san_luong'].get(), self.cv_vars['kg_pallet'].get(),
                            self.cv_vars['loai_bao'].get(), ma_hang=ma, mix=mix, phu=phu,
                            so_cont=self.cv_vars['so_cont'].get(), pallet_cont=pc or 20)
        if s:
            self.cv_vars['cach_xep'].set(s)
        # Xep lan: cap nhat Tong so Pallet = pallet day tung ma + pallet lan
        if mix and self.cv_vars.get('so_pallet'):
            from engine import xep_lan_pallet
            kl = xep_lan_pallet(self.cv_vars['san_luong'].get(), self.cv_vars['kg_pallet'].get(),
                                self.cv_vars['loai_bao'].get(), ma_hang=ma, phu=phu)
            if kl:
                self.cv_vars['so_pallet'].set(str(kl['tong_pallet']))

    def _cv_recalc(self):
        self._cv_calc_cont()
        self._cv_calc_pallet()
        self._cv_calc_xep()

    def _cv_load(self, don, ma):
        if not self.cv_vars:
            return
        sh, i = self._cv_find(don, ma)
        if sh is None or i is None:
            self._cv_calc_pallet()
            return
        r = sh['rows'][i]
        def g(idx):
            return r[idx] if len(r) > idx else ''
        if g(4) not in (None, ''):
            self.cv_vars['san_luong'].set(str(g(4)))
        if g(5) not in (None, ''):
            self.cv_vars['so_cont'].set(str(g(5)))
        if g(6) not in (None, ''):
            self.cv_vars['kg_pallet'].set(str(g(6)))
        if g(7):
            self.cv_vars['loai_bao'].set(str(g(7)))
        if g(8):
            self.cv_vars['quy_cach'].set(str(g(8)))
        if g(10) not in (None, ''):
            self.cv_vars['pallet_cont'].set(str(g(10)))
        if g(11) not in (None, '') and self.cv_vars.get('cach_xep'):
            self.cv_vars['cach_xep'].set(str(g(11)))
        # neu chua co so cont thi tu tinh (truoc, vi so pallet phu thuoc so cont)
        if g(5) in (None, ''):
            self._cv_calc_cont()
        if g(9) not in (None, ''):
            self.cv_vars['so_pallet'].set(str(g(9)))
        else:
            self._cv_calc_pallet()

    def _cv_save(self, don, ma, ma_sp, ma_san_pham):
        """Upsert 1 dong Container Volume theo (don, ma)."""
        if not self.cv_vars:
            return
        sh, i = self._cv_find(don, ma)
        if sh is None:
            return
        vals = [self.cv_vars['san_luong'].get().strip(), self.cv_vars['so_cont'].get().strip(),
                self.cv_vars['kg_pallet'].get().strip(), self.cv_vars['loai_bao'].get(),
                self.cv_vars['quy_cach'].get(), self.cv_vars['so_pallet'].get().strip(),
                self.cv_vars['pallet_cont'].get().strip(),
                self.cv_vars['cach_xep'].get().strip() if self.cv_vars.get('cach_xep') else '']
        if i is None:
            stt = len(sh['rows']) + 1
            row = [stt, don, ma, ma_san_pham or '', vals[0], vals[1], vals[2], vals[3],
                   vals[4], vals[5], vals[6], vals[7]]
            if CV_RAN_HEADER in sh.get('headers', []):
                row.append('')               # San luong da chay: don moi = chua chay
            sh['rows'].append(row)
        else:
            r = sh['rows'][i]
            while len(r) < 12:
                r.append('')
            r[4], r[5], r[6], r[7], r[8], r[9], r[10], r[11] = vals
            if CV_RAN_HEADER in sh.get('headers', []):
                while len(r) <= CV_RAN_IDX:
                    r.append('')
                # 'San luong da chay' giu nguyen (khong tu gan = san luong don nua,
                # neu khong don se bi coi la da chay du va an khoi o chon cua tab QC)
        # luu sheets.json + cap nhat index + tab
        try:
            save_sheets_file(self.app.sheets)
        except Exception as e:
            messagebox.showerror('Lỗi lưu', f'Không lưu được Container Volume:\n{e}', parent=self)
        if hasattr(self.app, 'rebuild_cv_index'):
            self.app.rebuild_cv_index()
        cvt = getattr(self.app, 'cv_tab', None)
        if cvt:
            try:
                cvt._fill()
            except Exception:
                pass

    def _save(self):
        import datetime
        ma_sp = self.v['ma_sp'].get().strip()
        if not ma_sp:
            messagebox.showwarning('Thiếu', 'Nhập Mã SP.', parent=self)
            return
        editing = (self.mode == 'edit' and self.edit_target is not None)
        if not editing and any(str(o.get('ma_sp')) == ma_sp for o in self.orders):
            messagebox.showwarning('Trùng', f'Mã SP {ma_sp} đã tồn tại.', parent=self)
            return
        if not editing and datastore.is_server():
            # nhieu may dung chung: may khac co the vua tao lenh cung Ma SP
            try:
                taken = datastore.order_key_taken(self.kind, ma_sp)
            except Exception as e:
                messagebox.showerror('Mất kết nối', f'Không kiểm tra được Mã SP trên máy chủ:\n{e}', parent=self)
                return
            if taken:
                new = datastore.next_free_ma_sp(self.kind, self.orders)
                if not messagebox.askyesno(
                        'Mã SP đã được dùng',
                        f'Mã SP {ma_sp} vừa được MÁY KHÁC dùng cho lệnh {self.kind} khác.\n\n'
                        f'Đổi sang Mã SP mới: {new} rồi lưu?', parent=self):
                    return
                ma_sp = new
                self.v['ma_sp'].set(new)
        mats = []
        thieu_ma = []      # co Ten nhung thieu Ma NVL
        thieu_ten = []     # co Ma nhung thieu Ten
        for idx, (e_ten, e_ma, e_kl, e_pct, e_klv, e_pctv) in enumerate(self.rows, 1):
            ten = e_ten.get().strip()
            ma = e_ma.get().strip()
            if not ten and not ma:
                continue
            if ten and not ma:
                thieu_ma.append(f'  • Dòng {idx}: {ten}')
                continue
            if ma and not ten:
                thieu_ten.append(f'  • Dòng {idx}: mã {ma}')
                continue
            try:
                kl = float(e_kl.get() or 0)
            except ValueError:
                kl = 0
            try:
                pct = float(e_pct.get() or 0) / 100.0
            except ValueError:
                pct = 0
            mats.append({'ten': ten, 'ma_nvl': ma, 'ten_day_du': '',
                         'mass_batch': kl, 'percentage': pct, 'nhom': None})
        # --- KIEM TRA: Ten NVL phai co trong danh sach Unit price & Ma phai dung ma cua ten do ---
        sai_ds = []
        for idx, (e_ten, e_ma, *_r) in enumerate(self.rows, 1):
            ten = e_ten.get().strip()
            ma = e_ma.get().strip()
            if not ten:
                continue
            codes = self.catalog.codes_for(ten)
            if not codes:
                sai_ds.append(f"  • Dòng {idx}: '{ten}' KHÔNG có trong danh sách Unit price")
                e_ten.configure(bg='#ffd6d6')
            elif ma and ma not in codes:
                sai_ds.append(f"  • Dòng {idx}: '{ten}' mã đúng là {' / '.join(codes)} "
                              f"(đang nhập {ma})")
        if sai_ds:
            messagebox.showerror(
                'Tên / Mã NVL không hợp lệ',
                'Không thể lưu — Nguyên vật liệu không khớp danh sách Unit price:\n\n'
                + '\n'.join(sai_ds) +
                '\n\nGõ tên rồi chọn trong danh sách gợi ý để Mã NVL tự cập nhật.',
                parent=self)
            return
        # --- KIEM TRA: Ten va Ma NVL phai di kem nhau ---
        if thieu_ma or thieu_ten:
            msg = 'Không thể lưu — Mã NVL chưa khớp với Tên NVL:\n'
            if thieu_ma:
                msg += '\nThiếu MÃ NVL cho các dòng có tên:\n' + '\n'.join(thieu_ma)
            if thieu_ten:
                msg += '\n\nThiếu TÊN NVL cho các dòng có mã:\n' + '\n'.join(thieu_ten)
            msg += '\n\nVui lòng nhập đủ cả Tên và Mã NVL trên cùng một dòng.'
            messagebox.showwarning('Thiếu Mã NVL', msg, parent=self)
            return
        if not mats:
            messagebox.showwarning('Thiếu', 'Nhập ít nhất 1 nguyên vật liệu.', parent=self)
            return
        # --- KIEM TRA: Tong % phai dung 100.000000000000% (khong thieu, khong vuot) ---
        from decimal import Decimal, InvalidOperation
        tong_pct = Decimal(0)
        for (e_ten, e_ma, e_kl, e_pct, e_klv, e_pctv) in self.rows:
            if not (e_ten.get().strip() and e_ma.get().strip()):
                continue
            try:
                tong_pct += Decimal((e_pct.get() or '0').strip().replace(',', '.'))
            except InvalidOperation:
                messagebox.showwarning('Sai tỉ lệ %',
                                       f"Tỉ lệ % của '{e_ten.get().strip()}' không phải là số.",
                                       parent=self)
                return
        # sai so lam tron so thuc (~1e-12) coi nhu bang 100
        if abs(tong_pct - Decimal(100)) > Decimal('1e-11'):
            lech = tong_pct - Decimal(100)
            tinh_trang = 'THIẾU' if lech < 0 else 'VƯỢT QUÁ'
            messagebox.showerror(
                'Tổng % không đúng',
                f'Không thể lưu — Tổng tỉ lệ % đang {tinh_trang} 100%.\n\n'
                f'Tổng hiện tại:  {tong_pct:.12f}%\n'
                f'Yêu cầu:          100.000000000000%\n'
                f'Chênh lệch:      {lech:+.12f}%\n\n'
                'Vui lòng chỉnh lại cột % cho đủ đúng 100% rồi lưu lại.',
                parent=self)
            return
        # --- KIEM TRA: ma hang <-> % bot da & PP/PE (P20/P10) ---
        errs, _info = check_order(self.v['ma_hang'].get(), self._rows_for_check(), self.catalog)
        if errs:
            messagebox.showerror(
                'Mã hàng không khớp công thức',
                'Không thể lưu — Mã hàng không khớp với tỉ lệ nguyên vật liệu:\n\n'
                + '\n\n'.join('• ' + e for e in errs) +
                '\n\nVí dụ: lệnh PO/26091/17.TBC, mã hàng HTP-P2015IM-3\n'
                '   P20 ⇒ PP nhiều hơn PE;  15 ⇒ bột đá = 100 − 15 = 85%\n'
                '   ⇒ nhập bột đá 86% là SAI, phải sửa về 85% (85 → dưới 86%)\n'
                '      hoặc đổi mã hàng cho đúng với công thức.',
                parent=self)
            return
        if not self._check_don_format():
            return
        # thoi gian lam lenh = luc bam Luu
        now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        ma_san_pham = self.v['ma_san_pham'].get().strip()
        for m in mats:            # gan Ma san pham vao tung NVL (cot 'nhom')
            m['nhom'] = ma_san_pham
        try:
            hao_hut = float(self.v['hao_hut'].get() or 0.7)
        except ValueError:
            hao_hut = 0.7
        fields = {'khach_hang': self.v['khach_hang'].get().strip(),
                  'ma_hang': self.v['ma_hang'].get().strip(),
                  'don_hang': self.v['don_hang'].get().strip(),
                  'ma_san_pham': ma_san_pham, 'hao_hut': hao_hut,
                  'loai': self._copied_loai, 'materials': mats, 'thoi_gian': now}

        if editing:
            self.edit_target.update(fields)
            order = self.edit_target
        else:
            order = {'ma_sp': ma_sp}
            order.update(fields)
            self.app.data['orders'].setdefault(self.kind, []).append(order)

        try:
            save_data_file(self.app.data)
        except Exception as e:
            if not editing and order in self.app.data['orders'].get(self.kind, []):
                self.app.data['orders'][self.kind].remove(order)     # chua luu duoc -> bo lenh moi
            messagebox.showerror('Lỗi lưu', f'Không lưu được lệnh sản xuất:\n{e}', parent=self)
            return
        # luu thong so vao Container Volume (Export/Domestic)
        self._cv_save(fields['don_hang'], fields['ma_hang'], ma_sp, ma_san_pham)
        # cap nhat giao dien
        if editing:
            self.dl_tab.replace_order_rows(order)
        else:
            self.dl_tab.append_order_rows(order)
        if getattr(self.app, 'order_tab', None):
            self.app.order_tab._reload_orders()
        messagebox.showinfo('Đã lưu',
                            ('Đã cập nhật' if editing else 'Đã thêm') +
                            f' lệnh sản xuất (Mã SP {ma_sp}).\nThời gian làm lệnh: {now}',
                            parent=self)
        self.orders = self.app.data['orders'].get(self.kind, [])
        self.src_cb['values'] = [self._olabel(o) for o in self.orders]
        if not editing:
            self.v['ma_sp'].set(self._next_masp())
            # FIX: truoc day chi Ma SP tu tang, Don hang/Khach hang giu so cu -> lenh tiep theo bi trung so.
            # Nay: tang luon Don hang & Khach hang theo lenh VUA LUU, va chon san lenh vua luu trong o
            # "Copy tu don co san" de bam Nap lan sau cung lay tu lenh moi nhat.
            try:
                new_don, new_khach = self._next_don_khach(fields['don_hang'], fields['khach_hang'])
                self.v['don_hang'].set(new_don)
                self.v['khach_hang'].set(new_khach)
            except Exception:
                pass
            try:
                self.src_cb.current(self.orders.index(order))
            except (ValueError, tk.TclError):
                pass


class SheetTab(tk.Frame):
    def __init__(self, master, sheet, app=None, edit_kind=None, row_edit=False):
        super().__init__(master, bg='#f4f6f9')
        self.sheet = sheet
        self.app = app
        self.edit_kind = edit_kind      # 'Export'/'Domestic'/'Sample' neu tab DL cho phep them
        self.row_edit = row_edit        # True: cho Them/Sua/Xoa dong bang thuong
        self._iid_idx = {}              # iid tree -> chi so dong that
        # Sheet Don gia vat tu (Unit price): hien dang NHOM +/- giong Outline Excel
        self.grouped = (sheet.get('title') in GROUPED_SHEETS)
        self.is_unit_price = (sheet.get('title') == UNIT_PRICE_TITLE)
        self._closed = set()            # cac nhom dang thu gon (giu khi ve lai)
        self._mcache = {}               # cache do rong chu (pixel) de can giua
        self._colw = []
        if self.is_unit_price and renumber_unit_price(sheet['rows']) and row_edit and app:
            # STT cu bi lon xon -> danh lai tu tren xuong va luu luon
            try:
                save_sheets_file(app.sheets)
            except Exception:
                pass
        self._build()
        self._fill()

    # ---- TOC DO: tab an (chua mo) thi KHONG ve bang ngay; ve khi nguoi dung mo tab ----
    def _fill(self):
        if self.winfo_ismapped():
            self._lazy_pending = False
            return self._fill_now()
        self._lazy_pending = True
        if not getattr(self, '_lazy_bound', False):
            self._lazy_bound = True
            self.bind('<Map>', self._lazy_on_map, add='+')

    def _lazy_on_map(self, ev=None):
        if ev is not None and ev.widget is not self:
            return
        if getattr(self, '_lazy_pending', False):
            self._lazy_pending = False
            self._fill_now()

    def _build(self):
        bar = tk.Frame(self, bg='#f4f6f9')
        bar.pack(fill='x', padx=12, pady=8)
        tk.Label(bar, text=(self.sheet['title'] if self.sheet.get('title') in VI_TAB_TITLE
                            else self.sheet.get('source') or self.sheet['title']), bg='#f4f6f9',
                 font=FONT_TITLE).pack(side='left')
        self.src_lbl = tk.Label(bar, text=f"(nguồn: {self.sheet['source']} — {len(self.sheet['rows'])} dòng)",
                                bg='#f4f6f9', fg='#666', font=FONT)
        self.src_lbl.pack(side='left', padx=10)
        if self.edit_kind:
            tk.Button(bar, text='✏️ Chỉnh sửa lệnh', font=FONT_B, bg='#e07a00', fg='white',
                      relief='flat', padx=12, pady=4,
                      command=self._edit_order).pack(side='right', padx=6)
            tk.Button(bar, text='➕ Thêm lệnh sản xuất', font=FONT_B, bg='#2d6a4f', fg='white',
                      relief='flat', padx=12, pady=4,
                      command=self._add_order).pack(side='right', padx=6)
        if self.edit_kind and self._is_admin():
            # ADMIN: chon nhieu dong de xoa + dan truc tiep tu Excel
            tk.Button(bar, text='📋 Dán từ Excel', font=FONT_B, bg='#1d6fa5', fg='white',
                      relief='flat', padx=10, pady=4,
                      command=self._dl_paste).pack(side='right', padx=4)
            tk.Button(bar, text='🗑️ Xóa các dòng đã chọn', font=FONT_B, bg='#c0392b', fg='white',
                      relief='flat', padx=10, pady=4,
                      command=self._dl_delete_selected).pack(side='right', padx=4)
        if self.row_edit and self.is_unit_price:
            tk.Button(bar, text='📦 Nhận diện tồn kho (ảnh / Excel)', font=FONT_B, bg='#6a4c93', fg='white',
                      relief='flat', padx=10, pady=4, command=self._stock_import).pack(side='right', padx=4)
        if self.row_edit:
            tk.Button(bar, text='🗑️ Xóa dòng', font=FONT_B, bg='#c0392b', fg='white',
                      relief='flat', padx=10, pady=4, command=self._row_del).pack(side='right', padx=4)
            tk.Button(bar, text='✏️ Sửa dòng', font=FONT_B, bg='#e07a00', fg='white',
                      relief='flat', padx=10, pady=4, command=self._row_edit).pack(side='right', padx=4)
            tk.Button(bar, text='➕ Thêm dòng', font=FONT_B, bg='#2d6a4f', fg='white',
                      relief='flat', padx=10, pady=4, command=self._row_add).pack(side='right', padx=4)
        self.stock_lbl = None
        if self.is_unit_price:
            try:
                import stock_import
                self.stock_lbl = tk.Label(self, text=stock_import.last_update_text(self.sheet), bg='#f4f6f9',
                                          fg='#6a4c93', font=('Segoe UI', 9, 'bold'), anchor='w')
                self.stock_lbl.pack(fill='x', padx=14, pady=(0, 4))
                self._update_stock_label()
            except Exception:
                self.stock_lbl = None
        tk.Label(bar, text='Tìm:', bg='#f4f6f9', font=FONT_B).pack(side='left', padx=(16, 4))
        self.search = tk.StringVar()
        e = tk.Entry(bar, textvariable=self.search, width=26, font=FONT)
        e.pack(side='left')
        e.bind('<KeyRelease>', lambda ev: self._fill())
        self.count = tk.Label(bar, text='', bg='#f4f6f9', fg='#666', font=FONT)
        self.count.pack(side='left', padx=10)

        wrap = tk.Frame(self, bg='#f4f6f9')
        wrap.pack(fill='both', expand=True, padx=12, pady=(0, 12))
        headers = self.sheet['headers']
        cols = [f'c{i}' for i in range(len(headers))]
        if self.grouped:
            self._build_outline_bar(wrap)
            self.tree = ttk.Treeview(wrap, columns=cols, show='tree headings')
            self.tree.heading('#0', text='Nhóm (+/−)')
            self.tree.column('#0', width=250, minwidth=120, stretch=False)
        else:
            self.tree = ttk.Treeview(wrap, columns=cols, show='headings')
        for i, (c, h) in enumerate(zip(cols, headers)):
            self.tree.heading(c, text=h)
            w = 60 if i == 0 else max(80, min(240, len(h) * 9 + 30))
            self.tree.column(c, width=w, anchor='w', stretch=False)
        self._cols = cols
        # Sheet Thong so van hanh may: 'Cot 6' chi la phan cuoi so don (07.KSP) - khong dung o dau
        # -> AN khoi bang (van giu trong du lieu de cac cot sau khong bi lech vi tri), tu dien khi luu.
        self.cot6_idx = None
        if self.sheet.get('title') == 'Thông số vận hành máy' and 'Cột 6' in headers:
            self.cot6_idx = headers.index('Cột 6')
            self.tree.configure(displaycolumns=[c for i, c in enumerate(cols) if i != self.cot6_idx])
        # keo gian cot xong -> tinh lai can giua
        self.tree.bind('<ButtonRelease-1>', lambda e: self.after_idle(self._check_resize), add='+')
        self.tree.tag_configure('odd', background='#f0f4f8')
        if self.grouped:
            # mau tieu de nhom giong Excel
            self.tree.tag_configure('g1', background='#1f3a5f', foreground='#ffd54f',
                                    font=('Segoe UI', 10, 'bold'))
            self.tree.tag_configure('g2', background='#bdd7ee', foreground='#1f3a5f',
                                    font=('Segoe UI', 10, 'bold'))
            self.tree.tag_configure('g3', background='#dde9f6', foreground='#1f3a5f',
                                    font=('Segoe UI', 10, 'bold'))
            self.tree.tag_configure('g4', background='#eef4fb', foreground='#2b4c7e',
                                    font=('Segoe UI', 10, 'bold italic'))
            self.tree.bind('<<TreeviewOpen>>', lambda e: self._on_toggle(True))
            self.tree.bind('<<TreeviewClose>>', lambda e: self._on_toggle(False))
        vs = ttk.Scrollbar(wrap, orient='vertical', command=self.tree.yview)
        hs = ttk.Scrollbar(wrap, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        self.tree.grid(row=0, column=0, sticky='nsew')
        vs.grid(row=0, column=1, sticky='ns')
        hs.grid(row=1, column=0, sticky='ew')
        wrap.grid_rowconfigure(0, weight=1)
        wrap.grid_columnconfigure(0, weight=1)
        if self.row_edit and not self.edit_kind:
            # bang thuong (Machine, Packaging, Unit price...): chon nhieu dong + phim Delete de xoa
            self.tree.configure(selectmode='extended')
            self.tree.bind('<Delete>', lambda e: self._row_del())
            self.tree.bind('<Control-a>', lambda e: (self.tree.selection_set(self.tree.get_children()), 'break')[1])
            self.tree.bind('<Control-A>', lambda e: (self.tree.selection_set(self.tree.get_children()), 'break')[1])
        if self.edit_kind and self._is_admin():
            self.tree.configure(selectmode='extended')
            self.tree.bind('<Delete>', lambda e: self._dl_delete_selected())
            self.tree.bind('<Control-a>', lambda e: (self.tree.selection_set(self.tree.get_children()), 'break')[1])
            self.tree.bind('<Control-A>', lambda e: (self.tree.selection_set(self.tree.get_children()), 'break')[1])
            self.tree.bind('<Control-v>', lambda e: (self._dl_paste(), 'break')[1])
            self.tree.bind('<Control-V>', lambda e: (self._dl_paste(), 'break')[1])
            self.tree.bind('<<TreeviewSelect>>', lambda e: self._dl_sel_count(), add='+')
            tk.Label(self, text='Admin: giữ Ctrl / Shift để chọn nhiều dòng, Ctrl+A chọn tất cả, '
                                'phím Delete để xoá. Copy vùng dữ liệu trong Excel (12 cột: Mã SP → '
                                'Thời gian làm lệnh) rồi bấm "Dán từ Excel" hoặc Ctrl+V trên bảng.',
                     bg='#f4f6f9', fg='#666', font=('Segoe UI', 9), anchor='w').pack(
                fill='x', padx=12, pady=(0, 6), before=wrap)

    # ---------------- Nhom +/- (Outline) cho sheet Don gia vat tu ----------------
    def _build_outline_bar(self, before):
        ob = tk.Frame(self, bg='#f4f6f9')
        ob.pack(fill='x', padx=12, pady=(0, 6), before=before)
        tk.Label(ob, text='Cấp nhóm:', bg='#f4f6f9', font=FONT_B).pack(side='left')
        for n in range(1, 6):
            tk.Button(ob, text=str(n), font=('Segoe UI', 9, 'bold'), width=2, relief='groove',
                      command=lambda k=n: self._show_level(k)).pack(side='left', padx=1)
        tk.Button(ob, text='⊞ Mở hết (+)', font=FONT, relief='groove',
                  command=lambda: self._show_level(5)).pack(side='left', padx=(12, 2))
        tk.Button(ob, text='⊟ Thu gọn hết (−)', font=FONT, relief='groove',
                  command=lambda: self._show_level(1)).pack(side='left', padx=2)
        tk.Label(ob, text='(bấm đúp / bấm mũi tên ở cột Nhóm để mở/đóng từng nhóm — '
                          'giống nút + / − của Excel)',
                 bg='#f4f6f9', fg='#666', font=('Segoe UI', 9)).pack(side='left', padx=10)

    @staticmethod
    def _sign(open_):
        return '[−] ' if open_ else '[+] '

    def _set_node_text(self, iid, open_):
        t = self.tree.item(iid, 'text') or ''
        if t.startswith('[−] ') or t.startswith('[+] '):
            t = t[4:]
        self.tree.item(iid, text=self._sign(open_) + t)

    def _on_toggle(self, open_):
        iid = self.tree.focus()
        if not iid or not self.tree.get_children(iid):
            return
        self._set_node_text(iid, open_)
        idx = self._iid_idx.get(iid)
        if idx is not None:
            (self._closed.discard(idx) if open_ else self._closed.add(idx))

    def _show_level(self, n):
        """Giong nut 1..5 cua Excel: hien den cap n (nhom cap < n duoc mo)."""
        self._closed = set()
        for iid, idx in self._iid_idx.items():
            if not self.tree.get_children(iid):
                continue
            lv = self._levels[idx]
            op = lv < n
            self.tree.item(iid, open=op)
            self._set_node_text(iid, op)
            if not op:
                self._closed.add(idx)

    def _fill_grouped(self, q):
        rows = self.sheet['rows']
        self._levels = group_levels(rows)
        # khi tim kiem: chi giu dong khop + cac nhom cha cua no
        keep = None
        if q:
            keep = set()
            stack = {}
            for idx, r in enumerate(rows):
                lv = self._levels[idx]
                if is_header_row(r):
                    stack = {k: v for k, v in stack.items() if k < lv}
                    stack[lv] = idx
                    if q in str(r[3]).lower():
                        keep.add(idx)
                        keep.update(stack.values())
                elif q in ' '.join(str(x) for x in r).lower():
                    keep.add(idx)
                    keep.update(stack.values())
        parents = {}          # level -> iid
        n = 0
        hdr_iids = []
        for idx, r in enumerate(rows):
            lv = self._levels[idx]
            if keep is not None and idx not in keep:
                if is_header_row(r):
                    parents = {k: v for k, v in parents.items() if k < lv}
                continue
            if is_header_row(r):
                parents = {k: v for k, v in parents.items() if k < lv}
                par = parents[max(parents)] if parents else ''
                op = bool(q) or idx not in self._closed
                iid = self.tree.insert(par, 'end', text=str(r[3]).strip(), values=[''] * len(r),
                                       open=op, tags=(f'g{min(lv, 4)}',))
                parents[lv] = iid
                hdr_iids.append((iid, op))
            else:
                par = parents[max(parents)] if parents else ''
                tag = 'odd' if n % 2 else ''
                iid = self.tree.insert(par, 'end', text='', values=self._disp(r), tags=(tag,))
                n += 1
            self._iid_idx[iid] = idx
        for iid, op in hdr_iids:
            if self.tree.get_children(iid):
                self._set_node_text(iid, op)
        self.count.config(text=f'{n} vật tư')

    # ---------------- Can le: ten ngan hon o -> can giua (giong tieu de) ----------------
    CELL_PAD = 6          # le trong o (px) moi ben

    def _font(self):
        if not hasattr(self, '_cfont'):
            f = ttk.Style(self).lookup('Treeview', 'font') or 'TkDefaultFont'
            try:
                self._cfont = tkfont.nametofont(f)
            except Exception:
                self._cfont = tkfont.Font(font=f)
            self._space_w = max(1, self._cfont.measure(' '))
        return self._cfont

    def _tw(self, t):
        w = self._mcache.get(t)
        if w is None:
            w = self._font().measure(t)
            self._mcache[t] = w
        return w

    def _compute_align(self):
        """Moi cot: neu MOI gia tri deu ngan hon do rong o -> anchor center (giua dung nhu tieu de).
        Cot co gia tri dai hon o -> anchor trai; cac gia tri ngan trong cot do se duoc
        chen khoang trang de van nam giua o (xem _disp)."""
        rows = self.sheet['rows']
        self._colw = [int(self.tree.column(c, 'width')) for c in self._cols]
        self._center = []
        for j, c in enumerate(self._cols):
            room = self._colw[j] - 2 * self.CELL_PAD
            fits = True
            for r in rows:
                if j < len(r):
                    t = _cell_disp(r[j])
                    if t and self._tw(t) > room:
                        fits = False
                        break
            self._center.append(fits)
            self.tree.column(c, anchor=('center' if fits else 'w'))

    def _disp(self, r):
        """Gia tri hien thi cua 1 dong: o ngan (trong cot can trai) duoc dem khoang trang de nam giua."""
        out = []
        for j in range(len(self._cols)):
            t = _cell_disp(r[j]) if j < len(r) else ''
            if t and j < len(self._center) and not self._center[j]:
                w = self._colw[j]
                tw = self._tw(t)
                if tw <= w - 2 * self.CELL_PAD:
                    lead = (w - tw) / 2.0 - 4          # 4px = le trai mac dinh cua Treeview
                    n = int(round(lead / self._space_w))
                    if n > 0:
                        t = ' ' * n + t
            out.append(t)
        return out

    def _check_resize(self):
        try:
            cur = [int(self.tree.column(c, 'width')) for c in self._cols]
        except Exception:
            return
        if cur != self._colw:
            sel = self._selected_index()
            y = self.tree.yview()[0]
            self._fill()
            self.tree.yview_moveto(y)
            if sel is not None:
                self._select_idx(sel, see=False)

    def _select_idx(self, idx, see=True):
        """Chon (va cuon toi) dong co chi so that idx."""
        for iid, i in self._iid_idx.items():
            if i == idx:
                par = self.tree.parent(iid)
                while par:
                    self.tree.item(par, open=True)
                    self._set_node_text(par, True)
                    self._closed.discard(self._iid_idx.get(par))
                    par = self.tree.parent(par)
                self.tree.selection_set(iid)
                self.tree.focus(iid)
                if see:
                    self.tree.see(iid)
                return True
        return False

    def _fill_now(self):
        q = self.search.get().strip().lower()
        self.tree.delete(*self.tree.get_children())
        self._compute_align()
        if self.grouped:
            self._iid_idx = {}
            self._fill_grouped(q)
            return
        self._iid_idx = {}
        n = 0
        for idx, r in enumerate(self.sheet['rows']):
            if q and q not in ' '.join(str(x) for x in r).lower():
                continue
            tag = 'odd' if n % 2 else ''
            iid = self.tree.insert('', 'end', values=self._disp(r), tags=(tag,))
            self._iid_idx[iid] = idx
            n += 1
        self.count.config(text=f'{n} dòng')

    # ---------------- Them/Sua/Xoa dong bang thuong ----------------
    def _selected_index(self):
        sel = self.tree.selection()
        if not sel:
            return None
        return self._iid_idx.get(sel[0])

    def _persist_sheets(self):
        try:
            save_sheets_file(self.app.sheets)
            if self.sheet.get('title') == 'Container Volume' and hasattr(self.app, 'rebuild_cv_index'):
                self.app.rebuild_cv_index()
            return True
        except Exception as e:
            messagebox.showerror('Lỗi lưu', f'Không lưu được dữ liệu sheet:\n{e}', parent=self)
            return False

    # --- rieng sheet Unit price: STT tu danh, Ma hoa khong trung ---
    UP_NOTE = ('• STT tự đánh lại theo thứ tự từ trên xuống (chèn thêm thì các dòng dưới tự lùi số).\n'
               '• "Mã hoá đưa vào SX" KHÔNG được trùng với mã đã có (không phân biệt hoa/thường).\n'
               '• Để trống cả "Mã hoá đưa vào SX" và "Mã vật tư" = dòng tiêu đề nhóm.')

    def _fill_cot6(self, vals):
        """Machine: 'Cot 6' = phan cuoi so Don Hang ('PO/24018/07.KSP' -> '07.KSP')."""
        i = getattr(self, 'cot6_idx', None)
        hd = self.sheet['headers']
        if i is None or 'Đơn Hàng' not in hd:
            return
        don = str(vals[hd.index('Đơn Hàng')] or '').strip() if len(vals) > hd.index('Đơn Hàng') else ''
        while len(vals) <= i:
            vals.append('')
        vals[i] = don.rsplit('/', 1)[-1].strip() if don else ''

    # --- rieng sheet NVL thay the (Replace): STT tu tang, go Ma hoa -> tu hien CODE ---
    # cot: 0 STT | 1 Ma hoa goc | 2 CODE goc | 3 Ma hoa thay the | 4 CODE thay the | 5 Ty le | 6 Uu tien
    def _is_replace(self):
        return self.sheet.get('title') == REPLACE_TITLE

    def _rp_next_stt(self):
        """STT lon nhat dang co + 1."""
        mx = 0
        for r in self.sheet.get('rows') or []:
            try:
                mx = max(mx, int(float(str(r[0]).strip().replace(',', '.'))))
            except (TypeError, ValueError, IndexError):
                pass
        return mx + 1

    def _rp_codes(self, ten):
        """Ten NVL (Ma hoa dua vao SX) -> danh sach Ma vat tu trong sheet Don gia vat tu."""
        try:
            sheets = self.app.sheets if self.app else []
        except Exception:
            sheets = []
        key = ' '.join(str(ten or '').split()).lower()
        out = []
        for sh in sheets:
            if sh.get('title') != UNIT_PRICE_TITLE:
                continue
            for r in sh.get('rows') or []:
                if len(r) <= UP_COL_MAVT:
                    continue
                if ' '.join(_cell_str(r[UP_COL_CODE]).split()).lower() == key:
                    c = _cell_disp(r[UP_COL_MAVT])
                    if c.endswith('.0') and c[:-2].isdigit():
                        c = c[:-2]
                    if c and c not in out:
                        out.append(c)
        return out

    def _rp_dialog_opts(self, adding):
        opts = {'code_lookup': self._rp_codes, 'code_pairs': [(1, 2), (3, 4)],
                'note': ('• STT tự tăng thêm 1 (so với số lớn nhất đang có).\n'
                         '• Gõ Mã hoá gốc / Mã hoá thay thế → CODE tự lấy theo tab Đơn giá vật tư '
                         '(tên có nhiều mã thì bấm ▼ để chọn).')}
        if adding:
            opts['defaults'] = {0: self._rp_next_stt()}
            opts['keep_on_load'] = {0}
        return opts

    def _up_dialog_opts(self):
        if getattr(self, 'cot6_idx', None) is not None:
            return {'readonly_cols': {self.cot6_idx: '(tự lấy từ Đơn Hàng)'}}
        if not self.is_unit_price:
            return {}
        return {'readonly_cols': {UP_COL_STT: '(tự đánh số)'}, 'note': self.UP_NOTE}

    def _up_check_code(self, vals, skip_idx, parent):
        """True neu Ma hoa hop le (khong trung). Sai -> bao loi, tra False."""
        code = vals[UP_COL_CODE] if len(vals) > UP_COL_CODE else ''
        dup = find_dup_code(self.sheet['rows'], code, skip_idx)
        if dup is None:
            return True
        _, r = dup
        stt = _cell_str(r[UP_COL_STT]) or '?'
        ten = _cell_str(r[UP_COL_TEN]) if len(r) > UP_COL_TEN else ''
        mavt = _cell_str(r[UP_COL_MAVT]) if len(r) > UP_COL_MAVT else ''
        messagebox.showerror(
            'Trùng mã',
            f"Mã hoá đưa vào SX '{_cell_str(code)}' đã tồn tại:\n\n"
            f"  STT {stt}  |  Mã vật tư {mavt}  |  {ten}\n\n"
            'Hãy nhập mã khác (không được trùng với mã cũ).', parent=parent)
        return False

    def _after_rows_changed(self):
        if self.is_unit_price:
            renumber_unit_price(self.sheet['rows'])
        self.src_lbl.config(text=f"(nguồn: {self.sheet['source']} — {len(self.sheet['rows'])} dòng)")

    def _row_add(self):
        headers = self.sheet['headers']

        pos = self._selected_index() if self.grouped else None
        dlg = []

        def save(vals):
            self._fill_cot6(vals)
            if self.is_unit_price:
                vals[UP_COL_STT] = ''
                if not self._up_check_code(vals, None, dlg[0] if dlg else self):
                    return False
            if self._is_replace() and vals:
                t = str(vals[0]).strip()
                if not t:
                    vals[0] = self._rp_next_stt()
                elif t.isdigit():
                    vals[0] = int(t)
            if pos is not None:
                new_idx = pos + 1
                self.sheet['rows'].insert(new_idx, vals)   # chen vao dung nhom dang chon
            else:
                new_idx = len(self.sheet['rows'])
                self.sheet['rows'].append(vals)
            self._after_rows_changed()
            if self._persist_sheets():
                self._fill()
                self._select_idx(new_idx)
        opts = self._rp_dialog_opts(True) if self._is_replace() else self._up_dialog_opts()
        dlg.append(RowEditor(self, f"Thêm dòng — {self.sheet['title']}", headers, None, save,
                             source_rows=self.sheet['rows'], **opts))

    def _row_edit(self):
        idx = self._selected_index()
        if idx is None:
            messagebox.showwarning('Chưa chọn', 'Hãy chọn 1 dòng để sửa.', parent=self)
            return
        headers = self.sheet['headers']
        cur = self.sheet['rows'][idx]
        dlg = []

        def save(vals):
            self._fill_cot6(vals)
            if self.is_unit_price:
                vals[UP_COL_STT] = cur[UP_COL_STT] if cur else ''
                old = cur[UP_COL_CODE] if len(cur) > UP_COL_CODE else ''
                # chi kiem tra khi Ma hoa bi doi (giu nguyen ma cu thi cho luu)
                if norm_code(vals[UP_COL_CODE]) != norm_code(old):
                    if not self._up_check_code(vals, idx, dlg[0] if dlg else self):
                        return False
            self.sheet['rows'][idx] = vals
            self._after_rows_changed()
            if self._persist_sheets():
                self._fill()
                self._select_idx(idx)
        opts = self._rp_dialog_opts(False) if self._is_replace() else self._up_dialog_opts()
        dlg.append(RowEditor(self, f"Sửa dòng — {self.sheet['title']}", headers, cur, save, **opts))

    def _row_del(self):
        """Xoa CAC dong dang chon (chon nhieu bang Ctrl / Shift, Ctrl+A = tat ca dong dang hien)."""
        idxs = sorted({self._iid_idx[i] for i in self.tree.selection() if i in self._iid_idx})
        if not idxs:
            messagebox.showwarning('Chưa chọn', 'Hãy chọn dòng cần xóa\n'
                                   '(giữ Ctrl hoặc Shift để chọn nhiều dòng).', parent=self)
            return
        rows = self.sheet['rows']

        def pv(r):
            t = ' | '.join(str(x) for x in r if str(x).strip())
            return (t[:90] + '…') if len(t) > 90 else (t or '(dòng trống)')
        preview = '\n'.join('• ' + pv(rows[i]) for i in idxs[:8])
        if len(idxs) > 8:
            preview += f'\n… và {len(idxs) - 8} dòng nữa'
        if not messagebox.askyesno('Xác nhận xóa', f'Xóa {len(idxs)} dòng đã chọn?\n\n{preview}\n\n'
                                   'Thao tác này KHÔNG hoàn tác được.', icon='warning', parent=self):
            return
        backup = list(rows)
        for i in reversed(idxs):                 # xoa tu duoi len -> chi so phia tren khong bi lech
            del rows[i]
        self._after_rows_changed()
        if not self._persist_sheets():
            rows[:] = backup                     # luu loi -> tra lai nhu cu
            self._after_rows_changed()
            return
        top = self.tree.yview()[0]
        self._fill()
        self.tree.yview_moveto(top)              # giu nguyen vi tri dang xem
        if rows:
            self._select_idx(min(idxs[0], len(rows) - 1), see=False)
        messagebox.showinfo('Đã xóa', f'Đã xóa {len(idxs)} dòng.', parent=self)

    # ---------------- Unit price: nhan dien ton kho tu anh / file Excel ----------------
    def _stock_import(self):
        try:
            from stock_ui import StockImportDialog
        except Exception as e:
            messagebox.showerror('Lỗi', f'Không nạp được stock_ui.py:\n{e}', parent=self)
            return

        def saved():
            ok = self._persist_sheets()
            if ok:
                self._fill()
            return ok
        dlg = StockImportDialog(self, self.sheet['rows'], saved, sheet=self.sheet)
        self.wait_window(dlg)
        self._update_stock_label()

    def _update_stock_label(self):
        if getattr(self, 'stock_lbl', None) is not None:
            import stock_import
            stale = bool(stock_import.stock_lock_reason(self.sheet))
            self.stock_lbl.config(text=stock_import.last_update_text(self.sheet)
                                  + ('   ⚠ QUÁ 5 NGÀY — đã khoá in/xuất lệnh' if stale else ''),
                                  fg=('#b00020' if stale else '#6a4c93'))
        # mo khoa / khoa lai cac nut o tab Lenh san xuat ngay sau khi nhan dien ton kho
        ot = getattr(self.app, 'order_tab', None) if getattr(self, 'app', None) else None
        if ot is not None:
            try:
                ot._refresh_stock_lock(repeat=False)
            except Exception:
                pass

    # ---------------- ADMIN: xoa nhieu dong / dan tu Excel (tab Data) ----------------
    def _is_admin(self):
        return bool(self.app) and self.app.user.get('tabs') == 'all'

    def _dl_sel_count(self):
        n = len(self.tree.selection())
        base = self.count.cget('text').split('  •')[0]
        self.count.config(text=base + (f'  •  đang chọn {n} dòng' if n > 1 else ''))

    def _dl_save_all(self):
        """Luu data.json + sheets.json, cap nhat tab Lenh san xuat."""
        try:
            save_data_file(self.app.data)
        except Exception as e:
            messagebox.showerror('Lỗi lưu', f'Không lưu được dữ liệu lệnh sản xuất:\n{e}', parent=self)
            return False
        if not self._persist_sheets():
            return False
        if getattr(self.app, 'order_tab', None):
            try:
                self.app.order_tab._reload_orders()
            except Exception:
                pass
        return True

    def _dl_delete_selected(self):
        import dl_tools as dt
        idxs = sorted({self._iid_idx[i] for i in self.tree.selection() if i in self._iid_idx})
        if not idxs:
            messagebox.showwarning('Chưa chọn', 'Hãy chọn các dòng cần xoá\n'
                                   '(giữ Ctrl hoặc Shift để chọn nhiều, Ctrl+A = tất cả).', parent=self)
            return
        rows = self.sheet['rows']
        new_rows, changed, removed = dt.delete_rows(rows, idxs)
        n_del = len(rows) - len(new_rows)
        if not changed and not removed:
            messagebox.showinfo('Không có gì để xoá',
                                'Chỉ chọn dòng "Tổng" — dòng Tổng được tính tự động, không xoá riêng.\n'
                                'Muốn xoá cả lệnh: chọn hết các dòng của lệnh đó.', parent=self)
            return
        msg = f'Xoá {len(idxs)} dòng đã chọn?\n'
        if removed:
            ds = ', '.join(sorted(removed, key=lambda x: (len(x), x))[:15])
            msg += f'\n• XOÁ HẲN {len(removed)} lệnh (Mã SP: {ds}{"..." if len(removed) > 15 else ""})'
        if changed:
            ds = ', '.join(sorted(changed, key=lambda x: (len(x), x))[:15])
            msg += f'\n• Bớt NVL ở {len(changed)} lệnh (Mã SP: {ds}) — dòng Tổng tự tính lại'
        msg += '\n\nThao tác này KHÔNG hoàn tác được.'
        if not messagebox.askyesno('Xác nhận xoá', msg, icon='warning', parent=self):
            return
        rows[:] = new_rows
        orders = self.app.data['orders'].setdefault(self.edit_kind, [])
        if removed:
            dt.remove_orders(orders, removed)
        for b in changed.values():
            dt.update_order_materials(orders, b)
        self._after_rows_changed()
        if self._dl_save_all():
            self._fill()
            messagebox.showinfo('Đã xoá', f'Đã xoá {len(idxs)} dòng'
                                + (f', {len(removed)} lệnh' if removed else '') + '.', parent=self)

    def _dl_paste(self):
        import dl_tools as dt
        try:
            text = self.clipboard_get()
        except Exception:
            text = ''
        if not text.strip():
            messagebox.showinfo('Clipboard trống',
                                'Trong Excel: bôi đen các dòng dữ liệu (12 cột: Mã SP, Đơn hàng, Mã hàng, '
                                'Khách hàng, Mã sản phẩm, Kí hiệu NVL, Mã NVL, Tên NVL, KL/mẻ, Tỉ lệ %, '
                                'Loại, Thời gian) → Ctrl+C → quay lại bấm "Dán từ Excel".', parent=self)
            return
        blocks = [b for b in dt.parse_paste(text, self.sheet.get('headers')) if b.rows]
        if not blocks:
            messagebox.showwarning('Không đọc được', 'Không tìm thấy dòng dữ liệu hợp lệ trong clipboard.',
                                   parent=self)
            return
        ncols = max(sum(1 for c in r if str(c).strip()) for b in blocks for r in b.rows)
        if ncols < 6:
            if not messagebox.askyesno(
                    'Dữ liệu ít cột',
                    f'Dữ liệu dán chỉ có tối đa {ncols} cột có giá trị (bảng Data cần 12 cột theo thứ tự:\n'
                    + ' | '.join(self.sheet['headers']) + ').\n\nVẫn tiếp tục?', parent=self):
                return
        PastePreview(self, blocks, self._dl_apply_paste)

    def _dl_apply_paste(self, blocks, mode):
        """mode: 'overwrite' = Ma SP trung thi ghi de; 'new' = Ma SP trung thi danh so moi."""
        import dl_tools as dt
        rows = self.sheet['rows']
        orders = self.app.data['orders'].setdefault(self.edit_kind, [])
        existing = {str(r[0]).strip() for r in rows if not dt.is_tong(r)}
        existing |= {str(o.get('ma_sp')).strip() for o in orders}
        n_new = n_over = 0
        for b in blocks:
            if not b.ma_sp or (b.ma_sp in existing and mode == 'new'):
                dt.set_block_ma_sp(b, dt.next_ma_sp(orders, rows))
            if not any(str(r[dt.C_TG]).strip() for r in b.rows):
                b.rows[0][dt.C_TG] = dt.now_str()
            if b.ma_sp in existing:
                n_over += 1
            else:
                n_new += 1
            dt.replace_block_in_rows(rows, b)
            dt.upsert_order(orders, b)
            existing.add(b.ma_sp)
        self._after_rows_changed()
        if self._dl_save_all():
            self._fill()
            try:
                self.tree.yview_moveto(1.0)
            except Exception:
                pass
            messagebox.showinfo('Đã dán', f'Đã thêm {n_new} lệnh mới'
                                + (f', ghi đè {n_over} lệnh' if n_over else '') + '.', parent=self)

    def _add_order(self):
        OrderEditor(self, self.app, self.edit_kind, self, mode='add')

    def _edit_order(self):
        OrderEditor(self, self.app, self.edit_kind, self, mode='edit')

    def _order_block_rows(self, order):
        rows = [m for m in order.get('materials', []) if (m.get('ten') or '').strip()]
        tong_mass = sum(float(m.get('mass_batch') or 0) for m in rows)
        tong_pct = sum(float(m.get('percentage') or 0) for m in rows)
        tg = order.get('thoi_gian') or ''
        msp = order.get('ma_san_pham') or ''
        out = []
        for m in rows:
            out.append([
                order.get('ma_sp'), order.get('don_hang'), order.get('ma_hang'),
                order.get('khach_hang'), msp, m.get('ten'), m.get('ma_nvl') or '',
                m.get('ten_day_du') or '', m.get('mass_batch') or 0,
                round(float(m.get('percentage') or 0) * 100, 3), order.get('loai') or '', tg])
        out.append(['Tổng', '', '', '', '', '', '', '',
                    round(tong_mass, 2), round(tong_pct * 100, 2), '', ''])
        return out

    def append_order_rows(self, order):
        """Them cac dong cua 1 don moi vao cuoi bang DL va cuon xuong."""
        self.sheet['rows'].extend(self._order_block_rows(order))
        self.src_lbl.config(text=f"(nguồn: {self.sheet['source']} — {len(self.sheet['rows'])} dòng)")
        self._persist_sheets()
        self._fill()
        try:
            self.tree.yview_moveto(1.0)
        except Exception:
            pass

    def replace_order_rows(self, order):
        """Sua don cu: xoa cac dong cu cua ma_sp roi chen block moi tai cho."""
        masp = str(order.get('ma_sp'))
        rows = self.sheet['rows']
        # tim vi tri dau tien cua ma_sp
        start = None
        for i, r in enumerate(rows):
            if str(r[0]) == masp:
                start = i
                break
        new_block = self._order_block_rows(order)
        if start is None:
            rows.extend(new_block)
        else:
            # xoa lien tiep cac dong ma_sp + dong 'Tong' ngay sau
            end = start
            while end < len(rows) and str(rows[end][0]) == masp:
                end += 1
            if end < len(rows) and str(rows[end][0]) == 'Tổng':
                end += 1
            rows[start:end] = new_block
        self.src_lbl.config(text=f"(nguồn: {self.sheet['source']} — {len(rows)} dòng)")
        self._persist_sheets()
        self._fill()


class PastePreview(tk.Toplevel):
    """Xem truoc du lieu dan tu Excel vao tab Data truoc khi ghi."""

    def __init__(self, tab, blocks, on_ok):
        import dl_tools as dt
        super().__init__(tab)
        self.blocks = blocks
        self.on_ok = on_ok
        self.title('Xem trước dữ liệu dán từ Excel')
        self.geometry('1180x560')
        self.configure(bg='#f4f6f9')
        self.transient(tab.winfo_toplevel())
        rows = tab.sheet['rows']
        orders = tab.app.data['orders'].get(tab.edit_kind, [])
        existing = {str(r[0]).strip() for r in rows if not dt.is_tong(r)} | \
                   {str(o.get('ma_sp')).strip() for o in orders}
        dup = [b.ma_sp for b in blocks if b.ma_sp and b.ma_sp in existing]
        nrow = sum(len(b.rows) for b in blocks)
        tk.Label(self, text=f'{len(blocks)} lệnh — {nrow} dòng NVL sẽ được thêm vào {tab.sheet["title"]}',
                 bg='#f4f6f9', font=('Segoe UI', 11, 'bold'), fg='#1d3557').pack(anchor='w', padx=12, pady=(10, 4))
        catalog = NvlCatalog.from_sheets(tab.app.sheets)
        wrap = tk.Frame(self, bg='#f4f6f9')
        wrap.pack(fill='both', expand=True, padx=12)
        headers = list(tab.sheet['headers']) + ['Kiểm tra']
        cols = [f'c{i}' for i in range(len(headers))]
        tv = ttk.Treeview(wrap, columns=cols, show='headings', height=16)
        for c, h in zip(cols, headers):
            tv.heading(c, text=h)
            tv.column(c, width=(70 if h in ('Mã SP', 'KL/mẻ (kg)', 'Tỉ lệ %') else 110), anchor='w')
        tv.column(cols[-1], width=220)
        tv.tag_configure('tong', background='#dbe5f1', font=('Segoe UI', 9, 'bold'))
        tv.tag_configure('bad', background='#ffd6d6')
        tv.tag_configure('dup', background='#fff3cd')
        n_bad_pct = 0
        for b in blocks:
            for r in b.rows:
                ten, ma = str(r[dt.C_TEN] or '').strip(), str(r[dt.C_MANVL] or '').strip()
                chk, tag = '', ''
                if ten and not catalog.codes_for(ten):
                    chk, tag = 'Tên NVL không có trong Unit price', 'bad'
                elif ten and ma and ma not in catalog.codes_for(ten):
                    chk, tag = f'Mã đúng: {"/".join(catalog.codes_for(ten))}', 'bad'
                elif b.ma_sp in dup:
                    tag = 'dup'
                tv.insert('', 'end', values=[_cell_str(x) for x in r] + [chk], tags=(tag,))
            t = dt.block_pct_total(b)
            ok = abs(t - 100) < 0.01
            n_bad_pct += 0 if ok else 1
            tr = b.tong_row()
            tv.insert('', 'end', values=[_cell_str(x) for x in tr] +
                      [('✔ 100%' if ok else f'✖ Tổng % = {t:.4f}')], tags=('tong',))
        vs = ttk.Scrollbar(wrap, orient='vertical', command=tv.yview)
        hs = ttk.Scrollbar(wrap, orient='horizontal', command=tv.xview)
        tv.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        tv.grid(row=0, column=0, sticky='nsew')
        vs.grid(row=0, column=1, sticky='ns')
        hs.grid(row=1, column=0, sticky='ew')
        wrap.grid_rowconfigure(0, weight=1)
        wrap.grid_columnconfigure(0, weight=1)
        notes = []
        if dup:
            notes.append(f'⚠ {len(dup)} Mã SP đã tồn tại (tô vàng): {", ".join(dup[:12])}'
                         + ('...' if len(dup) > 12 else ''))
        if n_bad_pct:
            notes.append(f'⚠ {n_bad_pct} lệnh có Tổng % khác 100%.')
        notes.append('Dòng Tổng được tính lại tự động. Tỉ lệ % hiểu theo dạng 83.5 (= 83,5%); '
                     'nếu Excel lưu 0,835 thì tự nhân 100.')
        tk.Label(self, text='\n'.join(notes), bg='#fff8e1', fg='#6d4c00', font=('Segoe UI', 9),
                 justify='left', anchor='w').pack(fill='x', padx=12, pady=6)
        self.mode = tk.StringVar(value='overwrite')
        bot = tk.Frame(self, bg='#e9edf2')
        bot.pack(fill='x', side='bottom')
        if dup:
            tk.Radiobutton(bot, text='Mã SP trùng: GHI ĐÈ lệnh cũ', variable=self.mode, value='overwrite',
                           bg='#e9edf2', font=FONT).pack(side='left', padx=(12, 4), pady=8)
            tk.Radiobutton(bot, text='Mã SP trùng: đánh Mã SP MỚI', variable=self.mode, value='new',
                           bg='#e9edf2', font=FONT).pack(side='left', padx=4)
        tk.Button(bot, text='✔ Ghi vào bảng', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat',
                  padx=14, pady=6, command=self._ok).pack(side='right', padx=12, pady=6)
        tk.Button(bot, text='Huỷ', font=FONT, relief='flat', padx=12, pady=6,
                  command=self.destroy).pack(side='right')
        self.grab_set()

    def _ok(self):
        mode = self.mode.get()
        self.destroy()
        self.on_ok(self.blocks, mode)


# =====================================================================
#  TAB TAI LIEU (CLAIM/TDS/COA/MSDS) - hien thi dang FORM + xuat PDF
# =====================================================================
class CellEditDialog(tk.Toplevel):
    """ADMIN sua noi dung 1 o cua TDS / COA / MSDS."""

    def __init__(self, master, doc_title, r, c, text, is_override, is_auto, on_save, on_reset):
        super().__init__(master)
        from openpyxl.utils import get_column_letter
        self.on_save, self.on_reset = on_save, on_reset
        addr = f'{get_column_letter(c)}{r}'
        self.title(f'Sửa ô {addr} — {doc_title}')
        self.configure(bg='#f4f6f9')
        self.transient(master)
        self.geometry('640x330')
        self.minsize(460, 260)
        tk.Label(self, text=f'✏  {doc_title}  •  Ô {addr}', bg='#f4f6f9', font=FONT_TITLE,
                 fg='#1f3a5f').pack(anchor='w', padx=14, pady=(10, 2))
        note = 'Xuống dòng bằng Enter. Ctrl+Enter = Lưu.'
        if is_auto:
            note += ('\n⚠ Ô này đang TỰ TÍNH từ thông số (Mã hàng / % bột đá / MFI / lot / ngày). '
                     'Sửa tay sẽ GIỮ CỐ ĐỊNH nội dung mới — bấm "Trả về tự tính" để tự tính lại.')
        if is_override:
            note += '\n• Ô này đã được sửa tay trước đó.'
        tk.Label(self, text=note, bg='#f4f6f9', fg='#8a5a00' if is_auto else '#555',
                 font=('Segoe UI', 9), justify='left', wraplength=600).pack(anchor='w', padx=14)
        self.txt = tk.Text(self, font=FONT, height=6, wrap='word', bg='#fffde7')
        self.txt.pack(fill='both', expand=True, padx=14, pady=8)
        self.txt.insert('1.0', text or '')
        self.txt.focus_set()
        self.txt.bind('<Control-Return>', lambda e: (self._save(), 'break')[1])
        bf = tk.Frame(self, bg='#f4f6f9')
        bf.pack(fill='x', padx=14, pady=(0, 12))
        tk.Button(bf, text='💾 Lưu', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat',
                  padx=16, pady=5, command=self._save).pack(side='right')
        tk.Button(bf, text='Hủy', font=FONT, relief='flat', padx=12, pady=5,
                  command=self.destroy).pack(side='right', padx=6)
        if is_override:
            tk.Button(bf, text='↺ Trả về ' + ('tự tính' if is_auto else 'nội dung gốc'), font=FONT,
                      relief='groove', padx=10, pady=4, command=self._reset).pack(side='left')
        self.bind('<Escape>', lambda e: self.destroy())
        self.grab_set()

    def _save(self):
        t = self.txt.get('1.0', 'end').rstrip('\n')
        self.destroy()
        self.on_save(t)

    def _reset(self):
        self.destroy()
        self.on_reset()


class DocTab(LayoutRenderMixin, tk.Frame):
    def __init__(self, master, title, layout_key, app=None):
        super().__init__(master, bg='#e9edf2')
        self.doc_title = title
        self.layout_key = layout_key
        self.app = app
        self.layout = None
        bar = tk.Frame(self, bg='#e9edf2')
        bar.pack(fill='x')
        tk.Label(bar, text=title, bg='#e9edf2', font=FONT_TITLE).pack(side='left', padx=12, pady=6)
        self._off = (0, 0)
        self._cell_rects = []
        self.edit_mode = tk.BooleanVar(value=False)
        tk.Button(bar, text='📄 Xuất PDF', font=FONT_B, bg='#b0413e', fg='white', relief='flat',
                  padx=12, pady=5, command=self._pdf).pack(side='right', padx=(4, 10), pady=6)
        if layout_key in ('tds', 'coa', 'msds'):
            tk.Button(bar, text='📊 Xuất Excel', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat',
                      padx=12, pady=5, command=self._excel).pack(side='right', padx=4, pady=6)
            tk.Button(bar, text='🖨️ In', font=FONT_B, bg='#1d3557', fg='white', relief='flat',
                      padx=14, pady=5, command=self._print).pack(side='right', padx=4, pady=6)
            if self._is_admin():
                tk.Button(bar, text='↺ Bỏ mọi sửa tay', font=FONT, relief='groove', padx=8,
                          command=self._clear_all_overrides).pack(side='right', padx=(4, 12), pady=6)
                tk.Checkbutton(bar, text='✏ Sửa trực tiếp (Admin)', variable=self.edit_mode,
                               font=FONT_B, bg='#ffe8cc', activebackground='#ffd8a8', relief='groove',
                               padx=8, command=self._toggle_edit).pack(side='right', padx=4, pady=6)
                self.lbl_edit = tk.Label(bar, text='', bg='#e9edf2', fg='#b35c00', font=('Segoe UI', 9))
                self.lbl_edit.pack(side='left', padx=6)
        self._build_editor()
        wrap = tk.Frame(self, bg='#e9edf2')
        wrap.pack(fill='both', expand=True)
        self.cv = tk.Canvas(wrap, bg='white', highlightthickness=0)
        vs = ttk.Scrollbar(wrap, orient='vertical', command=self.cv.yview)
        hs = ttk.Scrollbar(self, orient='horizontal', command=self.cv.xview)
        self.cv.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        vs.pack(side='right', fill='y')
        self.cv.pack(side='left', fill='both', expand=True)
        hs.pack(fill='x')
        self.cv.bind('<Enter>', lambda e: self.cv.bind_all(
            '<MouseWheel>', lambda ev: self.cv.yview_scroll(int(-ev.delta / 120), 'units')))
        self.cv.bind('<Leave>', lambda e: self.cv.unbind_all('<MouseWheel>'))
        self.cv.bind('<Button-1>', self._on_click)
        self.cv.bind('<Motion>', self._on_motion)
        self._doc_dirty = True                 # TOC DO: ve khi mo tab lan dau
        self.bind('<Map>', self._doc_on_map, add='+')

    def _doc_on_map(self, ev=None):
        if ev is not None and ev.widget is not self:
            return
        if getattr(self, '_doc_dirty', False):
            self._doc_dirty = False
            if not getattr(self, '_doc_rendered', False):
                self._doc_rendered = True
                self._render()                 # lan dau: giong ban cu (__init__ goi _render)
            else:
                self.refresh(force=True)

    # ---------------- ADMIN: sua truc tiep bat ky o nao (TDS / COA / MSDS) ----------------
    def _is_admin(self):
        return bool(self.app) and getattr(self.app, 'user', {}).get('tabs') == 'all'

    def _toggle_edit(self):
        on = self.edit_mode.get()
        if hasattr(self, 'lbl_edit'):
            self.lbl_edit.config(text=('Bấm vào ô bất kỳ trên tài liệu để sửa  •  ô viền cam = đã sửa tay'
                                       if on else ''))
        self.cv.config(cursor='hand2' if on else '')
        self._draw_overrides()

    def _cell_at(self, ev):
        x = self.cv.canvasx(ev.x) - self._off[0]
        y = self.cv.canvasy(ev.y) - self._off[1]
        best = None
        for (x0, y0, x1, y1, r, c) in self._cell_rects:
            if x0 <= x < x1 and y0 <= y < y1:
                area = (x1 - x0) * (y1 - y0)
                if best is None or area < best[0]:
                    best = (area, r, c, x0, y0, x1, y1)
        return best

    def _on_motion(self, ev):
        if not self.edit_mode.get():
            return
        self.cv.delete('hover')
        hit = self._cell_at(ev)
        if hit:
            _a, r, c, x0, y0, x1, y1 = hit
            dx, dy = self._off
            self.cv.create_rectangle(x0 + dx, y0 + dy, x1 + dx, y1 + dy, outline='#1d6fa5', width=2,
                                     tags='hover')

    def _on_click(self, ev):
        if not self.edit_mode.get() or not self._is_admin():
            return
        hit = self._cell_at(ev)
        if not hit:
            return
        _a, r, c = hit[:3]
        lay = self._load()
        cell = next((x for x in lay['cells'] if x['r'] == r and x['c'] == c), None)
        if cell is None:
            return
        cur = self._cell_meta(cell, self._dyn())[0]
        ovs = doc_params.get_overrides(self.layout_key)
        auto = doc_params.dynamic(self.layout_key, dict(doc_params.load(), overrides={}))
        CellEditDialog(self, self.doc_title, r, c, cur, (r, c) in ovs, (r, c) in auto,
                       on_save=lambda txt: self._save_cell(r, c, txt),
                       on_reset=lambda: self._reset_cell(r, c))

    def _save_cell(self, r, c, text):
        try:
            doc_params.set_override(self.layout_key, r, c, text)
        except Exception as e:
            messagebox.showerror('Lỗi lưu', f'Không lưu được doc_params.json:\n{e}', parent=self)
            return
        self._after_override()

    def _reset_cell(self, r, c):
        doc_params.clear_override(self.layout_key, r, c)
        self._after_override()

    def _clear_all_overrides(self):
        n = len(doc_params.get_overrides(self.layout_key))
        if not n:
            messagebox.showinfo('Không có', f'{self.doc_title} chưa có ô nào sửa tay.', parent=self)
            return
        if messagebox.askyesno('Bỏ sửa tay', f'Trả {n} ô đã sửa tay của {self.doc_title} về nội dung gốc?',
                               parent=self):
            doc_params.clear_override(self.layout_key)
            self._after_override()

    def _after_override(self):
        if self.app is not None and hasattr(self.app, 'refresh_docs'):
            self.app.refresh_docs()
        else:
            self.refresh()

    def _draw_overrides(self):
        self.cv.delete('ov')
        self.cv.delete('hover')
        if not self.edit_mode.get():
            return
        ovs = doc_params.get_overrides(self.layout_key)
        dx, dy = self._off
        for (x0, y0, x1, y1, r, c) in self._cell_rects:
            if (r, c) in ovs:
                self.cv.create_rectangle(x0 + dx + 1, y0 + dy + 1, x1 + dx - 1, y1 + dy - 1,
                                         outline='#e07a00', width=2, dash=(4, 2), tags='ov')

    # ---------------- In / Xuat Excel (TDS / COA / MSDS) ----------------
    def _print(self):
        try:
            from pdf_export import render_doc_pdf  # noqa: F401
        except Exception as e:
            messagebox.showerror('Lỗi', f'Cần reportlab để in:\n{e}', parent=self)
            return
        ma = str(doc_params.load().get('ma_hang') or '')
        PrintDialog(self, self.doc_title, None, 0, settings_owner=self.app,
                    make_pdf=lambda path: __import__('pdf_export').render_doc_pdf(
                        self.layout_key, path, self._dyn()),
                    landscape=False, heading=f'In {self.doc_title}',
                    subtitle=f'Tài liệu: {self.doc_title}   •   Mã hàng: {ma}   •   Khổ A4 dọc',
                    pdf_prefix=f'{self.doc_title}_')

    def _excel(self):
        try:
            import openpyxl  # noqa: F401
            import doc_tools
        except ImportError as e:
            messagebox.showerror('Thiếu thư viện', f'Cần cài openpyxl:\n\npip install openpyxl\n\n{e}',
                                 parent=self)
            return
        ma = str(doc_params.load().get('ma_hang') or '')
        path = filedialog.asksaveasfilename(defaultextension='.xlsx',
                                            initialfile=doc_tools.doc_file_name(self.doc_title, ma, '.xlsx'),
                                            filetypes=[('Excel', '*.xlsx')], parent=self)
        if not path:
            return
        try:
            doc_tools.export_doc_excel(self.layout_key, path, self._dyn())
        except Exception as e:
            messagebox.showerror('Lỗi xuất Excel', str(e), parent=self)
            return
        if messagebox.askyesno('Xong', f'Đã xuất Excel:\n{path}\n\nMở ngay?', parent=self):
            self._open(path)

    # ---------------- bang chinh sua (TDS / COA / MSDS) ----------------
    def _build_editor(self):
        k = self.layout_key
        if k not in ('tds', 'coa', 'msds'):
            return
        p = doc_params.load()
        box = tk.LabelFrame(self, text=' ✏ Chỉnh sửa thông số ', bg='#f4f6f9', font=FONT_B,
                            fg='#1f3a5f', padx=10, pady=6)
        box.pack(fill='x', padx=10, pady=(0, 6))
        YEL = '#fff59d'
        if k == 'tds':
            self.v_ma = tk.StringVar(value=p['ma_hang'])
            self.v_bot = tk.StringVar(value=doc_params.fmt(p['bot_da']))
            self.v_mfi_t = tk.StringVar(value=p['mfi_type'])
            self.v_mfi_v = tk.StringVar(value=doc_params.fmt(p['mfi_value']))
            g = tk.Frame(box, bg='#f4f6f9')
            g.pack(fill='x')
            tk.Label(g, text='a) Mã hàng:', bg='#f4f6f9', font=FONT_B).grid(row=0, column=0, sticky='e', padx=(0, 4), pady=3)
            e1 = tk.Entry(g, textvariable=self.v_ma, width=18, font=FONT, bg=YEL, fg='#c00000')
            e1.grid(row=0, column=1, sticky='w')
            tk.Label(g, text='b) % Bột đá (CaCO3):', bg='#f4f6f9', font=FONT_B).grid(row=0, column=2, sticky='e', padx=(18, 4))
            e2 = tk.Entry(g, textvariable=self.v_bot, width=8, font=FONT, bg=YEL)
            e2.grid(row=0, column=3, sticky='w')
            tk.Label(g, text='c) Melting flow index:', bg='#f4f6f9', font=FONT_B).grid(row=1, column=0, sticky='e', padx=(0, 4), pady=3)
            cb = ttk.Combobox(g, textvariable=self.v_mfi_t, values=doc_params.MFI_OPTIONS,
                              state='readonly', font=FONT, width=34)
            cb.grid(row=1, column=1, sticky='w')
            cb.bind('<<ComboboxSelected>>', lambda e: self._apply())
            tk.Label(g, text='Chỉ số (ô vàng):', bg='#f4f6f9', font=FONT_B).grid(row=1, column=2, sticky='e', padx=(18, 4))
            e3 = tk.Entry(g, textvariable=self.v_mfi_v, width=8, font=FONT, bg=YEL)
            e3.grid(row=1, column=3, sticky='w')
            tk.Button(g, text='✔ Áp dụng', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat',
                      padx=14, pady=4, command=self._apply).grid(row=0, column=4, rowspan=2, padx=(24, 0))
            tk.Label(g, text='→ COA và MSDS tự lấy theo (Mã hàng, % bột đá, MFI) giống file Excel.',
                     bg='#f4f6f9', fg='#666', font=('Segoe UI', 8)).grid(row=2, column=0, columnspan=5, sticky='w')
            for e in (e1, e2, e3):
                e.bind('<Return>', lambda ev: self._apply())
            self.lbl_calc = tk.Label(g, text='', bg='#f4f6f9', fg='#2d6a4f', font=('Segoe UI', 9, 'bold'))
            self.lbl_calc.grid(row=3, column=0, columnspan=5, sticky='w')
            self._update_calc_label(p)
        else:
            top = tk.Frame(box, bg='#f4f6f9')
            top.pack(fill='x')
            self.lbl_link = tk.Label(top, text='', bg='#f4f6f9', fg='#1f3a5f', font=FONT)
            self.lbl_link.pack(side='left')
            body = tk.Frame(box, bg='#f4f6f9')
            body.pack(fill='x', pady=(4, 0))
            left = tk.Frame(body, bg='#f4f6f9')
            left.pack(side='left', anchor='n')
            tk.Label(left, text='Lot number\n(mỗi dòng 1 lot):', bg='#f4f6f9', font=FONT_B,
                     justify='left').pack(anchor='w')
            if k == 'msds':
                self.v_same = tk.BooleanVar(value=bool(p.get('msds_same_coa', True)))
                tk.Checkbutton(left, text='Giống lot COA', variable=self.v_same, bg='#f4f6f9',
                               font=FONT, command=self._toggle_same).pack(anchor='w', pady=(4, 0))
            self.txt_lots = tk.Text(body, width=40, height=4, font=FONT, bg='#fff59d', wrap='none')
            self.txt_lots.pack(side='left', padx=8)
            lots = p.get('coa_lots') if k == 'coa' else (
                p.get('coa_lots') if p.get('msds_same_coa', True) else p.get('msds_lots'))
            self.txt_lots.insert('1.0', '\n'.join(lots or []))
            btns = tk.Frame(body, bg='#f4f6f9')
            btns.pack(side='left', anchor='n')
            tk.Button(btns, text='✔ Áp dụng', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat',
                      padx=14, pady=4, command=self._apply).pack(fill='x')
            tk.Button(btns, text='Xóa hết lot', font=FONT, relief='flat', padx=10, pady=3,
                      command=lambda: self.txt_lots.delete('1.0', 'end')).pack(fill='x', pady=(6, 0))
            self.txt_lots.bind('<Control-Return>', lambda ev: (self._apply(), 'break')[1])
            if k == 'msds':
                self._toggle_same(init=True)
            self._update_link_label(p)

    def _update_calc_label(self, p=None):
        if not hasattr(self, 'lbl_calc'):
            return
        p = p or doc_params.load()
        t = doc_params.tds_values(p)
        grp = doc_params.tds_variant_key(t['E3']) or 'mặc định (mã không thuộc nhóm VBA)'
        dec, tol = doc_params.mfi_format(t['F19'])
        self.lbl_calc.config(text=f"Tự tính (giống VBA):  Nhóm TDS = {grp}   •   Nhựa nền = {t['D28']}   •   "
                                  f"Density = {doc_params.nf(t['F21'], 2)}   •   D50 = {doc_params.nf(t['F22'], 1)}   •   "
                                  f"MFI = {doc_params.nf(t['F19'], dec, tol)}")

    def _update_link_label(self, p=None):
        if not hasattr(self, 'lbl_link'):
            return
        p = p or doc_params.load()
        self.lbl_link.config(text=f"Lấy từ TDS:  Mã hàng = {p['ma_hang']}   •   % bột đá = "
                                  f"{doc_params.fmt(p['bot_da'])}   •   {p['mfi_type']} = "
                                  f"{doc_params.fmt(p['mfi_value'])}")

    def _toggle_same(self, init=False):
        same = self.v_same.get()
        if same:
            p = doc_params.load()
            self.txt_lots.config(state='normal')
            self.txt_lots.delete('1.0', 'end')
            self.txt_lots.insert('1.0', '\n'.join(p.get('coa_lots') or []))
            self.txt_lots.config(state='disabled', bg='#eeeeee')
        else:
            self.txt_lots.config(state='normal', bg='#fff59d')
            if not init:
                p = doc_params.load()
                self.txt_lots.delete('1.0', 'end')
                self.txt_lots.insert('1.0', '\n'.join(p.get('msds_lots') or p.get('coa_lots') or []))
        if not init:
            self._apply()

    def _apply(self):
        p = doc_params.load()
        k = self.layout_key
        if k == 'tds':
            ma = self.v_ma.get().strip()
            if not ma:
                messagebox.showwarning('Thiếu', 'Nhập Mã hàng.', parent=self)
                return
            bot = doc_params.to_num(self.v_bot.get())
            if not isinstance(bot, (int, float)) or not (0 <= bot <= 100):
                messagebox.showwarning('Sai số liệu', '% Bột đá phải là số từ 0 đến 100.', parent=self)
                return
            old_ma = str(p.get('ma_hang') or '').strip()
            p['ma_hang'] = ma
            p['bot_da'] = bot
            p['mfi_type'] = self.v_mfi_t.get()
            var = doc_params.tds_variant(ma)
            if ma != old_ma and var:          # VBA: doi E3 -> change_Worksheet ghi lai B19
                p['mfi_type'] = var['B19']
                self.v_mfi_t.set(var['B19'])
            p['mfi_value'] = doc_params.to_num(self.v_mfi_v.get())
        else:
            if k == 'msds':
                p['msds_same_coa'] = bool(self.v_same.get())
                if not p['msds_same_coa']:
                    p['msds_lots'] = doc_params.parse_lots(self.txt_lots.get('1.0', 'end'))
            else:
                p['coa_lots'] = doc_params.parse_lots(self.txt_lots.get('1.0', 'end'))
        try:
            doc_params.save(p)
        except Exception as e:
            messagebox.showerror('Lỗi lưu', f'Không lưu được doc_params.json:\n{e}', parent=self)
            return
        if self.app is not None and hasattr(self.app, 'refresh_docs'):
            self.app.refresh_docs()
        else:
            self.refresh()

    def refresh(self, force=False):
        """Ve lai tai lieu (sau khi doi thong so). Tab dang an -> danh dau, ve khi mo tab."""
        if not force and not self.winfo_ismapped():
            self._doc_dirty = True
            return
        self._doc_dirty = False
        self._doc_rendered = True
        if getattr(self, 'layout_key', '') == 'msds' and hasattr(self, 'v_same') and self.v_same.get():
            self._toggle_same(init=True)
        self._update_link_label()
        self._update_calc_label()
        self.cv.delete('all')
        self._render()

    def _load(self):
        if self.layout is None:
            try:
                self.layout = json.load(open(resource_path('layout_%s.json' % self.layout_key),
                                             encoding='utf-8'))
                self.layout = doc_params.trim_layout(self.layout, self.layout_key)
            except Exception:
                self.layout = None
        return self.layout

    def _draw_page_bg(self):
        """Ve logo / giay tieu de (header an trong Excel) phia sau noi dung, ti le 100%."""
        spec = doc_params.page_spec(self.layout_key)
        if not spec:
            return
        try:
            from PIL import Image, ImageTk
        except Exception:
            return
        k = 1 / 0.75                              # pt -> px (Excel 100%)
        bl, bt, br, bb_ = spec['box']
        # noi dung tren man hinh (chu Segoe UI) co the cao/rong hon vung in -> PHONG TO trang giay
        # (header / footer / logo chim) cho vua noi dung, tranh header-footer de len chu.
        cb = self.cv.bbox('all')
        if cb:
            cw, ch = cb[2] - 8 + 4, cb[3] - 8 + 4
            k = max(k, cw / max(1.0, br - bl), ch / max(1.0, bb_ - bt))
        # bang hep hon vung in (trang bi phong to theo chieu cao) -> CAN GIUA theo chieu ngang
        dx = 0.0
        if cb:
            spare = (br - bl) * k - (cb[2] - 8 + 4)
            if spare > 1:
                dx = spare / 2
        self.cv.move('all', bl * k - 8 + dx, bt * k - 8)
        self._off = (bl * k - 8 + dx, bt * k - 8)
        pw, ph = doc_params.PAGE_W * k, doc_params.PAGE_H * k
        self._bg_imgs = []
        for (bf, bx, by, bw, bh) in spec['bg']:
            path = resource_path(bf)
            if not os.path.exists(path):
                continue
            try:
                im = Image.open(path).convert('RGBA').resize((max(1, int(bw * k)), max(1, int(bh * k))))
                ph_img = ImageTk.PhotoImage(im)
                self._bg_imgs.append(ph_img)
                self.cv.create_image(bx * k, by * k, image=ph_img, anchor='nw', tags='bg')
            except Exception:
                pass
        self.cv.create_rectangle(0, 0, pw, ph, outline='#999', tags='bg')
        self.cv.tag_lower('bg')
        bb = self.cv.bbox('all') or (0, 0, pw, ph)
        self.cv.configure(scrollregion=(0, 0, max(pw, bb[2]) + 20, max(ph, bb[3]) + 20))

    def _dyn(self):
        try:
            return doc_params.dynamic(self.layout_key)
        except Exception:
            return {}

    def _render(self):
        lay = self._load()
        if not lay:
            tk.Label(self.cv, text='Không tìm thấy layout_%s.json' % self.layout_key,
                     bg='white', fg='red', font=FONT_TITLE).pack(pady=30)
            return
        self._off = (0, 0)
        self._draw_layout(lay, self._dyn(), place_logo=False)
        self._draw_page_bg()
        self._draw_overrides()

    def _pdf(self):
        try:
            from pdf_export import render_doc_pdf
        except Exception as e:
            messagebox.showerror('Lỗi', f'Cần reportlab để xuất PDF:\n{e}', parent=self)
            return
        path = filedialog.asksaveasfilename(defaultextension='.pdf',
                                            initialfile=self._pdf_name(),
                                            filetypes=[('PDF', '*.pdf')], parent=self)
        if not path:
            return
        try:
            render_doc_pdf(self.layout_key, path, self._dyn())
            if messagebox.askyesno('Xong', f'Đã xuất PDF:\n{path}\n\nMở ngay?', parent=self):
                self._open(path)
        except Exception as e:
            messagebox.showerror('Lỗi xuất PDF', str(e), parent=self)

    def _pdf_name(self):
        # ten file: <TDS|COA|MSDS>_<Ma hang>_<ngay-thang-nam>_<gio> -> khong bao gio trung
        import doc_tools
        ma = str(doc_params.load().get('ma_hang') or '').strip() if self.layout_key in ('tds', 'coa', 'msds') else ''
        return doc_tools.doc_file_name(self.doc_title, ma, '.pdf')

    def _open(self, path):
        try:
            if sys.platform.startswith('win'):
                os.startfile(path)     # noqa
            elif sys.platform == 'darwin':
                import subprocess
                subprocess.Popen(['open', path])
            else:
                import subprocess
                subprocess.Popen(['xdg-open', path])
        except Exception:
            pass


# =====================================================================
#  APP
# =====================================================================
class ImagePathDialog(tk.Toplevel):
    """Nguon anh Tem / Bao bi: THEO MAY CHU (dung chung) hoac THEO MAY NAY (thu muc tren may).
    Co the tai anh tu may nay len may chu de may chu luu lai cho moi may dung."""
    def __init__(self, master, app):
        super().__init__(master)
        import image_sync
        self.IS = image_sync
        self.app = app
        self.title('Ảnh Tem / Bao bì')
        self.configure(bg='#f4f6f9')
        self.transient(master)
        st = getattr(app, 'settings', None) or load_settings()   # thư mục mục 2 là của MÁY NÀY (cục bộ)
        self.app.settings = st
        self.srv_ok = image_sync.server_available()
        self.v_mode = tk.StringVar(value=image_sync.mode(st))
        self.v_fallback = tk.BooleanVar(value=bool(st.get('img_local_fallback', True)))
        self.v_stamp = tk.StringVar(value=st.get('stamp_dir', DEFAULT_STAMP_DIR))
        self.v_pkg = tk.StringVar(value=st.get('pkg_dir', DEFAULT_PKG_DIR))
        BG = '#f4f6f9'

        # ---- 1. nguon anh ----
        lf = tk.LabelFrame(self, text=' 1. Lấy ảnh Tem / Bao bì khi làm lệnh theo ', bg=BG, font=FONT_B, fg='#1f3a5f')
        lf.grid(row=0, column=0, columnspan=2, sticky='we', padx=12, pady=(12, 4))
        self.rb_srv = tk.Radiobutton(lf, text='🖥  MÁY CHỦ — ảnh lưu trên máy chủ, mọi máy dùng chung',
                                     variable=self.v_mode, value='server', bg=BG, font=FONT,
                                     command=self._refresh_state)
        self.rb_srv.grid(row=0, column=0, sticky='w', padx=8)
        tk.Radiobutton(lf, text='💻  MÁY NÀY — đọc thư mục trên máy (đường dẫn bên dưới)',
                       variable=self.v_mode, value='local', bg=BG, font=FONT,
                       command=self._refresh_state).grid(row=1, column=0, sticky='w', padx=8)
        self.cb_fb = tk.Checkbutton(lf, text='Máy chủ chưa có ảnh → thử tìm thêm trong thư mục trên máy này',
                                    variable=self.v_fallback, bg=BG, font=('Segoe UI', 9))
        self.cb_fb.grid(row=2, column=0, sticky='w', padx=30, pady=(0, 4))
        if not self.srv_ok:
            tk.Label(lf, text='(Máy này đang không kết nối máy chủ dữ liệu — chỉ dùng được ảnh đã tải về trước đó.)',
                     bg=BG, fg='#b26a00', font=('Segoe UI', 9)).grid(row=3, column=0, sticky='w', padx=30, pady=(0, 4))

        # ---- 2. thu muc tren may ----
        lf2 = tk.LabelFrame(self, text=' 2. Thư mục ảnh trên máy này ', bg=BG, font=FONT_B, fg='#1f3a5f')
        lf2.grid(row=1, column=0, columnspan=2, sticky='we', padx=12, pady=4)
        lf2.columnconfigure(1, weight=1)
        for r, (lbl, var) in enumerate((('Tem (Stamp images):', self.v_stamp),
                                        ('Bao bì (Packaging Picture):', self.v_pkg))):
            tk.Label(lf2, text=lbl, bg=BG, font=FONT).grid(row=r, column=0, sticky='e', padx=(8, 4), pady=3)
            tk.Entry(lf2, textvariable=var, width=62, font=FONT).grid(row=r, column=1, sticky='we', pady=3)
            tk.Button(lf2, text='Chọn…', font=FONT,
                      command=lambda v=var: self._pick(v)).grid(row=r, column=2, padx=(4, 8))
        tk.Label(lf2, text='Ảnh được lấy theo: <thư mục>\\<Mã KH>\\<Mã KH>-<Mã hàng>.png/jpg   '
                           '(Mã KH = phần sau dấu chấm, VD 01.KSP → KSP)',
                 bg=BG, fg='#666', font=('Segoe UI', 9)).grid(row=2, column=0, columnspan=3, sticky='w', padx=8,
                                                             pady=(2, 6))

        # ---- 3. may chu ----
        lf3 = tk.LabelFrame(self, text=' 3. Ảnh trên máy chủ ', bg=BG, font=FONT_B, fg='#1f3a5f')
        lf3.grid(row=2, column=0, columnspan=2, sticky='we', padx=12, pady=4)
        self.lbl_srv = tk.Label(lf3, text='', bg=BG, font=FONT, fg='#1d3557', justify='left')
        self.lbl_srv.grid(row=0, column=0, columnspan=3, sticky='w', padx=8, pady=(4, 6))
        self.srv_btns = []

        def B(text, cmd, r, c, color='#1d6fa5'):
            b = tk.Button(lf3, text=text, font=FONT, bg=color, fg='white', relief='flat', padx=10, pady=4,
                          command=cmd)
            b.grid(row=r, column=c, sticky='we', padx=6, pady=3)
            self.srv_btns.append(b)
        B('⬆ Tải thư mục Tem lên máy chủ', lambda: self._upload_folder('stamp'), 1, 0)
        B('⬆ Tải thư mục Bao bì lên máy chủ', lambda: self._upload_folder('pkg'), 1, 1)
        B('⬆ Chọn ảnh lẻ tải lên…', self._upload_files, 1, 2)
        B('⬇ Tải tất cả ảnh từ máy chủ về máy này', self._download_all, 2, 0, '#2d6a4f')
        B('📋 Danh sách ảnh trên máy chủ', self._show_list, 2, 1, '#6a4c93')
        tk.Label(lf3, text='Tải lên chỉ gửi ảnh MỚI hoặc ĐÃ SỬA. Máy chủ lưu lại để mọi máy làm lệnh đều có ảnh.',
                 bg=BG, fg='#666', font=('Segoe UI', 9)).grid(row=3, column=0, columnspan=3, sticky='w', padx=8)
        self.prog = ttk.Progressbar(lf3, mode='determinate', length=420)
        self.prog.grid(row=4, column=0, columnspan=3, sticky='we', padx=8, pady=(4, 2))
        self.lbl_prog = tk.Label(lf3, text='', bg=BG, fg='#555', font=('Segoe UI', 9), anchor='w')
        self.lbl_prog.grid(row=5, column=0, columnspan=3, sticky='we', padx=8, pady=(0, 6))

        bot = tk.Frame(self, bg=BG)
        bot.grid(row=3, column=0, columnspan=2, sticky='e', padx=12, pady=(6, 12))
        tk.Button(bot, text='💾 Lưu', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat',
                  padx=16, pady=6, command=self._save).pack(side='right')
        tk.Button(bot, text='Đóng', font=FONT, relief='flat', padx=14, pady=6,
                  command=self.destroy).pack(side='right', padx=8)
        self.columnconfigure(0, weight=1)
        self._busy = False
        self._refresh_state()
        self._refresh_counts()
        self.grab_set()

    # ---------------- trang thai ----------------
    def _refresh_state(self):
        if not self.srv_ok:
            for b in self.srv_btns:
                b.config(state='disabled')
        self.cb_fb.config(state='normal' if self.v_mode.get() == 'server' else 'disabled')

    def _refresh_counts(self):
        if not self.srv_ok:
            self.lbl_srv.config(text='Không kết nối máy chủ. Bộ nhớ đệm trên máy này:\n' + self.IS.cache_root())
            return
        try:
            n = self.IS.server_counts()
            cfg = server_img_dirs()
            sdir = cfg.get('stamp_dir') or ''
            pdir = cfg.get('pkg_dir') or ''
            host = cfg.get('host') or ''
            txt = (f"Máy chủ đang lưu: {n.get('stamp', 0)} ảnh Tem   •   {n.get('pkg', 0)} ảnh Bao bì\n"
                   f"Nơi lưu: database máy chủ (htp_images) – mọi máy dùng chung.")
            if sdir or pdir:
                txt += (f"\nThư mục nguồn trên MÁY CHỦ{(' ' + host) if host else ''}:  "
                        f"Tem = {sdir or '(chưa đặt)'}   •   Bao bì = {pdir or '(chưa đặt)'}")
            else:
                txt += "\nThư mục nguồn trên máy chủ: (máy chủ chưa Lưu cấu hình thư mục)"
            txt += f"\nBộ nhớ đệm trên máy này: {self.IS.cache_root()}"
            self.lbl_srv.config(text=txt, fg='#1d3557')
        except Exception as e:
            self.lbl_srv.config(text=f'Không đọc được ảnh trên máy chủ: {e}', fg='#c0392b')

    def _pick(self, var):
        d = filedialog.askdirectory(initialdir=var.get() or None, parent=self, title='Chọn thư mục ảnh')
        if d:
            var.set(os.path.normpath(d))

    # ---------------- chay nen + thanh tien trinh ----------------
    def _run(self, title, work, done):
        if self._busy:
            return
        self._busy = True
        for b in self.srv_btns:
            b.config(state='disabled')
        self.prog['value'] = 0
        self.lbl_prog.config(text=title + '…')
        import threading

        def progress(i, total, what):
            self.after(0, lambda: (self.prog.configure(maximum=max(1, total), value=i),
                                   self.lbl_prog.config(text=f'{title}: {i}/{total}  {what}')))

        def job():
            try:
                res, err = work(progress), None
            except Exception as e:
                res, err = None, e
            self.after(0, lambda: finish(res, err))

        def finish(res, err):
            self._busy = False
            try:
                for b in self.srv_btns:
                    b.config(state='normal')
                self._refresh_counts()
            except Exception:
                return
            if err is not None:
                self.lbl_prog.config(text=f'Lỗi: {err}')
                messagebox.showerror(title, str(err), parent=self)
            else:
                done(res)
        threading.Thread(target=job, daemon=True).start()

    def _report_upload(self, res, what):
        sent, skipped, errs = res
        msg = f'{what}: đã gửi {sent} ảnh lên máy chủ, {skipped} ảnh giống hệt nên bỏ qua.'
        if errs:
            msg += f'\n\n{len(errs)} ảnh lỗi:\n' + '\n'.join(errs[:15])
        self.lbl_prog.config(text=msg.split('\n')[0])
        (messagebox.showwarning if errs else messagebox.showinfo)('Tải lên máy chủ', msg, parent=self)

    # ---------------- tai len ----------------
    def _upload_folder(self, kind):
        var = self.v_stamp if kind == 'stamp' else self.v_pkg
        root = var.get().strip()
        if not root or not os.path.isdir(root):
            root = filedialog.askdirectory(parent=self, title=f'Chọn thư mục ảnh {self.IS.KIND_LABEL[kind]} '
                                                              '(chứa các thư mục Mã KH)')
            if not root:
                return
            var.set(os.path.normpath(root))
        items = self.IS.scan_folder(root)
        if not items:
            messagebox.showwarning('Tải lên máy chủ',
                                   f'Không thấy ảnh nào trong:\n{root}\n\n'
                                   'Cấu trúc cần có: <thư mục>\\<Mã KH>\\<ảnh>.png/jpg', parent=self)
            return
        nkh = len({f.upper() for f, _n, _p in items})
        if not messagebox.askyesno('Tải lên máy chủ',
                                   f'Tải {len(items)} ảnh {self.IS.KIND_LABEL[kind]} ({nkh} khách hàng) '
                                   f'từ:\n{root}\n\nlên máy chủ?\n(Ảnh đã có và giống hệt sẽ được bỏ qua.)',
                                   parent=self):
            return
        self._run(f'Tải {self.IS.KIND_LABEL[kind]} lên', lambda pg: self.IS.upload_items(kind, items, pg),
                  lambda res: self._report_upload(res, self.IS.KIND_LABEL[kind]))

    def _upload_files(self):
        dlg = tk.Toplevel(self)
        dlg.title('Tải ảnh lẻ lên máy chủ')
        dlg.configure(bg='#f4f6f9')
        dlg.transient(self)
        dlg.grab_set()
        v_kind = tk.StringVar(value='stamp')
        tk.Label(dlg, text='Loại ảnh:', bg='#f4f6f9', font=FONT_B).grid(row=0, column=0, sticky='w', padx=12,
                                                                        pady=(12, 2))
        tk.Radiobutton(dlg, text='Tem', variable=v_kind, value='stamp', bg='#f4f6f9', font=FONT).grid(
            row=0, column=1, sticky='w', pady=(12, 2))
        tk.Radiobutton(dlg, text='Bao bì', variable=v_kind, value='pkg', bg='#f4f6f9', font=FONT).grid(
            row=0, column=2, sticky='w', pady=(12, 2))
        tk.Label(dlg, text='Thư mục Mã KH trên máy chủ:', bg='#f4f6f9', font=FONT_B).grid(
            row=1, column=0, sticky='w', padx=12, pady=4)
        v_kh = tk.StringVar()
        tk.Entry(dlg, textvariable=v_kh, width=14, font=FONT).grid(row=1, column=1, columnspan=2, sticky='w')
        tk.Label(dlg, text='Để trống = tự lấy theo tên file (KSP-HTP-P1014BF-41.png → KSP)',
                 bg='#f4f6f9', fg='#666', font=('Segoe UI', 9)).grid(row=2, column=0, columnspan=3, sticky='w',
                                                                    padx=12)

        def go():
            kind, kh = v_kind.get(), v_kh.get().strip()
            files = filedialog.askopenfilenames(
                parent=dlg, title=f'Chọn ảnh {self.IS.KIND_LABEL[kind]}',
                filetypes=[('Ảnh', '*.png *.jpg *.jpeg *.bmp *.gif'), ('Tất cả', '*.*')])
            if not files:
                return
            items = [((kh or self.IS.guess_folder(f)), os.path.basename(f), f) for f in files]
            preview = '\n'.join(f'  {fo}\\{n}' for fo, n, _p in items[:15])
            if len(items) > 15:
                preview += f'\n  … và {len(items) - 15} ảnh khác'
            if not messagebox.askyesno('Tải lên máy chủ', f'Sẽ lưu trên máy chủ ({self.IS.KIND_LABEL[kind]}):\n'
                                       f'{preview}\n\nTiếp tục?', parent=dlg):
                return
            dlg.destroy()
            self._run(f'Tải {self.IS.KIND_LABEL[kind]} lên', lambda pg: self.IS.upload_items(kind, items, pg),
                      lambda res: self._report_upload(res, self.IS.KIND_LABEL[kind]))
        bf = tk.Frame(dlg, bg='#f4f6f9')
        bf.grid(row=3, column=0, columnspan=3, sticky='e', padx=12, pady=12)
        tk.Button(bf, text='📂 Chọn ảnh & tải lên', font=FONT_B, bg='#1d6fa5', fg='white', relief='flat',
                  padx=12, pady=5, command=go).pack(side='right')
        tk.Button(bf, text='Đóng', font=FONT, relief='flat', padx=12, pady=5,
                  command=dlg.destroy).pack(side='right', padx=8)

    # ---------------- tai ve ----------------
    def _download_all(self):
        def work(pg):
            return self.IS.sync_all('stamp', pg) + self.IS.sync_all('pkg', pg)
        self._run('Tải ảnh về máy này', work,
                  lambda n: (self.lbl_prog.config(text=f'Đã tải {n} ảnh mới/đã sửa về máy này.'),
                             messagebox.showinfo('Tải về', f'Đã tải {n} ảnh mới/đã sửa về:\n{self.IS.cache_root()}',
                                                 parent=self)))

    # ---------------- danh sach ----------------
    def _show_list(self):
        try:
            rows = [('Tem',) + tuple(r) for r in self.IS.server_list('stamp')] + \
                   [('Bao bì',) + tuple(r) for r in self.IS.server_list('pkg')]
        except Exception as e:
            messagebox.showerror('Ảnh trên máy chủ', str(e), parent=self)
            return
        w = tk.Toplevel(self)
        w.title(f'Ảnh trên máy chủ ({len(rows)})')
        w.transient(self)
        w.geometry('820x460')
        cols = ('loai', 'kh', 'ten', 'kb', 'luc', 'ai')
        tv = ttk.Treeview(w, columns=cols, show='headings', selectmode='extended')
        for c, h, wd in (('loai', 'Loại', 70), ('kh', 'Mã KH', 90), ('ten', 'Tên file', 300),
                         ('kb', 'KB', 70), ('luc', 'Cập nhật', 150), ('ai', 'Người tải', 110)):
            tv.heading(c, text=h)
            tv.column(c, width=wd, anchor='e' if c == 'kb' else 'w')
        for loai, fo, n, size, at, who in rows:
            tv.insert('', 'end', values=(loai, fo, n, f'{(size or 0) / 1024:,.0f}',
                                         str(at or '')[:16], who or ''))
        sb = ttk.Scrollbar(w, orient='vertical', command=tv.yview)
        tv.configure(yscrollcommand=sb.set)
        tv.pack(side='left', fill='both', expand=True, padx=(8, 0), pady=8)
        sb.pack(side='left', fill='y', pady=8)

        def delete():
            sel = tv.selection()
            if not sel or not messagebox.askyesno('Xoá ảnh', f'Xoá {len(sel)} ảnh khỏi máy chủ?', parent=w):
                return
            for iid in sel:
                loai, fo, n = tv.item(iid, 'values')[:3]
                try:
                    self.IS.delete_image('stamp' if loai == 'Tem' else 'pkg', fo, n)
                    tv.delete(iid)
                except Exception as e:
                    messagebox.showerror('Xoá ảnh', str(e), parent=w)
                    break
            self._refresh_counts()
        if (getattr(self.app, 'user', None) or {}).get('tabs') == 'all':          # chi Admin duoc xoa
            tk.Button(w, text='🗑 Xoá ảnh đã chọn', font=FONT, bg='#b0413e', fg='white', relief='flat',
                      command=delete).pack(side='bottom', padx=8, pady=8)

    def _save(self):
        self.app.settings['img_source'] = self.v_mode.get()
        self.app.settings['img_local_fallback'] = bool(self.v_fallback.get())
        self.app.settings['stamp_dir'] = self.v_stamp.get().strip()
        self.app.settings['pkg_dir'] = self.v_pkg.get().strip()
        save_img_cfg(self.app.settings)     # ghi lên máy chủ (dùng chung) + máy này
        extra = '\n\n(Đã lưu lên máy chủ – các máy khác sẽ tự đồng bộ cài đặt này.)' \
            if datastore.is_server() else ''
        messagebox.showinfo('Đã lưu', 'Đã lưu cài đặt ảnh Tem / Bao bì.' + extra, parent=self)
        self.destroy()


# Container Volume: cot 'San luong da chay' (tam thoi = San luong don, gia tri tinh)
CV_RAN_HEADER = 'Sản lượng đã chạy (Kg)'
CV_RAN_IDX = 12

# Tab QC thanh pham (module qc_app.py): khoa phan quyen -> chu hien tren tab
QC_TABS = {'qc_run': ' 🧪 Thành phẩm (nhập) ',
           'qc_fp': ' 📊 Bảng tổng thành phẩm ',
           'qc_rd': ' 🧫 Mẫu thử (nhập) ',
           'qc_rdfp': ' 📋 Bảng tổng mẫu thử ',
           'qc_eos': ' 🖨 Báo cáo cuối ca ',
           'qc_resin': ' 🧴 3. Nhựa (NVL) ',
           'qc_caco3': ' 🪨 4. CaCO3 (NVL) ',
           'qc_additive': ' ⚗ 5. Phụ gia (NVL) ',
           'qc_bbnvl': ' 📝 Biên bản nhập NVL '}
# tab kiem tra NVL nhap kho -> Sheet tuong ung trong qc_app.py
QC_RM_SHEETS = {'qc_resin': 'RESIN', 'qc_caco3': 'CACO3', 'qc_additive': 'ADDITIVE'}


class ScrollTabBar(tk.Frame):
    """Dai tab cuon ngang thay cho hang tab cua ttk.Notebook: ten tab day du, co thanh cuon ngang,
    lan chuot (hoac Shift + lan) tren dai tab de cuon, tab dang chon to dam + tu cuon toi."""
    BG, TAB_BG, TAB_FG, SEL_BG, SEL_FG, HOV_BG = '#f4f6f9', '#dde3ea', '#33415c', '#ffffff', '#1f3a5f', '#c9d3df'

    def __init__(self, master, nb, can_move=False, on_reorder=None, on_reset=None, on_rename=None):
        super().__init__(master, bg=self.BG)
        self.nb = nb
        # ADMIN: keo tha tab (hoac chuot phai) de doi thu tu -> on_reorder(list ten tab)
        self.can_move = can_move
        self.on_reorder = on_reorder
        self.on_reset = on_reset
        self.on_rename = on_rename           # on_rename(tab, reset=False): admin doi ten hien thi cua tab
        self._drag = None
        st = ttk.Style(master)
        st.layout('Scroll.TNotebook.Tab', [])                  # an hang tab goc
        st.configure('Scroll.TNotebook', tabmargins=0, borderwidth=0)
        nb.configure(style='Scroll.TNotebook')
        self.cv = tk.Canvas(self, height=34, bg=self.BG, highlightthickness=0, bd=0)
        self.cv.pack(fill='x', side='top')
        self.hs = ttk.Scrollbar(self, orient='horizontal', command=self.cv.xview)
        self.hs.pack(fill='x', side='top')
        self.cv.configure(xscrollcommand=self.hs.set)
        self.inner = tk.Frame(self.cv, bg=self.BG)
        self.cv.create_window((0, 0), window=self.inner, anchor='nw')
        self.inner.bind('<Configure>', lambda e: self.cv.configure(scrollregion=self.cv.bbox('all')))
        for w in (self.cv, self.inner, self.hs):
            w.bind('<MouseWheel>', self._wheel)
            w.bind('<Shift-MouseWheel>', self._wheel)
            w.bind('<Button-4>', lambda e: self.cv.xview_scroll(-3, 'units'))
            w.bind('<Button-5>', lambda e: self.cv.xview_scroll(3, 'units'))
        self.btns = {}
        self._sig = None
        self._painted_sel = None
        nb.bind('<<NotebookTabChanged>>', lambda e: self.after_idle(lambda: self.sync(scroll=True)), add='+')
        self.after(100, self._poll)

    def _wheel(self, e):
        self.cv.xview_scroll(-3 if e.delta > 0 else 3, 'units')
        return 'break'

    def _poll(self):
        try:
            self.sync()
        except tk.TclError:
            return
        self.after(1500, self._poll)

    def sync(self, scroll=False):
        """Ve lai dai tab khi doi tab / them tab. CHI tu cuon toi tab dang chon khi nguoi dung
        vua CHUYEN tab (scroll=True) - khong keo thanh cuon ve khi dang xem cac tab khac."""
        tabs = list(self.nb.tabs())
        sig = [(t, self.nb.tab(t, 'text')) for t in tabs]
        if sig != self._sig:
            self._sig = sig
            for b in self.btns.values():
                b.destroy()
            self.btns = {}
            for t, text in sig:
                b = tk.Label(self.inner, text=(text or '').strip(), font=FONT_B, bg=self.TAB_BG, fg=self.TAB_FG,
                             padx=10, pady=5, cursor='hand2', bd=1, relief='raised')
                b.pack(side='left', padx=(0, 2), pady=(2, 0))
                b.bind('<Button-1>', lambda e, tt=t: self._press(e, tt))
                if self.can_move:
                    b.bind('<B1-Motion>', self._motion)
                    b.bind('<ButtonRelease-1>', self._release)
                    b.bind('<Button-3>', lambda e, tt=t: self._menu(e, tt))
                b.bind('<Enter>', lambda e, tt=t: self._hover(tt, True))
                b.bind('<Leave>', lambda e, tt=t: self._hover(tt, False))
                b.bind('<MouseWheel>', self._wheel)
                b.bind('<Button-4>', lambda e: self.cv.xview_scroll(-3, 'units'))
                b.bind('<Button-5>', lambda e: self.cv.xview_scroll(3, 'units'))
                self.btns[t] = b
        self._paint(scroll)

    # ---------------- ADMIN: doi thu tu tab ----------------
    def _press(self, e, t):
        self.nb.select(t)
        if self.can_move:
            self._drag = {'tab': t, 'x0': e.x_root, 'moving': False}

    def _motion(self, e):
        d = self._drag
        if not d:
            return
        if not d['moving'] and abs(e.x_root - d['x0']) > 10:
            d['moving'] = True
            self.configure(cursor='sb_h_double_arrow')
        if d['moving']:
            # tu cuon khi keo sat mep
            x = e.x_root - self.cv.winfo_rootx()
            if x < 30:
                self.cv.xview_scroll(-2, 'units')
            elif x > self.cv.winfo_width() - 30:
                self.cv.xview_scroll(2, 'units')
            self._mark(self._target_index(e.x_root))

    def _mark(self, idx):
        """Vach chen (mau cam) truoc tab thu idx."""
        if not hasattr(self, '_ins'):
            self._ins = tk.Frame(self.inner, bg='#e07a00', width=4, height=30)
        tabs = list(self.btns.values())
        try:
            if idx < len(tabs):
                self._ins.place(in_=tabs[idx], x=-3, y=0, relheight=1.0, width=4)
            else:
                self._ins.place(in_=tabs[-1], relx=1.0, x=-1, y=0, relheight=1.0, width=4)
            self._ins.lift()
        except tk.TclError:
            pass

    def _target_index(self, x_root):
        tabs = list(self.btns.values())
        for i, b in enumerate(tabs):
            if x_root < b.winfo_rootx() + b.winfo_width() / 2:
                return i
        return len(tabs)

    def _release(self, e):
        d, self._drag = self._drag, None
        self.configure(cursor='')
        if hasattr(self, '_ins'):
            self._ins.place_forget()
        if not d or not d['moving']:
            return
        tabs = list(self.nb.tabs())
        src = tabs.index(str(d['tab'])) if str(d['tab']) in tabs else None
        if src is None:
            return
        tgt = self._target_index(e.x_root)
        if tgt > src:
            tgt -= 1
        self.move(d['tab'], tgt)

    def move(self, t, idx):
        tabs = list(self.nb.tabs())
        idx = max(0, min(idx, len(tabs) - 1))
        if tabs.index(str(t)) == idx:
            return
        self.nb.insert(idx, t)
        self.nb.select(t)
        self.sync(scroll=True)
        if self.on_reorder:
            self.on_reorder([self.nb.tab(x, 'text').strip() for x in self.nb.tabs()])

    def _menu(self, e, t):
        tabs = list(self.nb.tabs())
        i = tabs.index(str(t)) if str(t) in tabs else 0
        n = len(tabs)
        m = tk.Menu(self, tearoff=0, font=FONT)
        m.add_command(label='◀  Chuyển sang trái', command=lambda: self.move(t, i - 1),
                      state='normal' if i > 0 else 'disabled')
        m.add_command(label='▶  Chuyển sang phải', command=lambda: self.move(t, i + 1),
                      state='normal' if i < n - 1 else 'disabled')
        m.add_command(label='⏮  Đưa lên đầu', command=lambda: self.move(t, 0),
                      state='normal' if i > 0 else 'disabled')
        m.add_command(label='⏭  Đưa xuống cuối', command=lambda: self.move(t, n - 1),
                      state='normal' if i < n - 1 else 'disabled')
        m.add_separator()
        if self.on_rename:
            m.add_command(label='✏  Đổi tên tab này…', command=lambda: self.on_rename(t))
            m.add_command(label='↩  Trả lại tên gốc của tab này', command=lambda: self.on_rename(t, reset=True))
            m.add_separator()
        m.add_command(label='↺  Khôi phục thứ tự tab mặc định', command=lambda: self.on_reset and self.on_reset())
        m.add_separator()
        m.add_command(label='(Admin: có thể KÉO THẢ tab để đổi chỗ)', state='disabled')
        try:
            m.tk_popup(e.x_root, e.y_root)
        finally:
            m.grab_release()

    def _hover(self, t, on):
        if t != self._cur():
            self.btns[t].configure(bg=self.HOV_BG if on else self.TAB_BG)

    def _cur(self):
        try:
            return self.nb.select()
        except tk.TclError:
            return None

    def _paint(self, scroll=False):
        cur = self._cur()
        for t, b in self.btns.items():
            sel = str(t) == str(cur)
            b.configure(bg=self.SEL_BG if sel else self.TAB_BG, fg=self.SEL_FG if sel else self.TAB_FG,
                        relief='solid' if sel else 'raised', pady=7 if sel else 5)
        b = self.btns.get(cur)
        changed = str(cur) != str(self._painted_sel)
        self._painted_sel = cur
        if b is not None and (scroll or changed):
            self.after_idle(lambda: self._see(b))

    def _see(self, b):
        """Cuon cho tab dang chon nam trong khung nhin."""
        try:
            self.update_idletasks()
            total = max(1, self.inner.winfo_reqwidth())
            vw = self.cv.winfo_width()
            x0, x1 = b.winfo_x(), b.winfo_x() + b.winfo_width()
            l, r = self.cv.xview()
            if x0 < l * total:
                self.cv.xview_moveto(max(0.0, (x0 - 20) / total))
            elif x1 > r * total:
                self.cv.xview_moveto(max(0.0, (x1 - vw + 20) / total))
        except tk.TclError:
            pass


class App(tk.Tk):
    def __init__(self, user):
        super().__init__()
        self.user = user
        self.logout_flag = False
        self.app_year = get_app_year()
        self.title(f'Phần mềm tạo Lệnh sản xuất - HitechPlast {self.app_year}')
        self.geometry('1360x780')
        self.configure(bg='#f4f6f9')

        datastore.set_user(user.get('hoten') or user.get('username'))
        if datastore.is_offline():                 # may chu chua bat -> ban sao offline
            self.data = datastore.load_data() or {'orders': {}}
            self.sheets = datastore.load_sheets() or []
        elif datastore.is_server():
            try:
                self.data = datastore.load_data() or {'orders': {}}
                self.sheets = datastore.load_sheets() or []
            except Exception as e:
                messagebox.showerror('Lỗi máy chủ', f'Không tải được dữ liệu từ máy chủ:\n{e}')
                self.destroy()
                raise SystemExit
            datastore.server().on_conflict = lambda msg: messagebox.showwarning(
                'Dữ liệu vừa được máy khác sửa', msg, parent=self)
        else:
            self.data = self._load('data.json', required=True)
            self.sheets = self._load('sheets.json', required=False) or []
        datastore.set_live_sheets(self.sheets)
        datastore.audit_baseline_data(self.data)
        self.settings = load_settings()
        try:                                  # danh sach day chuyen dung chung (line_config.py)
            import line_config
            line_config.reload()
        except Exception as e:
            print('Khong nap duoc line_config.py:', e)
        self._ensure_cv_ran_col()
        self._ensure_up_test_cols()
        self.cv_index = self._build_cont_index()

        allowed = user.get('tabs', 'all')
        if allowed == 'all':
            self._ensure_xuat_nvl_sheet()

        def can(tab_key):
            return allowed == 'all' or tab_key in allowed

        top = tk.Frame(self, bg='#1f3a5f', height=58)
        top.pack(fill='x')
        top.pack_propagate(False)
        self.lbl_brand = tk.Label(top, text=f'HITECHPLAST {self.app_year}  •  {APP_SUBTITLE}',
                                  bg='#1f3a5f', fg='white', font=FONT_TITLE)
        self.lbl_brand.pack(side='left', padx=18, pady=10)
        tk.Button(top, text='Đăng xuất', font=FONT, bg='#c0392b', fg='white',
                  relief='flat', padx=12, pady=4, command=self._logout).pack(side='right', padx=14)
        # 🌐 Tieng Viet / English - doi NGAY tren cua so dang mo (loi moi mo lai phan mem, giu tai khoan)
        i18n.lang_button(top, lambda lg: self._lang_live_done(), admin=(allowed == 'all'),
                         live=True, on_fail=lambda lg: self.restart()).pack(side='right', padx=4)
        # ⬆ Cap nhat phan mem: tai tinh nang moi do may khac dang len (updater.py)
        try:
            import updater
            updater.install_button(top, self).pack(side='right', padx=4)
        except Exception as e:
            print('Khong nap duoc updater.py:', e)
        self.lbl_sync = None
        self.btn_online = None
        # khung trang thai may chu (che do may chu) / OFFLINE - doi qua lai khi may chu bat / tat
        self._srvbar = tk.Frame(top, bg='#1f3a5f')
        self._srvbar.pack(side='right')
        if datastore.is_server():
            self._build_server_bar()
        elif datastore.is_offline():
            self._build_offline_bar()
        tk.Label(top, text=f"👤  {user.get('hoten') or user.get('username')}",
                 bg='#1f3a5f', fg='#9fc0e8', font=FONT_B).pack(side='right', padx=8)
        if allowed != 'all':
            tk.Button(top, text='⏳ Gia hạn', font=FONT, bg='#7b2cbf', fg='white',
                      relief='flat', padx=10, pady=4,
                      command=self._renew).pack(side='right', padx=4)
            tk.Button(top, text='🔓 Mở sheet', font=FONT, bg='#0b7285', fg='white',
                      relief='flat', padx=10, pady=4,
                      command=self._unlock_sheets).pack(side='right', padx=4)
            self.lbl_expire = tk.Label(top, text='', bg='#1f3a5f', fg='#ffd166', font=FONT)
            self.lbl_expire.pack(side='right', padx=6)
            self._update_expire_label()
        # Nut cap quyen: chi hien cho tai khoan toan quyen (admin)
        if allowed == 'all':
            tk.Button(top, text='🔑 Cấp quyền', font=FONT_B, bg='#e07a00', fg='white',
                      relief='flat', padx=12, pady=4,
                      command=self._open_user_admin).pack(side='right', padx=8)
            # ADMIN: xuat bang cua tab dang mo ra file Excel (Ctrl+Shift+E)
            tk.Button(top, text='📤 Xuất Excel tab này', font=FONT_B, bg='#217346', fg='white',
                      relief='flat', padx=10, pady=4,
                      command=self._export_tab_excel).pack(side='right', padx=4)
            self.bind_all('<Control-Shift-E>', lambda e: self._export_tab_excel())
            self.bind_all('<Control-Shift-e>', lambda e: self._export_tab_excel())
        # cac nut ben phai duoc uu tien cho truoc, tieu de tu co lai khi cua so hep
        self.lbl_brand.pack_forget()
        self.lbl_brand.pack(side='left', padx=18, pady=10, fill='x', expand=True, anchor='w')
        self.lbl_brand.configure(anchor='w')
        top.bind('<Configure>', lambda e: self._fit_brand(e.width))

        style = ttk.Style(self)
        try:
            style.theme_use('clam')
        except Exception:
            pass
        style.configure('TNotebook.Tab', font=FONT_B, padding=[8, 5])
        # theme 'clam' mac dinh THU NHO padding cua tab dang chon -> chu bi cat.
        style.map('TNotebook.Tab', padding=[('selected', [8, 5])], expand=[('selected', [1, 1, 1, 0])])

        nb = ttk.Notebook(self)
        nb.pack(fill='both', expand=True, padx=8, pady=8)
        self.main_nb = nb

        # Thu tu hien thi tab (theo dung thu tu trong Excel).
        # 'order' = tab Lenh san xuat (tinh toan). Cac ten khac = 'title' cua sheet.
        TAB_ORDER = [
            'order',
            'nvl_calc',      # Tinh nguyen vat lieu tu lenh SX / anh / list ma hang (nvl_calc.py)
            'prod_month',    # Tong hop san luong da chay trong thang (production.py)
            'prod_plan',     # Ke hoach san xuat 6 day chuyen (production.py)
            'Data Export',
            'Data Domestic',
            'Data Sample',
            'Thông số vận hành máy',
            'Tiêu chuẩn đóng gói',
            'Đơn giá vật tư',
            'NVL thay thế',
            'Xuất NVL',      # NVL da xuat theo don (tu Excel 15.XUAT NL) - xuat_nvl.py
            'Container Volume',
            'CLAIM',
            'TDS',
            'COA',
            'MSDS',
            'qc_run',        # QC: Thanh pham (nhap)      - tu qc_app.py
            'qc_fp',         # QC: Bang tong thanh pham   - tu qc_app.py
            'qc_rd',         # QC: Mau thu (nhap)         - lenh Data Sample
            'qc_rdfp',       # QC: Bang tong mau thu
            'qc_eos',        # QC: Bao cao cuoi ca (in)   - tu qc_app.py
            'qc_resin',      # QC: 3. Nhua nhap kho       - tu qc_app.py
            'qc_caco3',      # QC: 4. CaCO3 nhap kho      - tu qc_app.py
            'qc_additive',   # QC: 5. Phu gia nhap kho    - tu qc_app.py
            'qc_bbnvl',      # QC: Bien ban nhap NVL (BBCLĐV + LIST INPUT NVL) - tu qc_app.py
            'qc_report',     # Bao cao QC thang / quy / 6 thang / nam (qc_report.py)
            'claim_log',     # So theo doi claim khach hang + xuat bao cao claim (claim_log.py)
            'qc_log',        # QC: Nhat ky bat thuong QC  - CHI ADMIN (khong cap quyen duoc)
            'rmqr',          # Bao cao chat luong NVL (form Initial report + Tool NVL) - report_forms.py
            'tem_btw',       # Tem BarTender: nhan du lieu tem tu anh / Excel / Word -> xuat .btw (tem_ui.py)
        ]
        sheet_by_title = {sh['title']: sh for sh in self.sheets}
        # gom cac tab con lai chua liet ke trong TAB_ORDER (neu co) vao cuoi
        ordered = list(TAB_ORDER) + [t for t in sheet_by_title if t not in TAB_ORDER]

        # map tab DL -> loai don (de cho phep them lenh san xuat)
        DL_KIND = {'Data Export': 'Export',
                   'Data Domestic': 'Domestic',
                   'Data Sample': 'Sample'}
        # cac sheet cho phep Them/Sua/Xoa dong
        ROW_EDIT = {'Thông số vận hành máy', 'Tiêu chuẩn đóng gói',
                    'Đơn giá vật tư', 'NVL thay thế', 'Container Volume'}
        # cac tab TAI LIEU -> hien dang form (DocTab): title -> layout key
        DOC_TABS = {'CLAIM': 'claim', 'TDS': 'tds', 'COA': 'coa', 'MSDS': 'msds'}
        self.order_tab = None
        self.doc_tabs = {}
        self.dl_tabs = {}
        self.cv_tab = None
        self.sheet_tabs = {}    # title -> SheetTab (de lam moi khi module khac sua sheet)

        self.qc = None          # cau noi toi module QC (qc_app.py), tao khi can

        added = 0
        for key in ordered:
            if key in ('qc_report', 'claim_log'):
                if can(key):
                    lbl = ' 📊 Báo cáo QC ' if key == 'qc_report' else ' 📮 Claim khách hàng '
                    try:
                        if key == 'qc_report':
                            import qc_report
                            self.qc_report_tab = qc_report.QcReportTab(nb, self)
                            nb.add(self.qc_report_tab, text=lbl)
                        else:
                            import claim_log
                            self.claim_tab = claim_log.ClaimTab(nb, self)
                            nb.add(self.claim_tab, text=lbl)
                    except Exception as e:
                        fr = tk.Frame(nb, bg='#f4f6f9')
                        tk.Label(fr, text='Không mở được tab (qc_report.py / claim_log.py).\n\n' + str(e),
                                 bg='#f4f6f9', fg='#c0392b', font=FONT_B, justify='left').pack(padx=20, pady=20)
                        nb.add(fr, text=lbl)
                    added += 1
                continue
            if key in ('prod_month', 'prod_plan'):
                if can(key):
                    try:
                        import production
                        if key == 'prod_month':
                            self.prod_month_tab = production.ProdSummaryTab(nb, self)
                            nb.add(self.prod_month_tab, text=' 📈 Sản lượng tháng ')
                        else:
                            self.prod_plan_tab = production.ProdPlanTab(nb, self)
                            nb.add(self.prod_plan_tab, text=' 🗓 Kế hoạch SX ')
                    except Exception as e:
                        fr = tk.Frame(nb, bg='#f4f6f9')
                        tk.Label(fr, text='Không mở được tab (production.py).\n\n' + str(e), bg='#f4f6f9',
                                 fg='#c0392b', font=FONT_B, justify='left').pack(padx=20, pady=20)
                        nb.add(fr, text=' 📈 Sản lượng tháng ' if key == 'prod_month' else ' 🗓 Kế hoạch SX ')
                    added += 1
                continue
            if key == 'nvl_calc':
                # Tinh NVL: ai duoc cap quyen tab nay HOAC tab Lenh san xuat deu thay
                if can('nvl_calc') or can('order'):
                    try:
                        import nvl_calc
                        self.nvl_tab = nvl_calc.NvlCalcTab(nb, self)
                        nb.add(self.nvl_tab, text='🧾 Tính NVL ')
                        added += 1
                    except Exception as e:
                        print('Khong nap duoc nvl_calc.py:', e)
                continue
            if key == 'rmqr':
                # Bao cao chat luong NVL: chon loai -> hien form, nhap, chen anh, luu, xuat PDF/Excel.
                # Admin (toan quyen) duoc THIET KE / THEM / XOA form.
                if can('rmqr'):
                    try:
                        import report_forms
                        self.rmqr_tab = report_forms.ReportFormTab(
                            nb, app=self, admin=(allowed == 'all'),
                            user_name=user.get('hoten') or user.get('username'))
                        nb.add(self.rmqr_tab, text=' 📑 Báo cáo CL NVL ')
                    except Exception as e:
                        fr = tk.Frame(nb, bg='#f4f6f9')
                        tk.Label(fr, text='Không mở được tab Báo cáo CL NVL.\n'
                                          'Kiểm tra report_forms.py, report_forms_io.py, rmqr_forms.json, '
                                          'thư mục rmqr_media nằm cạnh phần mềm.\n\n' + str(e),
                                 bg='#f4f6f9', fg='#c0392b', font=FONT_B, justify='left').pack(padx=20, pady=20)
                        nb.add(fr, text=' 📑 Báo cáo CL NVL ')
                    added += 1
                continue
            if key == 'tem_btw':
                # Tem BTW: nhan du lieu tem (anh / Excel / Word) -> xuat file BarTender .btw theo mau
                if can('tem_btw'):
                    try:
                        import tem_ui
                        self.tem_tab = tem_ui.TemTab(nb, app=self)
                        nb.add(self.tem_tab, text=' 🏷 Tem BTW ')
                    except Exception as e:
                        fr = tk.Frame(nb, bg='#f4f6f9')
                        tk.Label(fr, text='Không mở được tab Tem BTW.\n'
                                          'Kiểm tra tem_ui.py, btw_core.py, tem_mau_AA06.btw nằm cạnh phần mềm.\n\n'
                                          + str(e), bg='#f4f6f9', fg='#c0392b', font=FONT_B,
                                 justify='left').pack(padx=20, pady=20)
                        nb.add(fr, text=' 🏷 Tem BTW ')
                    added += 1
                continue
            if key == 'qc_log':
                # Nhat ky bat thuong QC: CHI hien khi dang nhap bang tai khoan admin (toan quyen)
                if allowed == 'all' and self._add_qc_log_tab(nb):
                    added += 1
                continue
            # Bao cao cuoi ca: ai duoc nhap Thanh pham thi duoc in bao cao cua ca minh
            # Bien ban nhap NVL: ai duoc nhap 3. Nhua / 4. CaCO3 / 5. Phu gia thi duoc lap bien ban
            if not can(key) and not (key == 'qc_eos' and can('qc_run')) \
                    and not (key == 'qc_bbnvl' and any(can(k) for k in QC_RM_SHEETS)):
                continue
            if key in QC_TABS:
                if self._add_qc_tab(nb, key, admin=(allowed == 'all')):
                    added += 1
                continue
            if key == 'order':
                self.order_tab = OrderTab(nb, self)
                nb.add(self.order_tab, text='🧮 Lệnh sản xuất ')
                added += 1
            elif key in DOC_TABS:
                dt = DocTab(nb, key, DOC_TABS[key], app=self)
                self.doc_tabs[key] = dt
                nb.add(dt, text=f" {key} ")
                added += 1
            elif key in sheet_by_title:
                ek = DL_KIND.get(key)
                if key == 'Thông số vận hành máy':
                    try:                     # cot 'TC bot da (%)' (QC lap tieu chuan) hien san trong bang
                        import mfi_std
                        mfi_std.ensure_bd_column(sheet_by_title[key])
                    except Exception:
                        pass
                st = SheetTab(nb, sheet_by_title[key], app=self, edit_kind=ek,
                              row_edit=(key in ROW_EDIT))
                if ek:
                    self.dl_tabs[ek] = st
                if key == 'Container Volume':
                    self.cv_tab = st
                if key == 'Xuất NVL':
                    try:                     # nut 'Cap nhat tu Excel (15.XUAT NL)' - xuat_nvl.py
                        import xuat_nvl
                        xuat_nvl.install_tab(st, self, admin=(allowed == 'all'))
                    except Exception as e:
                        print('Khong nap duoc xuat_nvl.py:', e)
                self.sheet_tabs[key] = st
                if allowed == 'all' and (key in DL_KIND or key in ROW_EDIT):
                    # ADMIN: dai nut cam 'Chuyen tu Excel / Word / anh' (sheet_import.py)
                    try:
                        import sheet_import
                        sheet_import.install_main(st)
                    except Exception as e:
                        print('Khong nap duoc sheet_import.py:', e)
                # ten tab = ten sheet nguon Excel; rieng cac sheet trong VI_TAB_TITLE hien TEN TIENG VIET
                # (vd 'Unit price' -> 'Đơn giá vật tư'; khi chon English tu dich lai thanh 'Unit price')
                tab_txt = key if key in VI_TAB_TITLE else (sheet_by_title[key].get('source') or key)
                nb.add(st, text=f" {tab_txt} ")
                added += 1

        # --- THANH TAB CUON NGANG: an hang tab goc cua ttk.Notebook (bi ep hep, cat chu), thay bang
        #     1 dai nut co THANH CUON NGANG -> moi tab luon hien TEN DAY DU, nhieu tab thi keo / lan chuot.
        # --- THU TU TAB: admin keo tha / chuot phai de doi cho, luu vao settings.json ---
        self._default_tab_texts = [nb.tab(t, 'text').strip() for t in nb.tabs()]
        # ten GOC cua tung tab (de luu thu tu / ten moi theo ten goc, khong phu thuoc ten da doi)
        # (lay chu GOC luc tao tab - KHONG doan nguoc tu ban dich: 'Unit price' co 2 nghia trong tu dien
        #  'Đơn giá vật tư' / 'Đơn Giá' -> doan nguoc sai lam mat ten tab da doi)
        self._tab_default = {str(t): self._tab_src(nb, t) for t in nb.tabs()}
        self._migrate_tab_keys()
        self._default_tab_ids = list(nb.tabs())
        self._apply_tab_order(nb, (self.settings or {}).get('tab_order') or [])
        self._apply_tab_names(nb)
        self._tabbar = ScrollTabBar(self, nb, can_move=(allowed == 'all'),
                                    on_reorder=self._save_tab_order,
                                    on_reset=lambda: self._reset_tab_order(nb),
                                    on_rename=(lambda t, reset=False: self._rename_tab(nb, t, reset))
                                    if allowed == 'all' else None)
        self._tabbar.pack(fill='x', padx=8, pady=(6, 0), before=nb)
        # ban dich online vua ve (ten tab moi doi...) -> ve lai dai tab ngay, khong can mo lai phan mem
        i18n.on_refresh(lambda: self._tabbar.sync() if self._tabbar.winfo_exists() else None)
        nb.pack_configure(pady=(0, 8))

        # chuyen tab -> kiem tra lai khoa ton kho (qua 5 ngay) cua tab Lenh san xuat
        def _on_tab(ev=None):
            nt = getattr(self, 'nvl_tab', None)
            if nt is not None:
                try:
                    if nb.select() == str(nt):
                        nt.refresh_live()
                except Exception:
                    pass
            if self.order_tab is not None:
                try:
                    self.order_tab._refresh_stock_lock(repeat=False)
                except Exception:
                    pass
        nb.bind('<<NotebookTabChanged>>', _on_tab, add='+')

        if added == 0:
            tk.Label(nb, text='Tài khoản của bạn chưa được cấp quyền xem tab nào.\n'
                              'Liên hệ quản trị viên.',
                     bg='#f4f6f9', font=FONT_TITLE, fg='#c0392b').pack(expand=True)

        # hoi doi mat khau neu la lan dau dang nhap
        self.after(400, self._first_login_check)
        # vua cap nhat xong -> bao noi dung ban moi (1 lan)
        self.after(1200, self._after_update_notice)

    def _ensure_xuat_nvl_sheet(self):
        """ADMIN: tao san sheet 'Xuat NVL' (rong) neu chua co -> tab hien ngay, bam 'Cap nhat tu Excel' de nap."""
        try:
            import xuat_nvl
            _sh, created = xuat_nvl.ensure_sheet(self.sheets)
            if created:
                save_sheets_file(self.sheets)
        except Exception as e:
            print('Khong tao duoc sheet Xuat NVL:', e)

    def _after_update_notice(self):
        try:
            import updater
            updater.after_start(self)
        except Exception as e:
            print('updater.after_start loi:', e)

    @staticmethod
    def _apply_tab_order(nb, order):
        """Sap xep lai tab theo danh sach TEN tab da luu (tab moi chua co trong danh sach giu vi tri cuoi)."""
        if not order:
            return
        by_text = {}
        for t in nb.tabs():
            k = i18n.orig(nb.tab(t, 'text').strip())
            by_text.setdefault(TAB_KEY_ALIAS.get(k, k), t)
        pos = 0
        for name in order:
            k = i18n.orig(str(name).strip())
            t = by_text.pop(TAB_KEY_ALIAS.get(k, k), None)
            if t is None:
                continue
            try:
                nb.insert(pos, t)
                pos += 1
            except tk.TclError:
                pass

    def _save_tab_order(self, texts):
        dflt = getattr(self, '_tab_default', {})
        nb = getattr(self._tabbar, 'nb', None) if hasattr(self, '_tabbar') else None
        if nb is not None and dflt:
            self.settings['tab_order'] = [dflt.get(str(t), i18n.orig(nb.tab(t, 'text').strip())) for t in nb.tabs()]
        else:
            self.settings['tab_order'] = [i18n.orig(t) for t in texts]
        save_settings(self.settings)

    @staticmethod
    def _tab_src(nb, t):
        """Chu goc (tieng Viet / ten sheet) cua tab, khong phu thuoc ngon ngu dang hien."""
        try:
            st = i18n._ORIG.get((nb._w, 'tab', str(t)))
        except Exception:
            st = None
        txt = (st if isinstance(st, str) else nb.tab(t, 'text')).strip()
        if not i18n.is_en():
            return TAB_KEY_ALIAS.get(txt, txt)          # che do Viet: chu dang hien chinh la chu goc
        txt = i18n.orig(txt) if st is None else txt
        return TAB_KEY_ALIAS.get(txt, txt)

    def _migrate_tab_keys(self):
        """Ten tab / thu tu tab da luu theo khoa cu ('Unit price', 'Đơn Giá') -> khoa moi 'Đơn giá vật tư'."""
        changed = False
        names = (self.settings or {}).get('tab_names') or {}
        for old, new in TAB_KEY_ALIAS.items():
            if old in names:
                v = names.pop(old)
                if v and v not in TAB_KEY_ALIAS and v != new:
                    names.setdefault(new, v)
                changed = True
        order = (self.settings or {}).get('tab_order') or []
        if any(x in TAB_KEY_ALIAS for x in order):
            self.settings['tab_order'] = [TAB_KEY_ALIAS.get(x, x) for x in order]
            changed = True
        if changed:
            if not names:
                self.settings.pop('tab_names', None)
            try:
                save_settings(self.settings)
            except Exception:
                pass

    # ---------------- DOI TEN TAB (admin: chuot phai tab -> Doi ten) - luu settings.json 'tab_names' ----------------
    def _apply_tab_names(self, nb):
        names = (self.settings or {}).get('tab_names') or {}
        for t in nb.tabs():
            new = names.get(self._tab_default.get(str(t), ''))
            if new:
                nb.tab(t, text=f' {new} ')

    def _rename_tab(self, nb, t, reset=False):
        from tkinter import simpledialog
        key = self._tab_default.get(str(t)) or i18n.orig(nb.tab(t, 'text').strip())
        names = self.settings.setdefault('tab_names', {})
        if reset:
            names.pop(key, None)
            nb.tab(t, text=f' {key} ')
        else:
            cur = names.get(key) or key            # luon sua tren ten TIENG VIET (ke ca khi dang o English)
            new = simpledialog.askstring('Đổi tên tab', f'Tên gốc: {key}\n\nTên mới hiển thị trên tab '
                                         '(để trống = trả lại tên gốc):', initialvalue=cur, parent=self)
            if new is None:
                return
            new = new.strip()
            if not new or new == key:
                names.pop(key, None)
                nb.tab(t, text=f' {key} ')
            else:
                names[key] = new
                # hoi luon ten TIENG ANH -> chuyen English la hien dung ngay, khong phu thuoc mang
                en_old = i18n.user_translation(new)
                en = simpledialog.askstring(
                    'Tên tiếng Anh', f'Tên tiếng Anh cho tab:\n{new}\n\n'
                    '(để trống = tự dịch online khi có mạng)', initialvalue=en_old, parent=self)
                if en is not None and en.strip() != en_old:
                    i18n.set_user_translation(new, en.strip())
                elif not en_old:
                    i18n.request_online(new)
                nb.tab(t, text=f' {new} ')
        if not names:
            self.settings.pop('tab_names', None)
        save_settings(self.settings)
        try:
            self._tabbar.sync(scroll=True)
        except Exception:
            pass

    def _reset_tab_order(self, nb):
        if not messagebox.askyesno('Khôi phục thứ tự tab', 'Đưa các tab về thứ tự mặc định?', parent=self):
            return
        for pos, t in enumerate(getattr(self, '_default_tab_ids', []) or []):
            try:
                nb.insert(pos, t)
            except tk.TclError:
                pass
        if not getattr(self, '_default_tab_ids', None):
            self._apply_tab_order(nb, self._default_tab_texts)
        self.settings.pop('tab_order', None)
        save_settings(self.settings)
        try:
            self._tabbar.sync(scroll=True)
        except Exception:
            pass

    def _add_qc_tab(self, nb, key, admin=False):
        """Tab QC thanh pham (nhap / bang tong) - dung chung du lieu qc_data.db voi qc_app.py."""
        try:
            if self.qc is None:
                import qc_app
                self.qc = qc_app.HostBridge(
                    self, nb,
                    qc_name=self.user.get('hoten') or self.user.get('username'),
                    qc_can_edit=(self.user.get('tabs') == 'all'),
                    cv_rows=self.cv_rows_for_qc,
                    sample_orders=lambda: self.data.get('orders', {}).get('Sample', []))
                qc_app.CV_SAVE_FN = self.cv_saved_by_qc   # QC cong san luong -> luu Container Volume
                qc_app.MACHINE_SAVE_FN = self.machine_saved_by_qc   # QC lap tieu chuan -> luu sheet Machine
            if key == 'qc_run':
                self.qc.add_run_tab(text=QC_TABS[key], admin=admin)
            elif key == 'qc_eos':
                self.qc.add_eos_tab(text=QC_TABS[key])
            elif key == 'qc_rd':
                self.qc.add_rd_tab(text=QC_TABS[key])
            elif key == 'qc_rdfp':
                self.qc.add_rdfp_tab(text=QC_TABS[key])
            elif key == 'qc_bbnvl':
                self.qc.add_bb_tab(text=QC_TABS[key])
            elif key in QC_RM_SHEETS:
                import qc_app
                self.qc.add_sheet_tab(getattr(qc_app, QC_RM_SHEETS[key]), text=QC_TABS[key])
            else:
                self.qc.add_fp_tab(text=QC_TABS[key])
            return True
        except BaseException as ex:          # thieu file qc_app.py / loi DB -> khong lam hong ca phan mem
            if isinstance(ex, KeyboardInterrupt):
                raise
            fr = tk.Frame(nb, bg='#f4f6f9')
            tk.Label(fr, text='Không mở được tab QC.\n'
                              'Kiểm tra file qc_app.py nằm cùng thư mục phần mềm.\n\n' + str(ex),
                     bg='#f4f6f9', fg='#c0392b', font=FONT_B, justify='left').pack(padx=20, pady=20)
            nb.add(fr, text=QC_TABS[key])
            return True

    def _add_qc_log_tab(self, nb):
        """Tab 'Nhat ky bat thuong QC' (chi admin): QC chep tu don truoc / tu dien / che do luoi / sua gio."""
        if self.user.get('tabs') != 'all':
            return False
        try:
            import qc_app
            if self.qc is None:
                self.qc = qc_app.HostBridge(
                    self, nb,
                    qc_name=self.user.get('hoten') or self.user.get('username'),
                    qc_can_edit=True,
                    cv_rows=self.cv_rows_for_qc,
                    sample_orders=lambda: self.data.get('orders', {}).get('Sample', []))
                qc_app.CV_SAVE_FN = self.cv_saved_by_qc
                qc_app.MACHINE_SAVE_FN = self.machine_saved_by_qc
            self.qc.add_log_tab(text=' 🕵 Nhật ký bất thường QC ')
            return True
        except BaseException as ex:
            if isinstance(ex, KeyboardInterrupt):
                raise
            fr = tk.Frame(nb, bg='#f4f6f9')
            tk.Label(fr, text='Không mở được tab Nhật ký bất thường QC.\n\n' + str(ex),
                     bg='#f4f6f9', fg='#c0392b', font=FONT_B, justify='left').pack(padx=20, pady=20)
            nb.add(fr, text=' 🕵 Nhật ký bất thường QC ')
            return True

    def apply_year(self, year):
        """Doi nam hien thi ngay (goi tu cua so Cap quyen)."""
        self.app_year = int(year)
        self.title(f'Phần mềm tạo Lệnh sản xuất - HitechPlast {self.app_year}')
        self.lbl_brand.config(text=f'HITECHPLAST {self.app_year}  •  {APP_SUBTITLE}')
        try:
            self.update_idletasks()
            self._fit_brand(self.winfo_width())
        except Exception:
            pass

    def switch_year_live(self, year):
        """CHUYEN NAM KHONG MO LAI PHAN MEM: doi nam hien thi + nap du lieu nam moi + ve lai moi tab.
        Loi -> mo lai phan mem nhu cu."""
        try:
            self.config(cursor='watch')
            self.update_idletasks()
            self.apply_year(year)
            if datastore.is_server():
                self._reload_live(datastore.server(), force=True)
            else:                                   # che do file: du lieu nam moi da nap vao bo nho
                self._reload_live(None, force=True, load=False)
            self.config(cursor='')
            self._last_auto = time.time()
            if self.lbl_sync is not None:
                self.lbl_sync.config(text=f'✔ Đã chuyển sang năm {year} lúc {time.strftime("%H:%M:%S")}',
                                     fg='#8fd694')
            if datastore.is_server():
                self.after(3000, self._offline_cache_once)
            return True
        except Exception as e:
            print('Chuyen nam tai cho loi:', e)
            try:
                self.config(cursor='')
            except Exception:
                pass
            self.restart()
            return False

    def refresh_docs(self):
        """TDS doi -> ve lai TDS/COA/MSDS (cac o lien ket)."""
        for dt in self.doc_tabs.values():
            try:
                dt.refresh()
            except Exception:
                pass

    def _update_expire_label(self):
        exp = self.user.get('expire', '')
        if exp:
            self.lbl_expire.config(text=f"Hạn: {lic.show_date(exp)} (còn {lic.days_left(self.user)} ngày)")
        else:
            self.lbl_expire.config(text='Hạn: không giới hạn')

    def _renew(self):
        dlg = RenewDialog(self, self.user, f"Hạn hiện tại: {lic.show_date(self.user.get('expire', ''))}")
        self.wait_window(dlg)
        self._update_expire_label()

    def _fit_brand(self, width):
        """Thu nho chu tieu de khi thanh tieu de khong du cho."""
        try:
            used = sum(w.winfo_reqwidth() + 12 for w in self.lbl_brand.master.winfo_children() if w is not self.lbl_brand)
            room = width - used - 40
            full = f'HITECHPLAST {self.app_year}  •  {APP_SUBTITLE}'
            f = tkfont.Font(font=FONT_TITLE)
            txt = full if f.measure(i18n.tr(full)) <= room else f'HITECHPLAST {self.app_year}'
            if self.lbl_brand.cget('text') != i18n.tr(txt):
                self.lbl_brand.config(text=txt)
        except Exception:
            pass

    def perm_list(self):
        """[(khoa quyen, nhan)] moi tab co the cap quyen."""
        perms = [('order', '🧮 Lệnh sản xuất'), ('nvl_calc', '🧾 Tính NVL'), ('prod_month', '📈 Sản lượng tháng'),
                 ('prod_plan', '🗓 Kế hoạch SX'), ('rmqr', '📑 Báo cáo CL NVL'), ('tem_btw', '🏷 Tem BTW'),
                 ('qc_report', '📊 Báo cáo QC'), ('claim_log', '📮 Claim khách hàng')]
        for sh in self.sheets:
            perms.append((sh['title'], sh['title'] if sh['title'] in VI_TAB_TITLE
                          else sh.get('source') or sh['title']))
        for k, lbl in QC_TABS.items():
            perms.append((k, lbl.strip()))
        return perms

    def _export_tab_excel(self):
        """ADMIN: xuat cac bang dang hien o tab dang mo ra file Excel (excel_export.py)."""
        if self.user.get('tabs', 'all') != 'all':
            return
        nb = getattr(self, 'main_nb', None)
        if nb is None:
            return
        try:
            import excel_export
        except Exception as e:
            messagebox.showerror('Lỗi', f'Không nạp được excel_export.py:\n{e}', parent=self)
            return
        w = excel_export.current_tab(nb)
        if w is None:
            messagebox.showinfo('Chưa chọn tab', 'Hãy mở 1 tab có bảng dữ liệu.', parent=self)
            return
        name = excel_export.tab_title(nb, w) or 'Bang'
        excel_export.export_widget(self, w, name)

    def _open_user_admin(self):
        UserAdmin(self, self.perm_list())

    def _unlock_sheets(self):
        """Thanh vien nhap MA MO SHEET do admin gui -> mo them sheet, tai lai phan mem."""
        import sheet_code
        perms = self.perm_list()
        dlg = sheet_code.UnlockDialog(self, self.user, [k for k, _ in perms], perms,
                                      lambda un, ch: update_user_in_file(un, **ch))
        self.wait_window(dlg.win)
        if dlg.ok:
            self.restart()

    def _watch_kick(self):
        """Admin vua NGAT / KHOA may nay (presence.py) -> dong phan mem ngay; thong bao hien sau khi dong."""
        try:
            import presence
            k = presence.kicked()
        except Exception:
            return
        if k:
            self.restart_flag = False
            self.logout_flag = (k.get('kind') == 'logout')    # dang xuat -> quay lai man hinh dang nhap
            try:
                self.destroy()
            except Exception:
                pass
            return
        try:
            self.after(3000, self._watch_kick)
        except Exception:
            pass

    def _open_chat(self, to=None):
        import messenger
        try:
            users = [{'username': u.get('username'), 'hoten': u.get('hoten')} for u in load_users()
                     if u.get('username')]
        except Exception:
            users = []
        messenger.open_chat(self, {'username': self.user.get('username') or '',
                                   'hoten': self.user.get('hoten') or ''},
                            users, admin=(self.user.get('tabs') == 'all'), to=to)

    def _msg_tick(self):
        """Tin nhan moi -> thong bao goc man hinh + Windows; cap nhat so tin chua doc tren nut."""
        try:
            import messenger
            btn = getattr(self, 'btn_msg', None)
            if btn is None or not btn.winfo_exists():
                return
            for m in messenger.pop_new():
                messenger.notify(self, m, on_open=lambda who: self._open_chat(who))
            n = messenger.unread_count()
            btn.config(text=f'💬 Tin nhắn ({n})' if n else '💬 Tin nhắn', fg='#ff6b6b' if n else '#ffd166')
        except Exception:
            pass
        try:
            self.after(2000, self._msg_tick)
        except Exception:
            pass

    def _update_online(self):
        if self.btn_online is None:
            return
        try:
            if not self.btn_online.winfo_exists():
                return
            import presence
            n = presence.online_count()
            self.btn_online.config(text=f'🟢 {n} máy online' if n is not None else '🟢 …')
        except Exception:
            return
        self.after(10000, self._update_online)

    def _first_login_check(self):
        """Neu tai khoan CHUA doi mat khau -> hoi doi (moi lan dang nhap, den khi doi that)."""
        if self.user.get('tabs') == 'all':          # bo qua admin
            return
        if self.user.get('pw_changed', False):       # da doi mat khau roi -> thoi hoi
            return
        if messagebox.askyesno('Đổi mật khẩu',
                               'Bạn chưa đổi mật khẩu.\nBạn có muốn đổi mật khẩu bây giờ không?',
                               parent=self):
            ChangePassword(self, self.user)
        # neu chon Khong: KHONG danh dau -> lan sau van hoi

    def _logout(self):
        """Dang xuat. Tai khoan dat 'Dang xuat: tat ca cac may' -> dang xuat ca cac may khac."""
        un = self.user.get('username') or ''
        mode = self.user.get('logout_mode')
        try:
            fresh = next((x for x in load_users() if x.get('username') == un), None)
            if fresh is not None:
                mode = fresh.get('logout_mode')          # admin vua doi cai dat -> theo cai dat moi
        except Exception:
            pass
        if mode == 'all' and datastore.is_server():
            try:
                import presence
                import socket as _so
                presence.logout_user(un, by=self.user.get('hoten') or un,
                                     msg=f'Tài khoản vừa đăng xuất trên máy {_so.gethostname()} '
                                         '(cài đặt: đăng xuất tất cả các máy).')
            except Exception:
                pass
        self.logout_flag = True
        self.destroy()

    # ---------------- che do may chu / OFFLINE ----------------
    def _clear_srvbar(self):
        for w in self._srvbar.winfo_children():
            w.destroy()
        self.lbl_sync = None
        self.btn_online = None
        self.btn_msg = None

    def _build_server_bar(self):
        """Nut Tai lai + trang thai may chu + so may online; bat vong kiem tra may chu."""
        user = self.user
        self._clear_srvbar()
        tk.Button(self._srvbar, text='🔄 Tải lại', font=FONT, bg='#2d6a4f', fg='white', relief='flat',
                  padx=10, pady=4, command=self._reload_all).pack(side='right', padx=4)
        self.lbl_sync = tk.Label(self._srvbar, text='', bg='#1f3a5f',
                                 fg='#8fd694', font=('Segoe UI', 9))
        self.lbl_sync.pack(side='right', padx=6)
        self.after(AUTO_SYNC_MS, self._poll_server)
        # 🟢 so may dang dung phan mem (nhip tim presence.py)
        try:
            import presence
            presence.start(user.get('hoten') or user.get('username'), username=user.get('username') or '')
            self.btn_online = tk.Button(self._srvbar, text='🟢 …', font=FONT_B, bg='#1f3a5f', fg='#8fd694',
                                        activebackground='#1f3a5f', relief='flat', cursor='hand2',
                                        command=lambda: presence.open_window(
                                            self, admin=(user.get('tabs') == 'all'),
                                            by=user.get('hoten') or user.get('username') or ''))
            self.btn_online.pack(side='right', padx=4)
            self.after(3000, self._update_online)
            self.after(3000, self._watch_kick)
        except Exception as e:
            print('Khong chay duoc presence.py:', e)
        # 💬 nhan tin giua cac may / tai khoan (messenger.py) - yeu cau #3
        try:
            import messenger
            messenger.start(datastore.server().cfg, user.get('username') or '', user.get('hoten') or '')
            self.btn_msg = tk.Button(self._srvbar, text='💬 Tin nhắn', font=FONT_B, bg='#1f3a5f', fg='#ffd166',
                                     activebackground='#1f3a5f', relief='flat', cursor='hand2',
                                     command=lambda: self._open_chat())
            self.btn_msg.pack(side='right', padx=4)
            self.after(2000, self._msg_tick)
        except Exception as e:
            self.btn_msg = None
            print('Khong chay duoc messenger.py:', e)
        # luu ban sao offline (dung khi may chu tat luc mo phan mem)
        if not datastore.is_linked():
            self.after(15000, self._offline_cache_tick)
        # may chu du phong: dang o du phong -> cho may chinh bat; dang o may chinh -> giu du phong moi nhat
        if not getattr(self, '_sb_started', False):
            self._sb_started = True
            self.after(4000 if datastore.is_backup() else 60000, self._standby_tick)

    def _build_offline_bar(self):
        import offline
        self._clear_srvbar()
        tk.Button(self._srvbar, text='🔌 Kết nối lại', font=FONT, bg='#b26a00', fg='white', relief='flat',
                  padx=10, pady=4, command=lambda: self._offline_tick(manual=True)).pack(side='right', padx=4)
        self.lbl_sync = tk.Label(self._srvbar, text='', bg='#1f3a5f', fg='#ffb4a9', font=('Segoe UI', 9, 'bold'))
        self.lbl_sync.pack(side='right', padx=6)
        self._offline_msg(f'📴 OFFLINE – dữ liệu lưu lúc {offline.saved_at_text()} · tự kết nối khi máy chủ bật')
        self._going_online = False
        self.after(OFFLINE_RETRY_MS, self._offline_tick)

    def _offline_msg(self, t, fg='#ffb4a9'):
        try:
            if self.lbl_sync is not None:
                self.lbl_sync.config(text=t, fg=fg)
        except Exception:
            pass

    def _offline_cache_tick(self):
        """Dang online: cap nhat ban sao offline (sheet/don hang ngay, du lieu QC o luong nen)."""
        if not self.winfo_exists() or not datastore.is_server() or datastore.is_linked():
            return
        if datastore.is_backup():                      # ban sao offline chi lay tu MAY CHU CHINH
            self.after(OFFLINE_CACHE_MS, self._offline_cache_tick)
            return
        try:
            import offline
            srv = datastore.server()
            do_main, do_qc = offline_policy()          # admin co the tat (🔑 Cap quyen -> Ban offline)
            if (do_main or do_qc) and not self._ui_busy() and not getattr(self, '_year_switching', False):
                if (offline.save_main(srv, self.sheets, self.data) if do_main else True) and do_qc:
                    threading.Thread(target=lambda: self._safe(self._save_qc_own_conn), daemon=True).start()
        except Exception as e:
            print('Luu ban sao offline loi:', e)
        self.after(OFFLINE_CACHE_MS, self._offline_cache_tick)

    def _offline_cache_once(self):
        try:
            import offline
            srv = datastore.server()
            do_main, do_qc = offline_policy()
            if (offline.save_main(srv, self.sheets, self.data) if do_main else True) and do_qc:
                threading.Thread(target=lambda: self._safe(self._save_qc_own_conn), daemon=True).start()
        except Exception as e:
            print('Luu ban sao offline loi:', e)

    @staticmethod
    def _save_qc_own_conn():
        """TOC DO: chep TOAN BO du lieu QC bang KET NOI RIENG. Truoc day dung chung ket noi voi giao dien
        -> trong luc tai (vai giay - vai chuc giay qua Tailscale) moi thao tac / tu dong bo deu bi treo."""
        import offline
        srv = datastore.server()
        if srv is None:
            return

        class _Own:
            pass
        own = _Own()
        own.c = datastore.Conn(dict(srv.cfg))
        try:
            offline.save_qc(own)
        finally:
            own.c.close()

    @staticmethod
    def _safe(fn, *a):
        try:
            return fn(*a)
        except Exception as e:
            print('offline:', e)

    def _offline_tick(self, manual=False):
        """Dang OFFLINE: thu ket noi may chu (luong nen). Ket noi duoc -> dong bo + chuyen sang may chu."""
        if not self.winfo_exists() or not datastore.is_offline() or getattr(self, '_going_online', False):
            return
        self._going_online = True
        if manual:
            self._offline_msg('⏳ Đang thử kết nối máy chủ…', '#ffd166')
        import queue as _queue
        q = _queue.Queue()

        def work():
            try:
                q.put(datastore.try_connect())
            except Exception as e:
                q.put(e)
        threading.Thread(target=work, daemon=True).start()

        def wait():
            try:
                res = q.get_nowait()
            except _queue.Empty:
                self.after(300, wait)
                return
            if isinstance(res, Exception):
                self._going_online = False
                import offline
                self._offline_msg(f'📴 OFFLINE – dữ liệu lưu lúc {offline.saved_at_text()} · '
                                  f'thử kết nối lại lúc {time.strftime("%H:%M")}')
                if manual:
                    messagebox.showwarning('Chưa kết nối được máy chủ',
                                           f'{res}\n\nPhần mềm vẫn chạy OFFLINE và tự thử lại sau ít phút.',
                                           parent=self)
                self.after(OFFLINE_RETRY_MS, self._offline_tick)
                return
            self._go_online(res)
        wait()

    def _go_online(self, st):
        """May chu da bat: dong bo phan sua offline len may chu roi chuyen sang che do may chu (khong mo lai)."""
        import offline
        if not self.winfo_exists():
            return
        if self._ui_modal_busy() or getattr(self, '_year_switching', False):
            # chi hoan khi dang co HOP THOAI SUA (modal); cua so xem thuong van chuyen ngay
            self._offline_msg('🔌 Máy chủ đã bật – đang chờ đóng hộp thoại đang sửa để chuyển…', '#ffd166')
            self.after(2000, lambda: self._go_online(st))
            return
        pend = offline.pending_summary()
        win = tk.Toplevel(self)
        win.title('Đang kết nối máy chủ…')
        win.transient(self)
        win.resizable(False, False)
        win.protocol('WM_DELETE_WINDOW', lambda: None)
        tk.Label(win, text='🔌 Máy chủ đã bật – đang chuyển sang chế độ máy chủ…', font=('Segoe UI', 11, 'bold'),
                 padx=20, pady=12).pack(anchor='w')
        lbl = tk.Label(win, text=('Đồng bộ phần sửa offline: ' + pend) if pend else 'Không có phần sửa offline.',
                       font=('Segoe UI', 9), fg='#1d3557', padx=20, anchor='w', width=60)
        lbl.pack(anchor='w')
        pb = ttk.Progressbar(win, mode='indeterminate', length=420)
        pb.pack(padx=20, pady=(8, 14))
        win.update_idletasks()
        try:
            x = self.winfo_rootx() + (self.winfo_width() - win.winfo_width()) // 2
            y = self.winfo_rooty() + (self.winfo_height() - win.winfo_height()) // 3
            win.geometry(f'+{max(0, x)}+{max(0, y)}')
        except Exception:
            pass
        win.grab_set()
        pb.start(12)
        import queue as _queue
        q = _queue.Queue()

        def work():
            try:
                q.put(('done', offline.sync(st, progress=lambda t: q.put(('msg', t)))))
            except Exception as e:
                q.put(('done', e))
        threading.Thread(target=work, daemon=True).start()

        def poll():
            try:
                while True:
                    kind, val = q.get_nowait()
                    if kind == 'msg':
                        lbl.config(text=val)
                    else:
                        self._go_online_finish(st, val, win)
                        return
            except _queue.Empty:
                pass
            self.after(150, poll)
        poll()

    def _go_online_finish(self, st, report, win):
        import offline
        err = report if isinstance(report, Exception) else None
        if err is None:
            try:
                offline.sync_tail(st, report)               # phan vua sua trong luc dong bo
                datastore.go_online(st)
                try:                                        # du lieu QC: chuyen sang may chu
                    import qc_app
                    if qc_app.STORE is not None:
                        qc_app.STORE.con = datastore.qc_connection()
                        qc_app.STORE.cache.clear()
                        qc_app.STORE.idx_cache.clear()
                        qc_app.STORE.ver += 1
                except Exception:
                    pass
                st.on_conflict = lambda msg: messagebox.showwarning('Dữ liệu vừa được máy khác sửa', msg, parent=self)
                try:
                    load_store()                            # tai khoan: ban tren may chu (da gop)
                except Exception:
                    pass
                self.apply_year(get_app_year())
                self._reload_live(st, force=True)
                self._build_server_bar()
            except Exception as e:
                err = e
        try:
            win.grab_release()
            win.destroy()
        except Exception:
            pass
        if err is not None:
            if not datastore.is_server():
                try:
                    st.c.close()
                except Exception:
                    pass
                self._going_online = False
                self._offline_msg(f'⚠ Đồng bộ offline lỗi – vẫn chạy OFFLINE, thử lại lúc '
                                  f'{time.strftime("%H:%M")}')
                messagebox.showerror('Chưa đồng bộ được',
                                     f'Máy chủ đã bật nhưng chưa đồng bộ được phần sửa offline:\n{err}\n\n'
                                     'Dữ liệu offline vẫn giữ nguyên trên máy này, phần mềm sẽ tự thử lại.',
                                     parent=self)
                self.after(OFFLINE_RETRY_MS * 4, self._offline_tick)
            else:
                messagebox.showwarning('Đã kết nối máy chủ',
                                       f'Đã chuyển sang máy chủ nhưng vẽ lại giao diện lỗi:\n{err}\n\n'
                                       'Bấm 🔄 Tải lại nếu thấy dữ liệu chưa đúng.', parent=self)
            return
        msg = 'Máy chủ đã bật – phần mềm đã chuyển sang CHẾ ĐỘ MÁY CHỦ.'
        if report.get('year_changed'):
            msg += ('\n\n⚠ Quản trị viên đã CHUYỂN NĂM trong lúc máy này offline → phần sửa offline KHÔNG ghi lên '
                    'máy chủ (tránh lẫn năm).')
            if report.get('parked'):
                msg += f'\nĐã giữ lại bản sửa offline tại:\n{report["parked"]}'
        else:
            done = []
            if report.get('qc'):
                done.append(f"{report['qc']} dòng QC / cài đặt")
            if report.get('orders'):
                done.append(f"{report['orders']} lệnh sản xuất")
            if report.get('docs'):
                done.append(f"{report['docs']} tài liệu")
            if report.get('users'):
                done.append('tài khoản')
            if report.get('sheets'):
                done.append('các sheet (gộp theo dòng với máy khác)')
            if done:
                msg += '\nĐã đồng bộ phần sửa offline: ' + ', '.join(done) + '.'
            for n in report.get('notes') or []:
                msg += '\n\n⚠ ' + n
        messagebox.showinfo('Đã kết nối máy chủ', msg, parent=self)

    # ---------------- MAY CHU DU PHONG (standby.py) ----------------
    def _standby_tick(self):
        """Dang o du phong: 30 giay thu may chinh. Dang o may chinh: 10 phut sao chep sang du phong."""
        import standby
        try:
            if not self.winfo_exists():
                return
        except Exception:
            return
        if not datastore.is_server() or datastore.is_linked():
            return
        try:
            cfg = datastore.load_config()
        except Exception:
            cfg = None
        if not cfg or not standby.backup_cfg(cfg):
            return                                      # chua cau hinh du phong
        on_bk = datastore.is_backup()
        if on_bk:
            self._offline_msg('🛟 Đang dùng MÁY CHỦ DỰ PHÒNG (máy chính đang tắt) · tự chuyển về khi máy chính bật',
                              '#ffd166')
        if getattr(self, '_sb_busy', False):
            self.after(5000, self._standby_tick)
            return
        self._sb_busy = True
        res = {}

        def work():
            try:
                if on_bk:
                    if standby.main_reachable(cfg):
                        st = datastore.ServerStore(standby.main_cfg(cfg))
                        st.who = datastore._WHO
                        res['main'] = st
                else:
                    res['msg'] = standby.maintain(cfg)
            except Exception as e:
                res['err'] = e
        th = threading.Thread(target=work, daemon=True)
        th.start()

        def wait():
            try:
                if not self.winfo_exists():
                    return
            except Exception:
                return
            if th.is_alive():
                self.after(500, wait)
                return
            self._sb_busy = False
            if res.get('err') is not None:
                print('standby:', res['err'])
            if res.get('main') is not None:
                self._standby_failback(res['main'])
                return
            if res.get('msg'):
                print('standby:', res['msg'])
            self.after((standby.CHECK_MAIN_SEC if on_bk else standby.MIRROR_EVERY_SEC) * 1000, self._standby_tick)
        self.after(500, wait)

    def _standby_failback(self, main_st):
        """May chinh da bat: dong bo phan sua tren du phong -> may chinh, roi chuyen phan mem sang may chinh."""
        import standby
        if not self.winfo_exists():
            return
        if self._ui_busy() or getattr(self, '_year_switching', False):
            self._offline_msg('🔌 Máy chủ CHÍNH đã bật – đóng cửa sổ đang mở để chuyển về', '#ffd166')
            self.after(3000, lambda: self._standby_failback(main_st))
            return
        bcfg = standby.backup_cfg(datastore.load_config() or {})
        win = tk.Toplevel(self)
        win.title('Đang chuyển về máy chủ chính…')
        win.transient(self)
        win.resizable(False, False)
        win.protocol('WM_DELETE_WINDOW', lambda: None)
        tk.Label(win, text='🔌 Máy chủ CHÍNH đã bật – đang đưa dữ liệu dự phòng lên và chuyển về…',
                 font=('Segoe UI', 11, 'bold'), padx=20, pady=12).pack(anchor='w')
        lbl = tk.Label(win, text='Đang khoá máy dự phòng…', font=('Segoe UI', 9), fg='#1d3557', padx=20,
                       anchor='w', width=64)
        lbl.pack(anchor='w')
        pb = ttk.Progressbar(win, mode='indeterminate', length=440)
        pb.pack(padx=20, pady=(8, 14))
        win.grab_set()
        pb.start(12)
        import queue as _queue
        q = _queue.Queue()

        def work():
            try:
                q.put(('done', standby.failback(main_st, bcfg, progress=lambda t: q.put(('msg', t)))))
            except Exception as e:
                q.put(('done', e))
        threading.Thread(target=work, daemon=True).start()

        def poll():
            try:
                while True:
                    kind, val = q.get_nowait()
                    if kind == 'msg':
                        lbl.config(text=val)
                    else:
                        try:
                            win.grab_release()
                            win.destroy()
                        except Exception:
                            pass
                        self._standby_switch(main_st, val)
                        return
            except _queue.Empty:
                pass
            self.after(150, poll)
        poll()

    def _standby_switch(self, main_st, report):
        import standby
        if isinstance(report, Exception):
            try:
                main_st.c.close()
            except Exception:
                pass
            self._offline_msg(f'⚠ Chưa chuyển được về máy chính ({str(report)[:80]}) – thử lại sau', '#ffb4a9')
            self.after(standby.CHECK_MAIN_SEC * 2000, self._standby_tick)
            return
        try:                                            # dung nhip tim / tin nhan tren du phong
            import presence
            presence.stop()
        except Exception:
            pass
        try:
            import messenger
            messenger.stop()
        except Exception:
            pass
        old = datastore.server()
        datastore.go_online(main_st)
        try:
            if old is not None and old is not main_st:
                old.c.close()
        except Exception:
            pass
        try:
            import qc_app
            if qc_app.STORE is not None:
                qc_app.STORE.con = datastore.qc_connection()
                qc_app.STORE.cache.clear()
                qc_app.STORE.idx_cache.clear()
                qc_app.STORE.ver += 1
        except Exception:
            pass
        try:
            load_store()
        except Exception:
            pass
        messagebox.showinfo('Đã chuyển về máy chủ chính',
                            'Máy chủ CHÍNH đã bật – phần mềm chuyển về máy chủ chính.\n\n'
                            + standby.report_text(report) + '\n\nBấm OK để mở lại dữ liệu.', parent=self)
        self.restart()

    def destroy(self):
        try:                                        # dong phan mem: cap nhat ban sao offline (sheet / don hang)
            if datastore.is_server() and not datastore.is_linked() and not datastore.is_backup() \
                    and getattr(self, 'sheets', None) is not None and offline_policy()[0]:
                import offline
                offline.save_main(datastore.server(), self.sheets, self.data, light=True)
        except Exception:
            pass
        super().destroy()

    def _lang_live_done(self):
        """Sau khi doi ngon ngu tai cho: can lai tieu de theo do dai chu moi."""
        try:
            self.update_idletasks()
            self._fit_brand(self.winfo_width())
        except Exception:
            pass

    def restart(self):
        """Tai lai phan mem (vd sau khi chuyen nam) - giu nguyen tai khoan dang dang nhap."""
        self.restart_flag = True
        self.destroy()

    # ---------------- che do may chu: tai lai / thong bao may khac vua sua ----------------
    def _reload_all(self):
        """🔄 Tai lai: nap du lieu moi nhat tu may chu NGAY TRONG cua so dang mo (khong dong / mo lai).
        Chi khi bat buoc (chuyen nam, dang mo cua so con, loi) moi mo lai phan mem nhu cu."""
        if self._ui_busy():
            if messagebox.askyesno('Tải lại dữ liệu',
                                   'Đang có cửa sổ con mở (sửa lệnh, sửa dòng, hộp thoại...).\n\n'
                                   'Đóng cửa sổ đó rồi bấm 🔄 Tải lại để cập nhật ngay mà không phải mở lại.\n\n'
                                   'Bấm "Có" nếu vẫn muốn mở lại phần mềm ngay (cửa sổ con sẽ đóng).',
                                   parent=self):
                self.restart()
            return
        try:
            srv = datastore.server()
            ch = srv.changes()
        except Exception as e:
            messagebox.showerror('Tải lại dữ liệu', f'Không kết nối được máy chủ:\n{e}', parent=self)
            return
        if 'year' in ch:
            y = get_app_year()
            if self.switch_year_live(y):
                messagebox.showinfo('Tải lại dữ liệu', f'Quản trị viên vừa CHUYỂN NĂM trên máy khác.\n'
                                    f'Đã nạp dữ liệu năm {y}.', parent=self)
            return
        try:
            self.config(cursor='watch')
            self.update_idletasks()
            msg = self._reload_live(srv)
        except Exception as e:
            self.config(cursor='')
            if messagebox.askyesno('Tải lại dữ liệu', f'Không cập nhật tại chỗ được:\n{e}\n\n'
                                   'Mở lại phần mềm để tải lại toàn bộ?', parent=self):
                self.restart()
            return
        self.config(cursor='')
        self._last_auto = time.time()
        if self.lbl_sync is not None:
            self.lbl_sync.config(text=f'✔ Đã tải lại lúc {time.strftime("%H:%M:%S")}', fg='#8fd694')
        if msg:
            if messagebox.askyesno('Tải lại dữ liệu', msg + '\n\nMở lại phần mềm ngay để hiện đủ?', parent=self):
                self.restart()

    def _reload_live(self, srv, force=False, load=True):
        """Nap lai sheets + don hang + du lieu QC tu may chu va ve lai cac tab. Tra ve thong bao (neu co).
        force=True: khong hoi sheet chua luu (chuyen nam). load=False: du lieu da nap san (che do file) -> chi ve lai."""
        note = ''
        if not load:
            try:
                reset_sheet_caches()
            except Exception:
                pass
            self._ensure_cv_ran_col()
            self._ensure_up_test_cols()
            self.rebuild_cv_index()
        # --- sheets: co thay doi CHUA LUU cua minh -> hoi truoc khi ghi de
        base = getattr(srv, '_sheet_base', {}) if srv is not None else {}
        dirty = [] if (force or not load) else [
            sh.get('title') for sh in self.sheets
            if base.get(sh.get('title'), (0, None))[1] != json.dumps(sh, ensure_ascii=False)]
        take_sheets = True
        if dirty:
            take_sheets = messagebox.askyesno(
                'Tải lại dữ liệu', 'Các sheet sau có thay đổi CHƯA LƯU lên máy chủ:\n  ' + ', '.join(dirty[:8])
                + ('…' if len(dirty) > 8 else '') + '\n\nBỏ thay đổi đó và tải bản mới nhất từ máy chủ?\n'
                '(Bấm "Không": giữ nguyên sheet, chỉ tải lại đơn hàng + dữ liệu QC.)', parent=self)
        if take_sheets:
            old_titles = {sh.get('title') for sh in self.sheets}
            new = srv.load_sheets() if load else self.sheets
            new_titles = {sh.get('title') for sh in new}
            if load:
                self._apply_new_sheets(new)
            # ve lai MOI tab sheet (ke ca sheet khong doi) cho chac
            for child in self._all_sheet_tabs():
                try:
                    sel = child._selected_index()
                    y = child.tree.yview()[0]
                    child._after_rows_changed()
                    child._fill()
                    child.tree.yview_moveto(y)
                    if sel is not None:
                        child._select_idx(sel, see=False)
                except Exception:
                    pass
            added, gone = new_titles - old_titles, old_titles - new_titles
            if added or gone:
                note = ('Máy chủ vừa ' + ('THÊM sheet: ' + ', '.join(sorted(added)) if added else '')
                        + ('; ' if added and gone else '') + ('XOÁ sheet: ' + ', '.join(sorted(gone)) if gone else '')
                        + '.\nDữ liệu đã cập nhật, nhưng tab sheet chỉ thêm / bớt sau khi mở lại phần mềm.')
        # --- don hang
        if load:
            new = srv.load_data() or {'orders': {}}
            self.data.clear()
            self.data.update(new)                   # giu nguyen doi tuong -> cac tab dang tro toi van dung
        # --- du lieu QC
        try:
            import qc_app
            if qc_app.STORE is not None:
                qc_app.STORE.cache.clear()
                qc_app.STORE.idx_cache.clear()
                qc_app.STORE.ver += 1
            if srv is not None:
                srv.mark_seen(['qc'])
        except Exception:
            pass
        if self.qc is not None:
            try:
                self.qc.data_changed(None)
                self.qc.on_tab()                    # tab QC dang xem -> ve lai ngay
            except Exception:
                pass
        # --- cac tab khac: nap lai
        for attr, fn in (('order_tab', 'refresh_live'), ('nvl_tab', 'refresh_live'),
                         ('prod_month_tab', 'refresh'), ('prod_plan_tab', 'reload'), ('claim_tab', 'reload')):
            t = getattr(self, attr, None)
            if t is not None and hasattr(t, fn):
                try:
                    getattr(t, fn)()
                except Exception:
                    pass
        t = getattr(self, 'qc_report_tab', None)        # bao cao QC (nang): chi nap lai khi dang xem
        if t is not None:
            try:
                if t.winfo_viewable():
                    t.refresh()
            except Exception:
                pass
        try:
            self.refresh_docs()
        except Exception:
            pass
        return note

    def _poll_server(self):
        """TOC DO: hoi may chu o LUONG NEN (truoc day hoi ngay tren luong giao dien 5 giay/lan ->
        mang cham / Tailscale / rot mang la cua so dung hinh, giat)."""
        if not self.winfo_exists():
            return
        if getattr(self, '_year_switching', False):     # dang chuyen nam (luong nen dang ghi) -> doi
            self.after(3000, self._poll_server)
            return
        res = {}

        def work():
            try:
                res['ch'] = datastore.server().changes()
                res['ok'] = True
            except Exception:
                res['ch'], res['ok'] = [], False
        th = threading.Thread(target=work, daemon=True)
        th.start()

        def wait():
            try:
                if not self.winfo_exists():
                    return
            except Exception:
                return
            if th.is_alive():
                self.after(150, wait)
                return
            self._poll_server_done(res.get('ch', []), res.get('ok', False))
        self.after(150, wait)

    def _poll_server_done(self, ch, ok):
        if not self.winfo_exists():
            return
        if 'year' in ch:
            if self._ui_busy():                   # dang mo cua so con -> cho dong xong moi cap nhat
                if self.lbl_sync is not None:
                    self.lbl_sync.config(text='🔔 Quản trị viên vừa CHUYỂN NĂM – đóng cửa sổ đang mở để cập nhật',
                                         fg='#ffd166')
                self.after(5000, self._poll_server)
                return
            y = get_app_year()
            if self.switch_year_live(y):
                messagebox.showinfo('Đã chuyển năm', f'Quản trị viên vừa CHUYỂN NĂM phần mềm trên máy khác.\n'
                                    f'Đã tự nạp dữ liệu năm {y} (không cần mở lại).', parent=self)
                self.after(AUTO_SYNC_MS, self._poll_server)
            return
        if 'qc' in ch and self.qc is not None:
            try:                                  # du lieu QC: lam moi cache, cac tab QC ve lai khi mo
                import qc_app
                if qc_app.STORE is not None:
                    qc_app.STORE.cache.clear()
                    qc_app.STORE.idx_cache.clear()
                    qc_app.STORE.ver += 1
                self.qc.data_changed(None)
                datastore.server().mark_seen(['qc'])
                ch = [c for c in ch if c != 'qc']
            except Exception:
                pass
        updated = False
        if ok and ch:
            ch, updated = self._auto_refresh(ch)
        if self.lbl_sync is not None:
            if not ok:
                self.lbl_sync.config(text='⚠ Mất kết nối máy chủ', fg='#ff8a80')
            elif ch:
                self.lbl_sync.config(text='🔔 Máy khác vừa cập nhật - sẽ tự cập nhật khi đóng cửa sổ đang mở',
                                     fg='#ffd166')
            elif updated:
                self.lbl_sync.config(text=f'✔ Đã tự cập nhật lúc {time.strftime("%H:%M:%S")}', fg='#8fd694')
            elif not self._sync_msg_kept():
                self.lbl_sync.config(text='', fg='#8fd694')
        self.after(AUTO_SYNC_MS if ok else 15000, self._poll_server)

    def _sync_msg_kept(self):
        """Giu dong 'Da tu cap nhat luc ...' khoang 1 phut roi moi tro ve ten may chu."""
        t = getattr(self, '_last_auto', 0)
        return t and time.time() - t < 60

    def _ui_modal_busy(self):
        """CHI dang SUA (hop thoai modal / grab) -> hoan chuyen che do de khong mat thao tac.
        Cua so XEM (khong modal) KHONG chan -> chuyen sang may chu ngay, khong cần đóng."""
        if _MODAL_DEPTH[0] > 0:
            return True
        try:
            return self.grab_current() is not None
        except Exception:
            return False

    def _ui_busy(self):
        """Dang mo cua so con (sua lenh, sua dong, hop thoai...) -> CHUA tu cap nhat (tranh mat thao tac)."""
        if _MODAL_DEPTH[0] > 0:
            return True
        try:
            if self.grab_current() is not None:
                return True
        except Exception:
            pass

        def has_top(w):
            for c in w.winfo_children():
                if isinstance(c, tk.Toplevel):
                    try:
                        if c.winfo_exists() and c.winfo_viewable():
                            return True
                    except Exception:
                        pass
                if has_top(c):
                    return True
            return False
        try:
            return has_top(self)
        except Exception:
            return True

    def _auto_refresh(self, ch):
        """Tu nap du lieu may khac vua sua (sheets / don hang) NGAY TRONG cua so dang mo, khong khoi dong lai.
        Tra ve (cac muc con cho, da_cap_nhat)."""
        if self._ui_busy():
            return ch, False
        srv = datastore.server()
        done = []
        if 'sheets' in ch:
            try:
                # co thay doi CHUA LUU cua minh? -> khong ghi de, de nguoi dung bam Tai lai
                base = getattr(srv, '_sheet_base', {})
                dirty = any(base.get(sh.get('title'), (0, None))[1] != json.dumps(sh, ensure_ascii=False)
                            for sh in self.sheets)
                if not dirty:
                    new = srv.load_sheets()
                    self._apply_new_sheets(new)
                    done.append('sheets')
            except Exception:
                pass
        if 'orders' in ch:
            try:
                new = srv.load_data() or {'orders': {}}
                self.data.clear()
                self.data.update(new)            # giu nguyen doi tuong -> cac tab dang tro toi van dung
                if self.order_tab is not None:
                    self.order_tab.refresh_live()
                done.append('orders')
            except Exception:
                pass
        if done:
            self._last_auto = time.time()
        return [c for c in ch if c not in done], bool(done)

    def _apply_new_sheets(self, new):
        """Cap nhat cac sheet trong bo nho (giu nguyen doi tuong dict ma cac tab dang dung) va ve lai."""
        by_title = {sh.get('title'): sh for sh in self.sheets}
        changed = []
        for nsh in new:
            t = nsh.get('title')
            old = by_title.get(t)
            if old is None:
                self.sheets.append(nsh)          # sheet moi: se hien tab khi mo lai phan mem
                continue
            if old == nsh:
                continue
            old.clear()
            old.update(nsh)
            changed.append(t)
        if not changed:
            return
        try:
            reset_sheet_caches()
        except Exception:
            pass
        self._ensure_cv_ran_col()
        self._ensure_up_test_cols()
        self.rebuild_cv_index()
        for child in self._all_sheet_tabs():
            if child.sheet.get('title') not in changed:
                continue
            try:
                sel = child._selected_index()
                y = child.tree.yview()[0]
                child._after_rows_changed()
                child._fill()
                child.tree.yview_moveto(y)
                if sel is not None:
                    child._select_idx(sel, see=False)
            except Exception:
                pass
        if self.order_tab is not None:
            try:
                self.order_tab.refresh_live()
            except Exception:
                pass
        self.refresh_docs()
        if self.qc is not None:
            try:
                self.qc.data_changed(None)
            except Exception:
                pass

    def _all_sheet_tabs(self):
        out = []

        def walk(w):
            for c in w.winfo_children():
                if isinstance(c, SheetTab):
                    out.append(c)
                else:
                    walk(c)
        walk(self)
        return out

    def _load(self, name, required):
        path = resource_path(name)
        if not os.path.exists(path):
            if required:
                messagebox.showerror('Lỗi', f'Không tìm thấy {name} tại:\n{path}')
                self.destroy()
                raise SystemExit
            return None
        return json.load(open(path, encoding='utf-8'))

    def cv_sheet(self):
        for sh in self.sheets:
            if sh.get('title') == 'Container Volume':
                return sh
        return None

    def _build_cont_index(self):
        """Index Container Volume: (don_hang, ma_hang) -> dict 5 truong."""
        idx = {}
        sh = self.cv_sheet()
        if sh:
            for r in sh['rows']:
                # 0 STT,1 Don hang,2 Ma hang,3 Ma SP,4 San luong,5 So cont,6 Kg/Pallet,7 Loai bao,8 Quy cach
                don = str(r[1] if len(r) > 1 else '').strip()
                ma = str(r[2] if len(r) > 2 else '').strip()
                if not don:
                    continue
                idx[(don, ma)] = {
                    'san_luong': r[4] if len(r) > 4 else '',
                    'so_cont': r[5] if len(r) > 5 else '',
                    'kg_pallet': r[6] if len(r) > 6 else '',
                    'loai_bao': r[7] if len(r) > 7 else '',
                    'quy_cach': r[8] if len(r) > 8 else '',
                    'so_pallet': r[9] if len(r) > 9 else '',
                    'pallet_cont': r[10] if len(r) > 10 else '',
                    'cach_xep': r[11] if len(r) > 11 else '',
                }
        return idx

    def _ensure_up_test_cols(self):
        """Unit price: trai cot 'Phuong phap test' thanh cac cot ngang (Do nhot, Co hat, % Tro,
        Tap chat, Do chiu nhiet) sau Don Gia; cot Ton Kho luon o CUOI. MFI giu cot cu."""
        try:
            import test_std
            sh = next((x for x in self.sheets or [] if x.get('title') == UNIT_PRICE_TITLE), None)
            if sh and test_std.ensure_columns(sh):
                save_sheets_file(self.sheets)
        except Exception:
            pass

    def _ensure_cv_ran_col(self):
        """Container Volume: them cot 'San luong da chay' (cot thu 13, index 12) neu chua co.
        Tam thoi gan = San luong don (gia tri tinh, khong lien ket cong thuc)."""
        sh = self.cv_sheet()
        if not sh:
            return
        hd = sh.setdefault('headers', [])
        if CV_RAN_HEADER in hd:
            return
        while len(hd) < CV_RAN_IDX:
            hd.append('')
        hd.insert(CV_RAN_IDX, CV_RAN_HEADER)
        for r in sh.get('rows', []):
            while len(r) < CV_RAN_IDX:
                r.append('')
            r.insert(CV_RAN_IDX, r[4])
        try:
            save_sheets_file(self.sheets)
        except Exception:
            pass

    def cv_done_orders(self):
        """Tap (don_hang, ma_hang) DA CHAY XONG theo Container Volume:
        San luong da chay >= San luong don (> 0). Khoa da chuan hoa nhu qc_app.norm_key."""
        def num(v):
            t = str(v if v is not None else '').strip().replace(' ', '')
            if not t:
                return None
            if ',' in t and '.' in t:
                t = t.replace('.', '').replace(',', '.') if t.rfind(',') > t.rfind('.') else t.replace(',', '')
            else:
                t = t.replace(',', '')
            try:
                return float(t)
            except ValueError:
                return None
        try:
            from qc_app import norm_key
        except Exception:
            norm_key = lambda x: str(x or '').strip().upper()
        done = set()
        sh = self.cv_sheet()
        if not sh:
            return done
        for r in sh.get('rows', []):
            if len(r) <= CV_RAN_IDX:
                continue
            sl, ran = num(r[4]), num(r[CV_RAN_IDX])
            if sl and ran is not None and ran >= sl:
                done.add((norm_key(r[1]), norm_key(r[2])))
        return done

    def cv_rows_for_qc(self):
        """Cac dong sheet Container Volume (du lieu dang mo) cho tab QC:
        o 'Don hang' chi hien don co San luong da chay < San luong don."""
        sh = self.cv_sheet()
        return sh.get('rows', []) if sh else []

    def rebuild_cv_index(self):
        self.cv_index = self._build_cont_index()

    def machine_saved_by_qc(self):
        """Tab QC vua LAP TIEU CHUAN ma hang -> luu sheet 'Thong so van hanh may' + ve lai bang."""
        save_sheets_file(self.sheets)                 # loi -> qc_app bao cho nguoi dung
        st = getattr(self, 'sheet_tabs', {}).get('Thông số vận hành máy')
        if st:
            try:
                st._fill()
            except Exception:
                pass

    def cv_saved_by_qc(self):
        """Tab QC vua cong san luong vao cot 'San luong da chay' -> luu sheets.json + cap nhat bang."""
        try:
            save_sheets_file(self.sheets)
        except Exception as e:
            messagebox.showerror('Lỗi lưu', f'Không lưu được sản lượng vào Container Volume:\n{e}', parent=self)
        self.rebuild_cv_index()
        cvt = getattr(self, 'cv_tab', None)
        if cvt:
            try:
                cvt._fill()
            except Exception:
                pass


# =====================================================================
#  MAN HINH DANG NHAP
# =====================================================================
def load_users():
    try:
        d = load_store()
        users = d.get('users', [])
        _USERS_SRC[id(users)] = (users, int(d.get('rev', 0)))
        while len(_USERS_SRC) > 50:
            _USERS_SRC.pop(next(iter(_USERS_SRC)))
        return users
    except secure_store.StoreError as e:
        messagebox.showerror('Lỗi dữ liệu tài khoản', str(e))
        return []


class RecoveryWindow(tk.Tk):
    """Du lieu tai khoan bi xoa / sua / thay ban cu -> can MA KHOI PHUC cua admin."""
    def __init__(self, err):
        super().__init__()
        self.ok = False
        self.title('Khôi phục dữ liệu tài khoản')
        self.configure(bg='#f4f6f9')
        self.resizable(False, False)
        w, h = 520, 300
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        self.geometry(f'{w}x{h}+{(sw-w)//2}+{(sh-h)//2}')
        tk.Label(self, text='⚠ Dữ liệu tài khoản không hợp lệ', bg='#f4f6f9', fg='#c0392b',
                 font=FONT_TITLE).pack(anchor='w', padx=16, pady=(14, 4))
        tk.Label(self, text=str(err) + '\nPhần mềm bị khóa. Chỉ quản trị viên (có MÃ KHÔI PHỤC) '
                 'mới mở lại được.', bg='#f4f6f9', font=FONT, wraplength=480,
                 justify='left').pack(anchor='w', padx=16)
        tk.Label(self, text='Mã khôi phục:', bg='#f4f6f9', font=FONT_B).pack(anchor='w', padx=16, pady=(12, 2))
        self.v = tk.StringVar()
        e = tk.Entry(self, textvariable=self.v, font=('Consolas', 14, 'bold'), width=30, show='●',
                     justify='center')
        e.pack(padx=16, ipady=4)
        e.bind('<Return>', lambda ev: self._do())
        e.focus_set()
        self.msg = tk.Label(self, text='', bg='#f4f6f9', fg='#c0392b', font=FONT)
        self.msg.pack(pady=4)
        bt = tk.Frame(self, bg='#f4f6f9')
        bt.pack(fill='x', padx=16, pady=6)
        tk.Button(bt, text='Thoát', font=FONT, relief='flat', padx=12, pady=5,
                  command=self.destroy).pack(side='right')
        tk.Button(bt, text='Khôi phục', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat',
                  padx=14, pady=5, command=self._do).pack(side='right', padx=8)

    def _do(self):
        if not lic.check_recovery(self.v.get()):
            self.msg.config(text='✗ Mã khôi phục không đúng.')
            self.v.set('')
            return
        ext = os.path.join(secure_store.app_dir(), 'users.json')
        data, src = secure_store.best_available(), 'bản dữ liệu mã hóa hợp lệ gần nhất'
        if data is None and os.path.exists(ext):
            try:
                data, src = secure_store.import_plain_json(ext), 'file users.json'
            except Exception:
                data = None
        if data is None:
            data = {'users': [{'username': 'admin', 'password': 'admin123',
                               'hoten': 'Quản trị viên', 'tabs': 'all'}], 'license_log': []}
            src = 'tài khoản mặc định admin / admin123'
        try:
            secure_store.save(data)
        except Exception as e:
            self.msg.config(text=f'✗ Không ghi được: {e}')
            return
        if os.path.exists(ext):
            try:
                os.remove(ext)
            except Exception:
                pass
        messagebox.showinfo('Đã khôi phục', f'Đã khôi phục từ: {src}.', parent=self)
        self.ok = True
        self.destroy()


def ensure_store_ready():
    """Truoc man hinh dang nhap: kiem tra kho tai khoan; loi -> hoi ma khoi phuc."""
    try:
        load_store()
        return True
    except secure_store.StoreError as e:
        w = RecoveryWindow(e)
        w.mainloop()
        return w.ok


class LoginWindow(tk.Tk):
    """Man hinh dang nhap: chon ten + nhap mat khau."""
    def __init__(self, users):
        super().__init__()
        self.users = users
        self.result = None
        year = get_app_year()
        self.title(f'Đăng nhập - HitechPlast {year}')
        self.configure(bg='#1f3a5f')
        self.resizable(False, False)
        w, h = 420, 380
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        self.geometry(f'{w}x{h}+{(sw-w)//2}+{(sh-h)//2}')
        self.restart_login = False

        tb = tk.Frame(self, bg='#1f3a5f')
        tb.pack(fill='x')
        i18n.lang_button(tb, None, live=True, on_fail=self._lang_changed).pack(side='right', padx=6, pady=4)
        tk.Label(self, text=f'HITECHPLAST {year}', bg='#1f3a5f', fg='white',
                 font=('Segoe UI', 18, 'bold')).pack(pady=(4, 2))
        tk.Label(self, text=APP_SUBTITLE, bg='#1f3a5f',
                 fg='#9fc0e8', font=FONT).pack(pady=(0, 20))

        box = tk.Frame(self, bg='white')
        box.pack(fill='both', expand=True, padx=24, pady=(0, 24))

        tk.Label(box, text='Tên đăng nhập (chọn hoặc gõ):', bg='white', font=FONT_B).pack(anchor='w', padx=20, pady=(20, 2))
        self.name = tk.StringVar()
        # tai khoan toan quyen (admin) KHONG hien trong danh sach -> phai go truc tiep
        self.list_users = [u for u in users if u.get('tabs') != 'all']
        names = [f"{u.get('hoten') or u['username']}  ({u['username']})" for u in self.list_users]
        self.cb = ttk.Combobox(box, textvariable=self.name, values=names,
                               state='normal', font=FONT, width=32)
        self.cb.pack(padx=20)
        if names:
            self.cb.current(0)

        tk.Label(box, text='Mật khẩu:', bg='white', font=FONT_B).pack(anchor='w', padx=20, pady=(14, 2))
        self.pw = tk.StringVar()
        pe = tk.Entry(box, textvariable=self.pw, show='●', font=FONT, width=34)
        pe.pack(padx=20)
        pe.bind('<Return>', lambda e: self._login())

        self.msg = tk.Label(box, text='', bg='white', fg='#c0392b', font=FONT)
        self.msg.pack(pady=(8, 0))

        tk.Button(box, text='ĐĂNG NHẬP', bg='#2d6a4f', fg='white', font=FONT_B,
                  relief='flat', padx=20, pady=8, command=self._login).pack(pady=6)

        # che do du lieu + nhap ma ket noi may chu online (cho tai khoan thanh vien)
        lf = tk.Frame(self, bg='#1f3a5f')
        lf.pack(side='bottom', pady=(0, 6), before=box)
        self.lbl_mode = tk.Label(lf, text='', bg='#1f3a5f', fg='#9fc0e8', font=('Segoe UI', 8))
        self.lbl_mode.pack(side='left')
        tk.Button(lf, text='🔗 Nhập mã kết nối', font=('Segoe UI', 8), bg='#1f3a5f', fg='#ffd166',
                  activebackground='#1f3a5f', relief='flat', cursor='hand2',
                  command=self._enter_link).pack(side='left', padx=4)
        self._show_mode()

        pe.focus_set()
        self.protocol('WM_DELETE_WINDOW', self._cancel)

    def _lang_changed(self, lg):
        """Doi ngon ngu o man hinh dang nhap -> ve lai man hinh."""
        self.restart_login = True
        self.destroy()

    def _show_mode(self):
        try:
            import server_link
            if datastore.is_server():
                txt = ''          # an dia chi may chu (yeu cau #10)
            elif datastore.is_offline():
                import offline
                txt = f'📴 OFFLINE (máy chủ chưa bật) – dữ liệu lưu lúc {offline.saved_at_text()}'
            elif server_link.has_link():
                txt = 'Thành viên → máy chủ online'
            else:
                txt = 'Dữ liệu trên máy này'
        except Exception:
            txt = ''
        self.lbl_mode.config(text=txt)

    def _enter_link(self):
        import server_link
        dlg = server_link.ConnectDialog(self)
        self.wait_window(dlg.win)
        self._show_mode()

    def _typed_key(self):
        """Ten dang nhap nguoi dung go / chon (chu thuong)."""
        key = self.name.get().strip().lower()
        if key.endswith(')') and '(' in key:
            key = key[key.rfind('(') + 1:-1].strip()
        return key

    def _switch_to_server(self):
        """Tai khoan thanh vien + may co server_link.dat -> chuyen sang may chu online.
        -> True neu da chuyen (danh sach tai khoan lay tu may chu)."""
        import server_link
        cfg = server_link.load_link()
        if not cfg or datastore.is_server():
            return False
        self.msg.config(text='Đang kết nối máy chủ online…', fg='#1d3557')
        self.config(cursor='watch')
        self.update()
        try:
            server_link.activate(cfg)
            ok = True
        except Exception as e:
            ok = False
            datastore.reset()
            messagebox.showerror('Không kết nối được máy chủ online',
                                 f'{e}\n\nKiểm tra: máy chủ đã bật, cùng mạng LAN hoặc đã bật Tailscale.',
                                 parent=self)
        finally:
            self.config(cursor='')
            self.msg.config(text='', fg='#c0392b')
        if not ok:
            return None
        try:
            load_store()
        except secure_store.StoreError as e:
            messagebox.showerror('Lỗi dữ liệu tài khoản', str(e), parent=self)
            datastore.reset()
            return None
        self.users = load_users()
        self.list_users = [u for u in self.users if u.get('tabs') != 'all']
        self._show_mode()
        return True

    def _find_user(self):
        txt = self.name.get().strip()
        if not txt:
            return None
        i = self.cb.current()
        names = list(self.cb.cget('values'))
        if 0 <= i < len(self.list_users) and i < len(names) and txt == names[i]:
            return self.list_users[i]
        key = txt.lower()
        # go 'Ho ten  (username)' hoac chi 'username'
        if key.endswith(')') and '(' in key:
            key = key[key.rfind('(') + 1:-1].strip()
        for u in self.users:
            if str(u.get('username', '')).lower() == key:
                return u
        return None

    def _save_lock(self, u):
        update_user_in_file(u.get('username'), **{k: u.get(k) for k in lic.LOCK_FIELDS})

    def _login(self):
        u = self._find_user()
        # Tai khoan THANH VIEN (hoac ten chua co tren may nay) + co ma ket noi -> tu chay sang may chu online
        if (u is None or u.get('tabs') != 'all') and not datastore.is_server():
            try:
                import server_link
                linked = server_link.has_link()
            except Exception:
                linked = False
            if linked:
                key = self._typed_key()
                sw = self._switch_to_server()
                if sw is None:
                    self.pw.set('')
                    return
                u = next((x for x in self.users if str(x.get('username', '')).lower() == key), None)
                if u is not None and u.get('tabs') == 'all':
                    # admin -> van duoc dung may chu (vi da go ten admin sau khi chon che do thanh vien)
                    pass
        if u is None:
            self.msg.config(text='Tên đăng nhập hoặc mật khẩu không đúng.')
            self.pw.set('')
            return
        # doc lai trang thai khoa moi nhat trong file
        fresh = next((x for x in load_users() if x.get('username') == u.get('username')), None)
        if fresh:
            for k in lic.LOCK_FIELDS:
                u[k] = fresh.get(k)
        locked, until = lic.lock_info(u)
        if locked:
            self.msg.config(text='Tài khoản bị khóa VĨNH VIỄN. Liên hệ quản trị viên.' if until is None
                            else f'Tài khoản bị khóa đến {lic.lock_text(until)}.')
            self.pw.set('')
            return
        if str(self.pw.get()) == str(u.get('password', '')):
            if not self._check_expire(u):
                return
            if not self._check_single_login(u):
                return
            if u.get('fail_count') or u.get('lock_level') or u.get('locked_until'):
                lic.register_success(u)
                self._save_lock(u)
            self.result = u
            self.destroy()
        else:
            self.pw.set('')
            if u.get('tabs') == 'all':
                self.msg.config(text='Tên đăng nhập hoặc mật khẩu không đúng.')
                return
            just_locked, until, left, days = lic.register_fail(u)
            self._save_lock(u)
            if just_locked and until is None:
                self.msg.config(text='Tài khoản đã bị khóa VĨNH VIỄN.')
                messagebox.showerror('Tài khoản bị khóa vĩnh viễn',
                                     f'Bạn đã nhập sai mật khẩu {lic.MAX_FAIL} lần sau khi đã bị khóa 1 tháng.\n'
                                     f'Tài khoản "{u.get("username")}" bị khóa VĨNH VIỄN.\n\n'
                                     'Chỉ quản trị viên mới mở khóa được.', parent=self)
            elif just_locked:
                self.msg.config(text=f'Sai {lic.MAX_FAIL} lần → khóa {days} ngày.')
                messagebox.showerror('Tài khoản bị khóa',
                                     f'Bạn đã nhập sai mật khẩu {lic.MAX_FAIL} lần.\n'
                                     f'Tài khoản "{u.get("username")}" bị khóa {days} ngày, '
                                     f'đến {until.strftime("%H:%M %d/%m/%Y")}.\n\n'
                                     'Liên hệ quản trị viên nếu cần mở khóa sớm.', parent=self)
            else:
                self.msg.config(text=f'Sai mật khẩu. Còn {left} lần thử trước khi bị khóa.')

    def _check_single_login(self, u):
        """Tai khoan 'Chi 1 may': dang online o may khac -> hoi; dong y -> may kia tu dang xuat."""
        fresh = next((x for x in load_users() if x.get('username') == u.get('username')), None) or u
        if fresh.get('login_mode') != 'single' or not datastore.is_server():
            return True
        try:
            import presence
            others = presence.user_sessions(u.get('username'))
        except Exception:
            return True
        if not others:
            return True
        where = '\n'.join(f"• {o['pc_name']}  ({o['ip_lan'] or o['ip_ts']})" for o in others)
        if not messagebox.askyesno('Tài khoản đang dùng ở máy khác',
                                   f'Tài khoản "{u.get("username")}" chỉ được đăng nhập trên 1 máy và đang '
                                   f'đăng nhập ở:\n{where}\n\nĐăng nhập ở máy này sẽ ĐĂNG XUẤT máy kia '
                                   '(sau khoảng 20 giây). Tiếp tục?', icon='warning', parent=self):
            self.pw.set('')
            return False
        try:
            import socket as _so
            presence.logout_user(u.get('username'), by=u.get('hoten') or u.get('username'),
                                 msg=f'Tài khoản vừa đăng nhập trên máy {_so.gethostname()} '
                                     '(tài khoản chỉ được dùng 1 máy).')
        except Exception:
            pass
        return True

    def _check_expire(self, u):
        """Kiem tra han su dung. Het han -> hien hop nhap ma gia han.
        Ngay xet han = ngay TIN CAY (trusted_time.py), khong phai ngay tren dong ho Windows."""
        try:
            import trusted_time as TT
            self.config(cursor='watch')
            self.update_idletasks()
            tt = TT.check()
        except Exception:
            TT, tt = None, None
        finally:
            self.config(cursor='')
        if tt and tt['rolled_back']:
            info = (f"Ngày trên máy tính: {tt['local'].strftime('%d/%m/%Y')}\n"
                    f"Ngày thực tế: {tt['today'].strftime('%d/%m/%Y')} (theo {tt['source']})")
            if u.get('tabs') != 'all':
                messagebox.showerror('Sai ngày giờ',
                                     'Ngày giờ máy tính đang bị LÙI so với ngày thực tế:\n\n' + info +
                                     '\n\nHãy chỉnh lại ngày giờ đúng (Settings → Time & language → bật '
                                     '"Set time automatically") rồi đăng nhập lại.', parent=self)
                self.pw.set('')
                return False
            # admin: co the do dong ho TUNG bi chinh TIEN nham -> cho xoa moc
            if messagebox.askyesno('Ngày giờ máy tính bị lùi',
                                   info + '\n\nNếu ngày máy tính hiện tại là ĐÚNG (trước đây đồng hồ từng bị '
                                          'chỉnh tiến nhầm), bấm Yes để xoá mốc ngày đã lưu trên máy này.',
                                   parent=self):
                TT.reset_marks()
        ok, why, exp_d = lic.check_account(u)
        if not ok:
            if why == 'clock':
                messagebox.showerror('Sai ngày giờ',
                                     'Ngày giờ máy tính không đúng (bị lùi về trước lần đăng nhập '
                                     'gần nhất).\nHãy chỉnh lại ngày giờ rồi đăng nhập lại.', parent=self)
                return False
            if why == 'expired':
                reason = f'Tài khoản "{u.get("username")}" đã hết hạn từ ngày {exp_d.strftime("%d/%m/%Y")}.'
            else:
                reason = (f'Tài khoản "{u.get("username")}" chưa được cấp hạn sử dụng hợp lệ '
                          '(hoặc hạn đã bị sửa tay). Liên hệ quản trị viên để lấy mã.')
            dlg = RenewDialog(self, u, reason)
            self.wait_window(dlg)
            if not dlg.ok:
                self.msg.config(text='Tài khoản đã hết hạn. Liên hệ quản trị viên.')
                self.pw.set('')
                return False
        # ghi lai ngay dang nhap (chong lui ngay may tinh) - theo ngay tin cay
        today = lic.fmt_date(lic._today())
        u['last_login'] = today
        update_user_in_file(u.get('username'), last_login=today)
        left = lic.days_left(u)
        if u.get('tabs') != 'all' and left is not None and left <= 7:
            messagebox.showwarning('Sắp hết hạn',
                                   f'Tài khoản sẽ hết hạn vào {lic.show_date(u.get("expire"))} '
                                   f'(còn {left} ngày).\nLiên hệ quản trị viên để lấy mã gia hạn '
                                   '(nút "⏳ Gia hạn" trên thanh tiêu đề).', parent=self)
        return True

    def _cancel(self):
        self.result = None
        if datastore.is_linked():
            datastore.reset()
        self.destroy()


def _apply_offline_dir():
    """Thu muc ban sao offline nguoi dung da chon (settings.json 'offline_dir')."""
    try:
        import offline
        d = str(load_settings().get('offline_dir') or '').strip()
        offline.OVERRIDE_DIR = d if (d and os.path.isdir(d)) else ''
    except Exception:
        pass


def _verify_admin_local(username, pw):
    """Kiem tra tai khoan quan tri KHONG can may chu: doc users.dat canh phan mem / trong ban sao offline."""
    import json as _json
    cands = [os.path.join(secure_store.app_dir(), secure_store.FILE_NAME)]
    try:
        import offline
        cands.append(offline.path(secure_store.FILE_NAME))
        cands.append(os.path.join(offline.default_cache_dir(), secure_store.FILE_NAME))
    except Exception:
        pass
    for p in cands:
        try:
            d = _json.loads(secure_store.decrypt(open(p, 'rb').read()).decode('utf-8'))
        except Exception:
            continue
        for u in d.get('users') or []:
            if (str(u.get('username') or '').lower() == str(username or '').strip().lower()
                    and u.get('tabs') == 'all' and str(u.get('password', '')) == str(pw)):
                return True
    return False


def ensure_server_ready():
    """Co server.json -> ket noi may chu. Loi -> hoi Thu lai / Thoat."""
    _apply_offline_dir()
    while True:
        try:
            datastore.init()
            return True
        except datastore.ServerError as e:
            datastore.reset()
            if datastore.start_offline(str(e)):      # co ban sao -> chay OFFLINE, tu ket noi khi may chu bat
                return True
            choice = _server_error_dialog(str(e) + '\n\n(Máy này chưa có bản sao dữ liệu để chạy OFFLINE – '
                                          'cần kết nối máy chủ thành công ít nhất 1 lần với phiên bản này.)')
            if choice == 'setup':
                import server_setup
                server_setup.open_dialog()
            elif choice == 'offline':
                _apply_offline_dir()
                if datastore.start_offline(str(e)):
                    return True
                _r = tk.Tk()
                _r.withdraw()
                messagebox.showerror('Không mở được bản offline',
                                     'Thư mục đã chọn không có bản sao dữ liệu offline hợp lệ.\n'
                                     '(Cần các file meta.json, sheets.json, data.json, base.json, qc.db.)', parent=_r)
                _r.destroy()
            elif choice != 'retry':
                return False


def _server_error_dialog(err):
    """Hop thoai loi ket noi: 'retry' | 'setup' | None (thoat)."""
    r = tk.Tk()
    r.title('Không kết nối được máy chủ')
    r.configure(bg='#f4f6f9')
    r.resizable(False, False)
    res = {'v': None}
    tk.Label(r, text='⚠ Không kết nối được máy chủ dữ liệu', bg='#f4f6f9', fg='#c0392b',
             font=FONT_TITLE).pack(anchor='w', padx=16, pady=(14, 4))
    tk.Label(r, text=err + '\n\nKiểm tra:\n• Máy chủ đã bật và PostgreSQL đang chạy\n'
             '• Máy này cùng mạng LAN với máy chủ (ở ngoài: đã bật Tailscale / Radmin VPN)\n'
             f'• Địa chỉ / mật khẩu trong file {datastore.CONFIG_NAME}',
             bg='#f4f6f9', font=FONT, justify='left', wraplength=520).pack(anchor='w', padx=16)
    bf = tk.Frame(r, bg='#f4f6f9')
    bf.pack(fill='x', padx=16, pady=12)

    def pick(v):
        res['v'] = v
        r.destroy()
    tk.Button(bf, text='🔁 Thử lại', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat', padx=12, pady=5,
              command=lambda: pick('retry')).pack(side='left')
    show_setup = bool(cached_policy().get('show_server_setup', True))
    if show_setup:
        tk.Button(bf, text='🖧 Cài đặt máy chủ', font=FONT, relief='flat', padx=12, pady=5,
                  command=lambda: pick('setup')).pack(side='left', padx=8)
    else:
        # Admin tat nut Cai dat may chu -> may con chi duoc chon duong dan ban sao OFFLINE (yeu cau #2)
        tk.Button(bf, text='📂 Dùng offline', font=FONT_B, bg='#b26a00', fg='white', relief='flat', padx=12,
                  pady=5, command=lambda: pick('offline')).pack(side='left', padx=8)

        def admin_unlock():
            from tkinter import simpledialog
            un = simpledialog.askstring('Quản trị viên', 'Tên đăng nhập quản trị viên:', parent=r)
            if not un:
                return
            pw = simpledialog.askstring('Quản trị viên', 'Mật khẩu:', show='●', parent=r)
            if pw is None:
                return
            if _verify_admin_local(un, pw):
                pick('setup')
            else:
                messagebox.showerror('Sai thông tin', 'Sai tài khoản / mật khẩu quản trị viên\n'
                                     '(hoặc máy này chưa có dữ liệu tài khoản để kiểm tra).', parent=r)
        tk.Button(bf, text='🔑', font=FONT, relief='flat', padx=6, pady=5,
                  command=admin_unlock).pack(side='left')
        of = tk.Frame(r, bg='#f4f6f9')
        of.pack(fill='x', padx=16, pady=(0, 10), before=bf)
        tk.Label(of, text='Đường dẫn dữ liệu offline:', bg='#f4f6f9', font=FONT_B).pack(anchor='w')
        try:
            import offline
            cur = offline.cache_dir()
        except Exception:
            cur = os.path.join(datastore.app_dir(), 'Du_lieu_offline')
        v_dir = tk.StringVar(value=cur)
        row = tk.Frame(of, bg='#f4f6f9')
        row.pack(fill='x')
        tk.Entry(row, textvariable=v_dir, font=FONT, width=58, state='readonly').pack(side='left', fill='x',
                                                                                      expand=True)

        def choose():
            d = filedialog.askdirectory(title='Chọn thư mục bản sao dữ liệu offline (Du_lieu_offline)',
                                        initialdir=v_dir.get() or datastore.app_dir(), parent=r)
            if not d:
                return
            try:
                import offline
                ok = offline.has_cache_in(d)
            except Exception:
                ok = False
            if not ok:
                messagebox.showwarning('Không đúng thư mục', 'Thư mục này không có bản sao dữ liệu offline.\n'
                                       'Chọn thư mục "Du_lieu_offline" (có meta.json, data.json, qc.db...).',
                                       parent=r)
                return
            v_dir.set(d)
            st = load_settings()
            st['offline_dir'] = d
            save_settings(st)
            lbl_dir.config(text='✔ Đã chọn – bấm "📂 Dùng offline" để mở.', fg='#2d6a4f')
        tk.Button(row, text='Chọn thư mục…', font=FONT, relief='groove', command=choose).pack(side='left', padx=6)
        try:
            import offline
            has = offline.has_cache_in(cur)
        except Exception:
            has = False
        lbl_dir = tk.Label(of, text=('✔ Có bản sao offline tại đường dẫn này.' if has else
                                     '✖ Đường dẫn này chưa có bản sao offline – chọn thư mục khác '
                                     '(vd chép thư mục Du_lieu_offline từ máy khác sang).'),
                           bg='#f4f6f9', fg='#2d6a4f' if has else '#b36b00', font=('Segoe UI', 9))
        lbl_dir.pack(anchor='w')
    tk.Button(bf, text='Thoát', font=FONT, relief='flat', padx=12, pady=5,
              command=lambda: pick(None)).pack(side='right')
    r.mainloop()
    return res['v']


def _admin_cmd_notice(k):
    """Thong bao may nay bi quan tri vien ngat ket noi / khoa (presence.py)."""
    if k['kind'] == 'block':
        txt = ('Máy này đã bị QUẢN TRỊ VIÊN KHOÁ – không được dùng phần mềm.\n'
               'Liên hệ quản trị viên để được mở khoá.')
    elif k['kind'] == 'logout':
        txt = 'Tài khoản của bạn đã bị ĐĂNG XUẤT trên máy này.'
    else:
        txt = 'Quản trị viên đã NGẮT KẾT NỐI máy này. Phần mềm đã đóng.'
    if k.get('by'):
        txt += f"\n\nNgười thực hiện: {k['by']}"
    if k.get('msg'):
        txt += f"\nLời nhắn: {k['msg']}"
    r = tk.Tk()
    r.withdraw()
    title = {'block': 'Máy bị khoá', 'logout': 'Đã đăng xuất'}.get(k['kind'], 'Đã ngắt kết nối')
    (messagebox.showinfo if k['kind'] == 'logout' else messagebox.showerror)(title, txt, parent=r)
    r.destroy()


def _instance_blocked():
    """Cai dat 'Moi may chi mo 1 phan mem' + da co phan mem khac dang mo -> bao, dua cua so cu len, True."""
    try:
        import single_instance
        first = single_instance.acquire()          # luon giu khoa (du dang cho phep nhieu)
    except Exception:
        return False
    if first:
        return False
    try:                                           # chi doc ban ghi nho (khong ket noi may chu -> bao ngay)
        pol = load_settings().get(POLICY_KEY) or {}
        if isinstance(pol, dict) and pol.get('instance') == 'multi':
            return False
    except Exception:
        pass
    r = tk.Tk()
    r.withdraw()
    messagebox.showwarning('Phần mềm đang mở',
                           'Phần mềm HitechPlast ĐÃ ĐƯỢC MỞ trên máy này.\n'
                           'Mỗi máy chỉ được mở 1 phần mềm – hãy dùng cửa sổ đang mở\n'
                           '(xem trên thanh Taskbar).', parent=r)
    r.destroy()
    try:
        single_instance.focus_existing()
    except Exception:
        pass
    return True


if __name__ == '__main__':
    try:                                   # ma hoa ngay cac file du lieu cu con o dang chu thuong
        file_crypt.migrate_existing()
    except Exception:
        pass
    # kiem tra NGAY khi bam mo (truoc khi ket noi may chu) theo cai dat da ghi nho trong settings.json
    if _instance_blocked():
        sys.exit(0)
    _first_check = True
    while True:
        if not ensure_server_ready():
            break
        if _first_check:
            _first_check = False
            try:
                load_app_policy()          # cap nhat cai dat chung moi nhat tu may chu vao settings.json
            except Exception:
                pass
        if not ensure_store_ready():
            break
        try:
            import presence
            blk = presence.check_blocked()
        except Exception:
            blk = None
        if blk:
            _admin_cmd_notice(blk)
            break
        users = load_users()
        if not users:
            break
        login = LoginWindow(users)
        login.mainloop()
        if getattr(login, 'restart_login', False):      # vua doi ngon ngu -> ve lai man hinh dang nhap
            continue
        user = login.result
        if not user:
            break                      # nguoi dung dong cua so dang nhap -> thoat
        app = App(user)
        app.mainloop()
        while getattr(app, 'restart_flag', False):     # chuyen nam / doi ngon ngu / mo sheet -> mo lai
            app = App(user)
            app.mainloop()
        kicked = None
        try:
            import presence
            kicked = presence.kicked()
            presence.stop()            # bao may nay offline
            try:
                import messenger
                messenger.stop()
            except Exception:
                pass
        except Exception:
            pass
        if kicked and kicked.get('kind') == 'netlock':
            # bi KHOA MANG: dong ket noi, quay lai buoc ket noi -> hien hop thoai "Khong ket noi duoc may chu"
            linked = datastore.is_linked()
            datastore.reset()
            if linked:
                _r = tk.Tk()
                _r.withdraw()
                messagebox.showerror('Không kết nối được máy chủ', 'Mất kết nối máy chủ dữ liệu.', parent=_r)
                _r.destroy()
            continue
        if kicked:
            if datastore.is_linked():
                datastore.reset()
            _admin_cmd_notice(kicked)
            if kicked.get('kind') == 'logout':
                continue                   # bi dang xuat -> ve man hinh dang nhap
            break
        if datastore.is_linked():      # thanh vien dang xuat -> tro ve che do du lieu goc cua may
            datastore.reset()
        if not getattr(app, 'logout_flag', False):
            break                      # dong app binh thuong -> thoat
        # neu bam 'Dang xuat' -> quay lai vong lap hien login
