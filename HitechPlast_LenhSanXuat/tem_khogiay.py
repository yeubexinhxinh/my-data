# -*- coding: utf-8 -*-
"""
tem_khogiay.py - Ghép tem vào KHỔ GIẤY tiêu chuẩn (A3 / A4 / A5) để in.

- A3 / A4: 1 tem/trang, căn giữa, phóng vừa khổ (giữ tỉ lệ).
- A5 (hoặc bật "2 tem/trang"): xếp 2 tem trên 1 trang + KẺ 1 ĐƯỜNG Ở GIỮA
  để cắt -> được 2 tem bằng nhau.
- Xuất PDF (in trực tiếp) hoặc PNG.

Phần lõi (compose_page / export_page) KHÔNG phụ thuộc giao diện -> test được.
Hộp thoại chọn khổ (choose_and_export) dùng tkinter, gọi từ tab Làm Tem Nhanh.
"""
import os

# Kích thước khổ giấy theo chuẩn ISO 216 (mm): (rộng, cao) - khổ dọc
PAPERS_MM = {
    'A3': (297, 420),
    'A4': (210, 297),
    'A5': (148, 210),
}


def _mm2px(mm, dpi):
    return int(round(mm / 25.4 * dpi))


def _fit(img, bw, bh, allow_rotate=True):
    """Phóng ảnh vừa khung bw x bh (giữ tỉ lệ); tự xoay 90° nếu vừa hơn."""
    from PIL import Image
    if bw <= 0 or bh <= 0 or img.width <= 0 or img.height <= 0:
        return img
    cands = [img]
    if allow_rotate:
        cands.append(img.rotate(90, expand=True))
    best, best_s = img, -1.0
    for im in cands:
        s = min(bw / im.width, bh / im.height)
        if s > best_s:
            best, best_s = im, s
    w = max(1, int(im_round(best.width * best_s)))
    h = max(1, int(im_round(best.height * best_s)))
    return best.resize((w, h), Image.LANCZOS)


def im_round(x):
    return int(round(x))


def _dashed_line(draw, x0, y0, x1, y1, dash, gap, fill, width):
    """Vẽ đường nét đứt ngang hoặc dọc."""
    if y0 == y1:  # ngang
        x = x0
        while x < x1:
            draw.line([(x, y0), (min(x + dash, x1), y0)], fill=fill, width=width)
            x += dash + gap
    else:          # dọc
        y = y0
        while y < y1:
            draw.line([(x0, y), (x0, min(y + dash, y1))], fill=fill, width=width)
            y += dash + gap


def _scissor(draw, x, y, size, fill):
    """Ký hiệu kéo ✂ đơn giản ở mép đường cắt."""
    r = max(2, size // 2)
    draw.ellipse([x - r, y - r, x, y], outline=fill, width=max(1, size // 8))
    draw.ellipse([x - r, y, x, y + r], outline=fill, width=max(1, size // 8))
    draw.line([(x, y), (x + size, y)], fill=fill, width=max(1, size // 8))


def compose_page(label_img, paper='A4', dpi=300, two_up=None,
                 cut_line=True, margin_mm=5.0, gap_mm=4.0):
    """Trả về 1 ảnh PIL = trang giấy đã ghép tem.
       two_up=None -> tự bật khi paper == 'A5'."""
    from PIL import Image, ImageDraw
    paper = paper.upper()
    if paper not in PAPERS_MM:
        paper = 'A4'
    if two_up is None:
        two_up = (paper == 'A5')

    pw, ph = PAPERS_MM[paper]
    W, H = _mm2px(pw, dpi), _mm2px(ph, dpi)
    m = _mm2px(margin_mm, dpi)
    page = Image.new('RGB', (W, H), 'white')

    if two_up:
        # chia đôi theo CHIỀU DÀI (cao) -> 2 ô trên/dưới, cắt ngang ở giữa
        mid = H // 2
        half_gap = _mm2px(gap_mm, dpi)
        cell_w = W - 2 * m
        cell_h = mid - m - half_gap
        img = _fit(label_img, cell_w, cell_h)
        for cy in (0, mid):
            x = (W - img.width) // 2
            # căn giữa trong nửa trang
            y = cy + (mid - img.height) // 2
            page.paste(img, (x, max(cy + (m // 2), y)))
        if cut_line:
            d = ImageDraw.Draw(page)
            col = (110, 110, 110)
            wdt = max(1, _mm2px(0.3, dpi))
            _dashed_line(d, 0, mid, W, mid, _mm2px(4, dpi), _mm2px(2.5, dpi), col, wdt)
            _scissor(d, _mm2px(6, dpi), mid, _mm2px(5, dpi), col)
        return page

    # 1 tem/trang, căn giữa
    img = _fit(label_img, W - 2 * m, H - 2 * m)
    page.paste(img, ((W - img.width) // 2, (H - img.height) // 2))
    return page


def save_label_pdf(label_img, path, size_mm, dpi=300):
    """Xuất 1 tem ra PDF ĐÚNG kích thước tem (size_mm = (rộng, cao) mm)."""
    from PIL import Image
    w_mm, h_mm = size_mm
    W = max(1, int(round(w_mm / 25.4 * dpi)))
    H = max(1, int(round(h_mm / 25.4 * dpi)))
    im = label_img.convert('RGB').resize((W, H), Image.LANCZOS)
    im.save(path, 'PDF', resolution=float(dpi))
    return path


def save_label_excel(label_img, path, size_mm, dpi=200):
    """Xuất 1 tem vào file Excel (.xlsx) - ảnh đúng kích thước tem."""
    import os
    import tempfile
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image
    w_mm, h_mm = size_mm
    W = max(1, int(round(w_mm / 25.4 * dpi)))
    H = max(1, int(round(h_mm / 25.4 * dpi)))
    im = label_img.convert('RGB').resize((W, H), Image.LANCZOS)
    tmp = tempfile.NamedTemporaryFile(suffix='.png', delete=False)
    tmp.close()
    im.save(tmp.name)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Tem'
    xi = XLImage(tmp.name)
    xi.width, xi.height = W, H
    ws.add_image(xi, 'A1')
    wb.save(path)
    try:
        os.remove(tmp.name)
    except OSError:
        pass
    return path


def save_label_word(label_img, path, size_mm):
    """Xuất 1 tem vào file Word (.docx) - ảnh rộng đúng bằng tem. Cần python-docx."""
    try:
        import docx
        from docx.shared import Mm
    except Exception:
        raise RuntimeError("Chưa cài thư viện Word.\nChạy: pip install python-docx")
    import os
    import tempfile
    w_mm, h_mm = size_mm
    tmp = tempfile.NamedTemporaryFile(suffix='.png', delete=False)
    tmp.close()
    label_img.convert('RGB').save(tmp.name)
    doc = docx.Document()
    sec = doc.sections[0]
    try:
        sec.left_margin = sec.right_margin = Mm(5)
        sec.top_margin = sec.bottom_margin = Mm(5)
    except Exception:
        pass
    doc.add_picture(tmp.name, width=Mm(w_mm))
    doc.save(path)
    try:
        os.remove(tmp.name)
    except OSError:
        pass
    return path


def is_a5(size_mm, tol=3.0):
    """size_mm có phải khổ A5 (148x210) không, kể cả xoay ngang."""
    w, h = size_mm
    a, b = sorted((w, h))
    return abs(a - 148) <= tol and abs(b - 210) <= tol


def export_page(label_img, path, paper='A4', dpi=300, **kw):
    """Ghép rồi lưu .pdf (in được) hoặc .png."""
    page = compose_page(label_img, paper=paper, dpi=dpi, **kw)
    ext = os.path.splitext(path)[1].lower()
    if ext == '.pdf':
        page.save(path, 'PDF', resolution=float(dpi))
    else:
        page.save(path, dpi=(dpi, dpi))
    return path


# --------------------------------------------------------------------------
#  Hộp thoại chọn khổ giấy + xuất  (gọi từ tab Làm Tem Nhanh)
# --------------------------------------------------------------------------
def choose_and_export(parent, get_label_image, default_dir=None, default_name='Tem'):
    """get_label_image(): hàm trả về ảnh PIL của tem hiện tại (hoặc None)."""
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox

    try:
        img = get_label_image()
    except Exception as e:
        messagebox.showerror('Lỗi', 'Không tạo được ảnh tem:\n%s' % e, parent=parent)
        return
    if img is None:
        messagebox.showinfo('Khổ giấy', 'Chưa có tem để xuất.', parent=parent)
        return

    dlg = tk.Toplevel(parent)
    dlg.title('🖨 Xuất tem theo khổ giấy')
    dlg.transient(parent.winfo_toplevel())
    dlg.resizable(False, False)
    dlg.configure(bg='#f4f6f9')

    v_paper = tk.StringVar(value='A5')
    v_two = tk.BooleanVar(value=True)
    v_cut = tk.BooleanVar(value=True)
    v_fmt = tk.StringVar(value='PDF')

    frm = ttk.Frame(dlg, padding=12)
    frm.pack(fill='both', expand=True)
    ttk.Label(frm, text='Khổ giấy:', font=('Segoe UI', 10, 'bold')).grid(row=0, column=0, sticky='w')
    cb = ttk.Combobox(frm, textvariable=v_paper, state='readonly', width=8,
                      values=['A3', 'A4', 'A5'])
    cb.grid(row=0, column=1, sticky='w', padx=6, pady=4)

    chk2 = ttk.Checkbutton(frm, text='2 tem / trang (cắt ở giữa)', variable=v_two)
    chk2.grid(row=1, column=0, columnspan=2, sticky='w')
    chkc = ttk.Checkbutton(frm, text='Kẻ đường cắt ở giữa', variable=v_cut)
    chkc.grid(row=2, column=0, columnspan=2, sticky='w')

    ttk.Label(frm, text='Định dạng:', font=('Segoe UI', 10, 'bold')).grid(row=3, column=0, sticky='w', pady=(6, 0))
    fr = ttk.Frame(frm)
    fr.grid(row=3, column=1, sticky='w', pady=(6, 0))
    ttk.Radiobutton(fr, text='PDF (in)', value='PDF', variable=v_fmt).pack(side='left')
    ttk.Radiobutton(fr, text='PNG', value='PNG', variable=v_fmt).pack(side='left', padx=8)

    hint = ttk.Label(frm, foreground='#555', wraplength=320, justify='left',
                     text='A5: tự xếp 2 tem/trang + kẻ đường ở giữa để cắt thành 2 tem bằng nhau.\n'
                          'A3/A4: 1 tem/trang, căn giữa, phóng vừa khổ.')
    hint.grid(row=4, column=0, columnspan=2, sticky='w', pady=(8, 6))

    def _sync(*_a):
        # gợi ý: A5 mặc định 2 tem; A3/A4 mặc định 1 tem (vẫn cho bật tay)
        v_two.set(v_paper.get() == 'A5')
    cb.bind('<<ComboboxSelected>>', _sync)

    def do_export():
        paper = v_paper.get()
        ext = '.pdf' if v_fmt.get() == 'PDF' else '.png'
        init = (default_name or 'Tem') + '_' + paper + ext
        p = filedialog.asksaveasfilename(parent=dlg, defaultextension=ext,
                                         initialdir=default_dir or os.getcwd(),
                                         initialfile=init,
                                         filetypes=[('PDF', '*.pdf')] if ext == '.pdf' else [('PNG', '*.png')])
        if not p:
            return
        try:
            export_page(img, p, paper=paper, two_up=v_two.get(), cut_line=v_cut.get())
        except Exception as e:
            messagebox.showerror('Lỗi xuất', str(e), parent=dlg)
            return
        dlg.destroy()
        try:
            if os.name == 'nt':
                os.startfile(p)  # noqa
        except Exception:
            pass
        messagebox.showinfo('Xong', 'Đã xuất:\n%s' % p, parent=parent)

    bar = ttk.Frame(frm)
    bar.grid(row=5, column=0, columnspan=2, sticky='e', pady=(6, 0))
    ttk.Button(bar, text='Huỷ', command=dlg.destroy).pack(side='right', padx=4)
    tk.Button(bar, text='🖨 Xuất', bg='#0068ff', fg='white', relief='flat',
              font=('Segoe UI', 10, 'bold'), padx=14, command=do_export).pack(side='right')
    dlg.grab_set()
    dlg.focus_set()
