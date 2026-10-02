# -*- coding: utf-8 -*-
"""
report_forms.py - TAB "📑 Báo cáo CL NVL" (thay cho sheet 'Initial report' + add-in Tool NVL)

Chuc nang (tuong duong cac nut Ribbon trong Tool-QC-QA-RD-HitechPlast.xlam):
    • Chon LOAI bao cao (cay: Nhựa / Bột đá / Phụ gia / Chất độn...)  -> hien dung form do  (= Origins)
    • Nhap lieu truc tiep tren form: bam o -> go (Enter luu, Shift+Enter xuong dong), o co danh sach -> chon
      Nhap Ten NVL -> tu ghi Ngay tra KQ (giong Worksheet_Change)
    • Chen anh vao o dang chon (file / Ctrl+V dan tu clipboard / keo tha)  (= Chen_Anh_Vao_O_Da_Chon)
      Anh tu co gian vua o gop (= FitPic). Xoa anh o dang chon.
    • 💾 Luu bao cao (= ThemSheet: luu ban sao + ten = Ten NVL + ngay) -> tu xuat PDF (tuy chon)
    • Danh sach bao cao da luu: tim kiem, mo lai, sua, nhan ban, xoa, xuat PDF nhieu bao cao
    • 📄 Xuat PDF (= XuatPDF_RMQR)  •  📊 Xuat Excel  •  🖨 In
    • ADMIN - Thiet ke form: sua chu / dinh dang / mau / vien tung o, gop - tach o, them - xoa hang/cot,
      chinh do rong cot / cao hang, danh sach chon, anh mau; them form moi (trang / nhan ban / nhap tu Excel),
      doi ten, doi nhom, xoa, sap xep form.

Chay rieng:  python report_forms.py
Trong phan mem chinh: tab 'rmqr' (xem app_hitechplast.py)
"""
import copy
import os
import re
import sys
import tempfile
import tkinter as tk
import tkinter.font as tkfont
from tkinter import ttk, messagebox, filedialog, simpledialog, colorchooser

import report_forms_io as rio
from report_forms_io import key, unkey, FormModel

# ---------------------------------------------------------------------------
#  Bảng màu & font dùng chung cho giao diện (để chỉnh 1 chỗ là đổi toàn bộ)
# ---------------------------------------------------------------------------
FONT_FAMILY = 'Segoe UI'
FONT = (FONT_FAMILY, 10)
FONT_B = (FONT_FAMILY, 10, 'bold')
FONT_T = (FONT_FAMILY, 13, 'bold')
FONT_H = (FONT_FAMILY, 15, 'bold')          # tiêu đề header
FONT_SUB = (FONT_FAMILY, 10)                 # phụ đề header

# Tông màu xanh dương doanh nghiệp, thống nhất toàn giao diện
CLR = {
    'primary':   '#1d6fa5',   # xanh chính (nút chính, chọn)
    'primary_d': '#155a86',   # xanh đậm (hover/nhấn)
    'accent':    '#e07a00',   # cam nhấn (thiết kế / form)
    'success':   '#2d6a4f',   # xanh lá (lưu)
    'danger':    '#b0413e',   # đỏ (PDF/xoá)
    'ink':       '#1f3a5f',   # chữ tiêu đề đậm
    'muted':     '#5b6b7b',   # chữ phụ
    'surface':   '#ffffff',   # nền bảng/danh sách
    'panel':     '#eef2f6',   # nền panel nhạt
    'bg':        '#f4f6f9',   # nền chung
    'toolbar':   '#e3eaf1',   # nền thanh công cụ
    'border':    '#cdd7e1',   # viền
    'canvas':    '#aeb9c6',   # nền vùng vẽ form
    'header1':   '#17426b',   # nền header (đậm)
    'header_sub':'#cfe0f2',   # phụ đề trên header
    'row_alt':   '#f3f7fb',   # hàng lẻ (zebra) trong danh sách
}
BG = CLR['bg']
SEL = CLR['primary']


def _open_file(path, op=None):
    try:
        if sys.platform.startswith('win'):
            os.startfile(path, op) if op else os.startfile(path)   # noqa
        elif sys.platform == 'darwin':
            import subprocess
            subprocess.Popen(['open', path])
        else:
            import subprocess
            subprocess.Popen(['xdg-open', path])
    except Exception as e:
        messagebox.showerror('Lỗi', f'Không mở được file:\n{e}')


def _safe_name(s):
    return re.sub(r'[\\/:*?"<>|\n\r\t]+', ' ', str(s or '')).strip()[:120] or 'BaoCao'


def _norm_vi(s):
    """Bo dau tieng Viet + ve chu thuong (de do chuoi 'Nguoi lap...')."""
    import unicodedata
    t = unicodedata.normalize('NFD', str(s or ''))
    t = ''.join(ch for ch in t if unicodedata.category(ch) != 'Mn')
    return t.replace('đ', 'd').replace('Đ', 'D').lower()


# =====================================================================
#  Hop thoai sua o (THIET KE)
# =====================================================================
class CellDesignDialog(tk.Toplevel):
    def __init__(self, master, form, r, c, on_ok):
        super().__init__(master)
        self.form, self.r, self.c, self.on_ok = form, r, c, on_ok
        fm = FormModel(form)
        st = dict(fm.style(r, c))
        from openpyxl.utils import get_column_letter as L
        self.title(f'Thiết kế ô {L(c)}{r}')
        self.configure(bg=BG)
        self.transient(master)
        self.resizable(True, True)
        tk.Label(self, text=f'✏ Ô {L(c)}{r}   (vùng {self._area_txt(fm, r, c)})', bg=BG, font=FONT_T,
                 fg='#1f3a5f').grid(row=0, column=0, columnspan=6, sticky='w', padx=12, pady=(10, 4))
        tk.Label(self, text='Nội dung mẫu (chữ cố định trên form):', bg=BG, font=FONT_B).grid(
            row=1, column=0, columnspan=6, sticky='w', padx=12)
        self.txt = tk.Text(self, height=4, width=60, font=FONT, wrap='word', bg='#fffde7')
        self.txt.grid(row=2, column=0, columnspan=6, sticky='nsew', padx=12, pady=4)
        v = st.get('v')
        self.txt.insert('1.0', '' if v is None else FormModel.fmt_value(v, st.get('nf', '')))
        ref = st.get('ref')
        self.v_ref = tk.StringVar(value=(f'{L(ref[1])}{ref[0]}' if ref else ''))
        tk.Label(self, text='Lấy theo ô (VD C3, để trống = không):', bg=BG, font=FONT).grid(
            row=3, column=0, columnspan=3, sticky='w', padx=12)
        tk.Entry(self, textvariable=self.v_ref, width=8, font=FONT).grid(row=3, column=3, sticky='w')

        self.v_b = tk.BooleanVar(value=bool(st.get('b')))
        self.v_i = tk.BooleanVar(value=bool(st.get('i')))
        self.v_wrap = tk.BooleanVar(value=bool(st.get('wrap', 1)))
        self.v_sz = tk.StringVar(value=str(st.get('sz') or 12))
        self.v_al = tk.StringVar(value=st.get('al') or 'left')
        self.v_va = tk.StringVar(value=st.get('va') or 'center')
        self.fill = st.get('fill') or ''
        self.color = st.get('color') or '#000000'
        bd = st.get('bd') or [0, 0, 0, 0]
        self.v_bd = [tk.BooleanVar(value=bool(x)) for x in bd]
        f = tk.Frame(self, bg=BG)
        f.grid(row=4, column=0, columnspan=6, sticky='w', padx=12, pady=6)
        tk.Checkbutton(f, text='Đậm', variable=self.v_b, bg=BG, font=FONT_B).pack(side='left')
        tk.Checkbutton(f, text='Nghiêng', variable=self.v_i, bg=BG, font=('Segoe UI', 10, 'italic')).pack(side='left')
        tk.Checkbutton(f, text='Xuống dòng', variable=self.v_wrap, bg=BG, font=FONT).pack(side='left')
        tk.Label(f, text='  Cỡ chữ:', bg=BG, font=FONT).pack(side='left')
        ttk.Combobox(f, textvariable=self.v_sz, width=5,
                     values=['8', '9', '10', '11', '12', '13', '14', '16', '18', '20', '24']).pack(side='left')
        f2 = tk.Frame(self, bg=BG)
        f2.grid(row=5, column=0, columnspan=6, sticky='w', padx=12, pady=2)
        tk.Label(f2, text='Căn ngang:', bg=BG, font=FONT).pack(side='left')
        ttk.Combobox(f2, textvariable=self.v_al, width=8, state='readonly',
                     values=['left', 'center', 'right']).pack(side='left', padx=(2, 10))
        tk.Label(f2, text='Căn dọc:', bg=BG, font=FONT).pack(side='left')
        ttk.Combobox(f2, textvariable=self.v_va, width=8, state='readonly',
                     values=['top', 'center', 'bottom']).pack(side='left', padx=2)
        f3 = tk.Frame(self, bg=BG)
        f3.grid(row=6, column=0, columnspan=6, sticky='w', padx=12, pady=4)
        self.b_fill = tk.Button(f3, text='Màu nền', font=FONT, relief='groove', command=self._pick_fill)
        self.b_fill.pack(side='left')
        tk.Button(f3, text='✕', font=FONT, relief='flat', command=self._no_fill).pack(side='left', padx=(0, 10))
        self.b_col = tk.Button(f3, text='Màu chữ', font=FONT, relief='groove', command=self._pick_color)
        self.b_col.pack(side='left')
        tk.Label(f3, text='   Viền:', bg=BG, font=FONT).pack(side='left')
        for i, t in enumerate(('Trái', 'Trên', 'Phải', 'Dưới')):
            tk.Checkbutton(f3, text=t, variable=self.v_bd[i], bg=BG, font=FONT).pack(side='left')
        self._paint_btns()
        tk.Label(self, text='Danh sách chọn (mỗi dòng 1 mục, để trống = ô nhập tự do):', bg=BG,
                 font=FONT_B).grid(row=7, column=0, columnspan=6, sticky='w', padx=12, pady=(6, 0))
        self.lst = tk.Text(self, height=4, width=60, font=FONT)
        self.lst.grid(row=8, column=0, columnspan=6, sticky='nsew', padx=12, pady=4)
        self.lst.insert('1.0', '\n'.join((form.get('lists') or {}).get(key(r, c), [])))
        self.v_apply_area = tk.BooleanVar(value=False)
        tk.Checkbutton(self, text='Áp dụng định dạng (đậm, cỡ, căn, màu, viền) cho CẢ VÙNG đang chọn',
                       variable=self.v_apply_area, bg=BG, font=FONT).grid(row=9, column=0, columnspan=6,
                                                                          sticky='w', padx=12)
        b = tk.Frame(self, bg=BG)
        b.grid(row=10, column=0, columnspan=6, sticky='e', padx=12, pady=10)
        tk.Button(b, text='💾 Áp dụng', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat', padx=14,
                  command=self._ok).pack(side='right')
        tk.Button(b, text='Hủy', font=FONT, relief='flat', padx=10, command=self.destroy).pack(side='right', padx=6)
        self.columnconfigure(0, weight=1)
        self.rowconfigure(2, weight=1)
        self.bind('<Escape>', lambda e: self.destroy())
        self.txt.focus_set()
        self.grab_set()

    @staticmethod
    def _area_txt(fm, r, c):
        from openpyxl.utils import get_column_letter as L
        r0, c0, r1, c1 = fm.area(r, c)
        return f'{L(c0)}{r0}' + (f':{L(c1)}{r1}' if (r0, c0) != (r1, c1) else '')

    def _paint_btns(self):
        self.b_fill.config(bg=self.fill or '#ffffff')
        self.b_col.config(fg=self.color or '#000000')

    def _pick_fill(self):
        c = colorchooser.askcolor(self.fill or '#ffffff', parent=self)[1]
        if c:
            self.fill = c.upper()
            self._paint_btns()

    def _no_fill(self):
        self.fill = ''
        self._paint_btns()

    def _pick_color(self):
        c = colorchooser.askcolor(self.color, parent=self)[1]
        if c:
            self.color = c.upper()
            self._paint_btns()

    def _ok(self):
        from openpyxl.utils import column_index_from_string as CI
        text = self.txt.get('1.0', 'end').rstrip('\n')
        ref = None
        m = re.fullmatch(r'\s*([A-Za-z]{1,3})(\d+)\s*', self.v_ref.get() or '')
        if m:
            ref = [int(m.group(2)), CI(m.group(1).upper())]
        try:
            sz = float(self.v_sz.get())
        except ValueError:
            sz = 12.0
        style = {'b': int(self.v_b.get()), 'i': int(self.v_i.get()), 'sz': sz, 'color': self.color,
                 'al': self.v_al.get(), 'va': self.v_va.get(), 'wrap': int(self.v_wrap.get()),
                 'fill': self.fill, 'bd': [int(v.get()) for v in self.v_bd]}
        items = [x.strip() for x in self.lst.get('1.0', 'end').split('\n') if x.strip()]
        self.destroy()
        self.on_ok(text, ref, style, items, self.v_apply_area.get())


# =====================================================================
#  Hop thoai: Form moi / Nhap tu Excel / Doi ten
# =====================================================================
class FormInfoDialog(tk.Toplevel):
    def __init__(self, master, title, groups, name='', group='', mode='new', on_ok=None):
        super().__init__(master)
        self.on_ok, self.mode = on_ok, mode
        self.title(title)
        self.configure(bg=BG)
        self.transient(master)
        tk.Label(self, text=title, bg=BG, font=FONT_T, fg='#1f3a5f').grid(row=0, column=0, columnspan=3,
                                                                         sticky='w', padx=12, pady=(10, 6))
        tk.Label(self, text='Tên form (loại báo cáo):', bg=BG, font=FONT_B).grid(row=1, column=0, sticky='e', padx=12)
        self.v_name = tk.StringVar(value=name)
        e = tk.Entry(self, textvariable=self.v_name, width=46, font=FONT)
        e.grid(row=1, column=1, columnspan=2, sticky='we', padx=(0, 12), pady=3)
        tk.Label(self, text='Nhóm (cách nhau bằng " > "):', bg=BG, font=FONT_B).grid(row=2, column=0, sticky='e', padx=12)
        self.v_group = tk.StringVar(value=group)
        ttk.Combobox(self, textvariable=self.v_group, values=groups, width=44, font=FONT).grid(
            row=2, column=1, columnspan=2, sticky='we', padx=(0, 12), pady=3)
        r = 3
        self.v_base = tk.StringVar(value='copy')
        if mode == 'new':
            tk.Label(self, text='Tạo từ:', bg=BG, font=FONT_B).grid(row=r, column=0, sticky='e', padx=12)
            fb = tk.Frame(self, bg=BG)
            fb.grid(row=r, column=1, columnspan=2, sticky='w')
            tk.Radiobutton(fb, text='Nhân bản form đang mở', variable=self.v_base, value='copy', bg=BG,
                           font=FONT).pack(side='left')
            tk.Radiobutton(fb, text='Form trắng', variable=self.v_base, value='blank', bg=BG,
                           font=FONT).pack(side='left', padx=8)
            r += 1
        if mode == 'excel':
            tk.Label(self, text='File Excel:', bg=BG, font=FONT_B).grid(row=r, column=0, sticky='e', padx=12)
            self.v_path = tk.StringVar()
            tk.Entry(self, textvariable=self.v_path, width=40, font=FONT).grid(row=r, column=1, sticky='we', pady=3)
            tk.Button(self, text='Chọn…', font=FONT, command=self._pick).grid(row=r, column=2, padx=(4, 12))
            r += 1
            tk.Label(self, text='Sheet:', bg=BG, font=FONT_B).grid(row=r, column=0, sticky='e', padx=12)
            self.v_sheet = tk.StringVar()
            self.cb_sheet = ttk.Combobox(self, textvariable=self.v_sheet, width=30, font=FONT, state='readonly')
            self.cb_sheet.grid(row=r, column=1, sticky='w', pady=3)
            r += 1
            tk.Label(self, text='Vùng (VD A1:H39, trống = cả sheet):', bg=BG, font=FONT_B).grid(
                row=r, column=0, sticky='e', padx=12)
            self.v_rng = tk.StringVar()
            tk.Entry(self, textvariable=self.v_rng, width=14, font=FONT).grid(row=r, column=1, sticky='w', pady=3)
            r += 1
            tk.Label(self, text='Hỗ trợ .xlsx / .xlsm. File .xls/.xlsb: mở bằng Excel → Save As .xlsx trước.',
                     bg=BG, fg='#666', font=('Segoe UI', 9)).grid(row=r, column=0, columnspan=3, sticky='w', padx=12)
            r += 1
        bf = tk.Frame(self, bg=BG)
        bf.grid(row=r, column=0, columnspan=3, sticky='e', padx=12, pady=10)
        tk.Button(bf, text='✔ Đồng ý', font=FONT_B, bg='#2d6a4f', fg='white', relief='flat', padx=14,
                  command=self._ok).pack(side='right')
        tk.Button(bf, text='Hủy', font=FONT, relief='flat', padx=10, command=self.destroy).pack(side='right', padx=6)
        self.columnconfigure(1, weight=1)
        e.focus_set()
        self.bind('<Return>', lambda ev: self._ok())
        self.bind('<Escape>', lambda ev: self.destroy())
        self.grab_set()

    def _pick(self):
        p = filedialog.askopenfilename(parent=self, filetypes=[('Excel', '*.xlsx *.xlsm'), ('Tất cả', '*.*')])
        if not p:
            return
        self.v_path.set(p)
        try:
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True)
            self.cb_sheet['values'] = wb.sheetnames
            self.v_sheet.set(wb.sheetnames[0])
            wb.close()
        except Exception as e:
            messagebox.showerror('Lỗi', f'Không đọc được file:\n{e}', parent=self)
        if not self.v_name.get().strip():
            self.v_name.set(os.path.splitext(os.path.basename(p))[0])

    def _ok(self):
        name = self.v_name.get().strip()
        if not name:
            messagebox.showwarning('Thiếu', 'Nhập tên form.', parent=self)
            return
        group = [g.strip() for g in self.v_group.get().split('>') if g.strip()] or ['Form tự tạo']
        extra = {'base': self.v_base.get()}
        if self.mode == 'excel':
            if not self.v_path.get() or not self.v_sheet.get():
                messagebox.showwarning('Thiếu', 'Chọn file Excel và sheet.', parent=self)
                return
            extra.update(path=self.v_path.get(), sheet=self.v_sheet.get(), rng=self.v_rng.get().strip())
        self.destroy()
        self.on_ok(name, group, extra)


# =====================================================================
#  TAB CHINH
# =====================================================================
class ReportFormTab(tk.Frame):
    def __init__(self, master, app=None, admin=True, user_name=''):
        super().__init__(master, bg=BG)
        self.app = app
        self.admin = admin
        self.user_name = user_name or (getattr(app, 'user', {}) or {}).get('hoten') or ''
        self.forms = rio.FormStore()
        self.reports = rio.ReportStore()
        self.settings = rio.load_settings()
        self.rec = None             # bao cao dang mo
        self.tpl = None             # form dang THIET KE (ban sao)
        self.dirty = False
        self.tpl_dirty = False
        self.sel = None             # (r0,c0,r1,c1) vung chon
        self.zoom = float(self.settings.get('zoom', 0.8))
        self.design = tk.BooleanVar(value=False)
        self._undo = []
        self._imgcache = {}
        self._editor = None
        self._cellbuf = None        # bo nho Copy/Paste dinh dang o (thiet ke form)
        self._imgbuf = None         # bo nho Cat/Dan anh (sua cho anh add nham)
        self._build()
        self._fill_tree()
        self._fill_reports()
        last = self.settings.get('last_form')
        if last and self.forms.get(last):
            self._select_form_in_tree(last)
            self.new_report(last, ask=False)
        elif self.forms.forms:
            self.new_report(self.forms.forms[0]['id'], ask=False)

    # ------------------------------------------------------------ giao dien
    def _apply_style(self):
        """Tao style rieng (ten 'RF.*') cho cac widget ttk trong tab nay.
        Dung ten rieng de KHONG anh huong cac tab/widget khac cua app."""
        try:
            st = ttk.Style(self)
        except Exception:
            return
        # Danh sach / cay: chu ro rang, hang cao thoang, chon mau xanh chinh
        st.configure('RF.Treeview', font=FONT, rowheight=26,
                     background=CLR['surface'], fieldbackground=CLR['surface'],
                     borderwidth=0, relief='flat')
        st.map('RF.Treeview',
               background=[('selected', CLR['primary'])],
               foreground=[('selected', '#ffffff')])
        st.configure('RF.Treeview.Heading', font=FONT_B,
                     background=CLR['panel'], foreground=CLR['ink'],
                     relief='flat', padding=(6, 5))
        st.map('RF.Treeview.Heading', background=[('active', CLR['border'])])

    def _load_logo(self, target_h):
        """Nap logo cong ty (logo.png / icon) va thu nho vua chieu cao header."""
        for name in ('logo.png', 'hitechplast_icon_256.png', 'hitechplast.ico'):
            for base in (rio.resource_dir(), rio.writable_dir()):
                p = os.path.join(base, name)
                if not os.path.exists(p):
                    continue
                try:
                    img = tk.PhotoImage(file=p)
                    f = max(1, round(img.height() / target_h))
                    if f > 1:
                        img = img.subsample(f, f)
                    return img
                except Exception:
                    continue
        return None

    def _build_header(self):
        """Dai header: logo + ten cong ty + ten bao cao (dong dau tab)."""
        company = (self.settings.get('company_name') or 'HITECH PLAST')
        hdr = tk.Frame(self, bg=CLR['header1'], height=58)
        hdr.pack(side='top', fill='x')
        hdr.pack_propagate(False)

        self._logo_img = self._load_logo(42)
        if self._logo_img is not None:
            tk.Label(hdr, image=self._logo_img, bg=CLR['header1']).pack(
                side='left', padx=(14, 12))

        box = tk.Frame(hdr, bg=CLR['header1'])
        box.pack(side='left', fill='y')
        tk.Label(box, text=company, bg=CLR['header1'], fg='#ffffff',
                 font=FONT_H).pack(anchor='w', pady=(9, 0))
        tk.Label(box, text='BÁO CÁO CHẤT LƯỢNG NGUYÊN VẬT LIỆU',
                 bg=CLR['header1'], fg=CLR['header_sub'],
                 font=FONT_SUB).pack(anchor='w')

        # dai mau nhan mong duoi header
        tk.Frame(self, bg=CLR['accent'], height=3).pack(side='top', fill='x')

    def _build(self):
        self._apply_style()
        self._build_header()
        pw = ttk.PanedWindow(self, orient='horizontal')
        pw.pack(fill='both', expand=True)
        left = tk.Frame(pw, bg=BG, width=330)
        right = tk.Frame(pw, bg='#e9edf2')
        pw.add(left, weight=0)
        pw.add(right, weight=1)
        self.after(50, lambda: self._init_sash(pw))

        # --- nut thao tac bao cao da luu: pack o DAY truoc tien de luon hien,
        #     ke ca khi cua so thap (cac danh sach phia tren se tu co lai)
        bl = tk.Frame(left, bg=BG)
        bl.pack(side='bottom', fill='x', padx=8, pady=(2, 8))
        for i, (t, cmd) in enumerate((('📂 Mở', self.open_selected), ('⧉ Nhân bản', self.duplicate_selected),
                                      ('📄 PDF gộp', self.pdf_selected), ('🗑 Xoá', self.delete_selected))):
            tk.Button(bl, text=t, font=FONT, relief='groove', pady=2, command=cmd).grid(
                row=0, column=i, sticky='ew', padx=1)
            bl.grid_columnconfigure(i, weight=1, uniform='blbtn')

        # --- chon loai bao cao
        tk.Label(left, text='① Chọn loại báo cáo', bg=BG, font=FONT_T, fg='#1f3a5f').pack(anchor='w', padx=8, pady=(8, 2))
        sf = tk.Frame(left, bg=BG)
        sf.pack(fill='x', padx=8)
        tk.Label(sf, text='🔍', bg=BG).pack(side='left')
        self.v_fsearch = tk.StringVar()
        es = tk.Entry(sf, textvariable=self.v_fsearch, font=FONT)
        es.pack(side='left', fill='x', expand=True)
        self.v_fsearch.trace_add('write', lambda *a: self._fill_tree())
        tf = tk.Frame(left, bg=BG)
        tf.pack(fill='both', expand=True, padx=8, pady=4)
        self.tree = ttk.Treeview(tf, show='tree', selectmode='browse', height=8,
                                 style='RF.Treeview')
        self.tree.column('#0', width=260)
        ys = ttk.Scrollbar(tf, orient='vertical', command=self.tree.yview)
        self.tree.configure(yscrollcommand=ys.set)
        self.tree.pack(side='left', fill='both', expand=True)
        ys.pack(side='right', fill='y')
        self.tree.bind('<Double-1>', lambda e: self._tree_open())
        self.tree.bind('<Return>', lambda e: self._tree_open())
        self.tree.bind('<Button-3>', self._tree_menu)
        bt = tk.Frame(left, bg=BG)
        bt.pack(fill='x', padx=8)
        tk.Button(bt, text='📄 Mở báo cáo mới loại này', font=FONT_B, bg='#1d6fa5', fg='white', relief='flat',
                  command=self._tree_open).pack(side='left', fill='x', expand=True)
        if self.admin:
            mb = tk.Menubutton(bt, text='⚙ Form ▾', font=FONT_B, bg='#e07a00', fg='white', relief='flat',
                               padx=8)
            m = tk.Menu(mb, tearoff=0)
            m.add_command(label='➕ Thêm form mới…', command=self.form_new)
            m.add_command(label='📥 Nhập form từ file Excel…', command=self.form_from_excel)
            m.add_command(label='✏ Đổi tên / nhóm form đang chọn…', command=self.form_rename)
            m.add_command(label='⧉ Nhân bản form đang chọn', command=self.form_duplicate)
            m.add_command(label='⬆ Đưa lên', command=lambda: self._form_move(-1))
            m.add_command(label='⬇ Đưa xuống', command=lambda: self._form_move(1))
            m.add_separator()
            m.add_command(label='🗑 Xoá form đang chọn', command=self.form_delete)
            m.add_separator()
            m.add_command(label='📂 Thư mục lưu báo cáo…', command=self.pick_storage)
            m.add_command(label='↺ Khôi phục bộ form gốc (từ Excel)', command=self.form_reset_all)
            mb['menu'] = m
            mb.pack(side='left', padx=(4, 0))

        # --- bao cao da luu
        tk.Label(left, text='② Báo cáo đã lưu', bg=BG, font=FONT_T, fg='#1f3a5f').pack(anchor='w', padx=8, pady=(10, 2))
        rf = tk.Frame(left, bg=BG)
        rf.pack(fill='x', padx=8)
        tk.Label(rf, text='🔍', bg=BG).pack(side='left')
        self.v_rsearch = tk.StringVar()
        tk.Entry(rf, textvariable=self.v_rsearch, font=FONT).pack(side='left', fill='x', expand=True)
        self.v_rsearch.trace_add('write', lambda *a: self._fill_reports())
        self.v_only = tk.BooleanVar(value=False)
        tk.Checkbutton(left, text='Chỉ loại đang chọn', variable=self.v_only, bg=BG, font=FONT,
                       command=self._fill_reports).pack(anchor='w', padx=8)
        lf = tk.Frame(left, bg=BG)
        lf.pack(fill='both', expand=True, padx=8, pady=4)
        self.lst = ttk.Treeview(lf, columns=('title', 'form', 'date', 'user'), show='headings',
                                selectmode='extended', height=5, style='RF.Treeview')
        for c, t, w in (('title', 'Tên báo cáo', 130), ('form', 'Loại', 70), ('date', 'Ngày', 70),
                        ('user', 'Người lập', 60)):
            self.lst.heading(c, text=t)
            self.lst.column(c, width=w, stretch=(c == 'title'))
        # hang le/chan khac mau (zebra) cho de doc
        self.lst.tag_configure('odd', background=CLR['surface'])
        self.lst.tag_configure('even', background=CLR['row_alt'])
        ys2 = ttk.Scrollbar(lf, orient='vertical', command=self.lst.yview)
        self.lst.configure(yscrollcommand=ys2.set)
        self.lst.pack(side='left', fill='both', expand=True)
        ys2.pack(side='right', fill='y')
        self.lst.bind('<Double-1>', lambda e: self.open_selected())
        self.lst.bind('<Return>', lambda e: self.open_selected())
        self.lst.bind('<Delete>', lambda e: self.delete_selected())

        # --- thanh cong cu ben phai
        tb = tk.Frame(right, bg=CLR['toolbar'])
        tb.pack(fill='x')

        def btn(t, cmd, bg=CLR['surface'], fg=CLR['ink'], bold=False):
            b = tk.Button(tb, text=t, command=cmd, font=FONT_B if bold else FONT, bg=bg, fg=fg,
                          relief='flat', bd=0, padx=10, pady=5, cursor='hand2',
                          activebackground='#cfe0f5')
            b.pack(side='left', padx=2, pady=5)
            return b

        def sep():
            tk.Frame(tb, width=1, bg=CLR['border']).pack(side='left', fill='y', padx=5, pady=6)

        btn('📄 Mới', lambda: self.new_report(self._cur_form_id()))
        btn('💾 Lưu', self.save, bg=CLR['success'], fg='white', bold=True)
        btn('💾 Lưu bản mới', lambda: self.save(as_new=True))
        sep()
        btn('🖼 Chèn ảnh', self.insert_image_file)
        btn('📋 Dán ảnh', self.paste_image)
        btn('✖ Xoá ảnh', self.delete_image)
        btn('↶ Hoàn tác', self.undo)
        sep()
        btn('⚖ Kết luận', self.conclude, bg='#5a4b8a', fg='white', bold=True)
        sep()
        btn('📄 PDF', self.export_pdf, bg=CLR['danger'], fg='white', bold=True)
        btn('📊 Excel', self.export_excel, bg='#1e7145', fg='white', bold=True)
        btn('🖨 In', self.print_report, bg='#1d3557', fg='white', bold=True)
        self.v_zoom = tk.StringVar(value=f'{int(self.zoom * 100)}%')
        cz = ttk.Combobox(tb, textvariable=self.v_zoom, width=6, state='readonly',
                          values=['50%', '60%', '70%', '80%', '90%', '100%', '120%', '150%'])
        cz.pack(side='right', padx=6)
        cz.bind('<<ComboboxSelected>>', lambda e: self._set_zoom(int(self.v_zoom.get()[:-1]) / 100))
        tk.Label(tb, text='Zoom', bg=CLR['toolbar'], font=FONT).pack(side='right')

        # --- thong tin bao cao
        info = tk.Frame(right, bg='#e9edf2')
        info.pack(fill='x', padx=6, pady=(4, 0))
        self.lbl_form = tk.Label(info, text='', bg='#e9edf2', font=FONT_T, fg='#1f3a5f')
        self.lbl_form.pack(side='left')
        tk.Label(info, text='   Tên báo cáo:', bg='#e9edf2', font=FONT_B).pack(side='left')
        self.v_title = tk.StringVar()
        self.e_title = tk.Entry(info, textvariable=self.v_title, font=FONT, width=26)
        self.e_title.pack(side='left', padx=4)
        self.v_title.trace_add('write', self._title_changed)
        self.v_autopdf = tk.BooleanVar(value=bool(self.settings.get('auto_pdf', True)))
        tk.Checkbutton(info, text='Tự xuất PDF khi lưu', variable=self.v_autopdf, bg='#e9edf2', font=FONT,
                       command=self._save_settings).pack(side='left', padx=8)
        srow = tk.Frame(right, bg='#e9edf2')
        srow.pack(fill='x', padx=6)
        if self.admin:
            tk.Checkbutton(srow, text='✏ Thiết kế form', variable=self.design, font=FONT_B, bg='#ffe8cc',
                           activebackground='#ffd8a8', relief='groove', padx=6,
                           command=self._toggle_design).pack(side='right', padx=2, pady=2)
        self.lbl_state = tk.Label(srow, text='', bg='#e9edf2', font=('Segoe UI', 9, 'bold'), fg='#b35c00',
                                  anchor='w')
        self.lbl_state.pack(side='left', fill='x', expand=True, padx=2)

        # --- thanh thiet ke (an khi khong o che do thiet ke)
        self.dbar = tk.Frame(right, bg='#ffe8cc')
        row1 = tk.Frame(self.dbar, bg='#ffe8cc')
        row1.pack(fill='x')
        row2 = tk.Frame(self.dbar, bg='#ffe8cc')
        row2.pack(fill='x')
        for fr, items in ((row1, (('✏ Sửa ô', self.d_edit_cell), ('⊞ Gộp ô', self.d_merge),
                                  ('⊟ Tách ô', self.d_unmerge),
                                  ('➕ Hàng trên', lambda: self.d_rows(0, 1)), ('➕ Hàng dưới', lambda: self.d_rows(1, 1)),
                                  ('➖ Xoá hàng', lambda: self.d_rows(0, -1)),
                                  ('➕ Cột trái', lambda: self.d_cols(0, 1)), ('➕ Cột phải', lambda: self.d_cols(1, 1)),
                                  ('➖ Xoá cột', lambda: self.d_cols(0, -1)))),
                          (row2, (('↕ Cao hàng', self.d_row_height), ('↔ Rộng cột', self.d_col_width),
                                  ('🖼 Ảnh mẫu', self.d_tpl_image), ('✖ Xoá ảnh mẫu', self.d_del_tpl_image),
                                  ('⧉ Copy ô', self.d_copy_cell), ('📋 Paste ô', self.d_paste_cell),
                                  ('🏷 Đặt ô Tên NVL', self.d_name_cell), ('📅 Đặt ô Ngày trả KQ', self.d_date_cell),
                                  ('🧾 Đặt ô Kết luận', self.d_concl_cell)))):
            for t, cmd in items:
                tk.Button(fr, text=t, font=('Segoe UI', 9), relief='groove', bg='white',
                          command=cmd).pack(side='left', padx=1, pady=2)
        tk.Button(row2, text='💾 Lưu mẫu form', font=FONT_B, bg='#e07a00', fg='white', relief='flat',
                  command=self.d_save).pack(side='right', padx=4, pady=2)
        tk.Button(row2, text='Bỏ thay đổi', font=('Segoe UI', 9), relief='groove',
                  command=self.d_cancel).pack(side='right', padx=2)

        # --- canvas
        self.hint = tk.Label(right, text='Bấm ô để chọn • Nhấp đúp / gõ phím để nhập • Enter lưu ô, Shift+Enter xuống '
                                         'dòng • Ctrl+V dán ảnh vào ô • Shift+bấm = chọn vùng • Chuột phải = menu',
                             bg='#e9edf2', fg='#555', font=('Segoe UI', 9), anchor='w')
        self.hint.pack(fill='x', padx=6)
        cf = tk.Frame(right, bg='#e9edf2')
        cf.pack(fill='both', expand=True)
        self.cv = tk.Canvas(cf, bg=CLR['canvas'], highlightthickness=0, takefocus=1)
        vs = ttk.Scrollbar(cf, orient='vertical', command=self.cv.yview)
        hs = ttk.Scrollbar(right, orient='horizontal', command=self.cv.xview)
        self.cv.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        vs.pack(side='right', fill='y')
        self.cv.pack(side='left', fill='both', expand=True)
        hs.pack(fill='x')
        self.cv.bind('<Button-1>', self._click)
        self.cv.bind('<Shift-Button-1>', lambda e: self._click(e, extend=True))
        self.cv.bind('<B1-Motion>', self._drag)
        self.cv.bind('<Double-1>', self._dbl)
        self.cv.bind('<Button-3>', self._ctx_menu)
        self.cv.bind('<Key>', self._key)
        self.cv.bind('<Enter>', self._bind_wheel)
        self.cv.bind('<Leave>', self._unbind_wheel)
        for k in ('<Control-v>', '<Control-V>'):
            self.cv.bind(k, lambda e: (self.paste_image(from_key=True), 'break')[1])
        for k in ('<Control-s>', '<Control-S>'):
            self.cv.bind(k, lambda e: (self.save(), 'break')[1])
            self.bind_all(k, lambda e: (self.save(), 'break')[1] if self.winfo_ismapped() else None, add='+')
        for k in ('<Control-z>', '<Control-Z>'):
            self.cv.bind(k, lambda e: (self.undo(), 'break')[1])
        self.cv.bind('<Control-c>', lambda e: self._copy_text())
        # keo tha file anh (neu co windnd tren Windows)
        try:
            import windnd  # noqa
            windnd.hook_dropfiles(self.cv, func=self._on_drop)
        except Exception:
            pass

    @staticmethod
    def _init_sash(pw):
        try:
            pw.sashpos(0, 300)
        except Exception:
            pass

    # ------------------------------------------------------------ cay form
    def _fill_tree(self):
        q = (self.v_fsearch.get() if hasattr(self, 'v_fsearch') else '').strip().lower()
        self.tree.delete(*self.tree.get_children())
        nodes = {}
        for f in self.forms.forms:
            path = f.get('group') or []
            label = f.get('name', '')
            full = ' '.join(path + [label]).lower()
            if q and all(w not in full for w in q.split()):
                continue
            parent = ''
            acc = ()
            for g in path:
                acc = acc + (g,)
                if acc not in nodes:
                    nodes[acc] = self.tree.insert(parent, 'end', text='📁 ' + g, open=bool(q))
                parent = nodes[acc]
            self.tree.insert(parent, 'end', iid='F::' + f['id'], text='📄 ' + label)

    def _select_form_in_tree(self, fid):
        iid = 'F::' + fid
        if self.tree.exists(iid):
            p = self.tree.parent(iid)
            while p:
                self.tree.item(p, open=True)
                p = self.tree.parent(p)
            self.tree.selection_set(iid)
            self.tree.see(iid)

    def _tree_fid(self):
        s = self.tree.selection()
        if s and s[0].startswith('F::'):
            return s[0][3:]
        return None

    def _cur_form_id(self):
        if self.rec:
            return self.rec.get('form_id')
        return self._tree_fid()

    def _tree_open(self):
        fid = self._tree_fid()
        if not fid:
            s = self.tree.selection()
            if s:
                self.tree.item(s[0], open=not self.tree.item(s[0], 'open'))
            return
        self.new_report(fid)

    def _tree_menu(self, ev):
        iid = self.tree.identify_row(ev.y)
        if not iid:
            return
        self.tree.selection_set(iid)
        if not iid.startswith('F::'):
            return
        m = tk.Menu(self, tearoff=0)
        m.add_command(label='📄 Mở báo cáo mới loại này', command=self._tree_open)
        if self.admin:
            m.add_separator()
            m.add_command(label='✏ Đổi tên / nhóm…', command=self.form_rename)
            m.add_command(label='⧉ Nhân bản form', command=self.form_duplicate)
            m.add_command(label='⬆ Đưa lên', command=lambda: self._form_move(-1))
            m.add_command(label='⬇ Đưa xuống', command=lambda: self._form_move(1))
            m.add_command(label='🗑 Xoá form', command=self.form_delete)
        m.tk_popup(ev.x_root, ev.y_root)

    # ------------------------------------------------------------ danh sach bao cao
    def _fill_reports(self):
        if not hasattr(self, 'lst'):
            return
        q = self.v_rsearch.get().strip().lower()
        only = self.v_only.get()
        fid = self._cur_form_id()
        self.lst.delete(*self.lst.get_children())
        _i = 0
        for r in self.reports.list():
            if only and r.get('form_id') != fid:
                continue
            txt = ' '.join(str(r.get(k) or '') for k in ('title', 'form_name', 'user', 'created', 'material')).lower()
            if q and any(w not in txt for w in q.split()):
                continue
            d = (r.get('updated') or r.get('created') or '')[:16]
            self.lst.insert('', 'end', iid=r['id'], values=(r.get('title'), r.get('form_name'), d, r.get('user')),
                            tags=('even' if _i % 2 else 'odd',))
            _i += 1

    def _sel_report_ids(self):
        return list(self.lst.selection())

    def open_selected(self):
        ids = self._sel_report_ids()
        if not ids or not self._confirm_discard():
            return
        try:
            self.rec = self.reports.load(ids[0])
        except Exception as e:
            messagebox.showerror('Lỗi', f'Không mở được báo cáo:\n{e}', parent=self)
            return
        self._after_open()

    def duplicate_selected(self):
        ids = self._sel_report_ids()
        if not ids or not self._confirm_discard():
            return
        rec = self.reports.load(ids[0])
        for k in ('id', 'created', 'updated'):
            rec.pop(k, None)
        rec['title'] = (rec.get('title') or '') + ' (bản sao)'
        rec['user'] = self.user_name
        self.rec = rec
        self._after_open(dirty=True)

    def delete_selected(self):
        ids = self._sel_report_ids()
        if not ids:
            return
        if not messagebox.askyesno('Xoá báo cáo', f'Xoá {len(ids)} báo cáo đã chọn?\n(Không khôi phục được)',
                                   parent=self):
            return
        for i in ids:
            self.reports.delete(i)
            if self.rec and self.rec.get('id') == i:
                self.rec['id'] = None
                self.dirty = True
        self._fill_reports()
        self._update_state()

    def pdf_selected(self):
        ids = self._sel_report_ids()
        if not ids:
            messagebox.showinfo('Chọn báo cáo', 'Chọn 1 hoặc nhiều báo cáo trong danh sách (Ctrl/Shift + bấm).',
                                parent=self)
            return
        p = filedialog.asksaveasfilename(parent=self, defaultextension='.pdf', filetypes=[('PDF', '*.pdf')],
                                         initialfile=f'BaoCao_NVL_{len(ids)}_ban.pdf')
        if not p:
            return
        try:
            rio.export_pdf_many([self.reports.load(i) for i in ids], p)
        except Exception as e:
            messagebox.showerror('Lỗi xuất PDF', str(e), parent=self)
            return
        if messagebox.askyesno('Xong', f'Đã xuất {len(ids)} báo cáo:\n{p}\n\nMở ngay?', parent=self):
            _open_file(p)

    # ------------------------------------------------------------ bao cao
    def _confirm_discard(self):
        if self._editor:
            self._commit_editor()
        if self.design.get() and self.tpl_dirty:
            if not messagebox.askyesno('Chưa lưu mẫu', 'Mẫu form đang sửa chưa lưu. Bỏ thay đổi?', parent=self):
                return False
            self.design.set(False)
            self._toggle_design(silent=True)
        if self.dirty and self.rec:
            a = messagebox.askyesnocancel('Chưa lưu', 'Báo cáo đang mở chưa lưu. Lưu lại trước?', parent=self)
            if a is None:
                return False
            if a:
                return self.save()
        return True

    def new_report(self, fid, ask=True):
        if not fid:
            messagebox.showinfo('Chọn loại', 'Chọn 1 loại báo cáo trong cây bên trái.', parent=self)
            return
        f = self.forms.get(fid)
        if not f:
            return
        if ask and not self._confirm_discard():
            return
        if self.design.get():
            self.design.set(False)
            self._toggle_design(silent=True)
        self.rec = {'id': None, 'form_id': fid, 'form_name': f.get('name'), 'group': f.get('group'),
                    'title': '', 'user': self.user_name, 'form': copy.deepcopy(f), 'values': {}, 'images': [],
                    'hide_tpl': []}
        self.settings['last_form'] = fid
        self._save_settings()
        self._after_open(dirty=False)

    def _after_open(self, dirty=False):
        self._undo.clear()
        self.sel = None
        self.dirty = dirty
        self._title_lock = True
        self.v_title.set(self.rec.get('title') or '')
        self._title_lock = False
        self.lbl_form.config(text='📑 ' + ' › '.join((self.rec.get('group') or [])[-1:] +
                                                     [self.rec.get('form_name') or '']))
        self._select_form_in_tree(self.rec.get('form_id') or '')
        self._fill_reports()
        self.redraw()
        self.cv.yview_moveto(0)
        self.cv.xview_moveto(0)
        self._update_state()
        self.cv.focus_set()

    def _title_changed(self, *a):
        if getattr(self, '_title_lock', False) or not self.rec:
            return
        self.rec['title'] = self.v_title.get()
        self.rec['title_manual'] = True
        self._mark()

    def _auto_title(self):
        """Ten bao cao = Ten NVL + ngay (giong ThemSheet_RMQR: str & Format(Now,'dd mm yy'))."""
        if not self.rec or self.rec.get('title_manual'):
            return
        f = self.rec['form']
        nc = f.get('name_cell')
        mat = FormModel(f).text(*nc, self.rec['values']).strip() if nc else ''
        t = (mat or self.rec.get('form_name') or 'Báo cáo') + ' ' + rio.now_str('%d %m %y')
        self._title_lock = True
        self.v_title.set(t)
        self._title_lock = False
        self.rec['title'] = t
        self.rec['material'] = mat

    def _mark(self):
        self.dirty = True
        self._update_state()

    def _update_state(self):
        if self.design.get():
            self.lbl_state.config(text='● ĐANG THIẾT KẾ MẪU FORM' + (' (chưa lưu)' if self.tpl_dirty else ''),
                                  fg='#b35c00')
        elif self.rec:
            if not self.rec.get('id'):
                s = '● Báo cáo mới – chưa lưu' if self.dirty else '○ Báo cáo mới (trống)'
            else:
                s = ('● Có thay đổi chưa lưu' if self.dirty else '✔ Đã lưu') + f"   •   Mã: {self.rec['id']}"
            self.lbl_state.config(text=s, fg='#b35c00' if (self.dirty or not self.rec.get('id')) else '#2d6a4f')

    def save(self, as_new=False):
        if self._editor:
            self._commit_editor()
        if self.design.get():
            return self.d_save()
        if not self.rec:
            return False
        if as_new:
            self.rec = copy.deepcopy(self.rec)
            for k in ('id', 'created', 'updated'):
                self.rec.pop(k, None)
        if not (self.rec.get('title') or '').strip():
            self._auto_title()
        self.rec['user'] = self.rec.get('user') or self.user_name
        try:
            self.reports.save(self.rec)
        except Exception as e:
            messagebox.showerror('Lỗi lưu', f'Không lưu được báo cáo:\n{e}', parent=self)
            return False
        self.dirty = False
        self._fill_reports()
        if self.lst.exists(self.rec['id']):
            self.lst.selection_set(self.rec['id'])
            self.lst.see(self.rec['id'])
        msg = f"✔ Đã lưu: {self.rec.get('title')}   •   {rio.now_str('%H:%M:%S')}"
        if self.v_autopdf.get():
            try:
                p = os.path.join(rio.storage_dir(), 'pdf', _safe_name(self.rec.get('title')) + '.pdf')
                rio.export_pdf(self.rec, p)
                self.last_pdf = p
                msg += f'   •   PDF: {os.path.basename(p)} (thư mục BaoCaoNVL\\pdf)'
            except Exception as e:
                msg += f'   •   Không xuất được PDF: {e}'
        self._update_state()
        self.lbl_state.config(text=msg, fg='#2d6a4f')
        return True

    # ------------------------------------------------------------ du lieu dang ve
    def _form(self):
        return self.tpl if self.design.get() else (self.rec['form'] if self.rec else None)

    def _values(self):
        return {} if self.design.get() else (self.rec.get('values') if self.rec else {})

    def _images(self):
        if self.design.get():
            return [dict(im, tpl=i) for i, im in enumerate(self.tpl.get('images', []))]
        return rio.all_images(self.rec) if self.rec else []

    # ------------------------------------------------------------ ve
    def _set_zoom(self, z):
        self.zoom = max(0.3, min(2.5, z))
        self.v_zoom.set(f'{int(round(self.zoom * 100))}%')
        self.settings['zoom'] = self.zoom
        self._save_settings()
        self._imgcache.clear()
        self.redraw()

    def _font(self, st, scale=1.0):
        fam = (self._form() or {}).get('font') or 'Times New Roman'
        px = max(6, int(round(float(st.get('sz') or 11) * 96 / 72 * self.zoom * scale)))
        return tkfont.Font(family=fam, size=-px, weight='bold' if st.get('b') else 'normal',
                           slant='italic' if st.get('i') else 'roman')

    def redraw(self):
        if self._editor:
            self._commit_editor()
        cv = self.cv
        cv.delete('all')
        form = self._form()
        if not form:
            return
        fm = FormModel(form)
        self.fm = fm
        # Giãn cột/hàng khi hiển thị trên màn hình cho ô dữ liệu rộng rãi, dễ nhập.
        # (chỉ ảnh hưởng xem trên app; PDF/Excel xuất ra vẫn giữ kích thước gốc)
        cs = float(self.settings.get('grid_col_scale', 1.18))   # giãn ngang
        rs = float(self.settings.get('grid_row_scale', 1.15))   # giãn dọc
        xs, ys = rio.geometry(form, px_col=7.0 * cs, px_pad=5.0 * cs,
                              row_k=(96 / 72) * rs, zoom=self.zoom)
        M = 24
        self.xs = [x + M for x in xs]
        self.ys = [y + M for y in ys]
        X = lambda c: self.xs[min(max(c, 1), len(self.xs)) - 1]      # noqa
        Y = lambda r: self.ys[min(max(r, 1), len(self.ys)) - 1]      # noqa
        W, H = self.xs[-1], self.ys[-1]
        cv.create_rectangle(M - 6, M - 6, W + 6, H + 6, fill='white', outline='#666')
        design = self.design.get()
        if design:        # luoi mo + nhan hang/cot khi thiet ke
            from openpyxl.utils import get_column_letter as L
            for c in range(1, fm.ncol + 2):
                cv.create_line(X(c), M, X(c), H, fill='#eee')
            for r in range(1, fm.nrow + 2):
                cv.create_line(M, Y(r), W, Y(r), fill='#eee')
            for c in range(1, fm.ncol + 1):
                cv.create_text((X(c) + X(c + 1)) / 2, M - 12, text=L(c), fill='#b35c00', font=('Segoe UI', 8, 'bold'))
            for r in range(1, fm.nrow + 1):
                cv.create_text(M - 12, (Y(r) + Y(r + 1)) / 2, text=str(r), fill='#b35c00', font=('Segoe UI', 7))
        values = self._values()
        # nen
        for (r, c), st in fm.cells.items():
            if (r, c) in fm.covered or not st.get('fill'):
                continue
            r0, c0, r1, c1 = fm.area(r, c)
            cv.create_rectangle(X(c0), Y(r0), X(c1 + 1), Y(r1 + 1), fill=st['fill'], outline='')
        # o co danh sach chon / o nhap: to nhe khi nhap lieu
        if not design:
            for k_ in (form.get('lists') or {}):
                r, c = unkey(k_)
                if r <= fm.nrow and (r, c) not in fm.covered:
                    r0, c0, r1, c1 = fm.area(r, c)
                    cv.create_polygon(X(c1 + 1) - 9, Y(r1 + 1) - 1, X(c1 + 1) - 1, Y(r1 + 1) - 1,
                                      X(c1 + 1) - 1, Y(r1 + 1) - 9, fill='#7aa7d6', outline='')
        # anh
        self._photo_refs = []
        for im in self._images():
            self._draw_image(im, X, Y)
        # chu
        for r in range(1, fm.nrow + 1):
            for c in range(1, fm.ncol + 1):
                if (r, c) in fm.covered:
                    continue
                t = fm.text(r, c, values)
                if not t:
                    continue
                st = fm.style(r, c)
                r0, c0, r1, c1 = fm.area(r, c)
                x0, x1, y0, y1 = X(c0), X(c1 + 1), Y(r0), Y(r1 + 1)
                al = st.get('al') or ('right' if re.fullmatch(r'-?[\d.,]+%?', t.strip()) else 'left')
                va = st.get('va') or 'bottom'
                single = (r0 == r1 and c0 == c1)
                wrap = st.get('wrap', 1) or not single
                fnt = self._font(st)
                pad = 6                      # đệm trái/phải trong ô cho thoáng
                w = max(10, x1 - x0 - pad * 2)
                # chu qua cao so voi o -> thu nho (giong PDF)
                if wrap:
                    for sc in (1.0, 0.9, 0.8, 0.7, 0.6):
                        fnt = self._font(st, sc)
                        lines = sum(max(1, -(-fnt.measure(seg) // w)) for seg in t.split('\n'))
                        if lines * fnt.metrics('linespace') <= (y1 - y0) + 2:
                            break
                if al in ('center', 'centerContinuous', 'justify', 'distributed'):
                    ax, anchor_h, just = (x0 + x1) / 2, '', 'center'
                elif al == 'right':
                    ax, anchor_h, just = x1 - pad, 'e', 'right'
                else:
                    ax, anchor_h, just = x0 + pad, 'w', 'left'
                if va == 'top':
                    ay, anchor_v = y0 + 2, 'n'
                elif va == 'center':
                    ay, anchor_v = (y0 + y1) / 2, ''
                else:
                    ay, anchor_v = y1 - 2, 's'
                anchor = (anchor_v + anchor_h) or 'center'
                kw = dict(text=t, font=fnt, fill=st.get('color') or '#000', anchor=anchor, justify=just)
                if wrap:
                    kw['width'] = w
                cv.create_text(ax, ay, **kw)
        # vien
        for (r, c, side) in fm.border_lines():
            xa, xb, ya, yb = X(c), X(c + 1), Y(r), Y(r + 1)
            if side == 0:
                cv.create_line(xa, ya, xa, yb, fill='#000')
            elif side == 1:
                cv.create_line(xa, ya, xb, ya, fill='#000')
            elif side == 2:
                cv.create_line(xb, ya, xb, yb, fill='#000')
            else:
                cv.create_line(xa, yb, xb, yb, fill='#000')
        if design:
            for nm, cell, col in (('Tên NVL', form.get('name_cell'), '#2d6a4f'), ('Ngày KQ', form.get('date_cell'), '#7b2cbf'),
                                  ('Kết luận', form.get('concl_cell'), '#5a4b8a')):
                if cell:
                    r0, c0, r1, c1 = fm.area(*cell)
                    cv.create_rectangle(X(c0) + 2, Y(r0) + 2, X(c1 + 1) - 2, Y(r1 + 1) - 2, outline=col, dash=(3, 2), width=2)
                    cv.create_text(X(c1 + 1) - 4, Y(r0) + 3, text=nm, anchor='ne', fill=col, font=('Segoe UI', 7, 'bold'))
        cv.configure(scrollregion=(0, 0, W + M, H + M))
        self._draw_sel()

    def _draw_image(self, im, X, Y):
        p = rio.resolve_image(im.get('file'))
        x0, y0, x1, y1 = X(im['c0']), Y(im['r0']), X(im['c1'] + 1), Y(im['r1'] + 1)
        if not p:
            self.cv.create_rectangle(x0 + 2, y0 + 2, x1 - 2, y1 - 2, outline='#c00', dash=(3, 3))
            self.cv.create_text((x0 + x1) / 2, (y0 + y1) / 2, text='(mất ảnh)', fill='#c00')
            return
        try:
            from PIL import Image, ImageTk
            bw, bh = max(4, x1 - x0 - 4), max(4, y1 - y0 - 4)
            ck = (p, int(bw), int(bh))
            ph = self._imgcache.get(ck)
            if ph is None:
                pil = Image.open(p)
                pil.load()
                w, h, dx, dy = rio.fit_box(pil.width, pil.height, bw, bh)
                pil = pil.convert('RGBA').resize((max(1, int(w)), max(1, int(h))), Image.LANCZOS)
                ph = (ImageTk.PhotoImage(pil), dx, dy)
                self._imgcache[ck] = ph
            photo, dx, dy = ph
            self._photo_refs.append(photo)
            self.cv.create_image(x0 + 2 + dx, y0 + 2 + dy, image=photo, anchor='nw',
                                 tags=('img',))
        except Exception:
            pass

    def _draw_sel(self):
        self.cv.delete('sel')
        if not self.sel or not hasattr(self, 'xs'):
            return
        r0, c0, r1, c1 = self.sel
        x0, y0 = self.xs[c0 - 1], self.ys[r0 - 1]
        x1, y1 = self.xs[min(c1, len(self.xs) - 1)], self.ys[min(r1, len(self.ys) - 1)]
        self.cv.create_rectangle(x0, y0, x1, y1, outline=SEL, width=3, tags='sel')
        from openpyxl.utils import get_column_letter as L
        a = f'{L(c0)}{r0}' + (f':{L(c1)}{r1}' if (r0, c0) != (r1, c1) else '')
        self.cv.create_text(x0 + 2, y0 - 2, text=a, anchor='sw', fill=SEL, font=('Segoe UI', 8, 'bold'), tags='sel')

    # ------------------------------------------------------------ chon o
    def _hit(self, ev):
        if not hasattr(self, 'xs'):
            return None
        x, y = self.cv.canvasx(ev.x), self.cv.canvasy(ev.y)
        import bisect
        c = bisect.bisect_right(self.xs, x)
        r = bisect.bisect_right(self.ys, y)
        if 1 <= c <= self.fm.ncol and 1 <= r <= self.fm.nrow:
            return r, c
        return None

    def _click(self, ev, extend=False):
        self.cv.focus_set()
        if self._editor:
            self._commit_editor()
        hit = self._hit(ev)
        if not hit:
            return
        r0, c0, r1, c1 = self.fm.area(*hit)
        if extend and self.sel:
            a = self.sel
            self.sel = (min(a[0], r0), min(a[1], c0), max(a[2], r1), max(a[3], c1))
        else:
            self.sel = (r0, c0, r1, c1)
            self._anchor = (r0, c0)
        self._draw_sel()

    def _drag(self, ev):
        hit = self._hit(ev)
        if not hit or not getattr(self, '_anchor', None):
            return
        a0 = self.fm.area(*self._anchor)
        b0 = self.fm.area(*hit)
        self.sel = (min(a0[0], b0[0]), min(a0[1], b0[1]), max(a0[2], b0[2]), max(a0[3], b0[3]))
        self._draw_sel()

    def _dbl(self, ev):
        hit = self._hit(ev)
        if not hit:
            return
        self.sel = self.fm.area(*hit)
        if self.design.get():
            self.d_edit_cell()
        else:
            self._start_edit()

    def _key(self, ev):
        if not self.sel:
            return
        ks = ev.keysym
        r0, c0, r1, c1 = self.sel
        mv = {'Up': (-1, 0), 'Down': (1, 0), 'Left': (0, -1), 'Right': (0, 1), 'Tab': (0, 1)}
        if ks in mv:
            dr, dc = mv[ks]
            r = r0 + dr if dr < 0 else (r1 + dr if dr > 0 else r0)
            c = c0 + dc if dc < 0 else (c1 + dc if dc > 0 else c0)
            if 1 <= r <= self.fm.nrow and 1 <= c <= self.fm.ncol:
                self.sel = self.fm.area(r, c)
                self._draw_sel()
                self._see_sel()
            return 'break'
        if ks in ('Return', 'F2', 'KP_Enter'):
            if self.design.get():
                self.d_edit_cell()
            else:
                self._start_edit()
            return 'break'
        if ks in ('Delete', 'BackSpace') and not self.design.get():
            self._clear_sel()
            return 'break'
        if ev.char and ev.char.isprintable() and not (ev.state & 0x4) and not self.design.get():
            self._start_edit(initial=ev.char)
            return 'break'

    def _see_sel(self):
        if not self.sel:
            return
        r0, c0 = self.sel[:2]
        H = self.ys[-1] + 24
        W = self.xs[-1] + 24
        y = self.ys[r0 - 1]
        x = self.xs[c0 - 1]
        top, bot = self.cv.canvasy(0), self.cv.canvasy(self.cv.winfo_height())
        if y < top or y > bot - 40:
            self.cv.yview_moveto(max(0, (y - 60) / H))
        left, right = self.cv.canvasx(0), self.cv.canvasx(self.cv.winfo_width())
        if x < left or x > right - 60:
            self.cv.xview_moveto(max(0, (x - 60) / W))

    def _bind_wheel(self, e):
        self.cv.bind_all('<MouseWheel>', self._wheel)
        self.cv.bind_all('<Shift-MouseWheel>', lambda ev: self.cv.xview_scroll(int(-ev.delta / 120), 'units'))
        self.cv.bind_all('<Control-MouseWheel>', lambda ev: self._set_zoom(self.zoom + (0.1 if ev.delta > 0 else -0.1)))
        self.cv.bind_all('<Button-4>', lambda ev: self.cv.yview_scroll(-3, 'units'))
        self.cv.bind_all('<Button-5>', lambda ev: self.cv.yview_scroll(3, 'units'))

    def _unbind_wheel(self, e):
        for k in ('<MouseWheel>', '<Shift-MouseWheel>', '<Control-MouseWheel>', '<Button-4>', '<Button-5>'):
            self.cv.unbind_all(k)

    def _wheel(self, ev):
        self.cv.yview_scroll(int(-ev.delta / 40), 'units')

    # ------------------------------------------------------------ NHAP LIEU
    def _push_undo(self):
        if not self.rec:
            return
        self._undo.append((copy.deepcopy(self.rec.get('values', {})), copy.deepcopy(self.rec.get('images', [])),
                           list(self.rec.get('hide_tpl', []))))
        del self._undo[:-60]

    def undo(self):
        if self.design.get() or not self._undo or not self.rec:
            return
        v, im, h = self._undo.pop()
        self.rec['values'], self.rec['images'], self.rec['hide_tpl'] = v, im, h
        self._auto_title()
        self._mark()
        self.redraw()

    def _start_edit(self, initial=None):
        if not self.rec or not self.sel or self.design.get():
            return
        r0, c0, r1, c1 = self.fm.area(self.sel[0], self.sel[1])
        x0, y0, x1, y1 = self.xs[c0 - 1], self.ys[r0 - 1], self.xs[c1], self.ys[r1]
        cur = self.fm.text(r0, c0, self.rec['values'])
        items = (self.rec['form'].get('lists') or {}).get(key(r0, c0))
        st = self.fm.style(r0, c0)
        fnt = self._font(st)
        if items:
            var = tk.StringVar(value=cur)
            w = ttk.Combobox(self.cv, textvariable=var, values=items, font=fnt)
            w._var = var
            if initial:
                var.set(initial)
            w.bind('<<ComboboxSelected>>', lambda e: self._commit_editor())
            w.bind('<Return>', lambda e: (self._commit_editor(), 'break')[1])
        else:
            w = tk.Text(self.cv, font=fnt, wrap='word', bg='#fffde7', relief='solid', bd=1, undo=True)
            w.insert('1.0', initial if initial is not None else cur)
            w.bind('<Return>', lambda e: (self._commit_editor(move=True), 'break')[1])
            w.bind('<KP_Enter>', lambda e: (self._commit_editor(move=True), 'break')[1])
            w.bind('<Shift-Return>', lambda e: None)
            w.bind('<Alt-Return>', lambda e: (w.insert('insert', '\n'), 'break')[1])
            w.bind('<Tab>', lambda e: (self._commit_editor(), self._key(_Ev('Right')), 'break')[2])
        w.bind('<Escape>', lambda e: (self._cancel_editor(), 'break')[1])
        h = max(y1 - y0, fnt.metrics('linespace') + 8)
        self._editor = (w, r0, c0, self.cv.create_window(x0 + 1, y0 + 1, window=w, anchor='nw',
                                                          width=max(40, x1 - x0 - 2), height=h - 2))
        w.focus_set()
        if isinstance(w, tk.Text):
            w.mark_set('insert', 'end-1c')
            w.see('insert')

    def _cancel_editor(self):
        if self._editor:
            w, r, c, wid = self._editor
            self._editor = None
            self.cv.delete(wid)
            w.destroy()
            self.cv.focus_set()

    def _commit_editor(self, move=False):
        if not self._editor:
            return
        w, r, c, wid = self._editor
        self._editor = None
        val = w._var.get() if hasattr(w, '_var') else w.get('1.0', 'end').rstrip('\n')
        self.cv.delete(wid)
        w.destroy()
        self._set_value(r, c, val)
        if move and self.sel:
            rr = self.fm.area(r, c)[2] + 1
            if rr <= self.fm.nrow:
                self.sel = self.fm.area(rr, c)
        self.redraw()
        self.cv.focus_set()

    def _set_value(self, r, c, val):
        if not self.rec:
            return
        old = self.fm.text(r, c, self.rec['values'])
        if val == old:
            return
        self._push_undo()
        k = key(r, c)
        tpl_txt = FormModel(self.rec['form']).text(r, c, {})
        if val == tpl_txt:
            self.rec['values'].pop(k, None)
        else:
            self.rec['values'][k] = val
        f = self.rec['form']
        # Worksheet_Change: nhap Ten NVL -> ghi Ngay tra KQ = Now; xoa ten -> xoa ngay
        if f.get('name_cell') and [r, c] == list(f['name_cell']) and f.get('date_cell'):
            dk = key(*f['date_cell'])
            self.rec['values'][dk] = rio.now_str() if val.strip() else ''
        self._auto_title()
        self._mark()
        # Tu dong ghi KET LUAN khi nhap du lieu (tru khi dang ghi chinh o ket luan)
        if not getattr(self, '_in_concl', False):
            try:
                self._auto_conclude(create=True)
            except Exception:
                pass

    def _clear_sel(self):
        if not self.rec or not self.sel:
            return
        self._push_undo()
        r0, c0, r1, c1 = self.sel
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                if (r, c) in self.fm.covered:
                    continue
                if self.fm.text(r, c, self.rec['values']):
                    self.rec['values'][key(r, c)] = ''
        self._auto_title()
        self._mark()
        self.redraw()

    def _copy_text(self):
        if self.sel:
            t = self.fm.text(self.sel[0], self.sel[1], self._values())
            self.clipboard_clear()
            self.clipboard_append(t)

    # ------------------------------------------------------------ ANH
    def _target_area(self):
        if not self.sel:
            messagebox.showinfo('Chọn ô', 'Bấm chọn ô (vùng) cần chèn ảnh trước.', parent=self)
            return None
        return self.sel

    def _add_images(self, sources, tpl=False):
        area = self._target_area()
        if not area:
            return
        r0, c0, r1, c1 = area
        n = len(sources)
        if not tpl:
            self._push_undo()
        # nhieu anh -> chia deu theo chieu ngang vung chon (neu du cot), khong thi chong vao cung vung
        cols = list(range(c0, c1 + 1))
        for i, src in enumerate(sources):
            try:
                name = rio.import_image(src, sub='form_media' if tpl else 'images')
            except Exception as e:
                messagebox.showerror('Lỗi ảnh', f'Không đọc được ảnh:\n{e}', parent=self)
                continue
            if n > 1 and len(cols) >= n:
                per = len(cols) // n
                a = cols[i * per]
                b = cols[(i + 1) * per - 1] if i < n - 1 else c1
            else:
                a, b = c0, c1
            item = {'r0': r0, 'c0': a, 'r1': r1, 'c1': b, 'file': name}
            if tpl:
                self.tpl.setdefault('images', []).append(item)
                self.tpl_dirty = True
            else:
                self.rec.setdefault('images', []).append(item)
        if tpl:
            self._update_state()
        else:
            self._mark()
        self.redraw()

    def insert_image_file(self):
        if self.design.get():
            return self.d_tpl_image()
        if not self.rec or not self._target_area():
            return
        ps = filedialog.askopenfilenames(parent=self, title='Chọn ảnh chèn vào ô',
                                         filetypes=[('Ảnh', '*.png *.jpg *.jpeg *.bmp *.gif *.tif *.webp')])
        if ps:
            self._add_images(list(ps))

    def paste_image(self, from_key=False):
        if not self.rec and not self.design.get():
            return
        imgs = []
        try:
            from paste_image import clipboard_images
            imgs = [src for _n, src in clipboard_images(self)]
        except Exception:
            try:
                from PIL import ImageGrab, Image
                d = ImageGrab.grabclipboard()
                if isinstance(d, Image.Image):
                    imgs = [d]
                elif isinstance(d, (list, tuple)):
                    imgs = [p for p in d if str(p).lower().endswith(rio.IMG_EXT)]
            except Exception:
                imgs = []
        if imgs:
            self._add_images(imgs, tpl=self.design.get())
            return
        if from_key and self.sel and not self.design.get():      # clipboard la chu -> dan vao o
            try:
                t = self.clipboard_get()
            except Exception:
                t = ''
            if t:
                self._set_value(self.sel[0], self.sel[1], t.rstrip('\n'))
                self.redraw()
            return
        messagebox.showinfo('Không có ảnh', 'Clipboard chưa có ảnh.\nChụp màn hình (Win+Shift+S) hoặc Copy ảnh '
                                            'rồi bấm lại (hoặc Ctrl+V trên form).', parent=self)

    def _on_drop(self, files):
        paths = []
        for f in files:
            try:
                f = f.decode('gbk') if isinstance(f, bytes) else f
            except Exception:
                f = f.decode('utf-8', 'ignore') if isinstance(f, bytes) else f
            if str(f).lower().endswith(rio.IMG_EXT):
                paths.append(f)
        if paths:
            self.after(10, lambda: self._add_images(paths, tpl=self.design.get()))

    def _imgs_in_sel(self):
        if not self.sel:
            return []
        r0, c0, r1, c1 = self.sel
        out = []
        for im in self._images():
            if not (im['r1'] < r0 or im['r0'] > r1 or im['c1'] < c0 or im['c0'] > c1):
                out.append(im)
        return out

    def delete_image(self):
        if self.design.get():
            return self.d_del_tpl_image()
        if not self.rec:
            return
        ims = self._imgs_in_sel()
        if not ims:
            messagebox.showinfo('Không có ảnh', 'Chọn ô có ảnh cần xoá.', parent=self)
            return
        self._push_undo()
        drop = {im['idx'] for im in ims if im.get('tpl') is None}
        self.rec['images'] = [im for j, im in enumerate(self.rec.get('images', [])) if j not in drop]
        hide = set(self.rec.get('hide_tpl') or [])
        hide |= {im['tpl'] for im in ims if im.get('tpl') is not None}
        self.rec['hide_tpl'] = sorted(hide)
        self._mark()
        self.redraw()

    # ------------------------------------------------------------ menu chuot phai
    def _ctx_menu(self, ev):
        hit = self._hit(ev)
        if hit and (not self.sel or not (self.sel[0] <= hit[0] <= self.sel[2] and self.sel[1] <= hit[1] <= self.sel[3])):
            self.sel = self.fm.area(*hit)
            self._draw_sel()
        m = tk.Menu(self, tearoff=0)
        if self.design.get():
            m.add_command(label='✏ Sửa ô / định dạng…', command=self.d_edit_cell)
            m.add_command(label='⊞ Gộp vùng chọn', command=self.d_merge)
            m.add_command(label='⊟ Tách ô', command=self.d_unmerge)
            m.add_separator()
            m.add_command(label='➕ Chèn hàng phía trên', command=lambda: self.d_rows(0, 1))
            m.add_command(label='➕ Chèn hàng phía dưới', command=lambda: self.d_rows(1, 1))
            m.add_command(label='➖ Xoá hàng đang chọn', command=lambda: self.d_rows(0, -1))
            m.add_command(label='➕ Chèn cột bên trái', command=lambda: self.d_cols(0, 1))
            m.add_command(label='➕ Chèn cột bên phải', command=lambda: self.d_cols(1, 1))
            m.add_command(label='➖ Xoá cột đang chọn', command=lambda: self.d_cols(0, -1))
            m.add_separator()
            m.add_command(label='🖼 Chèn ảnh mẫu vào vùng…', command=self.d_tpl_image)
            m.add_command(label='📋 Dán ảnh mẫu (Ctrl+V)', command=self.paste_image)
            m.add_command(label='✖ Xoá ảnh mẫu trong vùng', command=self.d_del_tpl_image)
        else:
            m.add_command(label='✏ Nhập / sửa ô (Enter)', command=self._start_edit)
            m.add_command(label='⌫ Xoá nội dung (Delete)', command=self._clear_sel)
            m.add_command(label='↺ Trả ô về nội dung mẫu', command=self._reset_sel)
            m.add_separator()
            m.add_command(label='🖼 Chèn ảnh từ file vào ô…', command=self.insert_image_file)
            m.add_command(label='📋 Dán ảnh vào ô (Ctrl+V)', command=self.paste_image)
            m.add_command(label='✖ Xoá ảnh trong ô', command=self.delete_image)
            m.add_separator()
            m.add_command(label='✂ Cắt ảnh trong ô (để chuyển)', command=self.cut_image)
            m.add_command(label='📌 Dán ảnh đã cắt vào ô này', command=self.paste_cut_image,
                          state=('normal' if self._imgbuf else 'disabled'))
        m.tk_popup(ev.x_root, ev.y_root)

    def _reset_sel(self):
        if not self.rec or not self.sel:
            return
        self._push_undo()
        r0, c0, r1, c1 = self.sel
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                self.rec['values'].pop(key(r, c), None)
        self._auto_title()
        self._mark()
        self.redraw()

    # ------------------------------------------------------------ XUAT
    def _rec_for_output(self):
        if self._editor:
            self._commit_editor()
        if self.design.get():
            return {'form': self.tpl, 'values': {}, 'images': [], 'title': self.tpl.get('name')}
        if not (self.rec.get('title') or '').strip():
            self._auto_title()
        return self.rec

    def export_pdf(self):
        if not self.rec and not self.design.get():
            return
        rec = self._rec_for_output()
        p = filedialog.asksaveasfilename(parent=self, defaultextension='.pdf', filetypes=[('PDF', '*.pdf')],
                                         initialfile=_safe_name(rec.get('title')) + '.pdf')
        if not p:
            return
        try:
            rio.export_pdf(rec, p)
        except Exception as e:
            messagebox.showerror('Lỗi xuất PDF', f'{e}\n\n(Cần: pip install reportlab pillow)', parent=self)
            return
        if messagebox.askyesno('Xong', f'Đã xuất PDF:\n{p}\n\nMở ngay?', parent=self):
            _open_file(p)

    def export_excel(self):
        if not self.rec and not self.design.get():
            return
        rec = self._rec_for_output()
        p = filedialog.asksaveasfilename(parent=self, defaultextension='.xlsx', filetypes=[('Excel', '*.xlsx')],
                                         initialfile=_safe_name(rec.get('title')) + '.xlsx')
        if not p:
            return
        try:
            rio.export_excel(rec, p)
        except Exception as e:
            messagebox.showerror('Lỗi xuất Excel', f'{e}\n\n(Cần: pip install openpyxl pillow)', parent=self)
            return
        if messagebox.askyesno('Xong', f'Đã xuất Excel:\n{p}\n\nMở ngay?', parent=self):
            _open_file(p)

    def print_report(self):
        if not self.rec and not self.design.get():
            return
        rec = self._rec_for_output()
        p = os.path.join(tempfile.gettempdir(), _safe_name(rec.get('title')) + '_in.pdf')
        try:
            rio.export_pdf(rec, p)
        except Exception as e:
            messagebox.showerror('Lỗi', str(e), parent=self)
            return
        if sys.platform.startswith('win'):
            if messagebox.askyesno('In', 'Gửi thẳng tới máy in mặc định?\n(Không = mở PDF để xem rồi in)', parent=self):
                _open_file(p, 'print')
                return
        _open_file(p)

    # ------------------------------------------------------------ THIET KE FORM (ADMIN)
    def _toggle_design(self, silent=False):
        if self._editor:
            self._commit_editor()
        on = self.design.get()
        if on:
            fid = self._cur_form_id()
            f = self.forms.get(fid) if fid else None
            if not f:
                self.design.set(False)
                messagebox.showinfo('Chọn form', 'Mở 1 loại báo cáo trước khi thiết kế.', parent=self)
                return
            self.tpl = copy.deepcopy(f)
            self.tpl_dirty = False
            self.dbar.pack(fill='x', before=self.hint)
            self.hint.config(text='THIẾT KẾ: bấm / kéo chọn ô • nhấp đúp = sửa ô • Shift+bấm chọn vùng • '
                                  'chuột phải = menu • "💾 Lưu mẫu form" để áp dụng cho báo cáo mới')
        else:
            if self.tpl_dirty and not silent:
                a = messagebox.askyesnocancel('Lưu mẫu?', 'Lưu thay đổi của mẫu form?', parent=self)
                if a is None:
                    self.design.set(True)
                    return
                if a:
                    self.design.set(True)
                    self.d_save()
                    self.design.set(False)
            self.tpl = None
            self.tpl_dirty = False
            self.dbar.pack_forget()
            self.hint.config(text='Bấm ô để chọn • Nhấp đúp / gõ phím để nhập • Enter lưu ô, Shift+Enter xuống '
                                  'dòng • Ctrl+V dán ảnh vào ô • Shift+bấm = chọn vùng • Chuột phải = menu')
        self.sel = None
        self._imgcache.clear()
        self.redraw()
        self._update_state()

    def _need_design(self):
        if not self.design.get() or not self.tpl:
            return False
        if not self.sel:
            messagebox.showinfo('Chọn ô', 'Bấm chọn ô / vùng trên form trước.', parent=self)
            return False
        return True

    def _tpl_changed(self):
        self.tpl_dirty = True
        self._update_state()
        self.redraw()

    def _tpl_cell(self, r, c):
        for cell in self.tpl['cells']:
            if cell['r'] == r and cell['c'] == c:
                return cell
        cell = {'r': r, 'c': c, 'b': 0, 'i': 0, 'sz': 12.0, 'color': '#000000', 'al': 'left', 'va': 'center',
                'wrap': 1, 'fill': '', 'bd': [0, 0, 0, 0]}
        self.tpl['cells'].append(cell)
        return cell

    def d_edit_cell(self):
        if not self._need_design():
            return
        r, c = self.sel[0], self.sel[1]
        area = self.sel

        def ok(text, ref, style, items, whole):
            cell = self._tpl_cell(r, c)
            cell.pop('v', None)
            cell.pop('ref', None)
            cell.pop('nf', None)
            if ref:
                cell['ref'] = ref
            elif text != '':
                t = text
                if re.fullmatch(r'-?\d+', t):
                    cell['v'] = int(t)
                elif re.fullmatch(r'-?\d+\.\d+', t):
                    cell['v'] = float(t)
                else:
                    cell['v'] = t
            targets = [(rr, cc) for rr in range(area[0], area[2] + 1) for cc in range(area[1], area[3] + 1)] \
                if whole else [(r, c)]
            for (rr, cc) in targets:
                tc = self._tpl_cell(rr, cc)
                if whole and (rr, cc) != (r, c):
                    st = dict(style)
                    # vien: chi ve mep ngoai vung khi ap dung ca vung
                    st['bd'] = [style['bd'][0] if cc == area[1] else tc.get('bd', [0] * 4)[0],
                                style['bd'][1] if rr == area[0] else tc.get('bd', [0] * 4)[1],
                                style['bd'][2] if cc == area[3] else tc.get('bd', [0] * 4)[2],
                                style['bd'][3] if rr == area[2] else tc.get('bd', [0] * 4)[3]]
                    tc.update(st)
                else:
                    tc.update(style)
            lists = self.tpl.setdefault('lists', {})
            if items:
                lists[key(r, c)] = items
            else:
                lists.pop(key(r, c), None)
            self._tpl_changed()
        CellDesignDialog(self, self.tpl, r, c, ok)

    def d_copy_cell(self):
        """Copy toan bo dinh dang + noi dung mau cua o dang chon."""
        if not self._need_design():
            return
        r, c = self.sel[0], self.sel[1]
        src = dict(self._tpl_cell(r, c))
        src.pop('r', None)
        src.pop('c', None)
        lst = (self.tpl.get('lists') or {}).get(key(r, c))
        self._cellbuf = {'cell': src, 'list': list(lst) if lst else None}
        self._update_state()
        messagebox.showinfo('Copy ô', f'Đã copy định dạng ô (hàng {r}, cột {c}).\n'
                            'Chọn ô/vùng khác rồi bấm "📋 Paste ô".', parent=self)

    def d_paste_cell(self):
        """Paste dinh dang + noi dung da copy vao CA vung dang chon."""
        if not self._need_design():
            return
        if not self._cellbuf:
            messagebox.showinfo('Paste ô', 'Chưa copy ô nào. Chọn ô rồi bấm "⧉ Copy ô" trước.', parent=self)
            return
        src = self._cellbuf['cell']
        r0, c0, r1, c1 = self.sel
        lists = self.tpl.setdefault('lists', {})
        for rr in range(r0, r1 + 1):
            for cc in range(c0, c1 + 1):
                tc = self._tpl_cell(rr, cc)
                for kk in ('v', 'ref', 'nf'):
                    tc.pop(kk, None)
                for kk, vv in src.items():
                    tc[kk] = copy.deepcopy(vv)
                tc['r'], tc['c'] = rr, cc
                if self._cellbuf.get('list'):
                    lists[key(rr, cc)] = list(self._cellbuf['list'])
                else:
                    lists.pop(key(rr, cc), None)
        self._tpl_changed()

    def cut_image(self):
        """Cat anh (da chen) trong o dang chon vao bo nho - de chuyen sang o dung."""
        if self.design.get() or not self.rec:
            return
        ims = [im for im in self._imgs_in_sel() if im.get('tpl') is None]
        if not ims:
            messagebox.showinfo('Không có ảnh', 'Chọn ô có ảnh (ảnh bạn đã chèn) cần chuyển.', parent=self)
            return
        self._push_undo()
        self._imgbuf = [{'file': im['file']} for im in ims]
        drop = {im['idx'] for im in ims}
        self.rec['images'] = [im for j, im in enumerate(self.rec.get('images', [])) if j not in drop]
        self._mark()
        self.redraw()

    def paste_cut_image(self):
        """Dan anh da cat vao o dang chon (chuyen anh add nham ve dung cho)."""
        if self.design.get() or not self.rec:
            return
        if not self._imgbuf:
            messagebox.showinfo('Dán ảnh', 'Chưa cắt ảnh nào.\nChọn ô có ảnh sai → "✂ Cắt ảnh", '
                                'rồi chọn ô đúng → "📌 Dán ảnh".', parent=self)
            return
        if not self.sel:
            return
        self._push_undo()
        r0, c0, r1, c1 = self.sel
        for b in self._imgbuf:
            self.rec.setdefault('images', []).append(
                {'r0': r0, 'c0': c0, 'r1': r1, 'c1': c1, 'file': b['file']})
        self._imgbuf = None
        self._mark()
        self.redraw()

    def d_merge(self):
        if not self._need_design():
            return
        r0, c0, r1, c1 = self.sel
        if (r0, c0) == (r1, c1):
            messagebox.showinfo('Gộp ô', 'Kéo chuột (hoặc Shift+bấm) để chọn NHIỀU ô rồi bấm Gộp.', parent=self)
            return
        ms = [m for m in self.tpl['merges'] if m[2] < r0 or m[0] > r1 or m[3] < c0 or m[1] > c1]
        ms.append([r0, c0, r1, c1])
        self.tpl['merges'] = ms
        self._tpl_changed()

    def d_unmerge(self):
        if not self._need_design():
            return
        r0, c0, r1, c1 = self.sel
        self.tpl['merges'] = [m for m in self.tpl['merges'] if m[2] < r0 or m[0] > r1 or m[3] < c0 or m[1] > c1]
        self.sel = (r0, c0, r0, c0)
        self._tpl_changed()

    def d_rows(self, below, n):
        if not self._need_design():
            return
        r0, c0, r1, c1 = self.sel
        if n < 0:
            cnt = r1 - r0 + 1
            if cnt >= self.tpl['nrow']:
                return
            if not messagebox.askyesno('Xoá hàng', f'Xoá {cnt} hàng ({r0}–{r1})?', parent=self):
                return
            rio.shift_form(self.tpl, 'r', r0, -cnt)
            self.sel = None
        else:
            at = r1 + 1 if below else r0
            rio.shift_form(self.tpl, 'r', at, 1)
            self.sel = (at, c0, at, c0)
        self._tpl_changed()

    def d_cols(self, right, n):
        if not self._need_design():
            return
        r0, c0, r1, c1 = self.sel
        if n < 0:
            cnt = c1 - c0 + 1
            if cnt >= self.tpl['ncol']:
                return
            if not messagebox.askyesno('Xoá cột', f'Xoá {cnt} cột đang chọn?', parent=self):
                return
            rio.shift_form(self.tpl, 'c', c0, -cnt)
            self.sel = None
        else:
            at = c1 + 1 if right else c0
            rio.shift_form(self.tpl, 'c', at, 1)
            self.sel = (r0, at, r0, at)
        self._tpl_changed()

    def d_row_height(self):
        if not self._need_design():
            return
        r0, r1 = self.sel[0], self.sel[2]
        cur = self.tpl['rowh'][r0 - 1]
        v = simpledialog.askfloat('Cao hàng', f'Chiều cao hàng {r0}–{r1} (point, Excel):', initialvalue=cur,
                                  minvalue=3, maxvalue=600, parent=self)
        if v:
            for r in range(r0, r1 + 1):
                self.tpl['rowh'][r - 1] = v
            self._tpl_changed()

    def d_col_width(self):
        if not self._need_design():
            return
        c0, c1 = self.sel[1], self.sel[3]
        cur = self.tpl['cols'][c0 - 1]
        v = simpledialog.askfloat('Rộng cột', f'Độ rộng cột (đơn vị Excel):', initialvalue=cur,
                                  minvalue=0.5, maxvalue=200, parent=self)
        if v:
            for c in range(c0, c1 + 1):
                self.tpl['cols'][c - 1] = v
            self._tpl_changed()

    def d_tpl_image(self):
        if not self._need_design():
            return
        ps = filedialog.askopenfilenames(parent=self, title='Ảnh mẫu (hiện trên mọi báo cáo loại này)',
                                         filetypes=[('Ảnh', '*.png *.jpg *.jpeg *.bmp *.gif *.tif *.webp')])
        if ps:
            self._add_images(list(ps), tpl=True)

    def d_del_tpl_image(self):
        if not self._need_design():
            return
        ims = self._imgs_in_sel()
        if not ims:
            messagebox.showinfo('Không có ảnh', 'Vùng chọn không có ảnh mẫu.', parent=self)
            return
        drop = {im['tpl'] for im in ims}
        self.tpl['images'] = [im for i, im in enumerate(self.tpl.get('images', [])) if i not in drop]
        self._tpl_changed()

    def d_name_cell(self):
        if not self._need_design():
            return
        self.tpl['name_cell'] = [self.sel[0], self.sel[1]]
        self._tpl_changed()

    def d_date_cell(self):
        if not self._need_design():
            return
        self.tpl['date_cell'] = [self.sel[0], self.sel[1]]
        self._tpl_changed()

    @staticmethod
    def _merge_full_row(form, r):
        """Gop ca HANG r thanh 1 o full-width [r,1,r,ncol] + style cho o ket luan."""
        nc = int(form.get('ncol') or len(form.get('cols', [])))
        if nc < 1:
            return
        merges = form.setdefault('merges', [])
        merges[:] = [m for m in merges if not (m[0] <= r <= m[2])]   # go merge cu chong hang
        merges.append([r, 1, r, nc])
        cells = form.setdefault('cells', [])
        anchor = next((cc for cc in cells if cc['r'] == r and cc['c'] == 1), None)
        if anchor is None:
            anchor = {'r': r, 'c': 1}
            cells.append(anchor)
        anchor.update({'al': 'left', 'va': 'center', 'wrap': 1,
                       'sz': anchor.get('sz') or 11.0, 'bd': [1, 1, 1, 1]})

    def d_concl_cell(self):
        """ADMIN: dat o KET LUAN = gop ca hang dang chon thanh 1 o full-width."""
        if not self._need_design():
            return
        r = self.sel[0]
        self._merge_full_row(self.tpl, r)
        self.tpl['concl_cell'] = [r, 1]
        self._tpl_changed()

    def _sign_row(self, fm, values):
        """Tim hang chua 'Nguoi lap bao cao'. None neu khong thay."""
        for r in range(1, fm.nrow + 1):
            for c in range(1, fm.ncol + 1):
                if 'nguoi lap' in _norm_vi(fm.text(r, c, values)):
                    return r
        return None

    def _concl_anchor(self, create=False):
        """Tra (r,c) o Ket luan cua bao cao dang mo.
        create=True: neu chua co -> tu gop full-width hang TRONG ngay tren
        dong 'Nguoi lap bao cao'."""
        if not self.rec:
            return None
        form = self.rec['form']
        fm = getattr(self, 'fm', None) or FormModel(form)
        cell = form.get('concl_cell')
        if cell:
            a = fm.area(cell[0], cell[1])
            return a[0], a[1]
        if not create:
            return None
        values = self.rec.get('values') or {}
        sr = self._sign_row(fm, values)
        if not sr or sr < 2:
            return None
        tr = sr - 1
        if any(fm.text(tr, c, values) for c in range(1, fm.ncol + 1)):
            return None   # hang tren khong trong -> khong tu tao (tranh de len du lieu)
        self._merge_full_row(form, tr)
        form['concl_cell'] = [tr, 1]
        self.fm = FormModel(form)
        return tr, 1

    def _auto_conclude(self, create=False):
        """Ghi/ cap nhat ket luan tu dong. KHONG de len neu nguoi lap da tu sua."""
        if not self.rec or getattr(self, '_in_concl', False):
            return False
        values = self.rec.get('values') or {}
        fm0 = getattr(self, 'fm', None) or FormModel(self.rec['form'])
        # chi tu dong khi co it nhat 1 chi tieu SO so sanh duoc voi tieu chuan
        checked, _ = rio.suggest_conclusion(fm0, values)
        if checked == 0:
            return False
        anchor = self._concl_anchor(create=create)
        if not anchor:
            return False
        r, c = anchor
        fm = getattr(self, 'fm', None) or FormModel(self.rec['form'])
        cur = fm.text(r, c, values)
        last = self.rec.get('_concl_auto')
        # chi ghi khi o dang trong hoac dang la van ban tu dong truoc do
        if cur and cur != last:
            return False
        text = rio.conclusion_text(fm, values)
        if text == cur:
            return True
        self._in_concl = True
        try:
            self.fm = fm
            self._set_value(r, c, text)
            self.rec['_concl_auto'] = text
        finally:
            self._in_concl = False
        return True

    def conclude(self):
        """Nut '⚖ Ket luan': tu dat o ket luan (full-width, tren chu ky) + ghi ket luan."""
        if not self.rec:
            messagebox.showinfo('Kết luận', 'Hãy mở một báo cáo trước.', parent=self)
            return
        anchor = self._concl_anchor(create=True)
        if not anchor:
            messagebox.showinfo(
                'Kết luận',
                'Chưa tạo được ô Kết luận tự động.\n\n'
                'Cách thủ công (Admin): bật "✏ Thiết kế form", chọn HÀNG trống nằm '
                'phía trên dòng "Người lập báo cáo" rồi bấm "🧾 Đặt ô Kết luận".',
                parent=self)
            return
        r, c = anchor
        fm = getattr(self, 'fm', None) or FormModel(self.rec['form'])
        values = self.rec.get('values') or {}
        text = rio.conclusion_text(fm, values)
        self.fm = fm
        self._in_concl = True
        try:
            self._set_value(r, c, text)
            self.rec['_concl_auto'] = text
        finally:
            self._in_concl = False
        self.sel = fm.area(r, c)
        self.redraw()
        _, issues = rio.suggest_conclusion(fm, values)
        if issues:
            messagebox.showwarning(
                'Kết luận',
                f'Phát hiện {len(issues)} chỉ tiêu NGOÀI tiêu chuẩn.\n'
                'Đã ghi gợi ý vào ô Kết luận — người lập kiểm tra và sửa lại nếu cần.',
                parent=self)
        else:
            messagebox.showinfo(
                'Kết luận',
                'Đã ghi kết luận vào ô. Người lập kiểm tra / chỉnh lại nội dung nếu cần.',
                parent=self)

    def d_save(self):
        if not self.tpl:
            return False
        self.forms.update(self.tpl)
        self.tpl_dirty = False
        # bao cao dang mo chua co du lieu -> dung ngay mau moi
        if self.rec and self.rec.get('form_id') == self.tpl.get('id'):
            if not self.rec.get('values') and not self.rec.get('images'):
                self.rec['form'] = copy.deepcopy(self.tpl)
                self.rec['form_name'] = self.tpl.get('name')
            elif messagebox.askyesno('Áp dụng mẫu mới',
                                     'Đã lưu mẫu form.\n\nÁp dụng mẫu mới cho báo cáo ĐANG MỞ?\n'
                                     '(Nội dung đã nhập được giữ theo vị trí ô)', parent=self):
                self.rec['form'] = copy.deepcopy(self.tpl)
                self.rec['hide_tpl'] = []
                self._mark()
        self._fill_tree()
        self._select_form_in_tree(self.tpl.get('id'))
        self._update_state()
        messagebox.showinfo('Đã lưu', f"Đã lưu mẫu form: {self.tpl.get('name')}\n"
                                      'Báo cáo MỚI loại này sẽ dùng mẫu vừa sửa.\n'
                                      'Báo cáo đã lưu trước đó giữ nguyên mẫu cũ.', parent=self)
        return True

    def d_cancel(self):
        if not self.tpl:
            return
        if self.tpl_dirty and not messagebox.askyesno('Bỏ thay đổi', 'Bỏ mọi thay đổi chưa lưu của mẫu?', parent=self):
            return
        self.tpl = copy.deepcopy(self.forms.get(self.tpl['id']))
        self.tpl_dirty = False
        self.redraw()
        self._update_state()

    # ------------------------------------------------------------ QUAN LY FORM
    def _groups(self):
        return self.forms.groups()

    def _open_form_id(self, fid):
        self._fill_tree()
        self._select_form_in_tree(fid)
        self.new_report(fid, ask=False)

    def form_new(self):
        if not self._confirm_discard():
            return
        base_id = self._cur_form_id()
        base = self.forms.get(base_id) if base_id else None
        g = ' > '.join(base.get('group') or []) if base else 'Form tự tạo'

        def ok(name, group, extra):
            if extra.get('base') == 'copy' and base:
                f = copy.deepcopy(base)
                f['note'] = f"Nhân bản từ: {base.get('name')}"
            else:
                f = rio.blank_form(name, group)
            f['name'], f['group'] = name, group
            nf = self.forms.add(f, after_id=base_id)
            self._open_form_id(nf['id'])
            if self.admin:
                self.design.set(True)
                self._toggle_design()
        FormInfoDialog(self, '➕ Thêm form báo cáo mới', self._groups(), '', g, 'new', ok)

    def form_from_excel(self):
        if not self._confirm_discard():
            return

        def ok(name, group, extra):
            try:
                f = rio.form_from_excel(extra['path'], extra['sheet'], extra['rng'], name, group)
            except Exception as e:
                messagebox.showerror('Lỗi nhập', f'Không nhập được form:\n{e}', parent=self)
                return
            nf = self.forms.add(f, after_id=self._cur_form_id())
            self._open_form_id(nf['id'])
            messagebox.showinfo('Đã nhập', f"Đã tạo form '{name}' ({f['nrow']} hàng × {f['ncol']} cột, "
                                           f"{len(f['images'])} ảnh).\nBật '✏ Thiết kế form' để chỉnh thêm "
                                           '(ô Tên NVL, ô ngày, danh sách chọn...).', parent=self)
        FormInfoDialog(self, '📥 Nhập form từ file Excel', self._groups(), '', 'Form tự tạo', 'excel', ok)

    def form_rename(self):
        fid = self._tree_fid() or self._cur_form_id()
        f = self.forms.get(fid) if fid else None
        if not f:
            return

        def ok(name, group, extra):
            f['name'], f['group'] = name, group
            self.forms.update(f)
            if self.rec and self.rec.get('form_id') == fid:
                self.rec['form_name'], self.rec['group'] = name, group
                self.lbl_form.config(text='📑 ' + ' › '.join(group[-1:] + [name]))
            if self.tpl and self.tpl.get('id') == fid:
                self.tpl['name'], self.tpl['group'] = name, group
            self._fill_tree()
            self._select_form_in_tree(fid)
        FormInfoDialog(self, '✏ Đổi tên / nhóm form', self._groups(), f.get('name', ''),
                       ' > '.join(f.get('group') or []), 'rename', ok)

    def form_duplicate(self):
        fid = self._tree_fid() or self._cur_form_id()
        f = self.forms.get(fid) if fid else None
        if not f or not self._confirm_discard():
            return
        nf = copy.deepcopy(f)
        nf['name'] = f.get('name', '') + ' (bản sao)'
        nf['note'] = f"Nhân bản từ: {f.get('name')}"
        nf = self.forms.add(nf, after_id=fid)
        self._open_form_id(nf['id'])

    def form_delete(self):
        fid = self._tree_fid()
        f = self.forms.get(fid) if fid else None
        if not f:
            messagebox.showinfo('Chọn form', 'Chọn form cần xoá trong cây.', parent=self)
            return
        n = sum(1 for r in self.reports.list() if r.get('form_id') == fid)
        if not messagebox.askyesno('Xoá form', f"Xoá form '{f.get('name')}'?\n\n"
                                               f"{n} báo cáo đã lưu của form này VẪN GIỮ (mở/xuất bình thường).",
                                   parent=self):
            return
        self.forms.delete(fid)
        if self.design.get():
            self.design.set(False)
            self._toggle_design(silent=True)
        self._fill_tree()

    def _form_move(self, step):
        fid = self._tree_fid()
        if fid:
            self.forms.move(fid, step)
            self._fill_tree()
            self._select_form_in_tree(fid)

    def form_reset_all(self):
        if not messagebox.askyesno('Khôi phục', 'Bỏ MỌI chỉnh sửa / form tự thêm và dùng lại bộ form gốc '
                                                '(trích từ Excel)?\nBáo cáo đã lưu không bị ảnh hưởng.', parent=self):
            return
        try:
            if os.path.exists(self.forms.path_user):
                bk = self.forms.path_user + '.bak_' + rio.now_str('%Y%m%d_%H%M%S')
                os.replace(self.forms.path_user, bk)
        except Exception as e:
            messagebox.showerror('Lỗi', str(e), parent=self)
            return
        self.forms.load()
        self._fill_tree()
        messagebox.showinfo('Xong', 'Đã khôi phục bộ form gốc (bản cũ được đổi tên .bak_...).', parent=self)

    def pick_storage(self):
        d = filedialog.askdirectory(parent=self, title='Thư mục lưu báo cáo (có thể là thư mục chia sẻ mạng LAN)',
                                    initialdir=rio.storage_dir())
        if not d:
            return
        s = rio.load_settings()
        s['storage_dir'] = d
        rio.save_settings(s)
        self.settings = s
        self.forms = rio.FormStore()
        self.reports = rio.ReportStore()
        self._fill_tree()
        self._fill_reports()
        messagebox.showinfo('Đã đổi', f'Thư mục lưu báo cáo:\n{rio.storage_dir()}', parent=self)

    def _save_settings(self):
        s = rio.load_settings()
        s.update({k: v for k, v in self.settings.items() if k != 'storage_dir'})
        s['auto_pdf'] = bool(self.v_autopdf.get()) if hasattr(self, 'v_autopdf') else s.get('auto_pdf', True)
        rio.save_settings(s)


class _Ev:
    def __init__(self, keysym):
        self.keysym, self.char, self.state = keysym, '', 0


# =====================================================================
#  Chay rieng
# =====================================================================
def main():
    root = tk.Tk()
    root.title('Báo cáo chất lượng nguyên vật liệu - HitechPlast')
    root.geometry('1400x860')
    try:
        ttk.Style(root).theme_use('clam')
    except Exception:
        pass
    tab = ReportFormTab(root, admin=True)
    tab.pack(fill='both', expand=True)

    def on_close():
        if tab._confirm_discard():
            root.destroy()
    root.protocol('WM_DELETE_WINDOW', on_close)
    root.mainloop()


if __name__ == '__main__':
    main()
